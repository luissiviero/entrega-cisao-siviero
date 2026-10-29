"""Build cisao-siviero-esboco-v1.pptx (python-pptx). Numbers are typed from the
sources listed in fontes-dos-numeros.md; see that file for each origin."""
import re
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION

OUT = sys.argv[1]

DARK = RGBColor(0x1F, 0x3A, 0x2E)     # deep forest green
MID = RGBColor(0x4E, 0x7A, 0x5E)      # moss
TINT = RGBColor(0xEC, 0xF1, 0xED)     # card tint
WARN = RGBColor(0xB8, 0x5A, 0x22)     # terracotta: provisório / em aberto
WARN_TINT = RGBColor(0xFA, 0xEE, 0xE6)
TEXT = RGBColor(0x22, 0x2B, 0x27)
MUTED = RGBColor(0x5E, 0x6A, 0x64)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEAD = "Cambria"
BODY = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
prs.core_properties.title = "Cisão Siviero — esboço v1"
prs.core_properties.author = "Luís Siviero (rascunho)"

state = {"n": 0}


def nb(t):
    """Keep currency, figures and units together on one line."""
    t = t.replace("R$ ", "R$\u00a0")
    return re.sub(r"(\d) (mil|mi|ha)\b", "\\1\u00a0\\2", t)


def text(slide, x, y, w, h, paras, size=16, color=TEXT, bold=False, font=BODY,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=6, margin=0.05):
    """paras: str | list of (str | list of (text, {opts}))"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.space_after = Pt(space_after)
        runs = p if isinstance(p, list) else [(p, {})]
        for t, o in runs:
            r = para.add_run()
            r.text = nb(t)
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
    return tb


def box(slide, x, y, w, h, fill=TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def badge(slide, x, y, label, d=0.55, fill=DARK, size=18):
    c = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = HEAD
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = WHITE
    return c


def tag(slide, x, y, label="PROVISÓRIO", w=1.6):
    s = box(slide, x, y, w, 0.36, fill=WARN)
    s.adjustments[0] = 0.5
    tf = s.text_frame
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = BODY
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = WHITE


def new_slide(title, source=None, provisional=False, dark=False):
    state["n"] += 1
    s = prs.slides.add_slide(BLANK)
    if dark:
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = DARK
        return s
    text(s, 0.6, 0.35, 10.4 if provisional else 12.1, 0.9, title, size=28,
         color=DARK, bold=True, font=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    if provisional:
        tag(s, 11.13, 0.62)
    foot = "Rascunho v1 · 28/10/2026"
    if source:
        foot += " · Fonte: " + source
    text(s, 0.6, 7.0, 11.4, 0.35, foot, size=10, color=MUTED)
    text(s, 12.1, 7.0, 0.65, 0.35, str(state["n"]), size=10, color=MUTED,
         align=PP_ALIGN.RIGHT)
    return s


def table(slide, x, y, col_w, rows, size=13, row_h=0.42, header_fill=DARK,
          bold_last=False, warn_rows=(), header_h=None):
    nrows, ncols = len(rows), len(col_w)
    shp = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y),
                                 Inches(sum(col_w)), Inches(row_h * nrows))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("bandRow", "0")
    tblPr.set("firstRow", "1")
    for j, w in enumerate(col_w):
        tbl.columns[j].width = Inches(w)
    for i in range(nrows):
        tbl.rows[i].height = Inches(header_h if (i == 0 and header_h) else row_h)
        for j in range(ncols):
            cell = tbl.cell(i, j)
            val = rows[i][j]
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = header_fill
            elif i in warn_rows:
                cell.fill.fore_color.rgb = WARN_TINT
            else:
                cell.fill.fore_color.rgb = WHITE if i % 2 else TINT
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = ""
            num = isinstance(val, str) and (val.startswith("R$") or val[:1].isdigit()
                                            or val.startswith("−") or val.startswith("+"))
            p.alignment = PP_ALIGN.RIGHT if (num and j > 0) else PP_ALIGN.LEFT
            r = p.add_run()
            r.text = nb(val)
            r.font.name = BODY
            r.font.size = Pt(size)
            r.font.bold = i == 0 or (bold_last and i == nrows - 1)
            r.font.color.rgb = WHITE if i == 0 else TEXT
    return tbl


# ---------------------------------------------------------------- 1 capa
s = new_slide("", dark=True)
tag(s, 0.8, 0.8, "RASCUNHO — ESBOÇO v1", w=3.0)
text(s, 0.8, 2.0, 11.5, 1.6, "Cisão do grupo Siviero", size=48, color=WHITE,
     bold=True, font=HEAD)
text(s, 0.8, 3.25, 11.5, 1.2,
     "O ponto de vista do Luís, sócio de 1/6 que está saindo", size=26,
     color=RGBColor(0xC9, 0xDB, 0xCF), font=HEAD)
text(s, 0.8, 5.4, 11.5, 1.2, [
    "Versão de trabalho de 28/10/2026, feita a partir do material técnico em andamento.",
    "Números marcados como PROVISÓRIO ainda têm pendência e podem mudar.",
], size=16, color=WHITE)

# ---------------------------------------------------------------- 2 contexto
s = new_slide("Contexto: uma cisão e duas avaliações",
              source="00_PLANO.md (premissas P1 e P2, seção 2)")
cards = [
    ("1/6", "A quota do Luís",
     "O Luís tem 1/6 do grupo e está saindo pela cisão. Toda conta usa 1/6 em "
     "todas as empresas, e não os percentuais dos quadros societários."),
    ("1", "Inventário do grupo",
     "Feito pelos demais sócios, com data-base em 31.12.2024. Um documento por "
     "empresa; o documento 11 é a SEI, Siviero Empreendimentos Imobiliários."),
    ("2", "Estudo de Rafael",
     "Rafael, outro sócio que sai, apresentou o Estudo 2 (documentos assinados "
     "BELISE LTDA, de set/2026), com avaliações de mercado de parte dos bens."),
]
for i, (b, h, body) in enumerate(cards):
    x = 0.6 + i * 4.1
    box(s, x, 1.6, 3.8, 4.1)
    badge(s, x + 0.3, 1.9, b, d=0.75, size=20 if len(b) < 3 else 16)
    text(s, x + 0.3, 2.85, 3.3, 0.5, h, size=19, bold=True, color=DARK, font=HEAD)
    text(s, x + 0.3, 3.5, 3.3, 2.2, body, size=15)
text(s, 0.6, 6.0, 12.1, 0.8,
     [[("A pergunta deste trabalho: ", {"bold": True, "color": DARK}),
       ("para cada bem, qual dos dois valores existentes é tecnicamente "
        "defensável, e o que ainda falta esclarecer.", {})]], size=17)

# ---------------------------------------------------------------- 3 método
s = new_slide("Como o estudo foi feito", source="00_PLANO.md, premissas P2, P3, P5 e P11")
steps = [
    ("2018 é a referência do patrimônio entregue",
     "O inventário de 31.12.2018 mostra o patrimônio como ele foi entregue pela "
     "administração anterior. Serve para conferir se cada bem ainda existe ou tem "
     "destino documentado, não para dar valor."),
    ("Data-base em 31.12.2024",
     "Fatos posteriores (compras, vendas, avaliações de 2025 e 2026) são registrados "
     "um a um. Nada é ajustado em silêncio. Até agora só um tem tratamento decidido "
     "(o trator); os demais estão pendentes."),
    ("Sem avaliação própria",
     "Para cada bem adota-se um dos dois valores que já existem, o do grupo ou o de "
     "Rafael, com o argumento técnico registrado, por exemplo a faixa de mercado do "
     "próprio Rafael."),
    ("Revisão independente",
     "Desde 28/10/2026, cada pacote passa por um revisor com contexto novo antes de "
     "chegar ao Luís. Os pacotes 1 a 3 foram revisados depois de aprovados, e a "
     "revisão levou a duas correções."),
]
for i, (h, body) in enumerate(steps):
    y = 1.45 + i * 1.35
    badge(s, 0.7, y + 0.1, str(i + 1))
    text(s, 1.5, y, 10.9, 0.45, h, size=19, bold=True, color=DARK, font=HEAD)
    text(s, 1.5, y + 0.45, 10.9, 0.85, body, size=15)

# ---------------------------------------------------------------- 4 partes
s = new_slide("As 7 partes do estudo e a situação de cada uma",
              source="00_PLANO.md, seções 4 e 7 (estado em 28/10/2026)")
rows = [
    ["Parte", "O que analisa", "Situação em 28/10/2026"],
    ["0", "Fontes, extrações de 2018 e 2025, planilha-mestre", "Concluída"],
    ["3", "SEI: onde as duas avaliações concordam e discordam",
     "Pacotes 1 a 4 aprovados; falta a regra única de “possível APP” e o fechamento"],
    ["4", "SEI: contratos, indenizações, compromissos e destino de cada bem que saiu desde 2018",
     "Não iniciada; depende da explicação do método pelo Luís"],
    ["5", "Assuntos não abordados (por exemplo, loteamentos contratados)", "Não iniciada"],
    ["6", "Descontos em duplicidade e dupla contagem", "Não iniciada; 5 suspeitas já listadas"],
    ["7", "Observações do Luís nas planilhas Google", "Não iniciada"],
    ["1 e 2", "Demais empresas: concorda × discorda", "Começando agora"],
]
table(s, 0.6, 1.45, [1.0, 6.0, 5.1], rows, size=14, row_h=0.62)

# ---------------------------------------------------------------- 5 SEI números
s = new_slide("SEI: os números gerais", provisional=True,
              source="planilha-mestre-sei.xlsx, abas Resumo e Mestre")
stats = [
    ("R$ 152,80 mi", "bens de 2025 pelo inventário do grupo", DARK),
    ("R$ 153,97 mi", "soma dos valores adotados até agora", MID),
    ("+ R$ 1,17 mi", "diferença: adotado menos grupo", WARN),
]
for i, (v, lab, col) in enumerate(stats):
    y = 1.45 + i * 1.45
    text(s, 0.6, y, 4.3, 0.75, v, size=36, bold=True, color=col, font=HEAD)
    text(s, 0.6, y + 0.72, 4.3, 0.6, lab, size=15, color=MUTED)
cd = CategoryChartData()
cd.categories = ["Loteamentos", "Áreas urbanas", "Áreas rurais", "Lotes urbanos",
                 "Máquinas", "Veículos"]
cd.add_series("Grupo", (111.18, 26.93, 7.66, 3.29, 2.52, 1.22))
cd.add_series("Adotado", (112.33, 23.93, 10.92, 3.51, 2.07, 1.22))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(5.2), Inches(1.35),
                        Inches(7.5), Inches(4.4), cd)
ch = gf.chart
ch.has_title = True
ch.chart_title.text_frame.text = "Bens da SEI por bloco (R$ milhões)"
tp = ch.chart_title.text_frame.paragraphs[0]
tp.runs[0].font.size = Pt(14)
tp.runs[0].font.bold = True
tp.runs[0].font.name = BODY
tp.runs[0].font.color.rgb = TEXT
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(12)
ch.legend.font.name = BODY
ca = ch.category_axis
ca.reverse_order = True
ca.tick_labels.font.size = Pt(12)
ca.tick_labels.font.name = BODY
ca.format.line.color.rgb = RGBColor(0xBF, 0xC8, 0xC3)
va = ch.value_axis
va.visible = False
va.has_major_gridlines = False
plot = ch.plots[0]
plot.gap_width = 60
plot.overlap = -10
plot.has_data_labels = True
dl = plot.data_labels
dl.number_format = '#,##0.00'
dl.number_format_is_linked = False
dl.position = XL_LABEL_POSITION.OUTSIDE_END
dl.font.size = Pt(11)
dl.font.name = BODY
dl.font.color.rgb = TEXT
for ser, col in zip(plot.series, (DARK, RGBColor(0x9C, 0xBF, 0xA7))):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = col
box(s, 0.6, 5.8, 12.1, 1.1, fill=WARN_TINT)
text(s, 0.8, 5.83, 11.7, 1.05, [
    [("Ainda sem valor adotado: ", {"bold": True, "color": WARN}),
     ("Fracionamento Simioni e 7 loteamentos contratados (Parte 5). ", {}),
     ("Fora do total: ", {"bold": True, "color": WARN}),
     ("7 lotes em “possível APP” e 6 lotes “FAMÍLIA SIVIERO”. ", {}),
     ("Sem decisão lote a lote: ", {"bold": True, "color": WARN}),
     ("26 lotes de terceiros e 8 do acerto Rovílio (Parte 4). ", {}),
     ("São só os bens; compromissos e passivos ficam para a Parte 4.", {})],
], size=14, anchor=MSO_ANCHOR.MIDDLE, space_after=0)

# ---------------------------------------------------------------- 6 concordância
s = new_slide("SEI: onde os dois valores são iguais", provisional=True,
              source="log-decisoes.md, pacote 1 (regras R1, R2 e R3)")
cards = [
    ("R$ 103,46 mi", "395 lotes em 28 loteamentos",
     "Mesmo valor nas duas fontes. A soma dos lotes confere com o total de cada "
     "loteamento."),
    ("R$ 8,56 mi", "12 bens fora de loteamento",
     "Mesmo valor nas duas: 6 lotes urbanos, 5 áreas rurais e a Ford Ranger."),
    ("R$ 16,52 mi", "30 bens avaliados só pelo grupo",
     "Veículos, máquinas e 3 áreas urbanas, que o estudo de Rafael para a SEI não "
     "avalia: adota-se o valor do grupo."),
]
for i, (v, h, body) in enumerate(cards):
    x = 0.6 + i * 4.1
    box(s, x, 1.5, 3.8, 3.2)
    text(s, x + 0.3, 1.7, 3.3, 0.8, v, size=32, bold=True, color=DARK, font=HEAD)
    text(s, x + 0.3, 2.45, 3.3, 0.45, h, size=17, bold=True, color=TEXT)
    text(s, x + 0.3, 2.95, 3.3, 1.65, body, size=15)
text(s, 0.6, 5.15, 12.1, 0.6,
     [[("R$ 128,53 mi de R$ 152,80 mi (84%): ", {"bold": True, "color": DARK}),
       ("resolvidos em bloco, por valor igual ou avaliação única.", {})]], size=17)
text(s, 0.6, 5.7, 12.1, 1.25, [
    "Ressalvas: em alguns itens o valor de Rafael repete o balanço do grupo, e ele anota "
    "que fica abaixo da sua faixa de mercado. Área rural 06 (6.700 m² a mais nas parcelas), "
    "área rural 10 (conferir com as áreas contratadas), lote urbano 04 (hipoteca à Receita). "
    "Terrenos da BR-470 e da serraria, e a escavadeira 320D: dupla contagem, na Parte 6.",
], size=13, color=MUTED)

# ---------------------------------------------------------------- 7 divergências
s = new_slide("SEI: principais divergências", provisional=True,
              source="log-decisoes.md, pacote 2; planilha-mestre-sei.xlsx, aba Mestre")
rows = [
    ["Bem", "Grupo", "Rafael", "Adotado", "Argumento registrado"],
    ["Imóvel Oleoplan, BR-470", "R$ 13.631.378,60", "R$ 10.631.378,60",
     "R$ 10.631.378,60",
     "Valor declarado por Rafael. Ponto fraco: as parcelas que ele mesmo descreve "
     "somam R$ 13,63 mi; a diferença de R$ 3 mi parece erro de soma."],
    ["Lago Negro, 5 lotes da quadra 01", "R$ 4.021.000,00", "R$ 5.500.000,00",
     "R$ 5.170.000,00",
     "Valor adotado entre o do grupo e o de Rafael. Rafael cita lotes no mercado a "
     "partir de R$ 1,13 mi (p. 79)."],
    ["Casa da Nona, lote urbano 02", "R$ 250.000,00", "R$ 450.000,00", "R$ 450.000,00",
     "Topo da faixa de mercado do próprio Rafael: R$ 220 mil a R$ 450 mil (p. 3)."],
    ["Ouro Verde II, lote urbano 07", "R$ 130.000,00", "R$ 150.000,00", "R$ 150.000,00",
     "Valor do grupo abaixo do piso de Rafael (R$ 135 mil); comparáveis de R$ 152 mil "
     "a R$ 166 mil (p. 8 e 9)."],
    ["Trator de esteira D6N XL", "R$ 450.000,00", "não avalia", "R$ 0,00",
     "Comprado em set/2025, depois da data-base. A dívida de R$ 320 mil também sai: "
     "efeito líquido de menos R$ 130 mil."],
]
table(s, 0.6, 1.4, [2.3, 1.6, 1.6, 1.6, 5.0], rows, size=13, row_h=0.88,
      warn_rows=(1,), header_h=0.5)
text(s, 0.6, 6.45, 12.1, 0.45,
     "Áreas rurais: ver o próximo slide. Linha destacada: valor aprovado, mas sem "
     "argumento técnico forte e com risco de contestação.", size=13, color=MUTED)

# ---------------------------------------------------------------- 8 áreas rurais
s = new_slide("SEI: as 7 áreas rurais com valores diferentes", provisional=True,
              source="log-decisoes.md, pacote 3 (páginas do estudo de Rafael)")
rows = [
    ["Área rural", "Grupo", "Rafael (adotado)", "Faixa de mercado de Rafael"],
    ["03, lote 14, 7,5 ha", "R$ 220.000", "R$ 375.000", "R$ 320.000 a 550.000"],
    ["04, lote 12, 13 ha", "R$ 520.000", "R$ 650.000", "R$ 550.000 a 850.000"],
    ["05, lote B, 1,8 ha, gleba urbana", "R$ 1.000.000", "R$ 1.800.000", "R$ 1.800.000 a 9.000.000"],
    ["08, lote 89, 27,8 ha", "R$ 250.200", "R$ 1.500.000", "R$ 1.100.000 a 1.950.000"],
    ["09, lote 33, mato nativo, 15,1 ha", "R$ 40.000", "R$ 250.000", "R$ 180.000 a 350.000"],
    ["11, lote 47, 7,5 ha", "R$ 20.000", "R$ 380.000", "R$ 250.000 a 1.200.000"],
    ["12, lote 47, 7,5 ha", "R$ 30.000", "R$ 380.000", "R$ 250.000 a 1.200.000"],
    ["Total das 7 áreas", "R$ 2.080.200", "R$ 5.335.000", ""],
]
table(s, 0.6, 1.4, [3.3, 1.7, 1.9, 2.9], rows, size=14, row_h=0.5, bold_last=True)
box(s, 10.65, 1.4, 2.05, 4.5)
text(s, 10.8, 1.55, 1.8, 4.3, [
    [("Argumento", {"bold": True, "color": DARK, "size": 16})],
    "Em todas, o valor do grupo fica abaixo do piso da faixa de mercado de Rafael.",
    "Nas áreas 09, 11 e 12, fica abaixo até do valor fiscal da terra nua (VTN da "
    "Receita Federal).",
], size=14)
text(s, 0.6, 6.15, 12.1, 0.6,
     "Ressalvas: Rafael usa um único texto para as áreas 11 e 12. Área 04: pergunta 7 ao "
     "grupo (matrícula). Áreas 04, 08 e 09 estavam em 2018 no inventário da Agropecuária.",
     size=13, color=MUTED)

# ---------------------------------------------------------------- 9 2018 x 2025
s = new_slide("SEI: bens de 2018 que não aparecem em 2025", provisional=True,
              source="planilha-mestre-sei.xlsx, abas Mestre e Loteamentos 2018x2025")
text(s, 0.6, 1.45, 3.6, 0.8, "26 bens", size=40, bold=True, color=DARK, font=HEAD)
text(s, 0.6, 2.25, 3.6, 1.0,
     "existiam em 2018 e não têm par no inventário de 2025, somando R$ 6,05 mi pelos "
     "valores de 2018.", size=15)
text(s, 0.6, 3.45, 3.6, 0.8, "6 bens", size=40, bold=True, color=WARN, font=HEAD)
text(s, 0.6, 4.25, 3.6, 1.0,
     "sem destino documentado até agora, R$ 1,34 mi pelos valores de 2018.", size=15)
rows = [
    ["Sem destino documentado", "Valor 2018"],
    ["Lote 1543, BR-470 (Bissani)", "R$ 900.000"],
    ["Apartamento 303, Edifício Boa Esperança", "R$ 100.000"],
    ["2 lotes urbanos em Bento Gonçalves", "R$ 200.000"],
    ["Lote urbano na Rua Fioravante Boff", "R$ 70.000"],
    ["Lote 0024 (Santa Rita)", "R$ 70.000"],
]
table(s, 4.6, 1.45, [5.6, 2.5], rows, size=14, row_h=0.46)
text(s, 4.6, 4.4, 8.1, 1.3, [
    "Dos outros 20, 3 têm situação anotada (2 lotes já vendidos em 2018 e móveis fora "
    "de escopo). Os 17 restantes (8 loteamentos com saldo, 3 em execução, 5 veículos e "
    "uma máquina roubada) ainda não foram pesquisados: é tarefa da Parte 4, não iniciada. "
    "As perguntas 5 e 9 ao grupo pedem o destino dos bens sem documento.",
], size=14)
box(s, 0.6, 5.8, 12.1, 1.05, fill=TINT)
text(s, 0.8, 5.85, 11.7, 0.95, [
    [("Além disso, ", {"bold": True, "color": DARK}),
     ("vários loteamentos têm menos lotes em 2025 do que em 2018: Alto das Graças I, "
      "de 51 para 14; Altos da Julieta, de 40 para 10; Lago Negro, de 35 para 12. "
      "O destino de cada lote que saiu também é apurado na Parte 4.", {})],
], size=14, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 10 mudaram de empresa
s = new_slide("SEI: bens em outra empresa ou em duas",
              provisional=True,
              source="planilha-mestre-sei.xlsx, aba Mestre; parte6-candidatos-duplicidade.md")
rows = [
    ["Bem", "Em 2018", "Em 2025", "O que verificar"],
    ["Sede Campestre (mat. 1.941 e 7.104)", "Na lista da SEI, como bem de outra empresa",
     "Hotel Princesa dos Vales, R$ 5.560.857,50", "Que entra uma vez só no total do grupo"],
    ["Área rural lote 52, Bassani (mat. 7.391)", "Na lista da SEI, como bem de outra empresa",
     "Hotel Princesa dos Vales, R$ 422.800", "Idem"],
    ["Terrenos da BR-470, da antiga serraria e ao lado da Oleoplan",
     "Na lista da SEI, como bens da Rádio e da Oficina",
     "No total da SEI, R$ 25.729.394,60 pelo valor do grupo",
     "Se também constam no inventário da Rádio ou de outra empresa"],
    ["Áreas rurais 04, 08 e 09", "Fora do total da SEI, no inventário da Agropecuária",
     "No total da SEI, R$ 2.400.000 pelo valor adotado",
     "Se também constam em outra empresa"],
    ["Escavadeira Caterpillar 320D, 2009", "A SEI já tinha uma 320D",
     "Na SEI, R$ 200.000; no estudo de Rafael para a Ceval, R$ 240.000",
     "Se é a mesma máquina contada duas vezes"],
    ["Área rural lotes 24 e 24/E (Pierozan)", "Na lista da SEI, como bem de outra empresa",
     "Vendida em 2018 a Angeli, R$ 510.000 à vista", "O recebimento (Parte 4)"],
]
table(s, 0.6, 1.35, [3.2, 2.9, 3.2, 2.8], rows, size=13, row_h=0.7, header_h=0.45)
text(s, 0.6, 6.3, 12.1, 0.6,
     "Risco: o mesmo bem contado em duas empresas infla o patrimônio. A verificação é "
     "feita na Parte 6, depois do estudo das demais empresas.", size=14, color=MUTED)

# ---------------------------------------------------------------- 11 contratados
s = new_slide("Em aberto: loteamentos contratados", provisional=True,
              source="log-decisoes.md; planilha-mestre-sei.xlsx, aba Mestre, linhas 86 a 94")
text(s, 0.6, 1.45, 4.4, 0.9, "até R$ 34 mi", size=40, bold=True, color=WARN, font=HEAD)
text(s, 0.6, 2.35, 4.4, 2.9, [
    "é o valor de ativo que Rafael atribui a 7 loteamentos contratados.",
    "O inventário do grupo não lhes dá valor de ativo e lança multas e despesas: "
    "R$ 15,75 mi negativos no total dos contratados.",
], size=15)
rows = [
    ["Loteamento contratado", "Valor segundo Rafael"],
    ["Cecconello", "R$ 9.165.000"],
    ["Caminhos do Sol", "R$ 8.601.000"],
    ["Vale das Colinas", "R$ 5.019.600"],
    ["Vista da Colina", "R$ 3.384.000"],
    ["Pedrotti, Andreis e Lorandi", "R$ 3.384.000"],
    ["Bonetto (Rafael mantém a multa de R$ 782 mil)", "R$ 3.111.000"],
    ["Pasqual", "R$ 1.479.000"],
    ["Total", "R$ 34.143.600"],
]
table(s, 5.4, 1.4, [4.8, 2.5], rows, size=14, row_h=0.44, bold_last=True)
box(s, 0.6, 5.55, 12.1, 1.3, fill=WARN_TINT)
text(s, 0.8, 5.6, 11.7, 1.2, [
    [("Dr. Letti e Sol Nascente III: ", {"bold": True, "color": WARN}),
     ("Rafael avalia em zero; o grupo lança multas de R$ 1,65 mi e R$ 1,058 mi.", {})],
    "Decisão pendente: uma regra única para todos, na Parte 5. Uma das áreas rurais "
    "(2 ha, R$ 600 mil) fica dentro do Cecconello: conferir dupla contagem.",
], size=14, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 12 APP e Família
s = new_slide("Em aberto: APP, Simioni e “FAMÍLIA SIVIERO”", provisional=True,
              source="log-decisoes.md; perguntas-ao-grupo.md; parte6-candidatos-duplicidade.md")
cards = [
    ("Lago Negro: 7 lotes de APP confirmada", DARK,
     "Valor zero. Há decisão do Ministério Público acatada pela juíza. No estudo de "
     "Rafael, esses lotes aparecem com diferença de R$ 7,7 mi a favor do lado que sai. "
     "A empresa já deu 2 lotes em dação como indenização: possível desconto duplo (Parte 6)."),
    ("7 lotes em “possível APP”", WARN,
     "Em San Marino, Firminio Dalzóchio, Primavera e Industrial Vicentino. O grupo lista "
     "R$ 2,32 mi, fora do total. Valor a definir como % do lote cheio, com uma regra "
     "única, salvo documento que justifique outra. Pergunta 2 ao grupo."),
    ("Fracionamento Simioni: 4 lotes", WARN,
     "Em 2018 valiam R$ 1,0 mi (R$ 250 mil cada), com o mesmo texto de restrição. "
     "Em 2025 aparecem como APP e sem valor, sem processo citado. Rafael: R$ 520 mil. "
     "Reaberto."),
    ("6 lotes “FAMÍLIA SIVIERO”", WARN,
     "Loteamento Salvador, quadra 05, sem valor no inventário. Os vizinhos da quadra "
     "valem de R$ 374 mil a R$ 387 mil, cerca de R$ 2,29 mi no total. Pergunta 1 ao grupo."),
]
for i, (h, col, body) in enumerate(cards):
    x = 0.6 + (i % 2) * 6.15
    y = 1.45 + (i // 2) * 2.7
    box(s, x, y, 5.95, 2.5)
    text(s, x + 0.25, y + 0.15, 5.45, 0.5, h, size=18, bold=True, color=col, font=HEAD)
    text(s, x + 0.25, y + 0.7, 5.45, 1.75, body, size=15)

# ---------------------------------------------------------------- 13 perguntas
s = new_slide("Perguntas ao grupo: 10 aprovadas, nenhuma enviada",
              source="perguntas-ao-grupo.md")
qs = [
    ("1", "A quem se destinam os 6 lotes “FAMÍLIA SIVIERO” e por que estão sem valor"),
    ("2", "Situação do processo ou laudo de cada lote em “possível APP”"),
    ("3", "Cadeia de matrículas da área rural 02 e o aumento de 3,6 para 4,0 ha"),
    ("4", "Cadeia de matrículas da área rural 06 (Ansolin) e área da parcela Frainer"),
    ("5", "Destino, data e valor do lote 1543, BR-470 (Bissani)"),
    ("6", "Destino do lote de Michelle Tedesco, que saiu entre as versões de 2025"),
    ("7", "Matrícula 3.062, lote 12 (Sangali): possível dupla contagem em 2018"),
    ("8", "Datas de compra do lote urbano 06 e das matrículas 29.954 e 33.516"),
    ("9", "Destino de 5 bens de 2018 sem documento (Santa Rita, Boff, Ouro Verde, apto 303)"),
    ("10", "De quem é o apartamento 401 do Edifício Luiz Siviero"),
]
for i, (n, q) in enumerate(qs):
    col, row = i // 5, i % 5
    x = 0.6 + col * 6.15
    y = 1.5 + row * 0.98
    badge(s, x, y + 0.05, n, d=0.5, fill=MID, size=15)
    text(s, x + 0.65, y - 0.02, 5.3, 0.64, q, size=15, anchor=MSO_ANCHOR.MIDDLE,
         space_after=0)
text(s, 0.6, 6.45, 12.1, 0.45, "O envio das perguntas é decisão do Luís.",
     size=13, color=MUTED)

# ---------------------------------------------------------------- 14 demais empresas
s = new_slide("Demais empresas: estudo começando agora", provisional=True,
              source="01_fontes/outros/conferencia-discrepancias.md (26/10/2026)")
rows = [
    ["Empresa", "Inventário do grupo, 31.12.2024", "Situação"],
    ["Hotel Araçá", "R$ 58.076.823,20", "Conferido em 26/10"],
    ["Transportadora Irmãos Siviero", "R$ 16.374.986,00", "Conferido; nota de R$ 1,135 mi a explicar"],
    ["Rádio Comunidade", "R$ 2.698.000,00", "Conferido em 26/10"],
    ["Ceval Terraplenagem", "R$ 315.000,00", "Conferido em 26/10"],
    ["Associação Cultural, Turismo, Siviero Transportes", "R$ 0,00", "Empresas baixadas"],
    ["Hotel Princesa dos Vales, Abastecedora, São Crispin", "não conferido", "Inventários existem; ainda não conferidos"],
]
table(s, 0.6, 1.4, [4.6, 3.2, 4.3], rows, size=14, row_h=0.5, header_h=0.45)
text(s, 0.6, 5.2, 12.1, 0.45, "Pontos já identificados", size=17, bold=True,
     color=DARK, font=HEAD)
text(s, 0.6, 5.65, 12.1, 1.3, [
    "Fazenda São Crispin: Rafael avalia 407,53 ha (R$ 44,83 mi); o documento do grupo "
    "indica 411,43 ha restantes. Falta decidir a base de área.",
    "Transportadora: uma nota acrescenta R$ 1,135 mi ao capital referente a 3 itens que "
    "somam R$ 1,02 mi (diferença de R$ 115 mil).",
], size=14)

# ---------------------------------------------------------------- 15 próximos passos
s = new_slide("Próximos passos", source="00_PLANO.md, seção 7")
steps = [
    ("SEI, fechar a Parte 3",
     "Regra única para “possível APP”, com o Simioni; ficha de fechamento; nova "
     "revisão independente de tudo o que foi aprovado."),
    ("SEI, Partes 4 a 7",
     "Destino dos bens que saíram (após o método do Luís), não abordados, "
     "duplicidades e observações."),
    ("Demais empresas",
     "Mesmo método da SEI, empresa por empresa, partindo da conferência de 26/10."),
    ("Perguntas ao grupo",
     "Envio quando o Luís decidir."),
    ("Esta apresentação",
     "Novas versões (v2, v3...) conforme o trabalho técnico for aprovado."),
]
for i, (h, body) in enumerate(steps):
    y = 1.4 + i * 1.08
    badge(s, 0.7, y + 0.08, str(i + 1), d=0.6)
    text(s, 1.55, y, 10.9, 0.42, h, size=18, bold=True, color=DARK, font=HEAD)
    text(s, 1.55, y + 0.42, 10.9, 0.6, body, size=14)

# ---------------------------------------------------------------- 16 como ler
s = new_slide("O que ainda é provisório neste esboço",
              source="log-decisoes.md; planilha-mestre-sei.xlsx, aba Resumo")
items = [
    ("Soma adotada da SEI",
     "R$ 153,97 mi muda com o Simioni, os loteamentos contratados, a “possível APP”, "
     "os lotes “FAMÍLIA SIVIERO” e os lotes de terceiros."),
    ("Decisões já aprovadas",
     "Passam por nova revisão independente antes de fechar a Parte 3. Algumas dependem "
     "de perguntas ao grupo (3, 4, 7 e 8)."),
    ("Data-base",
     "As avaliações de Rafael são de 2026; o tratamento delas frente a 31.12.2024 "
     "ainda não foi decidido."),
    ("Bens de 2018 sem par",
     "O destino de cada um sai da Parte 4; parte pode ser só diferença de descrição."),
    ("Bens em duas empresas",
     "São suspeitas, não conclusões. A verificação é feita na Parte 6."),
    ("Demais empresas",
     "Totais da conferência de 26/10, parcialmente desatualizada."),
]
for i, (h, body) in enumerate(items):
    x = 0.6 + (i % 3) * 4.1
    y = 1.45 + (i // 3) * 2.75
    box(s, x, y, 3.85, 2.5, fill=WARN_TINT)
    text(s, x + 0.25, y + 0.2, 3.35, 0.5, h, size=17, bold=True, color=WARN, font=HEAD)
    text(s, x + 0.25, y + 0.75, 3.35, 1.7, body, size=14)

prs.save(OUT)
print("saved", OUT, "slides:", state["n"])
