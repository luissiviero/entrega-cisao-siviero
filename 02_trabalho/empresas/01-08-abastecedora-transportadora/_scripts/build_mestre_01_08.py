"""Master spreadsheet for companies 01 (Abastecedora) and 08 (Transportadora).

Re-runs the Phase 0 extraction (extract_fase0.py), then builds
`planilha-mestre-01-08.xlsx` with one row per asset (2025 items, 2018 items
without a 2025 pair, 2018 items that moved to the SEI), the credits and
contingencies written in the inventories and contracts, Luís's general notes,
a summary and a legend. Approved decisions come from `decisoes.json`; nothing
is typed by hand.

Usage: python -I build_mestre_01_08.py
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import fitz
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
import extract_fase0 as fase0

WORK = Path(__file__).resolve().parents[1]
FASE0 = WORK / "_extracao" / "fase0-itens-01-08.xlsx"
CONTRACTS = WORK / "_extracao" / "contratos-abastecedora.md"
DECISIONS = WORK / "decisoes.json"
OUT = WORK / "planilha-mestre-01-08.xlsx"

BLOCK_2025 = {
    ("01", "01"): "Imóvel urbano (posto)",
    ("08", "01"): "Imóvel urbano (complexo da oficina)",
    ("08", "02"): "Benfeitoria (complexo da oficina)",
    ("08", "03"): "Equipamentos (complexo da oficina)",
    ("08", "04"): "Benfeitoria (complexo da oficina)",
    ("08", "05"): "Benfeitoria (complexo da oficina)",
    ("08", "06"): "Imóvel urbano (Ed. Luiz Siviero)",
    ("08", "07"): "Imóvel rural (Cotiporã)",
    ("08", "08"): "Imóvel rural (Cotiporã)",
    ("08", "09"): "Imóvel urbano (Ed. Luiz Siviero)",
    ("08", "10"): "Veículo",
    ("08", "11"): "Veículo",
    ("08", "12"): "Veículo",
    ("08", "13"): "Veículo",
}
# cut-off registry lines (registro-data-base.md) that touch each row
CUTOFF = {
    "01-2025-01": "DB-02, DB-04, DB-06",
    "08-2025-03": "DB-07 (declaração de Benta, sem data de referência)",
    "08-2025-07": "DB-08 (contrato do Rovílio, 28/08/2025: créditos hereditários sobre a terra de São Brás)",
    "08-2025-08": "DB-08 (idem; este item foi comprado em 2008)",
    "08-2025-11": "DB-07 (Rafael: pesquisa de mercado de 2026)",
    "08-2025-12": "DB-07 (idem)",
    "08-2025-13": "DB-07 (idem)",
}
# Luís's notes in the 2_0 Google sheet, by row id (cells read from the sheet below)
NOTES = {
    "01-2025-01": [
        ("1.Abastecedora", "I7"),
        ("GERAL", "H7"),
        ("Resumo", "B72"),
        ("Resumo", "E72"),
    ],
    "08-2025-03": [("Resumo", "B196"), ("Resumo", "B198"), ("Resumo", "B199")],
    "08-2025-07": [("8.Transportadora", "C10"), ("GERAL", "C124")],
    "08-2025-08": [("8.Transportadora", "C11"), ("GERAL", "C125")],
    "01-2018-02": [("Resumo", "B73"), ("Resumo", "C73")],
}
NOTE_STATUS = {
    "08-2025-03": "Pendente: a nota de R$ 1.135.000 cita os itens 03–05, que somam R$ 1.020.000 (diferença de R$ 115.000 sem explicação); pergunta ao grupo",
    "08-2025-07": "Conferido: em 2018 o bem está na Agropecuária (doc 17 item 07), 'R$ NÃO DIVISÍVEL', sem valor e fora do total",
    "08-2025-08": "Conferido: em 2018 o bem está na Agropecuária (doc 17 item 08), 'R$ NÃO DIVISÍVEL', sem valor e fora do total",
    "01-2018-02": "Conferido: o rótulo de Resumo B73 ('Hotel Princesa dos Vales') está errado na planilha Google; o valor (C73 = GERAL!C8) é o caminhão-tanque da Abastecedora",
}
# Luís's general notes about 01/08 that do not belong to a single asset
GENERAL_NOTES = [
    ("Resumo", "B4", "Rovílio", "Respondido em rovilio-quotas.md (aguarda aprovação)"),
    ("Resumo", "B53", "Rovílio", "Idem"),
    (
        "Resumo",
        "C63",
        "Rovílio",
        "Conferido: a conta dá 35,53 % (preço acima da base); a base fica 26,22 % abaixo do preço",
    ),
    ("Resumo", "B67", "Rovílio", "Respondido em rovilio-quotas.md (aguarda aprovação)"),
    (
        "GERAL",
        "H874",
        "Rovílio",
        "Idem; nota fora das linhas 01/08 da GERAL (bloco da SEI)",
    ),
    (
        "Resumo",
        "B16",
        "Avaliações",
        "Pendente: 'valuations necessárias' inclui 01 Abastecedora e 08 Transportadora (B19, B23)",
    ),
    ("Resumo", "B19", "Avaliações", "Idem"),
    ("Resumo", "B23", "Avaliações", "Idem"),
]
# destination of the 2025 assets, by the 2018 inventory they come from
ORIGIN = {
    "2018 doc 12 Abastecedora": "Permanece na 01 (2018: doc 12 item {item})",
    "2018 doc 15 Transportadora": "Permanece na 08 (2018: doc 15 item {item})",
    "2018 doc 10 Oficina Mecânica": (
        "Oficina Mecânica (2018, inventário sem CNPJ, doc 10 item {item}) → 08 Transportadora (2025); "
        "sem documento de transferência; o grupo afirma os imóveis registrados em nome da 08 (arq. 1 p. 15)"
    ),
    "2018 doc 17 Agropecuária": (
        "Agropecuária (2018, doc 17 item {item}, 'R$ NÃO DIVISÍVEL') → 08 Transportadora (2025); "
        "'em nome das COMPRADORAS' segundo o contrato do Rovílio (arq. 10, cl. 3ª); sem documento de transferência"
    ),
}
ROW_NOTES = {
    "01-2025-01": (
        "Hipoteca em favor da Petrobras 'registrada no imóvel da Locadora' (arq. 4 p. 12, cl. 21ª § 3º), provavelmente a mat. 3.752: confirmar na matrícula. "
        "Imóvel locado à SIM de 01/01/2022 a 31/12/2033 + 36 meses (arq. 4 p. 1 e 3)."
    ),
    "08-2025-01": (
        "A SIM deve respeitar a faixa de acesso da mat. 3.753 (arq. 4 p. 7). O grupo diz que a 08 deu área em comodato à 01 para a locação SIM (arq. 1 p. 16); "
        "o contrato cita 2.500 m² da mat. 3.754, da SEI (arq. 4 p. 1). Ver aba Créditos e contingências."
    ),
    "08-2025-04": (
        "Locada a terceiros (inventário 08; arq. 1 p. 16). A locação SIM trata borracharia e rampa como espaços da locadora, a incluir por aditivo "
        "(arq. 4 p. 1 e 3). Localização não documentada."
    ),
    "08-2025-05": "Idem item 04.",
    "08-2025-07": "Valor 2018: 'R$ NÃO DIVISÍVEL' (doc 17), fora do total.",
    "08-2025-08": "Valor 2018: 'R$ NÃO DIVISÍVEL' (doc 17), fora do total. Comprado de Tercílio Dal Más em 25/09/2008.",
    "08-2018-07": "Hipótese a perguntar, sem par: o item 06 de 2025 (sala 7.393) passou a dizer 'CUB … + instalações'.",
}


def sheet_rows(wb, name):
    ws = wb[name]
    header = [c.value for c in ws[1]]
    return [
        dict(zip(header, r, strict=False))
        for r in ws.iter_rows(min_row=2, values_only=True)
    ]


def short(text: str | None, n: int = 260) -> str | None:
    if text is None:
        return None
    text = " ".join(str(text).split())
    return text if len(text) <= n else text[: n - 1] + "…"


def fmt_pct(v: float) -> str:
    return f"{v:.2f}".replace(".", ",") + " %"


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    fase0.main()  # re-derive every number from the sources
    wb = openpyxl.load_workbook(FASE0)
    items = sheet_rows(wb, "Itens inventários")
    adopted = {
        (r["Empresa"][:2], r["Item"]): r for r in sheet_rows(wb, "Rafael adotado")
    }
    pairs = sheet_rows(wb, "Pares 2018×2025")
    gsheet = openpyxl.load_workbook(fase0.GSHEET, data_only=True)
    decisions = json.loads(DECISIONS.read_text(encoding="utf-8"))

    def full_item(source: str, number: str) -> dict:
        return next(i for i in items if i["Fonte"] == source and i["Item"] == number)

    # why the station went from 12.216.851,89 (2018) to 9.661.414,00 (2025): compare the written parcels
    old_posto = full_item("2018 doc 12 Abastecedora", "01")["Descrição"]
    new_posto = full_item("2025 atual 01", "01")["Descrição"]
    labels = [
        "terreno e pavimentação",
        "área construída",
        "área coberta",
        "equipamentos",
    ]
    deltas = [
        n - o
        for o, n in zip(
            fase0.composition(old_posto), fase0.composition(new_posto), strict=True
        )
    ]
    deltas.append(fase0.thirty_percent(new_posto) - fase0.thirty_percent(old_posto))
    labels.append("30 % de vantagem da coisa feita")
    NOTE_STATUS["01-2025-01"] = (
        f"Conferido pela descrição: redução de {fase0.fmt(-sum(deltas))} = "
        + "; ".join(
            f"{lab} {fase0.fmt(d)}" for lab, d in zip(labels, deltas, strict=True)
        )
        + ". Por que os preços unitários caíram: pergunta ao grupo"
    )

    def note_text(row_id: str) -> str | None:
        cells = NOTES.get(row_id)
        if not cells:
            return None
        return " | ".join(
            f"{tab}!{cell}: {short(gsheet[tab][cell].value, 160)}"
            for tab, cell in cells
        )

    def decision_for(
        row_id: str, group: float | None = None, rafael: float | None = None
    ):
        """Approved decision for a row; a value of 'grupo' or 'rafael' points to that column."""

        def resolve(v):
            return (
                {"grupo": group, "rafael": rafael}.get(v, v)
                if isinstance(v, str)
                else v
            )

        d = decisions["itens"].get(row_id)
        if d:
            motive = (
                f"{d.get('motivo')} Aprovada em {d.get('aprovada_em')}."
                if d.get("aprovada_em")
                else d.get("motivo")
            )
            return resolve(d.get("valor_adotado")), d.get("decisao"), motive
        for rule_id, rule in decisions["regras"].items():
            if row_id in rule.get("itens", []):
                return (
                    resolve(rule.get("valor")),
                    rule["decisao"],
                    f"{rule_id}: {rule['motivo']} Aprovada em {rule['aprovada_em']}.",
                )
        return None, None, None

    master = []
    # -- 2025 items with their 2018 pair and Rafael's value
    for x in [i for i in items if i["Fonte"] in ("2025 atual 01", "2025 atual 08")]:
        co = x["Fonte"][-2:]
        row_id = f"{co}-2025-{x['Item']}"
        value = (
            x["Valor (R$)"]
            if x["Valor (R$)"] is not None
            else x["Valor recomposto (R$)"]
        )
        pair = next(
            (
                p
                for p in pairs
                if p["Fonte 2025"] == x["Fonte"] and p["Item 2025"] == x["Item"]
            ),
            None,
        )
        has_old = bool(pair and pair["Fonte 2018"])
        raf = adopted[(co, x["Item"])]
        raf_value = raf["Rafael: valor adotado, balanço executado por Valter (R$)"]
        diff = raf_value - value if raf_value is not None else None
        situation = (
            "Concorda"
            if diff == 0
            else ("Discorda" if diff is not None else "Só grupo")
        )
        old_value = pair["Valor 2018 (R$)"] if has_old else None
        adopted_value, decision, motive = decision_for(row_id, value, raf_value)
        notes = []
        if x["Valor (R$)"] is None:
            notes.append(
                "Célula de valor vazia na versão atual; valor recomposto pela descrição (= total do doc 01 e do doc 00)."
            )
        if row_id in ROW_NOTES:
            notes.append(ROW_NOTES[row_id])
        master.append(
            {
                "Empresa": "01 Abastecedora" if co == "01" else "08 Transportadora",
                "Bloco": BLOCK_2025[(co, x["Item"])],
                "ID": row_id,
                "Chave registral": pair["Chave"] if pair else None,
                "Descrição 2025": short(x["Descrição"]),
                "Valor grupo 2025 (R$)": value,
                "Existe em 2018?": "Sim" if has_old else "Não",
                "Item 2018": f"{pair['Fonte 2018']} item {pair['Item 2018']}"
                if has_old
                else None,
                "Tipo de par 2018×2025": pair["Tipo de par"] if pair else None,
                "Descrição 2018": short(pair["Descrição 2018"]) if has_old else None,
                "Valor 2018 (R$)": old_value
                if isinstance(old_value, (int, float))
                else None,
                "Existe em 2025 (grupo)?": f"Sim ({x['Fonte']} {x['Tabela.linha']})",
                "Existe no Estudo 2 (Rafael)?": "Sim",
                "Valor Rafael (R$)": raf_value,
                "Página Rafael": f"PDF {'ABASTECEDORA' if co == '01' else 'TRANSPORTADORA'} p. {raf['p.']}"
                if raf.get("p.")
                else None,
                "Diferença Rafael − grupo (R$)": diff,
                "Diferença %": (diff / value) if diff is not None and value else None,
                "Situação": situation,
                "Na lista de Rafael (virá para nós)?": "Não",
                "Critério C0": "Fica na empresa: maior valor"
                if situation == "Discorda"
                else None,
                "Valor adotado (R$)": adopted_value,
                "Decisão (P5)": decision,
                "Motivo da decisão": motive,
                "Destino 2018→2025": ORIGIN[pair["Fonte 2018"]].format(
                    item=pair["Item 2018"]
                )
                if has_old
                else None,
                "Registro data-base": CUTOFF.get(row_id),
                "Observação do Luís": note_text(row_id),
                "Situação da observação": NOTE_STATUS.get(row_id),
                "Fonte 2025": f"{x['Fonte']} {x['Tabela.linha']}",
                "Nota": " ".join(notes) or None,
            }
        )

    # -- 2018 items without a 2025 pair, and 2018 items that moved to the SEI
    for p in pairs:
        if not p["Fonte 2018"] or p["Fonte 2025"]:
            continue
        src, num = p["Fonte 2018"], p["Item 2018"]
        prefix = {
            "2018 doc 12 Abastecedora": "01",
            "2018 doc 15 Transportadora": "08",
            "2018 doc 10 Oficina Mecânica": "OFI",
        }[src]
        row_id = f"{prefix}-2018-{num}"
        moved = p["Tipo de par"].startswith("migrou para a SEI")
        adopted_value, decision, motive = decision_for(row_id)
        donated = (
            "doado" in full_item(src, num)["Descrição"]
        )  # full text: the donation is at the end of the description
        if moved:
            destination = p["Tipo de par"] + " — tratado pela sessão SEI"
        elif donated:
            destination = (
                "Doação a Geni Siviero Mattiuz declarada no inventário de 2018 ('conforme acordo dentre as partes') e deduzida do total de 2018; "
                "sem escritura nem registro nas fontes: pendente de certidão da mat. 7.395 (pacote 2)"
            )
        else:
            destination = "Sem destino documentado (pendência P4)"
        master.append(
            {
                "Empresa": {
                    "01": "01 Abastecedora",
                    "08": "08 Transportadora",
                    "OFI": "Oficina Mecânica (2018)",
                }[prefix],
                "Bloco": "Bem de 2018 hoje na SEI"
                if moved
                else (
                    "Bem de 2018 doado em 2018"
                    if donated
                    else "Bem de 2018 sem par em 2025"
                ),
                "ID": row_id,
                "Chave registral": p["Chave"],
                "Descrição 2025": None,
                "Valor grupo 2025 (R$)": None,
                "Existe em 2018?": "Sim",
                "Item 2018": f"{src} item {num}",
                "Tipo de par 2018×2025": p["Tipo de par"],
                "Descrição 2018": short(p["Descrição 2018"]),
                "Valor 2018 (R$)": p["Valor 2018 (R$)"],
                "Existe em 2025 (grupo)?": "Não na 01/08; na SEI (doc 11)"
                if moved
                else "Não",
                "Existe no Estudo 2 (Rafael)?": "Não",
                "Valor Rafael (R$)": None,
                "Página Rafael": None,
                "Diferença Rafael − grupo (R$)": None,
                "Diferença %": None,
                "Situação": "Não abordado — hoje na SEI"
                if moved
                else "Não abordado — só 2018",
                "Na lista de Rafael (virá para nós)?": None,
                "Critério C0": None,
                "Valor adotado (R$)": adopted_value,
                "Decisão (P5)": decision,
                "Motivo da decisão": motive,
                "Destino 2018→2025": destination,
                "Registro data-base": None,
                "Observação do Luís": note_text(row_id),
                "Situação da observação": NOTE_STATUS.get(row_id),
                "Fonte 2025": None,
                "Nota": ROW_NOTES.get(row_id),
            }
        )

    # -- credits, contingencies and liabilities written in the 2025 inventories, the replies and the contracts
    txt01 = fase0.docx_text(fase0.one(fase0.INV25, "01. "))
    txt08 = fase0.docx_text(fase0.one(fase0.INV25, "08. "))
    contracts = CONTRACTS.read_text(
        encoding="utf-8"
    )  # extraction checked on the page images
    with fitz.open(str(fase0.one(fase0.REPLIES, "1 - "))) as reply1:
        reply1_p4 = " ".join(reply1[3].get_text().split())
    total01 = full_item("2025 atual 01", "01")["Valor recomposto (R$)"]
    total08 = sum(i["Valor (R$)"] for i in items if i["Fonte"] == "2025 atual 08")

    def money_after(text: str, anchor: str) -> float:
        m = re.search(re.escape(anchor) + r".*?R\$\s*([\d.]+,\d{2})", text, re.DOTALL)
        return fase0.brl(m.group(1))

    def result_after(text: str, anchor: str) -> float:
        """Last amount of the first paragraph with R$ after the anchor (the result of a written calculation)."""
        tail = text[text.index(anchor) + len(anchor) :]
        para = next(p for p in tail.splitlines() if "R$" in p)
        return fase0.brl(fase0.MONEY.findall(para)[-1])

    def pct_after(text: str, anchor: str) -> float:
        m = re.search(re.escape(anchor) + r"\s*([\d,]+)%", text)
        return float(m.group(1).replace(",", "."))

    red01 = pct_after(txt01, "estima numa redução de")
    red08 = pct_after(txt08, "estima numa redução de")
    bonus = money_after(contracts, "| Valor antecipado (item III) |")
    rent = money_after(contracts, "| Aluguel inicial |")
    penalty = fase0.brl(
        re.search(r"de 9\.001 a 12\.000 m³ \*\*R\$\s*([\d.]+,\d{2})", contracts).group(
            1
        )
    )
    lucchetta = money_after(reply1_p4, "pretensão inicial")
    contingencies = [
        (
            "01 Abastecedora",
            "Redução estimada por regularização",
            f"{fmt_pct(red01)} do valor da empresa (calculado sobre o total)",
            round(total01 * red01 / 100, 2),
            "Redução (grupo)",
            "Inventário 2025 doc 01, observações",
            "Projeto do posto aprovado; averbação adiada até passar o prazo do INSS (março de 2027, arq. 1 p. 3–4)",
        ),
        (
            "01 Abastecedora",
            "Processo Lucchetta (ponto comercial e aluguéis)",
            "Pedido inicial, ajuizado em 31/07/2023",
            lucchetta,
            "Contingência",
            "Arq. 1, p. 4; inventário 2025 doc 01",
            "Sentença de improcedência em 14/08/2025; trânsito em julgado maio/2026 (DB-02). Sem valor no inventário",
        ),
        (
            "01 Abastecedora",
            "Bonificação antecipada Petrobras recebida por cessão",
            "Valor do item III do contrato",
            bonus,
            "Contingência (devolução proporcional ao volume não comprado + juros 1 % a.m. + multa 10 %)",
            "Arq. 3, p. 1–3",
            "Obrigação de restituir é do revendedor (Abastecedora Certo); os contratos da Certo 'são garantidos pela ora Locadora' (a 01) e a SIM 'obriga-se a assumir' (cessão não apresentada) (arq. 4 p. 11–12); hipoteca em 'o(s) imóvel(is)' da Locadora, matrícula não identificada; vigência até 30/09/2025, prorrogável (arq. 2 cl. 3.2–3.3; DB-04). Sem valor no inventário",
        ),
        (
            "01 Abastecedora",
            "Multa rescisória do contrato Petrobras",
            "Faixa de 9.001 a 12.000 m³ (volume contratado 10.847,10 m³)",
            penalty,
            "Contingência garantida por hipoteca",
            "Arq. 2, p. 4–5 (cl. 5.2 e 5.2.1)",
            "Obrigação do revendedor; decresce na proporção do volume já comprado; vigência até 30/09/2025 (DB-04)",
        ),
        (
            "01 Abastecedora",
            "Aluguel da locação à SIM",
            "Mensal, valor de 2022; valor em 31.12.2024 não informado (reajuste anual e revisão trienal)",
            rent,
            "Receita (não é bem no inventário)",
            "Arq. 4, p. 1 e 4",
            "Área locada inclui 2.500 m² da mat. 3.754 (SEI) como estacionamento; contrato até 31/12/2033 + 36 meses",
        ),
        (
            "01 Abastecedora",
            "Aluguel do restaurante/lancheria",
            "Valor não informado",
            None,
            "Receita ('Agrega ao valor do Posto, o aluguel da Lancharia/Restaurante')",
            "Inventário 2025 doc 01 item 01; arq. 4 p. 1 e 3",
            "Espaço locado a terceiros em 2021, com distrato previsto até 31/05/2022 e inclusão na locação SIM por aditivo; contexto do processo Lucchetta",
        ),
        (
            "08 Transportadora",
            "Redução estimada por regularização",
            f"{fmt_pct(red08)} do valor da empresa (calculado sobre o total)",
            round(total08 * red08 / 100, 2),
            "Redução (grupo)",
            "Inventário 2025 doc 08, observações",
            "Pavilhão da oficina, borracharia e rampa sem projeto (arq. 1 p. 15)",
        ),
        (
            "08 Transportadora",
            "INSS da obra do pavilhão da oficina",
            "Estimativa do grupo",
            money_after(txt08, "PAVILHÃO DA OFICINA MECÂNICA: APROVAR"),
            "Custo de regularização",
            "Inventário 2025 doc 08, observações",
            "Mais APPCI, ART, engenheiro, taxas",
        ),
        (
            "08 Transportadora",
            "INSS das obras da borracharia e da rampa",
            "Estimativa do grupo",
            money_after(txt08, "PAVILHÃO DA BORRACHARIA E RAMPA DE LAVAGEM: APROVAR"),
            "Custo de regularização",
            "Inventário 2025 doc 08, observações",
            "Mais APPCI, ART, engenheiro, taxas",
        ),
        (
            "08 Transportadora",
            "Aluguel da rampa de lavagem e da borracharia",
            "Valor não informado",
            None,
            "Receita",
            "Inventário 2025 doc 08 itens 04–05; arq. 1 p. 16",
            "Espaços locados a terceiros, equipamentos dos locatários; pedir os contratos",
        ),
        (
            "08 Transportadora",
            "Comodato de área da 08 à 01 para a locação SIM",
            "Sem valor; possível ônus sobre a área da 08 até 31/12/2033 + 36 meses",
            None,
            "Ônus (afirmado pelo grupo)",
            "Arq. 1 p. 16 × arq. 4 p. 1–2 e 7",
            "O grupo afirma o comodato; o contrato SIM cita 2.500 m² da mat. 3.754 (SEI) e a mat. 3.753 só como faixa de acesso. Pela diretriz do Luís não vale como ônus; A-05 retirada",
        ),
        (
            "08 Transportadora",
            "Cobrança contra Joveni Transportes × JBS × Frangosul",
            "Valor aproximado corrigido (2018: 150.000)",
            money_after(txt08, "cobrança de um valor de aproximadamente"),
            "Crédito (sem valor no total)",
            "Inventário 2025 doc 08; 2018 doc 15",
            "Ref. dez/2011 a abr/2012; acordo pode alterar o valor",
        ),
        (
            "08 Transportadora",
            "Processo Sottili — pensão",
            "1 SM × 12 × 20 anos (cálculo do grupo)",
            result_after(txt08, "PENSÃO"),
            "Contingência (cenário)",
            "Inventário 2025 doc 08",
            "Pedido cumulativo com o dano moral (pensão + um dos dois cenários de dano moral). Acordo depositado pela seguradora em ago/2020, de que a autora desistiu em fev/2022; não homologado (arq. 1 p. 15–16; DB-03)",
        ),
        (
            "08 Transportadora",
            "Processo Sottili — dano moral pedido",
            "1.000 SM",
            result_after(txt08, "DANO MORAL"),
            "Contingência (cenário)",
            "Inventário 2025 doc 08",
            "Alternativo ao cenário de 150 SM (os dois de dano moral não se somam); soma-se à pensão",
        ),
        (
            "08 Transportadora",
            "Processo Sottili — dano moral alternativo",
            "150 SM de 2001 com juros e correção",
            result_after(txt08, "teremos:"),
            "Contingência (cenário)",
            "Inventário 2025 doc 08",
            "Alternativo ao dano moral pedido; somado à pensão dá cerca de R$ 991 mil mais juros, de onde vem a 'condenação superior a R$ 1.000.000,00' do grupo",
        ),
        (
            "08 Transportadora",
            "Ônus de eventual dissolução em relação à parte de sócio",
            "Não quantificado",
            None,
            "Ônus (anunciado pelo grupo)",
            "Inventários 2025 docs 01 e 08, observações",
            "'haverá ônus de eventual dissolução da empresa em relação a parte de algum(ns) sócio(s)'; vale para a 01 e para a 08; dirige-se à saída de sócio",
        ),
        (
            "01 Abastecedora",
            "Processos do Posto Pértile Ltda. (078/1.14.0000166-1 e 078/1.13.0001323-4)",
            "Sem valor; 'Ação de Indenização = Ponto e pagamento de aluguéis'",
            None,
            "Contingência (2018, sem registro em 2025)",
            "Inventário 2018 doc 12, observações",
            "Não aparece em 2025; desfecho desconhecido. Mesmo objeto do processo Lucchetta (2023). Pergunta A-07 retirada pela diretriz",
        ),
        (
            "01 Abastecedora",
            "Pendências com a Transportes Zandoná (contrato de transporte de 22/03/2021 entre a 08 e a Zandoná)",
            "Sem valor",
            None,
            "Obrigação (assumida pela 01 no contrato SIM)",
            "Arq. 4, p. 12, cl. 22ª e § 1º",
            "A 01 'se compromete a arcar isoladamente com quaisquer pendências que eventualmente existir com a contratada Transportes Zandona Ltda.'; contrato assumido pela SIM até 21/03/2022",
        ),
        (
            "08 Transportadora",
            "Nota: acrescentado ao capital de bens (código 1060), itens 03–05",
            "Itens 03–05 somam 1.020.000",
            money_after(txt08, "FOI ACRESCENTADO O VALOR DO CAPITAL DE BENS"),
            "Nota contábil",
            "Inventário 2025 doc 08",
            "Diferença de 115.000 sem explicação (pendência 4 da conferência anterior); Luís: Resumo B196–B199 'Não entendi'",
        ),
    ]

    # ---------------------------------------------------------------- write workbook
    out = openpyxl.Workbook()
    ws = out.active
    ws.title = "Mestre"
    cols = list(master[0].keys())
    ws.append(cols)
    for r in master:
        ws.append([r[c] for c in cols])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if cols[c.column - 1].endswith("(R$)"):
                c.number_format = "#,##0.00"
            if cols[c.column - 1] == "Diferença %":
                c.number_format = "0.0%"
            if cols[c.column - 1] == "Situação" and c.value == "Discorda":
                c.fill = PatternFill("solid", fgColor="FCE4D6")

    ws = out.create_sheet("Créditos e contingências")
    ws.append(
        [
            "Empresa",
            "Item",
            "Base",
            "Valor (R$)",
            "Natureza (como o documento trata)",
            "Fonte",
            "Observação",
            "Tratamento",
        ]
    )
    for c in contingencies:
        ws.append([*c, "A decidir em pacote"])
    for row in ws.iter_rows(min_row=2):
        row[3].number_format = "#,##0.00"

    ws = out.create_sheet("Observações gerais do Luís")
    ws.append(["Aba", "Célula", "Tema", "Texto / valor", "Situação da observação"])
    for tab, cell, theme, status in GENERAL_NOTES:
        ws.append([tab, cell, theme, short(gsheet[tab][cell].value, 400), status])

    ws = out.create_sheet("Resumo")
    ws.append(["Indicador", "01 Abastecedora", "08 Transportadora", "Total"])

    def agg(company: str, key: str, cond=lambda r: True) -> float:
        return sum((r[key] or 0) for r in master if r["Empresa"] == company and cond(r))

    def cnt(company: str, cond) -> int:
        return sum(1 for r in master if r["Empresa"] == company and cond(r))

    is25 = lambda r: r["ID"][3:7] == "2025"
    gone = lambda r: r["Bloco"] == "Bem de 2018 sem par em 2025"
    donated = lambda r: r["Bloco"] == "Bem de 2018 doado em 2018"
    lines = [
        ("Itens 2025 (quantidade)", lambda co: cnt(co, is25)),
        ("Valor grupo 2025 (R$)", lambda co: agg(co, "Valor grupo 2025 (R$)", is25)),
        ("Valor Rafael (R$)", lambda co: agg(co, "Valor Rafael (R$)", is25)),
        (
            "Concorda (quantidade)",
            lambda co: cnt(co, lambda r: r["Situação"] == "Concorda"),
        ),
        (
            "Concorda (R$, grupo)",
            lambda co: agg(
                co, "Valor grupo 2025 (R$)", lambda r: r["Situação"] == "Concorda"
            ),
        ),
        (
            "Discorda (quantidade)",
            lambda co: cnt(co, lambda r: r["Situação"] == "Discorda"),
        ),
        (
            "Discorda (R$, grupo)",
            lambda co: agg(
                co, "Valor grupo 2025 (R$)", lambda r: r["Situação"] == "Discorda"
            ),
        ),
        (
            "Discorda (R$, Rafael)",
            lambda co: agg(
                co, "Valor Rafael (R$)", lambda r: r["Situação"] == "Discorda"
            ),
        ),
        (
            "Na lista de Rafael (quantidade)",
            lambda co: cnt(
                co, lambda r: r["Na lista de Rafael (virá para nós)?"] == "Sim"
            ),
        ),
        (
            "Itens com valor adotado (quantidade)",
            lambda co: cnt(
                co, lambda r: is25(r) and r["Valor adotado (R$)"] is not None
            ),
        ),
        (
            "Bens de 2018 sem par em 2025, não deduzidos em 2018 (quantidade)",
            lambda co: cnt(co, gone),
        ),
        (
            "Bens de 2018 sem par em 2025, não deduzidos em 2018 (R$ 2018)",
            lambda co: agg(co, "Valor 2018 (R$)", gone),
        ),
        (
            "Bem doado e já deduzido no próprio inventário de 2018 (R$ 2018)",
            lambda co: agg(co, "Valor 2018 (R$)", donated),
        ),
    ]
    for label, fn in lines:
        a, b = fn("01 Abastecedora"), fn("08 Transportadora")
        ws.append([label, a, b, a + b])
    for row in ws.iter_rows(min_row=2):
        for c in row[1:]:
            c.number_format = "#,##0.00" if "R$" in row[0].value else "0"

    ws = out.create_sheet("Legenda")
    ws.append(["Campo / aba", "Significado"])
    for k, v in [
        (
            "Mestre",
            "Uma linha por bem: 14 itens do inventário 2025 (01 e 08), bens de 2018 sem par em 2025 (inclusive o doado em 2018) e bens de 2018 que hoje estão na SEI. Gerada por _scripts/build_mestre_01_08.py; não editar à mão.",
        ),
        (
            "ID",
            "Empresa-ano-item. 'OFI' = Oficina Mecânica Siviero Ltda. (inventário 2018 próprio, sem CNPJ e sem inventário em 2025).",
        ),
        (
            "Chave registral",
            "Matrícula, registro ou placa usada no par 2018×2025 (lição 3: nunca só semelhança de texto).",
        ),
        (
            "Tipo de par 2018×2025",
            "'par por chave'; 'par por descrição — confirmar' (pavilhão, equipamentos, rampa e borracharia: sem matrícula própria e localização não documentada); 'migrou para a SEI'; 'sem par em 2025'.",
        ),
        (
            "Existe em 2025 (grupo)?",
            "Sim, com documento e tabela.linha do inventário 2025; 'Não'; ou 'na SEI' para bens da Oficina de 2018 que hoje estão no doc 11.",
        ),
        (
            "Valor Rafael",
            "Valor do 'Balanço patrimonial executado por Valter' (01 p. 3; 08 p. 13–15), que é o valor que Rafael adota. As outras seções do PDF estão em _extracao/fase0-itens-01-08.xlsx, aba 'Rafael adotado'.",
        ),
        (
            "Situação",
            "Concorda = mesmo valor; Discorda = valores diferentes; Só grupo = Rafael sem valor (não ocorre aqui); 'Não abordado — só 2018' = bem de 2018 sem par em 2025; 'Não abordado — hoje na SEI' = bem da Oficina de 2018 hoje na SEI.",
        ),
        (
            "Na lista de Rafael",
            "Bem escolhido por Rafael (ativos-avaliados-valter.xlsx). Nenhum bem da 01 ou da 08 está na lista.",
        ),
        (
            "Critério C0",
            "Quando há divergência, o C0 aponta o maior valor (bem que fica na empresa). A decisão só é preenchida depois da aprovação do Luís (decisoes.json).",
        ),
        (
            "Destino 2018→2025",
            "Para bens de 2025: de onde vieram (mesma empresa, Oficina ou Agropecuária) e o documento que sustenta a mudança, se houver. Para bens de 2018 sem par: destino documentado ou 'Sem destino documentado (pendência P4)'.",
        ),
        ("Registro data-base", "Linha DB-nn de registro-data-base.md desta pasta."),
        (
            "Observação do Luís / Situação da observação",
            "Texto das células da planilha Google 2_0 (aba!célula) e o que foi feito com ela: Conferido, Pendente ou Respondido. Notas que não são de um bem: aba 'Observações gerais do Luís'.",
        ),
        (
            "Créditos e contingências",
            "Valores escritos nas observações dos inventários 2025, nas respostas ao Dr. Pablo (arq. 1) e nos contratos (arq. 2, 3, 4, pela extração conferida nas imagens). As duas reduções são calculadas pelo script (percentual do grupo × total). Nenhum valor entra no total do inventário; tratamento a decidir em pacote.",
        ),
    ]:
        ws.append([k, v])

    for sheet in out.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = "A2"
    out.save(OUT)

    print(
        f"saved {OUT}: {len(master)} rows in Mestre, {len(contingencies)} contingencies, {len(GENERAL_NOTES)} general notes"
    )
    for row in out["Resumo"].iter_rows(min_row=2, values_only=True):
        print(" | ".join(fase0.fmt(v) if isinstance(v, float) else str(v) for v in row))


if __name__ == "__main__":
    main()
