"""Figures for packet 2 (exceptions) of companies 01 and 08.

Reads `planilha-mestre-01-08.xlsx` and the Rafael ranges in
`_extracao/fase0-itens-01-08.xlsx`; applies Luís's directive of 28/10/2026
(no asset of 01/08 comes to him: the higher value is recommended) and prints
the table, the totals and the effect on 1/6 used in pacotes/pacote-02.md.

Usage: python -I pacote_02.py
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import extract_fase0 as fase0

WORK = Path(__file__).resolve().parents[1]
MASTER = WORK / "planilha-mestre-01-08.xlsx"
FASE0 = WORK / "_extracao" / "fase0-itens-01-08.xlsx"


def rows_of(path: Path, sheet: str) -> list[dict]:
    ws = openpyxl.load_workbook(path)[sheet]
    header = [c.value for c in ws[1]]
    return [
        dict(zip(header, r, strict=False))
        for r in ws.iter_rows(min_row=2, values_only=True)
    ]


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    f = fase0.fmt
    rows = rows_of(MASTER, "Mestre")
    ranges = rows_of(FASE0, "Rafael faixas")

    decided = [r for r in rows if r["Valor adotado (R$)"] is not None]
    open_rows = [r for r in rows if r["Decisão (P5)"] is None]
    discord = [r for r in open_rows if r["Situação"] == "Discorda"]
    gone = [r for r in open_rows if r["Situação"] == "Não abordado — só 2018"]
    print(
        f"já decididos: {len(decided)} = {f(sum(r['Valor adotado (R$)'] for r in decided))}; pacote 2: {len(discord)} divergências + {len(gone)} bens de 2018"
    )

    print("Divergências (maior valor = recomendação pela diretriz):")
    for r in sorted(
        discord, key=lambda x: -max(x["Valor grupo 2025 (R$)"], x["Valor Rafael (R$)"])
    ):
        g, raf = r["Valor grupo 2025 (R$)"], r["Valor Rafael (R$)"]
        item = r["ID"][-2:]
        rr = [
            x
            for x in ranges
            if x["Empresa"].startswith(r["ID"][:2]) and x["Item"] == item
        ]
        rng = "; ".join(
            f"{f(x['Mínimo (R$)'])}–{f(x['Máximo (R$)']) if x['Máximo (R$)'] else '…'} (p. {x['Página']})"
            for x in rr
        )
        print(
            f"  {r['ID']} | grupo {f(g)} | Rafael {f(raf)} ({r['Página Rafael']}) | maior {f(max(g, raf))} | +{f(max(g, raf) - g)} | 2018 {f(r['Valor 2018 (R$)'])} | faixas: {rng or '-'}"
        )

    print("Bens de 2018 sem destino:")
    for r in sorted(gone, key=lambda x: -(x["Valor 2018 (R$)"] or 0)):
        print(
            f"  {r['ID']} | 2018 {f(r['Valor 2018 (R$)'])} | {r['Chave registral']} | {r['Destino 2018→2025'][:80]}"
        )

    for co in ("01", "08"):
        g = sum(
            r["Valor grupo 2025 (R$)"] for r in rows if r["ID"].startswith(f"{co}-2025")
        )
        rec = sum(
            (
                r["Valor adotado (R$)"]
                if r["Valor adotado (R$)"] is not None
                else max(r["Valor grupo 2025 (R$)"], r["Valor Rafael (R$)"])
            )
            for r in rows
            if r["ID"].startswith(f"{co}-2025")
        )
        print(
            f"{co}: grupo {f(g)}; com o pacote 2 pela diretriz {f(rec)}; +{f(rec - g)}; 1/6 +{f((rec - g) / 6)}"
        )
    g_all = sum(r["Valor grupo 2025 (R$)"] for r in rows if r["ID"][3:7] == "2025")
    # decided rows keep their adopted value; open rows take the higher value (directive)
    rec_all = sum(
        r["Valor adotado (R$)"]
        if r["Valor adotado (R$)"] is not None
        else max(r["Valor grupo 2025 (R$)"], r["Valor Rafael (R$)"])
        for r in rows
        if r["ID"][3:7] == "2025"
    )
    print(
        f"total 01+08: grupo {f(g_all)}; recomendado {f(rec_all)}; +{f(rec_all - g_all)}; 1/6 +{f((rec_all - g_all) / 6)}"
    )
    posto = next(r for r in rows if r["ID"] == "01-2025-01")
    alt = posto["Valor 2018 (R$)"]
    print(
        f"posto: alternativa 'posto sem o caminhão' (valor do grupo em 2018) {f(alt)}; diferença para Rafael {f(posto['Valor Rafael (R$)'] - alt)}; para o grupo 2025 +{f(alt - posto['Valor grupo 2025 (R$)'])}"
    )


if __name__ == "__main__":
    main()
