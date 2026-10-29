"""Phase 0 extraction for companies 01 (Abastecedora) and 08 (Transportadora).

Reads every source directly (docx, PDF text, xlsx), extracts the items with
their values, the value Rafael adopts for each item, the 2018 x 2025 pairs by
registry key (matrícula, registro, plate) and the Google-sheet values, and
checks each extracted total against the total printed in the documents.
Writes `_extracao/fase0-itens-01-08.xlsx` (labels in Portuguese) and
`_extracao/hashes.txt`, and prints the checks.

Usage: python -I extract_fase0.py
"""

from __future__ import annotations

import hashlib
import re
import sys
from pathlib import Path

import docx
import fitz
import openpyxl
from openpyxl.styles import Font

BASE = Path(__file__).resolve().parents[4]
WORK = Path(__file__).resolve().parents[1]
OUT = WORK / "_extracao" / "fase0-itens-01-08.xlsx"
HASHES = WORK / "_extracao" / "hashes.txt"
ROVILIO_TRANSCRIPTION = WORK / "_extracao" / "arq10-rovilio-transcricao.md"
DRIVE2018 = Path(
    "G:/Meu Drive/02_Patrimonio_e_Sociedades/Cisao_Siviero/Material Siviero/2018/inventário-2018"
)
INV25 = BASE / "01_fontes" / "inventario-2025"
INV25_OLD = BASE / "01_fontes" / "inventario-2025-antigo"
RAFAEL = BASE / "01_fontes" / "estudo-rafael"
REPLIES = BASE / "01_fontes" / "respostas-docs-2026-08-13"
GSHEET = (
    BASE
    / "01_fontes"
    / "planilhas-google"
    / "2_0_Comparativo Patrimonial Siviero 2018 x 2025.xlsx"
)
VALTER_LIST = BASE / "01_fontes" / "outros" / "ativos-avaliados-valter.xlsx"

MONEY = re.compile(r"R\$\s*([\d.]+,\d{2})(?!\d)")
# Rafael's summary tables write millions as "R$ 12.361,850" (comma as a thousands separator)
MONEY_SUMMARY = re.compile(r"R\$\s*(\d{1,3}(?:\.\d{3})*,\d{3})(?!\d)")


def md5_10(path: Path) -> str:
    return hashlib.md5(path.read_bytes()).hexdigest()[:10]


def one(folder: Path, prefix: str) -> Path:
    hits = sorted(p for p in folder.iterdir() if p.name.startswith(prefix))
    if len(hits) != 1:
        raise SystemExit(
            f"expected one file starting with {prefix!r} in {folder}, got {hits}"
        )
    return hits[0]


def brl(text: str) -> float:
    return float(text.replace(".", "").replace(",", "."))


def fmt(v: float | None) -> str:
    if v is None:
        return "-"
    s = f"{v:,.2f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


# ---------------------------------------------------------------- docx inventories


def docx_rows(path: Path):
    """Yield (table_index, row_index, [cell texts]) with merged cells collapsed."""
    doc = docx.Document(str(path))
    for t_idx, table in enumerate(doc.tables, 1):
        for r_idx, row in enumerate(table.rows, 1):
            cells, prev = [], None
            for cell in row.cells:
                if prev is not None and cell._tc is prev:
                    continue
                prev = cell._tc
                cells.append(" ".join(cell.text.split()))
            yield t_idx, r_idx, cells


def docx_text(path: Path) -> str:
    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return "\n".join(parts)


def inventory_items(path: Path, company: str, source: str):
    """Items are table rows whose first cell is a 2-digit number; the value is the last cell."""
    items, totals = [], []
    for t_idx, r_idx, cells in docx_rows(path):
        first = cells[0] if cells else ""
        # rows of the 2018 Oficina equipment checklist ('.... Compressores', no value) are not items
        if (
            re.fullmatch(r"\d{2}", first)
            and len(cells) >= 3
            and not cells[1].startswith("....")
        ):
            vals = MONEY.findall(cells[-1])
            items.append(
                {
                    "empresa": company,
                    "fonte": source,
                    "tabela": f"T{t_idx}.R{r_idx}",
                    "item": first,
                    "descricao": cells[1],
                    "valor": brl(vals[-1]) if vals else None,
                    "celula_valor": cells[-1],
                }
            )
        elif (
            "TOTAL GERAL DA EMPRESA" in first.upper()
            or "TOTAL LÍQUIDO" in first.upper()
        ):
            vals = [brl(v) for v in MONEY.findall(" ".join(cells[1:]))]
            label = first.strip(". ")
            # the 2018 Agropecuária repeats its total in a second table
            if not any(t["rotulo"] == label and t["valores"] == vals for t in totals):
                totals.append(
                    {
                        "fonte": source,
                        "tabela": f"T{t_idx}.R{r_idx}",
                        "rotulo": label,
                        "valores": vals,
                    }
                )
    return items, totals


def composition(description: str) -> list[float]:
    """Amounts written as '= R$ x' inside a description (partial values of one item)."""
    return [brl(v) for v in re.findall(r"=\s*R\$\s*([\d.]+,\d{2})", description)]


def thirty_percent(description: str) -> float:
    m = re.search(r"30,00%\s*x\s*R\$\s*([\d.]+,\d{2})", description)
    return brl(m.group(1)) if m else 0.0


def keys_of(description: str) -> set[str]:
    """Registry keys written in a description: matrículas, registro and vehicle plates."""
    keys: set[str] = set()
    for m in re.finditer(
        # greedy list capture ("459, 460, 461 e 558."), then split into numbers with thousands dots
        r"Matr[íi]culas?\s*n[º°o]?\s*\.?:?\s*(\d[\d.]*(?:\s*(?:,|\be\b)\s*\d[\d.]*)*)",
        description,
    ):
        for num in re.findall(r"\d{1,3}(?:\.\d{3})+|\d+", m.group(1)):
            keys.add("mat. " + num.replace(".", ""))
    for m in re.finditer(r"Registro\s*n[º°o]?\s*\.?:?\s*([\d.]+\d)", description):
        keys.add("registro " + m.group(1).replace(".", ""))
    for m in re.finditer(
        r"PLACA:?\s*=?\s*([A-Z]{3})\s*(\d{4})", description, re.IGNORECASE
    ):
        keys.add(f"placa {m.group(1).upper()} {m.group(2)}")
    return keys


# ---------------------------------------------------------------- Rafael PDFs


SECTION_MARKERS = [
    ("Balanço apresentado pelo grupo", "BALANÇO PATRIMONIAL APRESENTADO PELO GRUPO"),
    ("Avaliação executada por Valter", "AVALIAÇÃO PATRIMONIAL EXECUTADA POR VALTER"),
    ("Balanço executado por Valter", "BALANÇO PATRIMONIAL EXECUTADO POR VALTER"),
    ("Quadro-resumo", "QUADRO-RESUMO"),
]
RANGE_PATTERNS = [
    re.compile(r"entre\s+R\$\s*([\d.]+)\s+e\s+R\$\s*([\d.]+)", re.IGNORECASE),
    re.compile(r"R\$\s*([\d.]+)\s+a\s+R\$\s*([\d.]+)"),
    re.compile(r"(?:a partir de|começam em)\s+R\$\s*([\d.]+)()"),
]


def rafael_items(path: Path, company: str, single_item: bool):
    """Walk the PDF line by line, tracking section and item; collect values, ranges and declarations."""
    with fitz.open(str(path)) as doc:
        pages = [p.get_text() for p in doc]
    rows, totals, summary, declarations = [], [], [], []
    blocks: dict[tuple[str, str], list[tuple[int, str]]] = {}
    section, item = None, None
    for p_no, text in enumerate(pages, 1):
        lines = [ln.strip() for ln in text.splitlines()]
        for i, ln in enumerate(lines):
            hit = next(
                (label for label, marker in SECTION_MARKERS if marker in ln.upper()),
                None,
            )
            if hit:
                section, item = (
                    hit,
                    ("01" if single_item and hit != "Quadro-resumo" else None),
                )
                continue
            m_item = re.fullmatch(r"(\d{1,2})\.?", ln)
            if m_item and int(m_item.group(1)) <= 20 and section != "Quadro-resumo":
                item = f"{int(m_item.group(1)):02d}"
                continue
            if section and item:
                blocks.setdefault((section, item), []).append((p_no, ln))
            vals = MONEY.findall(ln) if re.fullmatch(r"R\$\s*[\d.]+,\d{2}", ln) else []
            if ln == "R$" and i + 1 < len(lines):
                vals = MONEY.findall("R$ " + lines[i + 1])
            if vals:
                v = brl(vals[0])
                context = " ".join(" ".join(lines[max(0, i - 8) : i]).split())
                if "TOTAL GERAL" in context.upper():
                    totals.append({"secao": section, "pagina": p_no, "valor": v})
                elif section != "Quadro-resumo":
                    rows.append(
                        {
                            "empresa": company,
                            "secao": section,
                            "pagina": p_no,
                            "item": item,
                            "valor": v,
                        }
                    )
            for m in MONEY_SUMMARY.finditer(ln):
                label = " ".join(" ".join(lines[max(0, i - 6) : i]).split())[-120:]
                summary.append(
                    {
                        "empresa": company,
                        "pagina": p_no,
                        "escrito": m.group(0),
                        "valor": float(re.sub(r"\D", "", m.group(1))),
                        "rotulo": label,
                    }
                )
            m_decl = re.search(
                r"aproximadamente\s+R\$\s*([\d.]+,\d{2})", " ".join(lines[i : i + 2])
            )
            if m_decl and "aproximadamente" in ln:
                declarations.append(
                    {
                        "empresa": company,
                        "secao": section,
                        "pagina": p_no,
                        "item": item,
                        "valor": brl(m_decl.group(1)),
                    }
                )
        # summary values written in full ("R$ 16.924.986,00") inside the quadro-resumo
        if section == "Quadro-resumo":
            for ln_i, ln in enumerate(lines):
                for m in MONEY.finditer(ln):
                    label = " ".join(" ".join(lines[max(0, ln_i - 6) : ln_i]).split())[
                        -120:
                    ]
                    summary.append(
                        {
                            "empresa": company,
                            "pagina": p_no,
                            "escrito": m.group(0),
                            "valor": brl(m.group(1)),
                            "rotulo": label,
                        }
                    )

    ranges = []
    for (sec, it), chunk in blocks.items():
        if sec != "Avaliação executada por Valter":
            continue
        text_by_page: dict[int, str] = {}
        for p_no, ln in chunk:
            text_by_page[p_no] = text_by_page.get(p_no, "") + " " + ln
        for p_no, txt in text_by_page.items():
            flat = " ".join(txt.split())
            seen = set()
            for pat in RANGE_PATTERNS:
                for m in pat.finditer(flat):
                    lo = float(m.group(1).replace(".", ""))
                    hi = float(m.group(2).replace(".", "")) if m.group(2) else None
                    if (lo, hi) in seen:
                        continue
                    seen.add((lo, hi))
                    ranges.append(
                        {
                            "empresa": company,
                            "item": it,
                            "pagina": p_no,
                            "min": lo,
                            "max": hi,
                            "trecho": flat[max(0, m.start() - 140) : m.end() + 30],
                        }
                    )
    return rows, totals, summary, declarations, ranges


# ---------------------------------------------------------------- main


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    sources = {
        "2025 atual 01": one(INV25, "01. "),
        "2025 atual 08": one(INV25, "08. "),
        "2025 atual 00": one(INV25, "00. "),
        "2025 atual 11 (SEI)": one(INV25, "11. "),
        "2025 antigo 00": one(INV25_OLD, "ANTIGO 00. "),
        "2025 antigo 01": one(INV25_OLD, "ANTIGO 01. "),
        "2025 antigo 08": one(INV25_OLD, "ANTIGO 08. "),
        "2018 doc 12 Abastecedora": one(DRIVE2018, "12. 1. "),
        "2018 doc 15 Transportadora": one(DRIVE2018, "15. 1. "),
        "2018 doc 10 Oficina Mecânica": one(DRIVE2018, "10. 1. "),
        "2018 doc 17 Agropecuária": one(DRIVE2018, "17. 1. "),
        "2018 doc 27 Resumo": one(DRIVE2018, "27. 1. "),
        "Rafael 01": one(RAFAEL, "ABASTECEDORA"),
        "Rafael 08": one(RAFAEL, "TRANSPORTADORA"),
        "Rafael Diretrizes": one(RAFAEL, "DIRETRIZES"),
        "Rafael Introdução": one(RAFAEL, "INTRODU"),
        "Lista de Rafael (Valter)": VALTER_LIST,
        "Planilha Google 2_0": GSHEET,
        "Respostas arq. 1": one(REPLIES, "1 - "),
        "Respostas arq. 2": one(REPLIES, "2 - "),
        "Respostas arq. 3": one(REPLIES, "3 - "),
        "Respostas arq. 4": one(REPLIES, "4 - "),
        "Respostas arq. 10": one(REPLIES, "10 - "),
    }

    items, totals = [], []
    plan = [
        ("2025 atual 01", "01 Abastecedora"),
        ("2025 atual 08", "08 Transportadora"),
        ("2025 antigo 01", "01 Abastecedora"),
        ("2025 antigo 08", "08 Transportadora"),
        ("2018 doc 12 Abastecedora", "01 Abastecedora"),
        ("2018 doc 15 Transportadora", "08 Transportadora"),
        ("2018 doc 10 Oficina Mecânica", "Oficina Mecânica (2018)"),
        ("2018 doc 17 Agropecuária", "02 Agropecuária (2018)"),
    ]
    for key, company in plan:
        it, tot = inventory_items(sources[key], company, key)
        if key == "2018 doc 17 Agropecuária":
            # only the Cotiporã lands that are items 07 and 08 of the Transportadora in 2025
            it = [
                x
                for x in it
                if x["item"] in ("07", "08") and "Cotipor" in x["descricao"]
            ]
        items.extend(it)
        totals.extend(tot)

    checks = []

    def check(
        label: str,
        computed: float,
        declared: float,
        note: str = "",
        doc_divergence: bool = False,
    ) -> None:
        ok = abs(computed - declared) < 0.005
        result = (
            "OK"
            if ok
            else ("DIVERGÊNCIA NO DOCUMENTO" if doc_divergence else "DIFERENÇA")
        )
        checks.append(
            {
                "conferencia": label,
                "calculado": computed,
                "declarado": declared,
                "ok": result,
                "nota": note,
            }
        )

    def total_of(source: str) -> list[float]:
        hits = [t for t in totals if t["fonte"] == source and t["valores"]]
        return hits[0]["valores"] if hits else []

    def item_of(source: str, number: str) -> dict:
        return next(x for x in items if x["fonte"] == source and x["item"] == number)

    # -- item sums against the totals printed in each document
    for key, _company in plan:
        if key == "2018 doc 17 Agropecuária":
            continue
        rows = [x for x in items if x["fonte"] == key]
        s = sum(x["valor"] or 0 for x in rows)
        tv = total_of(key)
        missing = [x["item"] for x in rows if x["valor"] is None]
        note = f"{len(rows)} itens"
        if key == "2018 doc 15 Transportadora":
            # total row: gross (-) apartment 102 donated = net
            check(f"{key}: soma dos itens = total bruto", s, tv[0], note)
            check(
                f"{key}: bruto − item 09 (apto 102 doado) = total",
                tv[0] - tv[1],
                tv[2],
                "linha de total: bruto (−) 220.000 = líquido",
            )
        elif missing:
            # value erased from the cell: rebuild it from the composition written in the description
            for x in rows:
                if x["valor"] is None:
                    x["valor_recomposto"] = sum(
                        composition(x["descricao"])
                    ) + thirty_percent(x["descricao"])
                    note += f"; célula de valor do item {x['item']} vazia; recomposto pela descrição = {fmt(x['valor_recomposto'])}"
            s2 = sum(
                (x["valor"] if x["valor"] is not None else x["valor_recomposto"])
                for x in rows
            )
            check(
                f"{key}: soma dos itens (com o item recomposto) = total",
                s2,
                tv[0],
                note,
            )
        else:
            check(f"{key}: soma dos itens = total", s, tv[0], note)

    # -- compositions written inside descriptions
    for x in items:
        if x["valor"] is None:
            continue  # the rebuilt value would be compared with itself
        parts = composition(x["descricao"])
        if x["empresa"].startswith("01") and x["item"] == "01":
            vant = thirty_percent(x["descricao"])
            check(
                f"{x['fonte']} item 01: 30 % × (construída + coberta + equipamentos)",
                round(0.30 * sum(parts[1:4]), 2),
                vant,
            )
            check(
                f"{x['fonte']} item 01: parcelas + 30 % = valor do item",
                sum(parts) + vant,
                x["valor"],
                f"parcelas {', '.join(fmt(p) for p in parts)}",
            )
        if x["fonte"].startswith("2018 doc 10") and x["item"] == "02":
            feita = brl(
                re.search(r"obra feita:\s*R\$\s*([\d.]+,\d{2})", x["descricao"]).group(
                    1
                )
            )
            check(
                f"{x['fonte']} item 02 (pavilhão): CUB + 30 % = valor",
                parts[0] + feita,
                x["valor"],
                f"parcelas {fmt(parts[0])} + {fmt(feita)}",
            )
        if (
            x["fonte"].startswith("2025")
            and x["empresa"].startswith("08")
            and x["item"] in ("07", "08")
        ):
            ha = float(
                re.search(r"([\d,]+)\s*hectares", x["descricao"])
                .group(1)
                .replace(",", ".")
            )
            per_ha = brl(
                re.search(r"valor médio de R\$\s*([\d.]+,\d{2})", x["descricao"]).group(
                    1
                )
            )
            check(
                f"{x['fonte']} item {x['item']}: {fmt(ha)} ha × R$ {fmt(per_ha)}/ha = valor",
                ha * per_ha,
                x["valor"],
                "a conta escrita no item não fecha",
                doc_divergence=True,
            )
        if (
            x["fonte"].startswith("2025")
            and x["empresa"].startswith("08")
            and x["item"] == "01"
        ):
            area = brl(re.search(r"área de ([\d.]+,\d{2})m²", x["descricao"]).group(1))
            per_m2 = brl(
                re.search(
                    r"valor de R\$\s*([\d.]+,\d{2}), ao metro", x["descricao"]
                ).group(1)
            )
            check(
                f"{x['fonte']} item 01: {fmt(area)} m² × R$ {fmt(per_m2)} = valor",
                round(area * per_m2, 2),
                x["valor"],
            )
        if (
            x["fonte"].startswith("2025")
            and x["empresa"].startswith("08")
            and x["item"] == "02"
        ):
            check(
                f"{x['fonte']} item 02: 1.298 m² × R$ 1.100 = valor (30 % não calculado)",
                parts[0],
                x["valor"],
            )

    # -- the R$ 1.135.000 note of the 2025 Transportadora
    for key in ("2025 atual 08", "2025 antigo 08"):
        txt = docx_text(sources[key])
        m = re.search(
            r"IMPORTÂNCIA DE R\$\s*([\d.]+,\d{2}).*?ITENS 03\s*–\s*04\s*–\s*05",
            txt,
            re.DOTALL,
        )
        s = sum(item_of(key, n)["valor"] for n in ("03", "04", "05"))
        check(
            f"{key}: nota 'acrescentado ao capital de bens' × itens 03 + 04 + 05",
            s,
            brl(m.group(1)),
            "diferença de 115.000 sem explicação no documento",
            doc_divergence=True,
        )

    # -- 2018 summary doc 27
    resumo = {}
    for _, _, cells in docx_rows(sources["2018 doc 27 Resumo"]):
        if len(cells) >= 2 and MONEY.search(cells[-1]):
            resumo[cells[0]] = brl(MONEY.findall(cells[-1])[-1])
    for key, label in [
        ("2018 doc 12 Abastecedora", "ABASTECEDORA"),
        ("2018 doc 15 Transportadora", "TRANSPORTADORA"),
        ("2018 doc 10 Oficina Mecânica", "OFICINA"),
    ]:
        r = next(v for k, v in resumo.items() if label in k.upper())
        check(f"{key}: total = doc 27 (resumo 2018)", total_of(key)[-1], r)

    # -- 2025 summaries doc 00 (current and old)
    marks = []
    for key in ("2025 atual 00", "2025 antigo 00"):
        found = {}
        for t_idx, r_idx, cells in docx_rows(sources[key]):
            if (
                len(cells) >= 2
                and re.match(r"\d{2}\.", cells[0])
                and re.search(r"R\$", cells[-1])
            ):
                vals = MONEY.findall(cells[-1])
                found.setdefault(
                    cells[0][:2],
                    (
                        brl(vals[-1]) if vals else None,
                        f"T{t_idx}.R{r_idx}",
                        cells[0],
                        cells[-1],
                    ),
                )
                if " x2" in cells[0] or re.search(r"\d,\d{3},\d{2}", cells[-1]):
                    marks.append(
                        {
                            "fonte": key,
                            "tabela": f"T{t_idx}.R{r_idx}",
                            "texto": f"{cells[0]} | {cells[-1]}",
                        }
                    )
        for company in ("01", "08"):
            src = (
                "2025 atual " + company
                if key == "2025 atual 00"
                else "2025 antigo " + company
            )
            check(f"{src}: total = {key[:-3]} 00", total_of(src)[-1], found[company][0])

    # -- partner shares written in the 2025 inventories
    for key in ("2025 atual 01", "2025 atual 08", "2025 antigo 01", "2025 antigo 08"):
        pct = []
        for _, _, cells in docx_rows(sources[key]):
            m = (
                re.match(r"(\d+,\d+)%\s*=\s*R\$\s*([\d.]+,\d{2})", cells[-1])
                if cells
                else None
            )
            if m and not cells[0].upper().startswith("TOTAL"):
                pct.append((cells[0], brl(m.group(1)), brl(m.group(2))))
        total = total_of(key)[-1]
        check(
            f"{key}: soma dos percentuais do quadro 'valor real de cada sócio'",
            sum(p[1] for p in pct),
            100.0,
            f"{len(pct)} linhas; o documento escreve TOTAL 100,00 %",
            doc_divergence=True,
        )
        for name, p, v in pct:
            if "ROVILIO" in name.upper():
                check(
                    f"{key}: parcela de Rovílio = 2 % do total",
                    round(total * p / 100, 2),
                    v,
                    "",
                    doc_divergence=True,
                )

    # -- Rafael
    raf_rows, raf_totals, raf_summary, raf_decl, raf_ranges = [], [], [], [], []
    for key, company, single in [
        ("Rafael 01", "01 Abastecedora", True),
        ("Rafael 08", "08 Transportadora", False),
    ]:
        r, t, s, d, g = rafael_items(sources[key], company, single)
        for coll in (r, t, s, d, g):
            for x in coll:
                x["fonte"] = key
        raf_rows += r
        raf_totals += t
        raf_summary += s
        raf_decl += d
        raf_ranges += g

    def rafael_value(key: str, section: str, number: str) -> tuple[float, int] | None:
        hit = next(
            (
                x
                for x in raf_rows
                if x["fonte"] == key and x["secao"] == section and x["item"] == number
            ),
            None,
        )
        return (hit["valor"], hit["pagina"]) if hit else None

    adopted = []
    for key, inv_key, company in [
        ("Rafael 01", "2025 atual 01", "01 Abastecedora"),
        ("Rafael 08", "2025 atual 08", "08 Transportadora"),
    ]:
        inv = [x for x in items if x["fonte"] == inv_key]
        for x in inv:
            grupo = x["valor"] if x["valor"] is not None else x["valor_recomposto"]
            shown = rafael_value(key, "Balanço apresentado pelo grupo", x["item"])
            evaluated = rafael_value(key, "Avaliação executada por Valter", x["item"])
            final = rafael_value(key, "Balanço executado por Valter", x["item"])
            decl = next(
                (d for d in raf_decl if d["fonte"] == key and d["item"] == x["item"]),
                None,
            )
            adopted.append(
                {
                    "empresa": company,
                    "item": x["item"],
                    "grupo": grupo,
                    "transcrito": shown[0] if shown else None,
                    "pag_transcrito": shown[1] if shown else None,
                    "avaliacao": evaluated[0] if evaluated else None,
                    "pag_avaliacao": evaluated[1] if evaluated else None,
                    "declaracao": decl["valor"] if decl else None,
                    "pag_declaracao": decl["pagina"] if decl else None,
                    "adotado": final[0] if final else None,
                    "pag_adotado": final[1] if final else None,
                }
            )
            if shown:
                check(
                    f"{key} item {x['item']}: balanço do grupo transcrito por Rafael = inventário 2025",
                    shown[0],
                    grupo,
                    f"p. {shown[1]}",
                )
    for key in ("Rafael 01", "Rafael 08"):
        for sec in ("Balanço apresentado pelo grupo", "Balanço executado por Valter"):
            per_item = {}
            for x in raf_rows:
                if x["fonte"] == key and x["secao"] == sec and x["item"]:
                    per_item.setdefault(x["item"], x["valor"])
            tot = [
                t["valor"]
                for t in raf_totals
                if t["fonte"] == key and t["secao"] == sec
            ]
            if tot and per_item:
                check(
                    f"{key} · {sec}: soma dos itens = total",
                    sum(per_item.values()),
                    tot[0],
                    f"{len(per_item)} itens",
                )
    a01 = next(a for a in adopted if a["empresa"].startswith("01"))
    check(
        "Rafael 01: valor adotado (p. 3) × total da empresa em 2018 (doc 12)",
        a01["adotado"],
        total_of("2018 doc 12 Abastecedora")[-1],
        "o total de 2018 inclui o caminhão-tanque ICW 8342 (145.000), que não existe em 2025",
        doc_divergence=True,
    )
    check(
        "Rafael 01: valor adotado (p. 3) × posto em 2018 (doc 12 item 01)",
        a01["adotado"],
        item_of("2018 doc 12 Abastecedora", "01")["valor"],
        "Rafael usa o total da empresa, não o valor do posto",
        doc_divergence=True,
    )
    check(
        "Rafael 01: valor adotado (p. 3) × parcelas da descrição que ele mesmo transcreve",
        a01["adotado"],
        a01["grupo"],
        "a descrição da p. 3 é a de 2025 e soma 9.661.414",
        doc_divergence=True,
    )
    for key in ("Rafael 01", "Rafael 08"):
        vals = [s for s in raf_summary if s["fonte"] == key]
        total = max(v["valor"] for v in vals)
        sixth = sorted({v["valor"] for v in vals if v["valor"] < total})
        if sixth:
            check(
                f"{key}: quadro-resumo, 1/6 do patrimônio",
                round(total / 6, 2),
                sixth[-1],
                f"escrito {next(v['escrito'] for v in vals if v['valor'] == sixth[-1])}; total {fmt(total)}",
                doc_divergence=True,
            )

    # -- Rovílio contract: price × sum of the payment items (from the transcription checked on the page images)
    tr = ROVILIO_TRANSCRIPTION.read_text(encoding="utf-8")
    price = brl(re.search(r"importância de \*\*<u>R\$\s*([\d.]+,\d{2})", tr).group(1))
    pay = [
        brl(v)
        for v in re.findall(r"^\| R\$\s*([\d.]+,\d{2}) \| Item \d+", tr, re.MULTILINE)
    ]
    check(
        "arq. 10 (Rovílio): soma dos itens 1 a 10 = preço da cláusula 3ª",
        sum(pay),
        price,
        f"{len(pay)} itens (transcrição conferida nas p. 4–6)",
    )

    # -- 2018 x 2025 pairs by registry key
    old = [x for x in items if x["fonte"].startswith("2018")]
    new = [
        x
        for x in items
        if x["fonte"] == "2025 atual 01" or x["fonte"] == "2025 atual 08"
    ]
    sei_text = docx_text(sources["2025 atual 11 (SEI)"])
    sei_keys = {}
    for _, _, cells in docx_rows(sources["2025 atual 11 (SEI)"]):
        if cells and re.fullmatch(r"\d{2}", cells[0]) and len(cells) >= 3:
            for k in keys_of(cells[1]):
                sei_keys.setdefault(k, (cells[0], cells[1][:90], cells[-1]))
            m = re.search(
                r"oriunda da unificação das Matrículas nº ([\d.]+) e ([\d.]+)", cells[1]
            )
            if m:
                unified = re.search(r"Matrícula nº ([\d.]+)\s*–\s*\(oriunda", cells[1])
                # overwrite: the old matrículas were already added above without the unified number
                for num in m.groups():
                    sei_keys["mat. " + num.replace(".", "")] = (
                        cells[0],
                        cells[1][:90],
                        cells[-1],
                        unified.group(1).replace(".", "") if unified else None,
                    )
    pairs = []
    no_key_pairs = {
        ("2018 doc 10 Oficina Mecânica", n): ("2025 atual 08", n)
        for n in ("02", "03", "04", "05")
    }
    for o in old:
        ok = keys_of(o["descricao"])
        match = [n for n in new if ok & keys_of(n["descricao"])]
        if match:
            n = match[0]
            pairs.append(
                {
                    "o": o,
                    "n": n,
                    "chave": ", ".join(sorted(ok & keys_of(n["descricao"]))),
                    "tipo": "par por chave",
                }
            )
        elif (o["fonte"], o["item"]) in no_key_pairs:
            src, num = no_key_pairs[(o["fonte"], o["item"])]
            pairs.append(
                {
                    "o": o,
                    "n": item_of(src, num),
                    "chave": "sem chave; mesma posição na lista; localização não documentada (arq. 1 p. 16 × arq. 4 p. 1 e 3)",
                    "tipo": "par por descrição — confirmar",
                }
            )
        else:
            sei = [sei_keys[k] for k in ok if k in sei_keys]
            if sei:
                pairs.append(
                    {
                        "o": o,
                        "n": None,
                        "chave": ", ".join(sorted(k for k in ok if k in sei_keys))
                        + (
                            f" → mat. {sei[0][3]} (unificação)"
                            if len(sei[0]) > 3 and sei[0][3]
                            else ""
                        ),
                        "tipo": f"migrou para a SEI (doc 11 item {sei[0][0]}, {sei[0][2]})",
                    }
                )
            else:
                elsewhere = ok & keys_of(sei_text)
                pairs.append(
                    {
                        "o": o,
                        "n": None,
                        "chave": ", ".join(sorted(ok)) or "sem chave",
                        "tipo": "sem par em 2025"
                        + (" (chave citada no doc 11)" if elsewhere else ""),
                    }
                )
    for n in new:
        if not any(p["n"] is n for p in pairs):
            pairs.append(
                {
                    "o": None,
                    "n": n,
                    "chave": ", ".join(sorted(keys_of(n["descricao"]))),
                    "tipo": "sem par em 2018",
                }
            )

    # -- Google sheet: formulas and cached values, with headers
    wf = openpyxl.load_workbook(GSHEET, data_only=False)
    wv = openpyxl.load_workbook(GSHEET, data_only=True)
    gs_rows = []
    for tab, rng in [
        ("1.Abastecedora", [1] + list(range(3, 10))),
        ("8.Transportadora", list(range(1, 20))),
        ("GERAL", [1, 3] + list(range(5, 10)) + list(range(116, 134))),
    ]:
        for r in rng:
            for c in wf[tab][r]:
                if c.value is None:
                    continue
                cached = wv[tab][c.coordinate].value
                gs_rows.append(
                    {
                        "aba": tab,
                        "celula": c.coordinate,
                        "formula": c.value
                        if isinstance(c.value, str) and c.value.startswith("=")
                        else None,
                        "valor": cached,
                    }
                )
    gv = lambda tab, cell: wv[tab][cell].value
    check(
        "Planilha 1.Abastecedora C5 (2018 apresentado) = doc 12 (2018)",
        gv("1.Abastecedora", "C5"),
        total_of("2018 doc 12 Abastecedora")[-1],
    )
    check(
        "Planilha 1.Abastecedora F5/G5 (2025) = doc 00",
        gv("1.Abastecedora", "G5"),
        total_of("2025 atual 01")[-1],
    )
    check(
        "Planilha GERAL E5 (2025) = doc 00",
        gv("GERAL", "E5"),
        total_of("2025 atual 01")[-1],
    )
    check(
        "Planilha 8.Transportadora F2/G2 (2025) = doc 00",
        gv("8.Transportadora", "G2"),
        total_of("2025 atual 08")[-1],
    )
    check(
        "Planilha GERAL E116 (2025) = doc 00",
        gv("GERAL", "E116"),
        total_of("2025 atual 08")[-1],
    )
    oficina_01_05 = sum(
        item_of("2018 doc 10 Oficina Mecânica", n)["valor"]
        for n in ("01", "02", "03", "04", "05")
    )
    check(
        "Planilha 8.Transportadora C2 (2018 apresentado) = Oficina itens 01–05 + Transportadora 2018 (líquido)",
        oficina_01_05 + total_of("2018 doc 15 Transportadora")[-1],
        gv("8.Transportadora", "C2"),
        "base composta: não é o total de nenhum documento de 2018; a coluna D ('2018: atualizado') repete o valor de 2025 nos itens que existem",
    )
    author_notes = [
        ("1.Abastecedora", "I7"),
        ("GERAL", "H7"),
        ("8.Transportadora", "C10"),
        ("8.Transportadora", "C11"),
        ("GERAL", "C124"),
        ("GERAL", "C125"),
        ("GERAL", "H874"),
        ("GERAL", "E874"),
        ("Resumo", "C63"),
        ("Resumo", "B67"),
        ("Resumo", "B73"),
        ("Resumo", "C73"),
    ]

    # ---------------------------------------------------------------- write workbook
    out = openpyxl.Workbook()
    ws = out.active
    ws.title = "Fontes"
    ws.append(["fonte", "arquivo", "md5 (10)"])
    lines = []
    for k, p in sources.items():
        h = md5_10(p)
        ws.append([k, str(p), h])
        lines.append(f"{h}|{p.stat().st_size}|{p}")
    HASHES.write_text("\n".join(lines) + "\n", encoding="utf-8")

    ws = out.create_sheet("Itens inventários")
    cols = [
        "empresa",
        "fonte",
        "tabela",
        "item",
        "valor",
        "valor_recomposto",
        "celula_valor",
        "descricao",
    ]
    ws.append(
        [
            "Empresa",
            "Fonte",
            "Tabela.linha",
            "Item",
            "Valor (R$)",
            "Valor recomposto (R$)",
            "Texto da célula de valor",
            "Descrição",
        ]
    )
    for x in items:
        ws.append([x.get(c) for c in cols])

    ws = out.create_sheet("Totais declarados")
    ws.append(["Fonte", "Tabela.linha", "Rótulo", "Valores na linha de total (R$)"])
    for t in totals:
        ws.append(
            [
                t["fonte"],
                t["tabela"],
                t["rotulo"],
                " | ".join(fmt(v) for v in t["valores"]),
            ]
        )

    ws = out.create_sheet("Rafael adotado")
    ws.append(
        [
            "Empresa",
            "Item",
            "Grupo 2025 (R$)",
            "Rafael: balanço do grupo transcrito (R$)",
            "p.",
            "Rafael: seção 'avaliação' (R$)",
            "p.",
            "Rafael: declaração (R$)",
            "p.",
            "Rafael: valor adotado, balanço executado por Valter (R$)",
            "p.",
            "Diferença adotado − grupo (R$)",
        ]
    )
    for a in adopted:
        diff = (a["adotado"] - a["grupo"]) if a["adotado"] is not None else None
        ws.append(
            [
                a["empresa"],
                a["item"],
                a["grupo"],
                a["transcrito"],
                a["pag_transcrito"],
                a["avaliacao"],
                a["pag_avaliacao"],
                a["declaracao"],
                a["pag_declaracao"],
                a["adotado"],
                a["pag_adotado"],
                diff,
            ]
        )

    ws = out.create_sheet("Rafael itens")
    ws.append(["Empresa", "Fonte", "Seção", "Página", "Item", "Valor (R$)"])
    for x in raf_rows:
        ws.append(
            [x["empresa"], x["fonte"], x["secao"], x["pagina"], x["item"], x["valor"]]
        )
    for t in raf_totals:
        ws.append(["", t["fonte"], t["secao"], t["pagina"], "TOTAL GERAL", t["valor"]])
    for d in raf_decl:
        ws.append(
            [
                d["empresa"],
                d["fonte"],
                f"{d['secao']} (declaração 'aproximadamente')",
                d["pagina"],
                d["item"],
                d["valor"],
            ]
        )

    ws = out.create_sheet("Rafael quadro-resumo")
    ws.append(
        [
            "Empresa",
            "Fonte",
            "Página",
            "Como está escrito",
            "Valor lido (R$)",
            "Texto anterior",
        ]
    )
    for s in raf_summary:
        ws.append(
            [
                s["empresa"],
                s["fonte"],
                s["pagina"],
                s["escrito"],
                s["valor"],
                s["rotulo"],
            ]
        )

    ws = out.create_sheet("Rafael faixas")
    ws.append(["Empresa", "Item", "Página", "Mínimo (R$)", "Máximo (R$)", "Trecho"])
    for g in raf_ranges:
        ws.append(
            [g["empresa"], g["item"], g["pagina"], g["min"], g["max"], g["trecho"]]
        )

    ws = out.create_sheet("Pares 2018×2025")
    ws.append(
        [
            "Fonte 2018",
            "Item 2018",
            "Valor 2018 (R$)",
            "Descrição 2018",
            "Chave",
            "Tipo de par",
            "Fonte 2025",
            "Item 2025",
            "Valor 2025 (R$)",
            "Descrição 2025",
        ]
    )
    for p in pairs:
        o, n = p["o"], p["n"]
        ws.append(
            [
                o["fonte"] if o else None,
                o["item"] if o else None,
                (o["valor"] if o["valor"] is not None else o["celula_valor"])
                if o
                else None,
                o["descricao"][:160] if o else None,
                p["chave"],
                p["tipo"],
                n["fonte"] if n else None,
                n["item"] if n else None,
                (n["valor"] if n["valor"] is not None else n.get("valor_recomposto"))
                if n
                else None,
                n["descricao"][:160] if n else None,
            ]
        )

    ws = out.create_sheet("Planilha Google")
    ws.append(["Aba", "Célula", "Fórmula", "Valor (cache da exportação)"])
    for r in gs_rows:
        ws.append([r["aba"], r["celula"], r["formula"], r["valor"]])

    ws = out.create_sheet("Notas do autor (planilha)")
    ws.append(["Aba", "Célula", "Fórmula", "Valor / texto"])
    for tab, cell in author_notes:
        f = wf[tab][cell].value
        ws.append(
            [
                tab,
                cell,
                f if isinstance(f, str) and f.startswith("=") else None,
                wv[tab][cell].value,
            ]
        )

    ws = out.create_sheet("Marcas doc 00")
    ws.append(["Fonte", "Tabela.linha", "Texto"])
    for m in marks:
        ws.append([m["fonte"], m["tabela"], m["texto"]])

    ws = out.create_sheet("Conferências")
    ws.append(
        [
            "Conferência",
            "Calculado (R$ ou %)",
            "Declarado (R$ ou %)",
            "Resultado",
            "Nota",
        ]
    )
    for c in checks:
        ws.append(
            [c["conferencia"], c["calculado"], c["declarado"], c["ok"], c["nota"]]
        )

    for sheet in out.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
    out.save(OUT)

    print(f"saved {OUT}")
    print(
        f"items: {len(items)}; rafael values: {len(raf_rows)}; summary values: {len(raf_summary)}; declarations: {len(raf_decl)}; ranges: {len(raf_ranges)}; pairs: {len(pairs)}; google cells: {len(gs_rows)}"
    )
    for c in checks:
        print(
            f"[{c['ok']}] {c['conferencia']}: {fmt(c['calculado'])} x {fmt(c['declarado'])} {c['nota']}"
        )


if __name__ == "__main__":
    main()
