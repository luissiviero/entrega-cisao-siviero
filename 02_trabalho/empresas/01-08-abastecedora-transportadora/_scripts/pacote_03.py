"""Figures for packet 3 (credits and contingencies) of companies 01 and 08.

Checks that neither the group nor Rafael deducts any contingency, reduction
or liability from the totals of 01 and 08 (the totals equal the sum of the
assets), lists the rows of the 'Créditos e contingências' sheet by proposed
rule and prints the effect on 1/6 of the options shown to Luís.

Usage: python -I pacote_03.py
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import extract_fase0 as fase0

WORK = Path(__file__).resolve().parents[1]
MASTER = WORK / "planilha-mestre-01-08.xlsx"
FASE0 = WORK / "_extracao" / "fase0-itens-01-08.xlsx"


def rows_of(path: Path, sheet: str) -> list[dict]:
    ws = openpyxl.load_workbook(path)[sheet]
    header = [c.value for c in ws[1]]
    return [
        dict(zip(header, r, strict=False))
        for r in ws.iter_rows(min_row=2, values_only=True)
    ]


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    f = fase0.fmt
    checks = rows_of(FASE0, "Conferências")
    for c in checks:
        if c["Conferência"].startswith(
            (
                "2025 atual 01: soma",
                "2025 atual 08: soma",
                "Rafael 08 · Balanço executado",
            )
        ):
            print(
                f"[{c['Resultado']}] {c['Conferência']}: {f(c['Calculado (R$ ou %)'])} = {f(c['Declarado (R$ ou %)'])}"
            )

    cont = rows_of(MASTER, "Créditos e contingências")
    rules = {
        "C1 (nada deduzido)": [
            r
            for r in cont
            if r["Natureza (como o documento trata)"].startswith(
                ("Redução", "Contingência", "Custo", "Ônus", "Obrigação")
            )
        ],
        "C2 (créditos e receitas sem valor)": [
            r
            for r in cont
            if r["Natureza (como o documento trata)"].startswith(("Crédito", "Receita"))
        ],
        "C3 (nota contábil)": [
            r for r in cont if r["Natureza (como o documento trata)"].startswith("Nota")
        ],
    }
    covered = 0
    for rule, rows in rules.items():
        covered += len(rows)
        print(f"{rule}: {len(rows)} linhas")
        for r in rows:
            print(
                f"   {r['Empresa'][:2]} | {r['Item'][:60]} | {f(r['Valor (R$)']) if r['Valor (R$)'] is not None else 'sem valor'} | {r['Natureza (como o documento trata)'][:50]}"
            )
    print(f"linhas cobertas: {covered} de {len(cont)}")

    joveni = next(r for r in cont if r["Item"].startswith("Cobrança contra Joveni"))
    print(
        f"opção de incluir o crédito Joveni (caso extremo): +{f(joveni['Valor (R$)'])}; 1/6 +{f(joveni['Valor (R$)'] / 6)}"
    )
    red = sum(r["Valor (R$)"] for r in cont if r["Item"].startswith("Redução estimada"))
    inss = sum(r["Valor (R$)"] for r in cont if r["Item"].startswith("INSS"))
    print(
        f"reduções anunciadas sobre os totais do grupo (0,70 % + 1,5 %): {f(red)}; 1/6 {f(red / 6)}; INSS estimado: {f(inss)}; 1/6 {f(inss / 6)} (INSS e 1,5 % tratam do mesmo fato: não somar)"
    )
    # the same percentages on the values adopted in packets 1 and 2
    master = rows_of(MASTER, "Mestre")
    adopted = {
        co: sum(
            r["Valor adotado (R$)"] or 0
            for r in master
            if r["ID"].startswith(f"{co}-2025")
        )
        for co in ("01", "08")
    }
    pct = {"01": 0.70, "08": 1.50}
    red_adopted = {co: round(adopted[co] * pct[co] / 100, 2) for co in pct}
    total_red = sum(red_adopted.values())
    print(
        f"reduções sobre os valores adotados: 01 {f(red_adopted['01'])} (0,70 % × {f(adopted['01'])}); 08 {f(red_adopted['08'])} (1,5 % × {f(adopted['08'])}); soma {f(total_red)}; 1/6 {f(total_red / 6)}"
    )
    # 30 % "vantagem da coisa feita" left out of the unregularized works of the 08 (same method as the station of the 01)
    works = [r for r in master if r["ID"] in ("08-2025-02", "08-2025-04", "08-2025-05")]
    base = sum(r["Valor adotado (R$)"] for r in works)
    print(
        f"30 % fora das obras da 08 (pavilhão, rampa, borracharia): 30 % × {f(base)} = {f(0.30 * base)}"
    )


if __name__ == "__main__":
    main()
