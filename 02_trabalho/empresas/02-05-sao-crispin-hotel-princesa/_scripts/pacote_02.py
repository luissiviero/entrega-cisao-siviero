"""Write pacotes/pacote-02.md (exceptions) from planilha-mestre-02-05.xlsx.

Exceptions of the Hotel, largest value first: farm (item 05) with the 4 ha
(item 06) and the R$ 400.000 regularization, hotel building (item 01),
truck (item 08). Values, pages and quotes come from the master spreadsheet,
the Phase 0 extraction and the sources; labels are typed.

v2 (28/10/2026) applies `_diario/revisao-pacote-02.md` (A1-A3, M1-M5, B1-B5):
options side by side with the effect on the exiting 1/6.

Run: PYTHONIOENCODING=utf-8 python -W ignore pacote_02.py
"""

from __future__ import annotations

import contextlib
import io
import os
import re
import sys
from decimal import ROUND_HALF_UP, Decimal

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_mestre_02_05 as bm
import fase0_extracao as f0

OUT = os.path.join(bm.DIR, "pacotes", "pacote-02.md")
LOG_0108 = os.path.join(
    f0.ROOT, r"02_trabalho\empresas\01-08-abastecedora-transportadora\log-decisoes.md"
)
SIXTH = Decimal(1) / Decimal(6)


def money(v) -> str:
    return "—" if v in (None, "") else f"R$ {f0.br(Decimal(str(v)))}"


def q2(v: Decimal) -> Decimal:
    return v.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def effect_stays(delta: Decimal) -> Decimal:
    """Asset that stays: the company total rises by delta; the quota by delta/6."""
    return q2(delta * SIXTH)


def effect_comes(delta_lower: Decimal) -> Decimal:
    """Asset that comes to us valued delta lower: we receive it cheaper (+delta)
    but the company total falls by delta (-delta/6): net +5/6 delta."""
    return q2(delta_lower - delta_lower * SIXTH)


def millions(text: str) -> list[Decimal]:
    return [f0.num(x) * 1_000_000 for x in re.findall(r"R\$ ([\d,]+) milhões", text)]


def main() -> None:
    with contextlib.redirect_stdout(io.StringIO()):
        bm.main()
    wbm = openpyxl.load_workbook(bm.OUT, data_only=True)
    hdr = [c.value for c in wbm["Mestre"][1]]
    rows = [
        dict(zip(hdr, r)) for r in wbm["Mestre"].iter_rows(min_row=2, values_only=True)
    ]
    col = {k: t for k, t, _ in bm.COLS}

    def g(r, k):
        v = r[col[k]]
        return Decimal(str(v)) if isinstance(v, (int, float)) else v

    by = {g(r, "item"): r for r in rows if str(g(r, "chave")).startswith("05-2025")}
    adopted = {
        k: g(r, "valor_adotado")
        for k, r in by.items()
        if g(r, "valor_adotado") is not None
    }
    faixas = [
        (f[1], f[2], f[3])
        for f in openpyxl.load_workbook(f0.OUT_XLSX, data_only=True)[
            "Rafael 05 faixas"
        ].iter_rows(min_row=4, values_only=True)
    ]
    passivo = next(
        wbm["Passivos e deduções"].iter_rows(min_row=2, max_row=2, values_only=True)
    )
    reduction = Decimal(str(passivo[2]))
    r05, r06, r01, r02, r08 = by["05"], by["06"], by["01"], by["02"], by["08"]

    def page(path: str, n: int) -> str:
        return " ".join(f0.pymupdf.open(path)[n - 1].get_text().split())

    raf = f0.SRC["rafael_05"]
    p13, p14, p15 = page(raf, 13), page(raf, 14), page(raf, 15)
    prelim = f0.need(
        r"os valores atualmente apresentados deverão ser considerados parâmetros preliminares de referência[^.]*\.",
        p14,
        "p. 14",
    ).group(0)
    inferior = f0.need(
        r"tendo como resultado valores inferiores[^.]*\.", p13, "p. 13"
    ).group(0)
    arrend = f0.need(
        r"eventual valor econômico decorrente de exploração mercantil por arrendamento",
        p13 + " " + p14,
        "arrendamento",
    ).group(0)
    conc06 = f0.need(
        r"de acordo como citado nos itens anteriores \(02\., 03\., 04\.,05\.\)",
        p14,
        "p. 14 remissão",
    ).group(0)
    t05 = " ".join(f0.docx_paragraphs(f0.SRC["inv2025_05"]))
    nao_liquido = f0.need(
        r"O valor patrimonial lançado sofrerá redução pela regularização dos imóveis[^.]*?R\$ "
        + f0.MONEY
        + r"[^.]*\.",
        t05,
        "redução",
    ).group(0)
    quem_paga = f0.need(
        r"que deverão ser satisfeitas pela atual credora e proprietária das áreas de terras HOTEL PRINCESA DOS VALES LTDA",
        t05,
        "quem paga",
    ).group(0)
    rate05 = f0.need(
        r"R\$ (" + f0.MONEY + r") = R\$", g(r05, "descricao"), "R$/ha"
    ).group(1)
    a1 = f0.SRC_LATER["respostas_arq01"]
    pablo = f0.need(
        r"Esclarecer o motivo para o aumento do preço para R\$110\.000,00/hectare, uma vez que ainda estamos recebendo o pagamento de áreas vendidas por R\$ [\d.,]+/hectare\.",
        page(a1, 5),
        "Pablo",
    ).group(0)
    sem_prov = f0.need(
        r"ainda não se deu início a qualquer providência nesse sentido",
        page(a1, 6),
        "p. 6",
    ).group(0)
    fipe_grupo = f0.need(
        r"preço médio de R\$ (" + f0.MONEY + ")", page(a1, 9), "FIPE grupo"
    ).group(1)
    sales, _ = f0.sale_rows(f0.SRC["inv2018_17"])
    rigo = next(s for s in sales if s["comprador"] == "RIGO")
    rigo_ha = f0.brl(
        f0.need(
            r"ao valor de R\$ (" + f0.MONEY + "), ao hectare",
            rigo["texto"],
            "R$/ha Rigo",
        ).group(1)
    )
    fl18 = f0.farm_lines(
        next(
            x
            for x in f0.item_rows(f0.SRC["inv2018_17"], "02", "2018")[0]
            if x["seq"] == 1
        )["descricao"]
    )
    rate18 = sorted({x[2] for x in fl18}, reverse=True)
    with open(LOG_0108, encoding="utf-8") as fh:
        prec_0108 = next(
            ln for ln in fh if "ônus de dissolução" in ln and "Nada é descontado" in ln
        )
    built = f0.num(
        f0.need(r"de terreno: ([\d.,]+)m²", g(r01, "descricao"), "720").group(1)
    )

    # ---- item 01: Rafael's own range, restricted to the 720 m² of item 01
    m2txt = next(
        t for i, p, t in faixas if i == "01" and "por m²" in t and "esquina" in t
    )
    m2_lo, m2_hi = (
        f0.num(x) for x in re.findall(r"R\$ ([\d.]+) e R\$ ([\d.]+) por m²", m2txt)[0]
    )
    build_txt = next(
        t for i, p, t in faixas if i == "01" and "milhões" in t and "2.040" in t
    )
    b_lo, b_hi = millions(build_txt)
    item01_lo, item01_hi = built * m2_lo + b_lo, built * m2_hi + b_hi
    g01, raf01 = g(r01, "valor_grupo"), g(r01, "valor_rafael")
    ratio = (raf01 / g01).quantize(Decimal("0.0001"))

    # ---- effects on the exiting 1/6
    e01 = effect_stays(raf01 - g01)
    e08 = effect_comes(g(r08, "valor_rafael") - g(r08, "valor_grupo"))
    e400_farm = effect_comes(reduction)
    e400_liab = -q2(reduction * SIXTH)
    gross_rec = (
        sum(adopted.values())
        + g(r05, "valor_grupo")
        + g(r06, "valor_grupo")
        + raf01
        + g(r08, "valor_grupo")
    )
    total_group = sum(g(r, "valor_grupo") for r in by.values())
    total_raf = sum(g(r, "valor_rafael") for r in by.values())

    L: list[str] = []
    w = L.append
    w("# Pacote 2 — Hotel Princesa dos Vales (05): exceções")
    w("")
    w(
        "Rascunho revisado (revisão independente em `_diario/revisao-pacote-02.md`). Números gerados por `_scripts/pacote_02.py`. **Nada aqui está aprovado.**"
    )
    w("")
    w(
        'Critérios usados: **C0** — bem da lista de Rafael (vem para nós) fica com o menor valor; bem que fica na empresa, com o maior. **C1** — faixa de mercado do próprio Rafael, como argumento técnico. **P5** — adota-se um dos dois valores; um terceiro só como "caso extremo", com justificativa escrita.'
    )
    w("")
    w(
        '"Efeito na cota" = quanto a decisão muda o que cabe ao Luís (1/6): num bem que fica, 1/6 da diferença; num bem que vem para nós, recebê-lo por um valor menor rende 5/6 da diferença.'
    )
    w("")
    w("## Resumo")
    w("")
    w("| Ponto | Grupo | Rafael | Opções | Minha recomendação | Efeito na cota |")
    w("|---|---|---|---|---|---|")
    w(
        f"| Fazenda (05) + 4 ha (06), na lista de Rafael | {money(g(r05, 'valor_grupo'))} + {money(g(r06, 'valor_grupo'))} | iguais | Adotar o valor comum; pedir a avaliação independente que Rafael recomenda | Adotar o valor comum e pedir a avaliação, aceitando o resultado em qualquer direção | — (até haver avaliação) |"
    )
    w(
        f"| Redução de {money(reduction)} para regularizar a fazenda | anunciada, não abatida | repete o texto | (a) não abater; (b) abater como redução da fazenda (caso extremo) | (a) não abater, por coerência com as empresas 01 e 08 | (a) zero; (b) + {money(e400_farm)} |"
    )
    w(
        f"| Prédio do hotel (01), fica na empresa | {money(g01)} | {money(raf01)} | C0: Rafael; C1: grupo | Rafael (C0), sabendo que é frágil | + {money(e01)} com Rafael |"
    )
    w(
        f"| Caminhão Mercedinha (08), na lista de Rafael | {money(g(r08, 'valor_grupo'))} | {money(g(r08, 'valor_rafael'))} | C0 e FIPE: grupo | Grupo | + {money(e08)} com o grupo |"
    )
    w("")
    w(
        f"Com as recomendações e o pacote 1, o Hotel fica em **{money(gross_rec)}** (grupo {money(total_group)}; Rafael {money(total_raf)}; todos sem abater a redução)."
    )
    w("")
    w("## 1. Fazenda São Crispin: item 05 (407,53 ha) e item 06 (4 ha)")
    w("")
    w(
        f"- **Valores:** grupo e Rafael dão o mesmo valor: {money(g(r05, 'valor_grupo'))} e {money(g(r06, 'valor_grupo'))}, a R$ {rate05} por hectare para toda a área, produtiva ou não. Os dois estão na lista de Rafael; pelo C0 valeria o menor, mas não há valor menor em nenhum dos dois estudos."
    )
    w(
        f'- **O que Rafael escreve:** não há bloco de concordância logo após o item 05, mas as concordâncias dos itens 06 e 07 remetem a ele ("{conc06}", p. 14 e 15), e o quadro-resumo dele usa esse valor. Ao mesmo tempo, diz que "{prelim}" (p. 14), que a avaliação por uso do solo teve "{inferior[len("tendo como resultado ") :]}" (p. 13) e recomenda avaliação técnica independente, que considere também o "{arrend}" (p. 13–14).'
    )
    w(
        "- **Indícios de que o valor pode ser alto** (não são faixa de mercado de Rafael; servem para pedir a avaliação):"
    )
    w(
        f"  - a venda real aos Rigo em 2018 foi a {money(rigo_ha)} por hectare (inventário 2018);"
    )
    w(
        f"  - o inventário de 2018 diferenciava {money(rate18[0])}/ha para área produtiva e {money(rate18[-1])}/ha para não produtiva; o de 2025 usa R$ {rate05}/ha para tudo; a versão antiga do inventário 2025 usava {money(g(r05, 'valor_antigo') / Decimal('407.53'))}/ha ({money(g(r05, 'valor_antigo'))} no item 05);"
    )
    w(
        f'  - o Dr. Pablo perguntou ao grupo (arq. 1, p. 5): "{pablo}" O grupo respondeu que é "uma média por hectare" (p. 6);'
    )
    w(
        '  - suas notas na planilha: "Avaliação por produtividade?" e "Melhor avaliação é verificada pelo preço do arrendo".'
    )
    w(
        f'- **Recomendação:** adotar {money(g(r05, "valor_grupo"))} e {money(g(r06, "valor_grupo"))} ("adoto o valor do grupo, igual ao de Rafael") e pedir ao grupo a avaliação independente que o próprio Rafael recomenda. Antes de pedir, decida se aceita o laudo **em qualquer direção**: o Rafael quer que ele considere o arrendamento e a capacidade produtiva, o que pode subir o valor. Usar o laudo só se vier menor seria uma assimetria visível. Se o laudo não vier, o valor adotado fica definitivo. Falta combinar quem encomenda e paga (o grupo disse que prefere evitar despesas nesta fase, arq. 1, p. 6).'
    )
    w(
        "- **Ressalvas:** a área avaliada (407,53 + 4,0 = 411,53 ha) difere em 0,1 ha dos 411,4 ha do texto e dos 411,43 ha do inventário de 2018 (cerca de R$ 11 mil); a matrícula 00.760 aparece duas vezes (48 e 100 ha) — o grupo explica que são duas compras e que restam 64,53 ha (arq. 1, p. 9), a confirmar (Q3); a 11.090 também está numa parcela vendida aos Rigo (Q1); o ITR de 2024 pago pela SEI declara 259 ha."
    )
    w("")
    w(f"## 2. Redução de {money(reduction)} para regularizar a fazenda")
    w("")
    w(
        f'- **O que o grupo escreve (docs 02 e 05; Rafael repete):** "{nao_liquido}" E também que as despesas "{quem_paga}". O texto aparece nos dois inventários e nos dois PDFs de Rafael, mas é uma redução só. Nenhum dos lados a abate: o total do grupo ({money(total_group)}), a cota por sócio do próprio doc 05, o doc 00 e o total de Rafael são brutos.'
    )
    w(
        f'- **Precedente:** nas empresas 01 e 08 você aprovou: "{prec_0108.split("|")[3].strip()}" para as reduções e ônus anunciados pelo grupo do mesmo jeito.'
    )
    w(
        "- **Opção (a), não abater:** coerente com as empresas 01 e 08 e com a P5 (nenhum dos lados aplica). O fato fica registrado. Efeito na cota: zero."
    )
    w(
        f"- **Opção (b), abater como redução da fazenda:** a fazenda vem para nós {money(reduction)} mais barata e o total cai o mesmo: efeito na cota **+ {money(e400_farm)}**. Exige justificativa de caso extremo (é o único caso em que o grupo dá um valor em reais). O grupo vai apontar a assimetria com as empresas 01 e 08 e o risco de benefício duplo, se depois a empresa ainda pagar a transferência da fazenda para nós."
    )
    w(
        f"- Se a redução fosse tratada como dívida da empresa (o texto diz que quem paga é o Hotel), o efeito seria o contrário: **menos {money(-e400_liab)}** na cota."
    )
    w(
        f'- O custo é de uma operação que não começou ("{sem_prov}", arq. 1, p. 6): transferir as escrituras do Hotel e de terceiros para a Agropecuária. Se a fazenda vier para nós, essa transferência vira outra, com custos próprios.'
    )
    w(
        '- **Recomendação:** (a) não abater; tratar quem paga a regularização e a transferência da fazenda no instrumento da cisão. Os "ônus de eventual dissolução" citados no mesmo texto ficam como aviso para a negociação, como nas empresas 01 e 08.'
    )
    w("")
    w("## 3. Item 01 — Prédio do hotel e terreno sob o prédio")
    w("")
    w(
        f'- **Valores:** grupo {money(g01)}, com memória de cálculo (terreno de {f0.br(built, 0)} m², prédio, subsolo, reforma, móveis e 25 % de "coisa feita"); Rafael {money(raf01)} no balanço executado (p. 16), **sem cálculo em nenhum documento dele**: os componentes impressos logo abaixo são os do grupo, e na seção de avaliação (p. 7) ele repete o valor do grupo. O valor dele entra no total ({money(total_raf)}) e na cota de 1/6 do quadro-resumo (p. 19–20). É quase exatamente 1,30 × o valor do grupo (razão {f0.br(ratio, 4)}) — hipótese de acréscimo de 30 %, sem documento.'
    )
    w(
        f"- **C0:** o bem fica na empresa → vale o maior: **Rafael**. Efeito na cota: + {money(e01)}."
    )
    w(
        f'- **C1 aponta o contrário:** a faixa do próprio Rafael para o item 01 — os {f0.br(built, 0)} m² de terreno a R$ {f0.br(m2_lo, 0)}–{f0.br(m2_hi, 0)}/m² (p. 8) mais o prédio a {money(b_lo)}–{money(b_hi)} (p. 9) — vai de {money(item01_lo)} a {money(item01_hi)}. O valor do grupo está dentro; o de Rafael fica {money(raf01 - item01_hi)} acima do teto. Se a faixa do prédio ("valor estimado de mercado para o Hotel") já inclui o terreno, o teto cai para {money(b_hi)}.'
    )
    w(
        f"- **Outros riscos:** a mesma frase de Rafael (R$ {f0.br(m2_lo, 0)}–{f0.br(m2_hi, 0)}/m²) pesa contra os R$ 7.500/m² do item 02, já aprovado ({money(g(r02, 'valor_grupo'))}); as faixas são de 2026 (data-base, RDB-4)."
    )
    w(
        '- **Precedente Oleoplan (SEI):** você adotou o valor de Rafael mesmo com o risco de "erro de soma". Diferenças: lá o bem vinha para nós e o valor adotado foi o menor; o valor de Rafael estava em dois documentos (PDF do empreendimento e Diretrizes); aqui está numa página só e a faixa dele não o sustenta.'
    )
    w(
        f"- **Recomendação:** é uma escolha real. Pelo C0, Rafael ({money(raf01)}, + {money(e01)} na cota), sabendo que é o valor mais frágil do Hotel. Pelo C1, grupo ({money(g01)}), mais fácil de defender. Recomendo Rafael, pela regra que você fixou (C0 prevalece, como no Oleoplan), registrando os riscos acima."
    )
    w("")
    w("## 4. Item 08 — Caminhão Mercedinha MB 709 (1991)")
    w("")
    w(
        f"- **Valores:** grupo {money(g(r08, 'valor_grupo'))}; Rafael {money(g(r08, 'valor_rafael'))} no balanço executado (p. 19); na avaliação (p. 15) ele repete o valor do grupo."
    )
    w(
        f"- **C0:** está na lista de Rafael (vem para nós) → vale o menor: **grupo**. Efeito na cota: + {money(e08)}."
    )
    fipe_raf = f0.need(
        r"Aproximadamente R\$ (" + f0.MONEY + ")", p15, "FIPE Rafael"
    ).group(1)
    w(
        f'- **C1:** Rafael cita a Tabela FIPE "somente o chassi" em aproximadamente R$ {fipe_raf} (p. 15); o grupo cita FIPE de out/2025 em R$ {fipe_grupo} para um veículo similar (arq. 1, p. 9). Os {money(g(r08, "valor_grupo"))} dão valor zero à carroceria boiadeiro; o grupo não pode contestar o próprio número. As duas referências FIPE são posteriores à data-base (RDB-2).'
    )
    w(
        '- **Risco:** Rafael escreve que o caminhão com boiadeiro vale de R$ 85.000 a R$ 130.000 "para unidades reformadas" (p. 15); o estado do caminhão não está descrito.'
    )
    w(f"- **Recomendação:** grupo, {money(g(r08, 'valor_grupo'))}.")
    w("")
    w("## Perguntas ao grupo ligadas a este pacote (rascunho)")
    w("")
    w(
        "- Avaliação independente da fazenda: quem encomenda, quem paga, em que prazo, e se as partes aceitam o resultado em qualquer direção."
    )
    w(
        "- Certidão atualizada da matrícula 1.412 (a certidão entregue pelo grupo, arq. 11, é antiga; o inventário avalia 2.040 m² de prédio mais 510 m² de subsolo)."
    )
    w("")
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))
    print(
        f"bruto recomendado {gross_rec}; e01 {e01}; e08 {e08}; e400 farm {e400_farm} liab {e400_liab}; faixa item01 {item01_lo}-{item01_hi}; ratio {ratio}"
    )


if __name__ == "__main__":
    main()
