"""Phase 0 extraction for companies 02 (Agropecuaria Fazenda Sao Crispin) and
05 (Hotel Princesa dos Vales).

Reads every source by script, writes `fase0-extracoes-02-05.xlsx` in the
session folder and prints the checks. Every number comes from the documents;
declared totals and declared areas are compared with what the items add up to.

v2 (28/10/2026) applies the independent review in `_diario/revisao-fase0.md`
(A1-A2, M1-M7, B1-B6).

Run from anywhere: python fase0_extracao.py
"""

from __future__ import annotations

import difflib
import glob
import hashlib
import os
import re
from decimal import Decimal

import docx
import openpyxl
import pymupdf
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = r"C:\entrega-cisao-siviero"
OUT_DIR = os.path.join(ROOT, r"02_trabalho\empresas\02-05-sao-crispin-hotel-princesa")
OUT_XLSX = os.path.join(OUT_DIR, "fase0-extracoes-02-05.xlsx")
DRIVE_2018 = (
    r"G:\Meu Drive\02_Patrimonio_e_Sociedades\Cisao_Siviero"
    r"\Material Siviero\2018\inventário-2018"
)


def one(pattern: str) -> str:
    hits = glob.glob(pattern)
    if len(hits) != 1:
        raise SystemExit(f"expected one file for {pattern!r}, got {hits}")
    return hits[0]


SRC = {
    "inv2025_02": one(os.path.join(ROOT, r"01_fontes\inventario-2025\02. *.docx")),
    "inv2025_05": one(os.path.join(ROOT, r"01_fontes\inventario-2025\05. *.docx")),
    "inv2025_areas": one(
        os.path.join(ROOT, r"01_fontes\inventario-2025\HOTEL PRINCESA*.docx")
    ),
    "inv2025_00": one(os.path.join(ROOT, r"01_fontes\inventario-2025\00. *.docx")),
    "antigo_05": one(
        os.path.join(ROOT, r"01_fontes\inventario-2025-antigo\ANTIGO 05*.docx")
    ),
    "antigo_00": one(
        os.path.join(ROOT, r"01_fontes\inventario-2025-antigo\ANTIGO 00*.docx")
    ),
    "inv2018_13": one(os.path.join(DRIVE_2018, "13. 1. HOTEL*.docx")),
    "inv2018_17": one(os.path.join(DRIVE_2018, "17. 1. AGROPECU*.docx")),
    "inv2018_27": one(os.path.join(ROOT, r"01_fontes\inventario-2018\27. *.docx")),
    "inv2018_18x": one(os.path.join(ROOT, r"01_fontes\inventario-2018\18. *.xlsx")),
    "rafael_05": one(
        os.path.join(ROOT, r"01_fontes\estudo-rafael\HOTEL PRINCESA*.pdf")
    ),
    "rafael_02": one(os.path.join(ROOT, r"01_fontes\estudo-rafael\AGROPECUARIA*.pdf")),
    "diretrizes": one(os.path.join(ROOT, r"01_fontes\estudo-rafael\DIRETRIZES*.pdf")),
    "lista_rafael": os.path.join(
        ROOT, r"01_fontes\outros\ativos-avaliados-valter.xlsx"
    ),
    "planilha_2_0": one(os.path.join(ROOT, r"01_fontes\planilhas-google\2_0_*.xlsx")),
    "observacoes": os.path.join(ROOT, r"02_trabalho\observacoes-planilhas-google.xlsx"),
}

# Other 2025 inventories: searched only for where 2018 assets went (by matrícula).
OTHER_2025 = {
    f"inv2025_{os.path.basename(p)[:2]}": p
    for p in sorted(
        glob.glob(os.path.join(ROOT, r"01_fontes\inventario-2025\[0-9][0-9]. *.docx"))
    )
    if os.path.basename(p)[:2] not in ("00", "02", "05")
}
OLD_2025 = {
    f"antigo_{os.path.basename(p)[7:9]}": p
    for p in sorted(
        glob.glob(
            os.path.join(
                ROOT, r"01_fontes\inventario-2025-antigo\ANTIGO [0-9][0-9]. *.docx"
            )
        )
    )
    if os.path.basename(p)[7:9] not in ("00", "05")
}

# Sources read later (themes and packets); hashed now so the journal has them.
SRC_LATER = {
    "respostas_arq01": one(
        os.path.join(ROOT, r"01_fontes\respostas-docs-2026-08-13\1 - *.pdf")
    ),
    "respostas_arq07": one(
        os.path.join(ROOT, r"01_fontes\respostas-docs-2026-08-13\7 - *.pdf")
    ),
    "respostas_arq08": one(
        os.path.join(ROOT, r"01_fontes\respostas-docs-2026-08-13\8 - *.pdf")
    ),
    "respostas_arq09": one(
        os.path.join(ROOT, r"01_fontes\respostas-docs-2026-08-13\9 - *.pdf")
    ),
    "respostas_arq11": one(
        os.path.join(ROOT, r"01_fontes\respostas-docs-2026-08-13\11 - *.pdf")
    ),
    "conferencia_md": os.path.join(
        ROOT, r"01_fontes\outros\conferencia-discrepancias.md"
    ),
    "conferencia_xlsx": os.path.join(
        ROOT, r"01_fontes\outros\inventario-2025-avaliacoes-corrigido.xlsx"
    ),
    "diff_antigo": os.path.join(
        ROOT, r"02_trabalho\diff-inventario-2025-antigo-x-atual.md"
    ),
    "pendencias": os.path.join(ROOT, r"02_trabalho\pendencias-conferencia-anterior.md"),
    "parte6_sei": os.path.join(ROOT, r"02_trabalho\parte6-candidatos-duplicidade.md"),
    "achados_acertos": os.path.join(
        ROOT, r"02_trabalho\analises\acertos-e-areas-vendidas\achados-para-empresas.md"
    ),
    "log_acertos": os.path.join(
        ROOT, r"02_trabalho\analises\acertos-e-areas-vendidas\log-decisoes.md"
    ),
}

ROLE = {
    "inv2025_02": "Inventário 2025 atual, empresa 02",
    "inv2025_05": "Inventário 2025 atual, empresa 05",
    "inv2025_areas": "Docx do Hotel: áreas vendidas, soja, arrendamento, áreas contratadas da SEI",
    "inv2025_00": "Doc 00 2025: capital por empresa",
    "antigo_05": "Inventário 2025 antigo, empresa 05",
    "antigo_00": "Doc 00 2025 antigo",
    "inv2018_13": "Inventário 2018 do Hotel (Drive)",
    "inv2018_17": "Inventário 2018 da Agropecuária (Drive)",
    "inv2018_27": "Doc 27 2018: resumo do capital",
    "inv2018_18x": "Inventário 2018 da SEI (xlsx), aba Bens em outras empresas",
    "rafael_05": "Estudo 2 de Rafael, Hotel (24 p.)",
    "rafael_02": "Estudo 2 de Rafael, Agropecuária (4 p.)",
    "diretrizes": "PDF Diretrizes (lista de Rafael)",
    "lista_rafael": "Lista de Rafael transcrita do Diretrizes",
    "planilha_2_0": "Planilha Google 2_0 exportada",
    "observacoes": "Observações extraídas pela sessão SEI",
    "respostas_arq01": "Respostas ao Dr. Pablo, arquivo 1 (tem texto)",
    "respostas_arq07": "Arrendamento Lorenzet 2017-2020 (escaneado)",
    "respostas_arq08": "Arrendamento Lorenzet 2020-2023 (escaneado)",
    "respostas_arq09": "Arrendamento Lorenzet 2023-2026 (escaneado)",
    "respostas_arq11": "Matrícula 1.412 (escaneado)",
    "conferencia_md": "Conferência de 26/10 (não cobre 02 e 05)",
    "conferencia_xlsx": "Planilha da conferência de 26/10",
    "diff_antigo": "Diff 2025 antigo x atual (sessão SEI)",
    "pendencias": "Pendências da conferência anterior",
    "parte6_sei": "Candidatos à Parte 6 da SEI (só leitura)",
    "achados_acertos": "Achados da análise de ACERTOS (E1-E3 aprovados; E4-E7 não)",
    "log_acertos": "Log de decisões da análise de ACERTOS (registro da aprovação)",
}
ROLE.update(
    {
        k: f"Inventário 2025, empresa {k[-2:]} (busca de destino por matrícula)"
        for k in OTHER_2025
    }
)
ROLE.update(
    {
        k: f"Inventário 2025 antigo, empresa {k[-2:]} (busca de destino por matrícula)"
        for k in OLD_2025
    }
)

# Luís's notes in the Resumo tab of the 2_0 sheet about companies 02/05
# (listed cell by cell; labels and numbers of his tables are not notes)
RESUMO_OBS = {
    "B20",  # "05. Hotel Princesa dos Vales" in the list "VALUATIONS NECESSÁRIAS"
    "B80", "E80", "B87", "B88", "E88", "F88", "F91", "B92", "E92", "C94",
    "B96", "E96", "B100", "B191",
}

MONEY = r"\d{1,3}(?:\.\d{3})*,\d{2,3}"


def md5(path: str) -> str:
    h = hashlib.md5()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def brl(text: str) -> Decimal:
    return Decimal(text.replace(".", "").replace(",", "."))


def num(text: str) -> Decimal:
    """Brazilian number with optional thousands dots, e.g. 1.800,00 or 64,53."""
    return Decimal(text.replace(".", "").replace(",", "."))


def money_values(text: str) -> list[Decimal]:
    return [brl(m) for m in re.findall(r"R\$\s*(" + MONEY + ")", text)]


def need(pattern: str, text: str, what: str, flags: int = 0) -> re.Match:
    m = re.search(pattern, text, flags)
    if not m:
        raise SystemExit(f"pattern not found ({what}): {pattern}")
    return m


# ---------------------------------------------------------------- docx


def docx_tables(path: str) -> list[list[list[str]]]:
    doc = docx.Document(path)
    tables = []
    for child in doc.element.body.iterchildren():
        if child.tag.rsplit("}", 1)[-1] != "tbl":
            continue
        rows = []
        for row in Table(child, doc).rows:
            cells, prev = [], None
            for cell in row.cells:
                if prev is not None and cell._tc is prev:
                    continue
                prev = cell._tc
                cells.append(" ".join(cell.text.split()))
            rows.append(cells)
        tables.append(rows)
    return tables


def docx_paragraphs(path: str) -> list[str]:
    doc = docx.Document(path)
    out = []
    for child in doc.element.body.iterchildren():
        if child.tag.rsplit("}", 1)[-1] == "p":
            t = " ".join(Paragraph(child, doc).text.split())
            if t:
                out.append(t)
    return out


def matriculas(desc: str) -> list[str]:
    found = []
    for m in re.finditer(r"(?:Matr[íi]culas?|Registro)\s*n[ºo°]\s*\.?\s*:?\s*", desc):
        tail = desc[m.end() : m.end() + 60]
        for tok in re.finditer(r"(\d{1,3}(?:\.\d{3})+|\d+)|(\s*(?:,|e)\s*)|(.)", tail):
            if tok.group(1):
                found.append(tok.group(1))
            elif tok.group(2):
                continue
            else:
                break
    uniq = []
    for f in found:
        if f not in uniq:
            uniq.append(f)
    return uniq


def area_text(desc: str) -> str:
    plus = re.search(r"área (?:total )?de ([\d,]+) hectares \+ ([\d,]+) hectares", desc)
    if plus:  # e.g. lote 12 in 2018: "13,0 hectares + 7,5 hectares"
        total = num(plus.group(1)) + num(plus.group(2))
        return f"{str(total).replace('.', ',')} ha ({plus.group(1)} + {plus.group(2)})"
    m = re.search(r"área (?:total )?de ([\d.,]+)\s*(hectares|m²|ha)", desc)
    if m:
        unit = "ha" if m.group(2).startswith("h") else "m²"
        return f"{m.group(1)} {unit}"
    return ""


def area_ha(text: str) -> Decimal | None:
    m = re.match(r"([\d.,]+) ha\b", text)
    return num(m.group(1)) if m else None


def seller(desc: str) -> str:
    m = re.search(
        r"adquirid[ao] d[oa]s?\s+(?:Senhora?|Senhor)?\s*(.+?)(?:,|\s+área|\s+considerando)",
        desc,
    )
    return m.group(1).strip() if m else ""


def buyer(value_cell: str) -> str:
    m = re.search(r"Fam[íi]lia\s+([A-ZÁ-Úa-zá-ú]+)", value_cell)
    return m.group(1).upper() if m else ""


def status(value_cell: str) -> str:
    out = []
    if re.search(r"V\s?E\s?N\s?D\s?I\s?D\s?A", value_cell):
        part = re.search(r"[Ss]omente ([\d,]+) ?ha", value_cell)
        out.append(f"VENDIDA PARCIAL ({part.group(1)} ha)" if part else "VENDIDA")
    if "NÃO DIVISÍVEL" in value_cell:
        out.append("NÃO DIVISÍVEL")
    return "; ".join(out)


def item_rows(path: str, company: str, version: str) -> tuple[list[dict], list[dict]]:
    """Return (items, totals) from the asset tables of an inventory docx.

    Items listed before a 'Valor total de todas as áreas' row are flagged as
    part of the farm block (2018 doc 17)."""
    items, totals = [], []
    seq = 0
    for t_idx, rows in enumerate(docx_tables(path), 1):
        if not rows or rows[0][:2] != ["Item", "Descrição do(s) bem(ns)"]:
            continue
        for r_idx, cells in enumerate(rows[1:], 2):
            ref = f"T{t_idx}R{r_idx}"
            if len(cells) == 2:  # total row: label | value
                vals = money_values(cells[1])
                totals.append(
                    {
                        "rotulo": cells[0].rstrip("."),
                        "valor": vals[-1] if vals else None,
                        "ref": ref,
                    }
                )
                continue
            item, desc, val = cells[0], cells[1], cells[2]
            vals = money_values(val)
            if desc.startswith("Valor total de todas as áreas"):
                totals.append(
                    {"rotulo": desc.rstrip("."), "valor": vals[-1], "ref": ref}
                )
                for x in items:
                    x["bloco_fazenda"] = "sim"
                continue
            seq += 1
            evaluated = re.search(r"([\d.,]+) ha\s+R\$", val)
            items.append(
                {
                    "empresa": company,
                    "versao": version,
                    "seq": seq,
                    "item": item,
                    "descricao": desc,
                    "matriculas": ", ".join(matriculas(desc)),
                    "area": area_text(desc),
                    "area_avaliada": f"{evaluated.group(1)} ha" if evaluated else "",
                    "adquirido_de": seller(desc),
                    "situacao": status(val),
                    "comprador": buyer(val),
                    "texto_valor": val,
                    "valor": vals[-1] if vals else None,
                    "ref": ref,
                    "bloco_fazenda": "",
                }
            )
    return items, totals


def sale_rows(path: str) -> tuple[list[dict], list[dict]]:
    """Sales table (2018 doc 17 and the Hotel 'areas vendidas' docx)."""
    sales, totals = [], []
    for t_idx, rows in enumerate(docx_tables(path), 1):
        flat = " ".join(" ".join(r) for r in rows)
        if (
            "TOTAL RECEBIDO DAS ÁREAS VENDIDAS" not in flat
            or "Vendido em 2018" not in flat
        ):
            continue
        for r_idx, cells in enumerate(rows, 1):
            ref = f"T{t_idx}R{r_idx}"
            vals = money_values(cells[-1])
            if cells[0].startswith("TOTAL"):
                totals.append(
                    {"rotulo": cells[0].rstrip("."), "valor": vals[-1], "ref": ref}
                )
                continue
            text = cells[0]
            ha = re.search(
                r"área (?:de terras de |total de )?([\d.,]+)\s*hectares", text
            )
            fam = re.search(r"para a FAMILIA\s+(\w+)|para o Senhor\s+(\w+)", text)
            sales.append(
                {
                    "comprador": (fam.group(1) or fam.group(2)) if fam else "",
                    "area_ha": num(ha.group(1)) if ha else None,
                    "valor": vals[-1],
                    "texto": text,
                    "ref": ref,
                    "falta_receber": ", ".join(
                        re.findall(r"MAIO DE (\d{4})\.\s*\(FALTA RECEBER\)", text)
                    ),
                }
            )
    return sales, totals


def declared_areas(path: str) -> dict[str, Decimal]:
    """The '686,75 ha – ÁREA TOTAL … 411,43 ha – TOTAL' paragraphs."""
    out = {}
    for p in docx_paragraphs(path):
        m = re.match(r"(\(-\))?\s*([\d.,]+) ha\s*[–-]\s*(.+)$", p)
        if not m:
            continue
        label = m.group(3).upper()
        key = (
            "total"
            if "ÁREA TOTAL" in label
            else "rigo"
            if "RIGO" in label
            else "pellegrini"
            if "PELLEGRINI" in label
            else "remanescente"
            if label.startswith("TOTAL DE TERRAS")
            else label
        )
        out[key] = num(m.group(2))
    return out


def lease_2025(path: str) -> tuple[Decimal, list[tuple[str, Decimal]]]:
    paras = docx_paragraphs(path)
    text = " ".join(paras)
    total = num(need(r"área total de ([\d,]+) ha \(", text, "411,4 ha").group(1))
    parts = []
    for p in paras:
        m = re.match(r"([a-d])\.\) ([\d,]+) hectares", p)
        if m:
            parts.append((m.group(1), num(m.group(2))))
    return total, parts


def soy_credits(path: str) -> list[dict]:
    """Soy credits in the OBSERVAÇÕES paragraphs of the 'areas vendidas' docx."""
    rows, block = [], None
    paras: list[str] = []
    for p in docx_paragraphs(path):
        # "...2.606 SACAS PARA" continues in the next paragraph
        if paras and paras[-1].endswith(" PARA"):
            paras[-1] = f"{paras[-1]} {p}"
        else:
            paras.append(p)
    for p in paras:
        if "FAMÍLIA RIGO" in p:
            block = "RIGO"
        elif "FAMÍLIA LORENZET" in p:
            block = "LORENZET"
        elif p.startswith("CONTRATO PARTICULAR"):
            block = None
        if block is None:
            continue
        if m := re.match(r"([\d.]+) SACAS DE SOJA em MAIO DE (\d{4})\.(.*)$", p):
            sold = re.search(
                r"VENDIDO EM (\d{4}) A QUANTIA DE ([\d.]+) SACAS", m.group(3)
            )
            rows.append(
                {
                    "origem": block,
                    "tipo": "safra",
                    "ano": m.group(2),
                    "sacas": num(m.group(1)),
                    "empresa": "",
                    "texto": p,
                }
            )
            if sold:
                rows.append(
                    {
                        "origem": block,
                        "tipo": "vendida",
                        "ano": sold.group(1),
                        "sacas": -num(sold.group(2)),
                        "empresa": "",
                        "texto": p,
                    }
                )
        elif m := re.match(
            r"([\d.]+) SACAS DE SOJA em\s+REFERENTE AO ANO DE (\d{4})\.(\S+)", p
        ):
            rows.append(
                {
                    "origem": block,
                    "tipo": "safra",
                    "ano": m.group(2),
                    "sacas": num(m.group(1)),
                    "empresa": m.group(3),
                    "texto": p,
                }
            )
        elif m := re.match(r"([\d.]+) sacas de soja X R\$", p):
            rows.append(
                {
                    "origem": block,
                    "tipo": "crédito declarado",
                    "ano": "",
                    "sacas": num(m.group(1)),
                    "empresa": "",
                    "texto": p,
                }
            )
        elif m := re.match(r"(-\s*)?([\d.]+) sacas de soja$", p):
            kind = "dedução declarada" if m.group(1) else "subtotal declarado"
            rows.append(
                {
                    "origem": block,
                    "tipo": kind,
                    "ano": "",
                    "sacas": num(m.group(2)),
                    "empresa": "",
                    "texto": p,
                }
            )
    return rows


def capital_rows(path: str, label: str) -> list[dict]:
    out = []
    for t_idx, rows in enumerate(docx_tables(path), 1):
        for r_idx, cells in enumerate(rows, 1):
            line = " | ".join(cells)
            up = line.upper()
            if not any(n in up for n in ("CRISPIN", "HOTEL PRINCESA")) and not cells[
                0
            ].upper().startswith(("TOTAL", "PARCIAL")):
                continue
            vals = money_values(line)
            kind = (
                "cota por sócio"
                if "COTA DO CAPITAL" in up
                else "total do grupo"
                if cells[0].upper().startswith(("TOTAL", "PARCIAL"))
                else "total da empresa"
            )
            out.append(
                {
                    "fonte": label,
                    "tipo": kind,
                    "texto": line[:200],
                    # cota rows print the per-partner value first, then the company total
                    "valor": (vals[0] if kind == "cota por sócio" else vals[-1])
                    if vals
                    else None,
                    "ref": f"T{t_idx}R{r_idx}",
                }
            )
    return out


# ------------------------------------------------------------- Rafael PDF

SECTION_HEADS = [
    ("Balanço apresentado pelo grupo", "BALANÇO PATRIMONIAL APRESENTADO PELO GRUPO"),
    ("Avaliação executada por Valter", "AVALIAÇÃO PATRIMONIAL EXECUTADA POR VALTER"),
    ("Balanço executado por Valter", "BALANÇO PATRIMONIAL EXECUTADO POR VALTER"),
    ("Quadro-resumo", "QUADRO-RESUMO"),
    ("Bens escolhidos", "BENS MÓVEIS E IMÓVEIS ESCOLHIDOS PELA VALTER"),
]

# Last label whose key is found in the context wins (order matters).
QUADRO_LABELS = [
    ("100% do patrimônio", "PATRIMÔNIO TOTAL"),
    ("Participação atribuída", "1/6 ATRIBUÍDO A VALTER"),
    ("escolhidos", "SUBTOTAL ESCOLHIDOS"),
    ("pelos bens imóveis desta empresa", "DIFERENÇA A COMPLEMENTAR"),
]

# A sentence may contain dots between digits (1.800 m², R$ 9,2 milhões).
SENT = r"(?:[^.]|\.(?=\d))*"
RANGE_PATTERNS = [
    SENT
    + r"\bentre R\$ ?[\d.,]+(?: milhões)? e R\$ ?[\d.,]+(?: milhões)?"
    + SENT
    + r"\.",
    SENT + r"R\$ ?[\d.,]+ a R\$ ?[\d.,]+" + SENT + r"\.",
    SENT + r"\bna faixa de R\$ ?[\d.,]+" + SENT + r"\.",
    r"Aproximadamente R\$ ?[\d.,]+",
]


def rafael_sections(path: str) -> tuple[list[dict], list[dict], list[dict], dict]:
    doc = pymupdf.open(path)
    lines = []  # (page, text)
    for p_idx, page in enumerate(doc, 1):
        for ln in page.get_text().splitlines():
            lines.append((p_idx, ln.rstrip()))
    values, ranges, notes = [], [], []
    blocks: dict[tuple[str, int | None], list[str]] = {}
    section, marker, pending = None, None, None
    full_by_section: dict[str, list[tuple[int, int | None, str]]] = {}
    i = 0
    while i < len(lines):
        page, ln = lines[i]
        for name, head in SECTION_HEADS:
            heading = head in ln.upper().replace("  ", " ") or (
                head.split()[0] in ln.upper()
                and head in " ".join(x[1] for x in lines[i : i + 3]).upper()
            )
            if heading and section != name:
                section, marker = name, None
        s = ln.strip()
        if re.fullmatch(r"0?([1-9]|1[0-9])\.?", s):
            marker = int(s.rstrip("."))
        if "TOTAL GERAL" in s.upper() or "(04.,05.,06. e 08)" in s:
            pending = "TOTAL" if "TOTAL" in s.upper() else "SUBTOTAL ESCOLHIDOS"
        if s.upper().startswith("DA CONCORDÂNCIA") and section:
            follow = " ".join(x[1].strip() for x in lines[i : i + 3])
            cited = re.search(r"\(([\d., e]+)\)", follow)
            notes.append(
                {
                    "item": f"{marker:02d}" if marker else "?",
                    "pagina": page,
                    "tipo": "concordância expressa",
                    "texto": " ".join(follow.split())[:200],
                    "itens_citados": cited.group(1) if cited else "",
                }
            )
        if "AVALIAÇÃO TÉCNICA DA ÁREA DE TERRAS" in s.upper():
            notes.append(
                {
                    "item": "05",
                    "pagina": page,
                    "tipo": "recomenda avaliação independente (p. 13–14)",
                    "texto": "Valores atuais = 'parâmetros preliminares de referência', sujeitos a laudo técnico.",
                    "itens_citados": "",
                }
            )
        if "DA INCORPORAÇÃO E DA REGULARIDADE" in s.upper():
            notes.append(
                {
                    "item": "—",
                    "pagina": page,
                    "tipo": "incorporação da Agropecuária 'pendente de validação formal'",
                    "texto": s,
                    "itens_citados": "",
                }
            )
        val = None
        m = re.fullmatch(r"R\$\s*(" + MONEY + r")", s)
        if m:
            val = brl(m.group(1))
        elif (
            s == "R$"
            and i + 1 < len(lines)
            and re.fullmatch(MONEY, lines[i + 1][1].strip())
        ):
            val = brl(lines[i + 1][1].strip())
            i += 1
        if val is not None and section:
            label = pending or (f"{marker:02d}" if marker else "?")
            ctx = ""
            if section == "Quadro-resumo":
                back = [
                    x[1].strip()
                    for x in lines[max(0, i - 12) : i]
                    if re.search(r"[A-Za-zÀ-ú]{4}", x[1])
                ]
                ctx = " ".join(back[-6:])
                for key, name in QUADRO_LABELS:
                    if key in ctx:
                        label = name
            values.append(
                {
                    "secao": section,
                    "item": label,
                    "valor": val,
                    "pagina": page,
                    "contexto": ctx,
                }
            )
            pending = None
        if section:
            full_by_section.setdefault(section, []).append((page, marker, s))
            blocks.setdefault((section, marker), []).append(s)
        i += 1
    for section, rows in full_by_section.items():
        joined: dict[tuple[int, int | None], str] = {}
        for page, marker, s in rows:
            joined[(page, marker)] = joined.get((page, marker), "") + " " + s
        for (page, marker), text in joined.items():
            text = " ".join(text.split())
            found = []
            for pat in RANGE_PATTERNS:
                for m in re.finditer(pat, text):
                    found.append(re.sub(r"^R\$\s+", "", m.group(0).strip()))
            for f in dict.fromkeys(found):
                if any(f != g and f in g for g in found):
                    continue  # contained in a longer sentence
                ranges.append(
                    {
                        "secao": section,
                        "item": f"{marker:02d}" if marker else "?",
                        "pagina": page,
                        "texto": f,
                    }
                )
    block_text = {k: " ".join(v) for k, v in blocks.items()}
    return values, ranges, notes, block_text


# ---------------------------------------------------------------- checks

checks: list[dict] = []
avisos: list[dict] = []


def check(
    name: str, extracted, declared, source: str, tol=Decimal("0.01"), nota: str = ""
) -> bool:
    ok = (
        extracted is not None
        and declared is not None
        and abs(Decimal(extracted) - Decimal(declared)) <= tol
    )
    checks.append(
        {
            "conferencia": name,
            "extraido": extracted,
            "declarado": declared,
            "ok": "OK" if ok else "DIVERGE",
            "fonte": source,
            "nota": nota,
        }
    )
    return ok


def aviso(assunto: str, texto: str, fonte: str) -> None:
    avisos.append({"assunto": assunto, "texto": texto, "fonte": fonte})


def br(v: Decimal, places: int = 2) -> str:
    """Brazilian number format for labels and notes: 1.234.567,89."""
    s = f"{Decimal(v):,.{places}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def farm_lines(desc: str) -> list[tuple[str, Decimal, Decimal, Decimal]]:
    out = []
    for m in re.finditer(
        r"Área ([\w/ ]+?)\s*[–-]\s*parte[^=]*=\s*[\d,]+%\s*=\s*(?:[\d,]+\s*h[aá]\s*\(-\)\s*[\d,]+ha\s*=\s*)?([\d,]+)\s*h[aá]\s*[–-]\s*R\$\s*("
        + MONEY
        + r")\s*=\s*R\$\s*("
        + MONEY
        + ")",
        desc,
    ):
        out.append(
            (m.group(1).strip(), num(m.group(2)), brl(m.group(3)), brl(m.group(4)))
        )
    return out


def component_checks(items25: list[dict], items18: list[dict]) -> None:
    """Internal arithmetic of the composite items (review B3)."""
    by = {x["item"]: x for x in items25}
    d = by["01"]["descricao"]
    land = brl(
        need(r"720,00m² x R\$ [\d.,]+ = R\$ (" + MONEY + ")", d, "01 terreno").group(1)
    )
    deps = [
        brl(x)
        for x in re.findall(
            r"depreciação de [\d,]+%\s*[.\s]*\(-\s*R\$ ?"
            + MONEY
            + r"\)\s*=\s*R\$ ("
            + MONEY
            + ")",
            d,
        )
    ]
    extra = need(
        r"= ([\d,]+)% \([^)]*\) da obra feita: R\$ ("
        + MONEY
        + r") \([^)]*\)\s*=\s*R\$ ("
        + MONEY
        + ")",
        d,
        "01 coisa feita",
    )
    pct, base, res = num(extra.group(1)), brl(extra.group(2)), brl(extra.group(3))
    check("2025 item 01: 4 componentes depreciados lidos", len(deps), 4, "doc 05 2025")
    check(
        "2025 item 01: base da 'coisa feita' = soma dos 4 depreciados",
        sum(deps),
        base,
        "doc 05 2025",
    )
    check(
        f"2025 item 01: {br(pct)}% × base = valor da 'coisa feita'",
        (base * pct / 100).quantize(Decimal("0.01")),
        res,
        "doc 05 2025",
    )
    check(
        "2025 item 01: terreno + 4 depreciados + coisa feita = valor do item",
        land + sum(deps) + res,
        by["01"]["valor"],
        "doc 05 2025",
    )

    d = by["03"]["descricao"]
    rate = brl(need(r"= R\$ (" + MONEY + ") ao metro", d, "03 R$/m²").group(1))
    m2 = num(need(r"área de ([\d.,]+)m²", d, "03 área").group(1))
    check(
        "2025 item 03: 750 m² × R$/m² = valor",
        m2 * rate,
        by["03"]["valor"],
        "doc 05 2025",
        nota="O documento arredonda: 750 × 3.333,34 = 2.500.005,00.",
    )
    d = by["02"]["descricao"]
    m = need(r"= ([\d.,]+)m² X R\$ (" + MONEY + ")", d, "02")
    check(
        "2025 item 02: m² × R$/m² = valor",
        num(m.group(1)) * brl(m.group(2)),
        by["02"]["valor"],
        "doc 05 2025",
    )

    for x, label in (
        (by["04"], "2025"),
        (
            next(
                i
                for i in items18
                if i["empresa"].startswith("05") and i["item"] == "04"
            ),
            "2018",
        ),
    ):
        d = x["descricao"]
        parts = {
            "terra": brl(
                need(r"ao hectare = R\$ (" + MONEY + ")", d, "04 terra").group(1)
            ),
            "pavilhão": brl(
                need(
                    r"891,00m² x R\$ [\d.,]+ = R\$ (" + MONEY + ")", d, "04 pavilhão"
                ).group(1)
            ),
            "casa": brl(
                need(r"área construída: R\$ (" + MONEY + ")", d, "04 casa").group(1)
            ),
        }
        q = re.search(r"Quiosque: R\$ (" + MONEY + ")", d)
        if q:
            parts["quiosque"] = brl(q.group(1))
        extra = need(
            r"= ([\d,]+)% \(([^)]*)\) da obra feita: R\$ ("
            + MONEY
            + r")(?: \([^)]*\)\s*=\s*R\$ ("
            + MONEY
            + "))?",
            d,
            "04 obra",
        )
        pct, words = num(extra.group(1)), extra.group(2)
        if extra.group(4):
            base, res = brl(extra.group(3)), brl(extra.group(4))
        else:  # 2018: the amount printed is the result
            res, base = (
                brl(extra.group(3)),
                parts["pavilhão"] + parts["casa"] + parts.get("quiosque", 0),
            )
        if extra.group(4):  # in 2018 the base is not printed: nothing to check
            check(
                f"{label} item 04: base da obra = pavilhão + casa + quiosque",
                parts["pavilhão"] + parts["casa"] + parts.get("quiosque", 0),
                base,
                f"doc {label}",
            )
        check(
            f"{label} item 04: {br(pct)}% × base = obra feita",
            (base * pct / 100).quantize(Decimal("0.01")),
            res,
            f"doc {label}",
        )
        check(
            f"{label} item 04: soma dos componentes = valor do item",
            sum(parts.values()) + res,
            x["valor"],
            f"doc {label}",
        )
        if "Vinte por cento" in words and pct != 20:
            aviso(
                "Sede Campestre",
                f"{label}: '{br(pct)}% ({words})' — número e extenso divergem; o cálculo usa {br(pct)}%.",
                f"doc {label}, item 04",
            )

    h01 = next(
        i for i in items18 if i["empresa"].startswith("05") and i["item"] == "01"
    )
    d = h01["descricao"]
    comps = {
        "4 andares": brl(
            need(
                r"2\.040,00m² x R\$ [\d.,]+ = R\$ (" + MONEY + ")", d, "18 01 andares"
            ).group(1)
        ),
        "subsolo": brl(
            need(
                r"510,00m² x R\$ [\d.,]+ = R\$ (" + MONEY + ")", d, "18 01 subsolo"
            ).group(1)
        ),
        "terreno": brl(
            need(
                r"720,00m² x R\$ [\d.,]+ = R\$ (" + MONEY + ")", d, "18 01 terreno"
            ).group(1)
        ),
        "reforma": brl(
            need(r"48 x R\$ [\d.,]+ = R\$ (" + MONEY + ")", d, "18 01 reforma").group(1)
        ),
        "móveis": brl(need(r"etc\. = R\$ (" + MONEY + ")", d, "18 01 móveis").group(1)),
    }
    m = need(
        r"= ([\d,]+)% \([^)]*\) da obra feita: R\$ (" + MONEY + ")",
        d,
        "18 01 coisa feita",
    )
    pct, res = num(m.group(1)), brl(m.group(2))
    base = comps["4 andares"] + comps["subsolo"] + comps["reforma"] + comps["móveis"]
    check(
        f"2018 item 01 Hotel: {br(pct)}% × (andares + subsolo + reforma + móveis) = coisa feita",
        (base * pct / 100).quantize(Decimal("0.01")),
        res,
        "doc 13 2018",
    )
    check(
        "2018 item 01 Hotel: soma dos componentes = valor do item",
        sum(comps.values()) + res,
        h01["valor"],
        "doc 13 2018",
    )

    for x in items25:
        if x["item"] in ("06", "07"):
            m = need(
                r"([\d,]+) (?:ha|hectares) (?:–|X) R\$ (" + MONEY + ")",
                x["descricao"],
                f"item {x['item']}",
            )
            check(
                f"2025 item {x['item']}: ha × R$/ha = valor",
                num(m.group(1)) * brl(m.group(2)),
                x["valor"],
                "doc 05 2025",
            )


def regularization_checks(
    path05: str, path02: str, total05: Decimal, items05: list[dict]
) -> None:
    t05 = " ".join(docx_paragraphs(path05))
    t02 = " ".join(docx_paragraphs(path02))
    adm = brl(
        need(
            r"totalizando aproximadamente um custo de R\$ (" + MONEY + ")", t05, "adm"
        ).group(1)
    )
    ha = num(
        need(
            r"FAZENDA SÃO CRISPIN é de ([\d,]+) ha ao custo médio de R\$ ("
            + MONEY
            + ")",
            t05,
            "ITBI base",
        ).group(1)
    )
    rate = brl(
        need(r"ao custo médio de R\$ (" + MONEY + ")", t05, "ITBI R$/ha").group(1)
    )
    base = brl(
        need(r"importância de R\$ (" + MONEY + ")", t05, "ITBI base R$").group(1)
    )
    pct = num(need(r"percentual do imposto de ([\d,]+)", t05, "ITBI %").group(1))
    itbi = brl(
        need(
            r"incidindo um total de aproximadamente R\$ (" + MONEY + ")", t05, "ITBI"
        ).group(1)
    )
    notarial = brl(
        need(r"Serviço Notarial[^R]*R\$ (" + MONEY + ")", t05, "notarial").group(1)
    )
    registral = brl(
        need(r"Serviço Registral[^R]*R\$ (" + MONEY + ")", t05, "registral").group(1)
    )
    total = brl(
        need(r"um TOTAL de R\$ (" + MONEY + ")", t05, "total regularização").group(1)
    )
    check(
        f"Regularização: {br(ha, 1)} ha × R$ {br(rate)} = base do ITBI",
        ha * rate,
        base,
        "doc 05 2025",
    )
    check(
        f"Regularização: base × {br(pct)}% = ITBI",
        base * pct / 100,
        itbi,
        "doc 05 2025",
    )
    check(
        f"Regularização: custos administrativos R$ {br(adm)} + ITBI + notarial + registral = TOTAL declarado",
        adm + itbi + notarial + registral,
        total,
        "doc 05 2025",
        nota=f"O documento arredonda as despesas para R$ {br(total)}.",
    )
    reduction = brl(
        need(r"estima numa redução de R\$ (" + MONEY + ")", t05, "redução").group(1)
    )
    check(
        "Regularização: redução estimada no valor da fazenda = TOTAL das despesas",
        reduction,
        total,
        "doc 05 2025",
    )

    def block(text: str) -> str:
        a, b = text.find("É necessário transmitir"), text.find("Veranópolis – RS, 31")
        if a < 0 or b < 0:
            raise SystemExit("regularization block anchors not found")
        return re.sub(r"\s+", " ", text[a:b])

    b05, b02 = block(t05), block(t02)
    diff = [
        (" ".join(b05.split()[i1:i2]), " ".join(b02.split()[j1:j2]))
        for op, i1, i2, j1, j2 in difflib.SequenceMatcher(
            a=b05.split(), b=b02.split(), autojunk=False
        ).get_opcodes()
        if op != "equal"
    ]
    check(
        "Regularização: diferenças de redação entre os docs 02 e 05",
        len(diff),
        0,
        "docs 02 e 05 2025",
        tol=Decimal(0),
        nota="; ".join(f"doc 05 '{a}' × doc 02 '{b}'" for a, b in diff),
    )
    aviso(
        "Regularização",
        f"A redução de R$ {br(reduction)} é descrita no texto do doc 05 mas não abatida do TOTAL GERAL "
        f"(R$ {br(total05)} = soma dos itens); o doc 02 repete o mesmo texto com total R$ 0,00.",
        "docs 02 e 05 2025",
    )


# ---------------------------------------------------------------- main


def main() -> None:
    fontes = []
    for key, path in (
        list(SRC.items())
        + list(OTHER_2025.items())
        + list(OLD_2025.items())
        + list(SRC_LATER.items())
    ):
        fontes.append(
            {"chave": key, "caminho": path, "md5": md5(path), "papel": ROLE[key]}
        )

    # ---- 2025
    i25_05, t25_05 = item_rows(SRC["inv2025_05"], "05 Hotel Princesa", "2025 atual")
    i25a_05, t25a_05 = item_rows(SRC["antigo_05"], "05 Hotel Princesa", "2025 antigo")
    i25_02, t25_02 = item_rows(SRC["inv2025_02"], "02 São Crispin", "2025 atual")
    for t_idx, rows in enumerate(docx_tables(SRC["inv2025_02"]), 1):
        for r_idx, cells in enumerate(rows, 1):
            if cells[0].startswith("VALOR INSCRITO"):
                t25_02.append(
                    {
                        "rotulo": cells[0],
                        "valor": money_values(cells[-1])[-1],
                        "ref": f"T{t_idx}R{r_idx}",
                    }
                )

    def total(ts):
        return next(
            t["valor"]
            for t in ts
            if t["rotulo"].upper().startswith(("TOTAL GERAL", "VALOR INSCRITO"))
        )

    check(
        "2025 atual 05: soma dos itens = TOTAL GERAL",
        sum(x["valor"] for x in i25_05),
        total(t25_05),
        "doc 05 2025",
    )
    check(
        "2025 antigo 05: soma dos itens = TOTAL GERAL",
        sum(x["valor"] for x in i25a_05),
        total(t25a_05),
        "ANTIGO 05",
    )
    check(
        "2025 atual 02: soma dos itens = valor inscrito",
        sum(x["valor"] or 0 for x in i25_02),
        total(t25_02),
        "doc 02 2025",
    )
    check(
        "2025 atual 02: linhas com bem descrito",
        sum(1 for x in i25_02 if x["descricao"] not in ("-.-", "")),
        0,
        "doc 02 2025",
    )
    nums05 = [int(x["item"]) for x in i25_05]
    check(
        "2025 atual 05: itens numerados 01 a N sem falha",
        int(nums05 == list(range(1, len(nums05) + 1))),
        1,
        "doc 05 2025",
        nota=f"N = {len(nums05)}",
    )

    old_by = {x["item"]: x for x in i25a_05}
    diff_rows = []
    for b in i25_05:
        a = old_by.get(b["item"])
        diff_rows.append(
            {
                "item": b["item"],
                "antigo": a["valor"] if a else None,
                "atual": b["valor"],
                "diferenca": b["valor"] - (a["valor"] if a else 0),
                "texto_mudou": "item novo"
                if not a
                else ("sim" if a["descricao"] != b["descricao"] else "não"),
            }
        )
    check(
        "Antigo → atual: itens do antigo todos presentes no atual",
        len(set(old_by) - {x["item"] for x in i25_05}),
        0,
        "docs 05",
    )
    check(
        "Antigo → atual: diferença total",
        sum(r["diferenca"] for r in diff_rows),
        total(t25_05) - total(t25a_05),
        "docs 05",
    )

    # ---- capital summaries
    resumo = (
        capital_rows(SRC["inv2025_00"], "doc 00 2025")
        + capital_rows(SRC["antigo_00"], "doc 00 2025 antigo")
        + capital_rows(SRC["inv2018_27"], "doc 27 2018")
    )

    def find(src, needle, kind="total da empresa"):
        return next(
            r["valor"]
            for r in resumo
            if r["fonte"] == src and r["tipo"] == kind and needle in r["texto"].upper()
        )

    check(
        "Doc 00 2025: Hotel = total do doc 05",
        find("doc 00 2025", "05. HOTEL"),
        total(t25_05),
        "doc 00",
    )
    check(
        "Doc 00 2025: Agropecuária = 0",
        find("doc 00 2025", "02. AGROPECU"),
        0,
        "doc 00",
    )
    check(
        "Doc 00 antigo: Hotel = total do ANTIGO 05",
        find("doc 00 2025 antigo", "05. HOTEL"),
        total(t25a_05),
        "ANTIGO 00",
    )
    cota = find("doc 00 2025", "05. HOTEL", "cota por sócio")
    check(
        "Doc 00 2025: 6 × cota por sócio do Hotel = total",
        6 * cota,
        total(t25_05),
        "doc 00",
        tol=Decimal("0.05"),
        nota="Rótulo '16,66%'; o valor é 1/6 arredondado.",
    )

    # ---- 2018
    i18_13, t18_13 = item_rows(SRC["inv2018_13"], "05 Hotel Princesa", "2018")
    i18_17, t18_17 = item_rows(SRC["inv2018_17"], "02 São Crispin", "2018")
    v18, vt18 = sale_rows(SRC["inv2018_17"])
    totals18 = []
    seen = {}
    for t_idx, rows in enumerate(docx_tables(SRC["inv2018_17"]), 1):
        for r_idx, cells in enumerate(rows, 1):
            if (
                cells[0].startswith(("TOTAL GERAL", "TOTAL RECEBIDO", "TOTAL LÍQUIDO"))
                or cells[1:2]
                and cells[1].startswith("Valor total de todas")
            ):
                label = (cells[1] if cells[0] == "" else cells[0]).rstrip(".")
                vals = money_values(cells[-1])
                key = (label, vals[-1] if vals else None)
                if key in seen:
                    seen[key]["ref"] += f", T{t_idx}R{r_idx}"
                    continue
                seen[key] = {
                    "empresa": "02 São Crispin",
                    "rotulo": label,
                    "valor": key[1],
                    "ref": f"T{t_idx}R{r_idx}",
                }
                totals18.append(seen[key])
    totals18 += [dict(t, empresa="05 Hotel Princesa") for t in t18_13]
    check(
        "2018 Hotel: soma dos itens = TOTAL GERAL",
        sum(x["valor"] for x in i18_13),
        total(t18_13),
        "doc 13 2018",
    )
    check(
        "Doc 27 2018: Hotel = total do doc 13",
        find("doc 27 2018", "HOTEL PRINCESA"),
        total(t18_13),
        "doc 27",
    )
    fazenda = [x for x in i18_17 if x["bloco_fazenda"]]
    sub = next(
        t["valor"] for t in t18_17 if t["rotulo"].startswith("Valor total de todas")
    )
    check(
        f"2018 Agro: subtotal Fazenda São Crispin ({len(fazenda)} linhas antes do subtotal)",
        sum(x["valor"] for x in fazenda),
        sub,
        "doc 17 2018",
    )
    bruto = sum(x["valor"] or 0 for x in i18_17)
    check(
        "2018 Agro: soma dos itens = TOTAL GERAL", bruto, total(t18_17), "doc 17 2018"
    )
    recebido = next(t["valor"] for t in vt18 if "RECEBIDO" in t["rotulo"])
    check(
        "2018 Agro: vendas somam TOTAL RECEBIDO",
        sum(s["valor"] for s in v18),
        recebido,
        "doc 17 2018",
    )
    liquido = next(
        t["valor"] for t in totals18 if t["rotulo"].startswith("TOTAL LÍQUIDO")
    )
    check(
        "2018 Agro: TOTAL GERAL − RECEBIDO = TOTAL LÍQUIDO",
        total(t18_17) - recebido,
        liquido,
        "doc 17 2018",
    )
    check(
        "Doc 27 2018: Agropecuária = líquido do doc 17",
        find("doc 27 2018", "CRISPIN"),
        liquido,
        "doc 27",
    )

    # ---- sales docx (2025) x doc 17 (2018)
    va, vat = sale_rows(SRC["inv2025_areas"])
    check(
        "Docx áreas vendidas: vendas somam TOTAL RECEBIDO",
        sum(s["valor"] for s in va),
        vat[0]["valor"],
        "docx Hotel 2025",
    )
    check(
        "Docx áreas vendidas = vendas de 2018 (doc 17)",
        sum(s["valor"] for s in va),
        recebido,
        "docx x doc 17",
    )
    sales_diff = []
    va_by = {s["comprador"]: s for s in va}
    check(
        "Docx áreas × doc 17: mesmos compradores",
        int(sorted(va_by) == sorted(s["comprador"] for s in v18)),
        1,
        "docx x doc 17",
    )
    for a in v18:
        b = va_by[a["comprador"]]
        wa, wb_ = a["texto"].split(), b["texto"].split()
        for op, i1, i2, j1, j2 in difflib.SequenceMatcher(
            a=wa, b=wb_, autojunk=False
        ).get_opcodes():
            if op != "equal":
                sales_diff.append(
                    {
                        "comprador": a["comprador"],
                        "operacao": {
                            "replace": "trocado",
                            "delete": "só no doc 17",
                            "insert": "só no docx 2025",
                        }[op],
                        "doc17": " ".join(wa[i1:i2]),
                        "docx2025": " ".join(wb_[j1:j2]),
                    }
                )
    check(
        "Docx áreas × doc 17: diferenças de texto nas 3 vendas",
        len(sales_diff),
        0,
        "docx x doc 17",
        tol=Decimal(0),
        nota="Ver aba 'Vendas docx x doc 17'.",
    )
    falta = [s for s in va if s["falta_receber"]]
    price = brl(
        need(
            r"considerado foi de R\$ (" + MONEY + ")", va[0]["texto"], "preço saca"
        ).group(1)
    )
    parcel = num(
        need(
            r"([\d.]+) SACAS DE SOJA em MAIO DE 2021", va[0]["texto"], "parcela"
        ).group(1)
    )
    n_falta = sum(len(s["falta_receber"].split(", ")) for s in falta)
    check(
        "Docx: parcelas Rigo marcadas FALTA RECEBER dentro do TOTAL RECEBIDO (R$, à cotação de out/2018)",
        n_falta * parcel * price,
        0,
        "docx Hotel 2025, T1R1",
        nota=(
            f"{n_falta} × {br(parcel, 0)} sacas × R$ {br(price)} (cotação de out/2018) dentro do 'TOTAL RECEBIDO' "
            f"de R$ {br(vat[0]['valor'])}. Leitura em aberto: a tabela diz 'FALTA RECEBER', mas as OBSERVAÇÕES dizem "
            "'ainda não foram vendidos' e registram venda parcial em 2024 — pode ser soja em depósito, não crédito "
            "contra os Rigo. Saldo declarado: ver aba 'Créditos de soja'."
        ),
    )

    # ---- soy credits (A1)
    soy = soy_credits(SRC["inv2025_areas"])
    rigo_safras = [r for r in soy if r["origem"] == "RIGO" and r["tipo"] == "safra"]
    rigo_sub = next(
        r for r in soy if r["origem"] == "RIGO" and r["tipo"] == "subtotal declarado"
    )
    rigo_ded = next(
        r for r in soy if r["origem"] == "RIGO" and r["tipo"] == "dedução declarada"
    )
    rigo_cred = next(
        r for r in soy if r["origem"] == "RIGO" and r["tipo"] == "crédito declarado"
    )
    rigo_sold = sum(
        -r["sacas"] for r in soy if r["origem"] == "RIGO" and r["tipo"] == "vendida"
    )
    check(
        "Soja Rigo: safras não vendidas somam o subtotal declarado",
        sum(r["sacas"] for r in rigo_safras),
        rigo_sub["sacas"],
        "docx Hotel 2025",
    )
    check(
        "Soja Rigo: vendida em 2024 = dedução declarada",
        rigo_sold,
        rigo_ded["sacas"],
        "docx Hotel 2025",
    )
    check(
        "Soja Rigo: subtotal − dedução = crédito declarado",
        rigo_sub["sacas"] - rigo_ded["sacas"],
        rigo_cred["sacas"],
        "docx Hotel 2025",
    )
    lor = [r for r in soy if r["origem"] == "LORENZET" and r["tipo"] == "safra"]
    lor_cred = next(
        r for r in soy if r["origem"] == "LORENZET" and r["tipo"] == "crédito declarado"
    )
    check(
        "Soja Lorenzet: safras somam o crédito declarado",
        sum(r["sacas"] for r in lor),
        lor_cred["sacas"],
        "docx Hotel 2025",
    )
    agro_soy = sum(r["sacas"] for r in lor if r["empresa"].startswith("AGROPECU"))
    check(
        f"Soja atribuída à AGROPECUÁRIA × inventário 2025 da 02 (R$ {br(total(t25_02))})",
        agro_soy,
        0,
        "docx Hotel 2025 x doc 02 2025",
        nota="Crédito citado em documento do grupo para empresa com inventário zerado.",
    )
    sold_row = next(r for r in soy if r["tipo"] == "vendida")
    if "ESCAVADEIRA" in sold_row["texto"].upper():
        exc = [
            (t_idx, r_idx, " ".join(cells))
            for t_idx, rows in enumerate(docx_tables(OTHER_2025["inv2025_11"]), 1)
            for r_idx, cells in enumerate(rows, 1)
            if re.search(r"Escavadeira Hidráulica, ano 2024", " ".join(cells))
        ]
        where_exc = (
            "; ".join(
                f"doc 11 T{t}R{r}: {txt[:120]} (R$ {br(money_values(txt)[-1])})"
                for t, r, txt in exc
            )
            or "nenhuma escavadeira 2024 achada no doc 11"
        )
        aviso(
            "Soja → escavadeira",
            f"{br(-sold_row['sacas'], 0)} sacas do crédito Rigo (Hotel) vendidas em {sold_row['ano']} 'para aquisição da "
            f"escavadeira hidráulica'. Candidata: {where_exc}. Hipótese de fluxo Hotel → SEI, a confirmar (P4).",
            "docx Hotel 2025 x doc 11 2025",
        )
    if rigo_sold in [r["sacas"] for r in lor]:
        aviso(
            "Soja",
            f"O número {rigo_sold} aparece como soja Rigo vendida em 2024 e como safra Lorenzet; coincidência ou erro de digitação.",
            "docx Hotel 2025",
        )

    # ---- farm areas (M2: declared numbers read from the documents)
    areas_rows = []
    farm25 = next(x for x in i25_05 if x["item"] == "05")
    pairs = [
        (m, num(h))
        for m, h in re.findall(
            r"Matr[íi]cula nº ([\d.]+)\s*=\s*([\d.,]+) hectares", farm25["descricao"]
        )
    ]
    for m, h in pairs:
        areas_rows.append(
            {
                "bloco": "2025 item 05: matrículas listadas",
                "chave": m,
                "ha": h,
                "valor": None,
                "ref": farm25["ref"],
            }
        )
    farm_total25 = num(
        need(r"área total de ([\d,]+) hectares", farm25["descricao"], "491").group(1)
    )
    sold25 = num(need(r"\(-\)\s*([\d,]+)\s*ha", farm25["descricao"], "83,47").group(1))
    eval25 = area_ha(farm25["area_avaliada"])
    check(
        "2025 item 05: soma das matrículas = área total declarada",
        sum(h for _, h in pairs),
        farm_total25,
        "doc 05 2025",
    )
    dups = sorted({m for m, _ in pairs if [p for p, _ in pairs].count(m) > 1})
    check(
        "2025 item 05: matrículas repetidas na lista",
        len(dups),
        0,
        "doc 05 2025",
        tol=Decimal(0),
        nota=f"Repetidas: {', '.join(dups)}",
    )
    fl25 = farm_lines(farm25["descricao"])
    for nome, ha, rate, val in fl25:
        areas_rows.append(
            {
                "bloco": "2025 item 05: áreas avaliadas",
                "chave": nome,
                "ha": ha,
                "valor": val,
                "ref": farm25["ref"],
            }
        )
        check(
            f"2025 item 05 {nome}: {br(ha)} ha × R$ {br(rate)} = valor",
            ha * rate,
            val,
            "doc 05 2025",
        )
    check(
        "2025 item 05: soma das áreas avaliadas = área avaliada declarada",
        sum(x[1] for x in fl25),
        eval25,
        "doc 05 2025",
    )
    check(
        "2025 item 05: área total − vendida = área avaliada",
        farm_total25 - sold25,
        eval25,
        "doc 05 2025",
    )
    check(
        "2025 item 05: soma dos valores = valor do item",
        sum(x[3] for x in fl25),
        farm25["valor"],
        "doc 05 2025",
    )
    jp = next(x for x in fl25 if x[0].startswith("Jane"))
    m00760 = sum(h for m, h in pairs if m == "00.760")
    check(
        "2025: matrícula 00.760 (48 + 100 ha) − vendida = área Jane/João Paulo",
        m00760 - sold25,
        jp[1],
        "doc 05 2025",
    )
    fa = next(x for x in i25a_05 if x["item"] == "05")
    check(
        "Antigo item 05: soma dos valores",
        sum(x[3] for x in farm_lines(fa["descricao"])),
        fa["valor"],
        "ANTIGO 05",
    )

    fl18 = farm_lines(fazenda[0]["descricao"])
    farm_total18 = num(
        need(
            r"área total de ([\d,]+) hectares", fazenda[0]["descricao"], "491 2018"
        ).group(1)
    )
    check(
        "2018 Fazenda 491 ha: soma das áreas avaliadas = área total (avaliava também a parte vendida)",
        sum(x[1] for x in fl18),
        farm_total18,
        "doc 17 2018",
    )
    for x in fazenda:
        lines = farm_lines(x["descricao"])
        for nome, ha, rate, val in lines:
            check(
                f"2018 parcela {nome}: {br(ha, 3)} ha × R$ {br(rate)} = valor",
                ha * rate,
                val,
                "doc 17 2018",
            )
        if lines:
            check(
                f"2018 parcela {x['seq']} ({x['area']}): soma das linhas = valor da parcela",
                sum(v for *_, v in lines),
                x["valor"],
                "doc 17 2018",
            )
        x["area_ha"] = area_ha(x["area"])
        one_line = re.findall(
            r"Matr[íi]cula nº[^=]*=\s*([\d,]+) hectares", x["descricao"]
        )
        if (
            len(one_line) == 1
            and x["area_ha"] is not None
            and num(one_line[0]) != x["area_ha"]
        ):
            aviso(
                "Área de parcela 2018",
                f"Seq. {x['seq']}: cabeçalho {x['area']}, linha da matrícula '= {one_line[0]} hectares'.",
                f"doc 17 2018, {x['ref']}",
            )
    decl18 = declared_areas(SRC["inv2018_17"])
    decl_docx = declared_areas(SRC["inv2025_areas"])
    check(
        "Doc 17: áreas declaradas idênticas no docx de 2025",
        int(decl18 == decl_docx),
        1,
        "doc 17 x docx Hotel 2025",
    )
    check(
        "Doc 17: total − Rigo − Pellegrini = remanescente (aritmética declarada)",
        decl18["total"] - decl18["rigo"] - decl18["pellegrini"],
        decl18["remanescente"],
        "doc 17 2018",
    )
    total_farm_ha = sum(x["area_ha"] for x in fazenda if x["area_ha"])
    check(
        "2018 Fazenda: soma das parcelas = área total declarada",
        total_farm_ha,
        decl18["total"],
        "doc 17 2018",
    )
    rigo = [x for x in fazenda if x["comprador"] == "RIGO"]
    pell = [x for x in fazenda if x["comprador"] == "PELLEGRINI"]
    sold18 = num(
        need(r"[Ss]omente ([\d,]+) ?ha", rigo[0]["texto_valor"], "83,47 2018").group(1)
    )
    rigo_ha = sold18 + sum(x["area_ha"] for x in rigo[1:])
    check(
        "2018: Rigo = parte vendida da mat. 00.760 + parcela mat. 11.090 = declarado",
        rigo_ha,
        decl18["rigo"],
        "doc 17 2018",
    )
    pell_ha = sum(x["area_ha"] for x in pell)
    check(
        "2018: parcelas Pellegrini somam a área declarada",
        pell_ha,
        decl18["pellegrini"],
        "doc 17 2018",
    )
    remaining = total_farm_ha - rigo_ha - pell_ha
    check(
        "Remanescente pelas parcelas × remanescente declarado",
        remaining,
        decl18["remanescente"],
        "doc 17 2018",
        nota="A diferença vem dos 85,32 ha declarados para a Pellegrini.",
    )
    items_06 = next(x for x in i25_05 if x["item"] == "06")
    check(
        "Remanescente pelas parcelas × itens 05 + 06 de 2025",
        remaining,
        eval25 + area_ha(items_06["area"]),
        "docs 17 2018 e 05 2025",
    )
    lease_total, lease_parts = lease_2025(SRC["inv2025_05"])
    check(
        "2025: arrendamento a)+b)+c)+d) = área declarada no texto",
        sum(h for _, h in lease_parts),
        lease_total,
        "doc 05 2025",
    )
    check(
        "2025: itens 05 + 06 × área declarada no texto",
        eval25 + area_ha(items_06["area"]),
        lease_total,
        "doc 05 2025",
        nota=f"O texto (ITBI e arrendamento) usa {br(lease_total, 1)} ha.",
    )
    for x in fazenda:
        areas_rows.append(
            {
                "bloco": "2018 Fazenda: parcelas",
                "chave": f"{x['matriculas'] or 's/ matrícula'} ({x['adquirido_de'] or '—'})",
                "ha": x["area_ha"],
                "valor": x["valor"],
                "ref": x["ref"],
                "situacao": x["situacao"],
                "comprador": x["comprador"],
            }
        )
    for k, v in decl18.items():
        areas_rows.append(
            {
                "bloco": "2018 doc 17: áreas declaradas",
                "chave": k,
                "ha": v,
                "valor": None,
                "ref": "parágrafos após T5",
            }
        )
    mat_count: dict[str, list[str]] = {}
    for x in i18_17:
        for m in x["matriculas"].split(", "):
            if m:
                mat_count.setdefault(m, []).append(
                    f"seq. {x['seq']} ({x['area']}, {x['situacao'] or 'não vendida'})"
                )
    for m, where in mat_count.items():
        if len(where) > 1:
            aviso(
                "Matrícula em mais de uma parcela",
                f"Mat. {m}: {'; '.join(where)}. Em 2025 a mat. {m} é o item 06 do Hotel."
                if m == "11.090"
                else f"Mat. {m}: {'; '.join(where)}.",
                "doc 17 2018",
            )

    component_checks(i25_05, i18_13)
    regularization_checks(SRC["inv2025_05"], SRC["inv2025_02"], total(t25_05), i25_05)

    # ---- Rafael
    rv, rr, rnotes, rblocks = rafael_sections(SRC["rafael_05"])

    def by(sec):
        return [v for v in rv if v["secao"] == sec]

    grp = {v["item"]: v["valor"] for v in by("Balanço apresentado pelo grupo")}
    exe: dict[str, Decimal] = {}
    for v in by("Balanço executado por Valter"):
        if v["item"] in exe:
            v["observacao"] = (
                "linha repetida no PDF (mesmo item, mesmo valor)"
                if exe[v["item"]] == v["valor"]
                else "REPETIDA COM OUTRO VALOR"
            )
        else:
            exe[v["item"]] = v["valor"]
    ava = {v["item"]: v["valor"] for v in by("Avaliação executada por Valter")}
    esc = {v["item"]: v["valor"] for v in by("Bens escolhidos")}
    q = {v["item"]: v["valor"] for v in by("Quadro-resumo")}
    items05 = [x["item"] for x in i25_05]
    check(
        "Rafael, balanço do grupo: soma dos itens = TOTAL",
        sum(grp.get(k, 0) for k in items05),
        grp.get("TOTAL"),
        "PDF Hotel p. 1–4",
    )
    check(
        "Rafael, balanço do grupo = doc 05 2025 (item a item)",
        sum(abs(grp.get(x["item"], 0) - x["valor"]) for x in i25_05),
        0,
        "PDF Hotel p. 1–4",
    )
    check(
        "Rafael, balanço executado: soma dos itens = TOTAL",
        sum(exe.get(k, 0) for k in items05),
        exe.get("TOTAL"),
        "PDF Hotel p. 16–19",
    )
    check(
        "Rafael, avaliação executada: itens lidos = itens do doc 05",
        len([k for k in items05 if k in ava]),
        len(items05),
        "PDF Hotel p. 7–15",
    )
    check(
        "Rafael, bens escolhidos: soma = subtotal",
        sum(v for k, v in esc.items() if k != "SUBTOTAL ESCOLHIDOS"),
        esc.get("SUBTOTAL ESCOLHIDOS"),
        "PDF Hotel p. 22–23",
    )
    check(
        "Quadro-resumo: patrimônio total = TOTAL executado",
        q.get("PATRIMÔNIO TOTAL"),
        exe.get("TOTAL"),
        "PDF Hotel p. 19–20",
    )
    check(
        "Quadro-resumo: 1/6 do total",
        (q["PATRIMÔNIO TOTAL"] / 6).quantize(Decimal("0.001")),
        q.get("1/6 ATRIBUÍDO A VALTER"),
        "PDF Hotel p. 20",
        tol=Decimal("0.001"),
    )
    check(
        "Quadro-resumo: escolhidos = subtotal p. 23",
        q.get("SUBTOTAL ESCOLHIDOS"),
        esc.get("SUBTOTAL ESCOLHIDOS"),
        "PDF Hotel p. 20",
    )
    check(
        "Quadro-resumo: escolhidos − 1/6 = diferença a complementar",
        q["SUBTOTAL ESCOLHIDOS"] - q["1/6 ATRIBUÍDO A VALTER"],
        q.get("DIFERENÇA A COMPLEMENTAR"),
        "PDF Hotel p. 20",
        tol=Decimal("0.10"),
        nota="Rafael arredonda para R$ 37.665.498,00.",
    )
    # item 01 in the executed balance: are the printed components the group's?
    g01 = next(x for x in i25_05 if x["item"] == "01")["descricao"]
    comp_amounts = re.findall(r"=\s*R\$\s*(" + MONEY + r")(?![\d])", g01)
    exe01 = rblocks.get(("Balanço executado por Valter", 1), "")
    missing = [a for a in comp_amounts if a not in exe01]
    check(
        "Rafael executado item 01: componentes impressos = os do grupo",
        len(missing),
        0,
        "PDF Hotel p. 16–17",
        tol=Decimal(0),
        nota=f"{len(comp_amounts)} valores do grupo procurados; o total impresso (R$ {br(exe['01'])}) não tem memória de cálculo.",
    )
    exe01_is_group = not missing
    rv2, _, rnotes2, _ = rafael_sections(SRC["rafael_02"])
    check(
        "Rafael Agropecuária: valores de item/total no PDF",
        len(rv2),
        0,
        "PDF Agropecuária (4 p.)",
    )

    # ---- Rafael list and Diretrizes
    wb = openpyxl.load_workbook(SRC["lista_rafael"], data_only=True)
    lista = []
    for row in wb["Ativos"].iter_rows(min_row=1, values_only=True):
        if isinstance(row[1], str) and row[1].startswith(("05 -", "02 -")):
            lista.append(
                {
                    "n": row[0],
                    "empresa": row[1],
                    "categoria": row[2],
                    "item": row[3],
                    "descricao": row[4],
                    "matriculas": row[8],
                    "area": row[9],
                    "valor": Decimal(str(row[10])),
                    "pagina_diretrizes": row[13],
                    "obs": row[14],
                }
            )
    chosen = sorted(k for k in esc if k != "SUBTOTAL ESCOLHIDOS")
    check(
        "Lista de Rafael (xlsx): itens da 05 = escolhidos no PDF do Hotel",
        int(
            sorted(x["item"] for x in lista if x["empresa"].startswith("05")) == chosen
        ),
        1,
        "ativos-avaliados-valter.xlsx x PDF p. 22–23",
        nota=f"Escolhidos: {', '.join(chosen)}",
    )
    check(
        "Lista de Rafael: itens da empresa 02",
        len([x for x in lista if x["empresa"].startswith("02")]),
        0,
        "ativos-avaliados-valter.xlsx",
    )
    dtext = [p.get_text() for p in pymupdf.open(SRC["diretrizes"])]
    start = next(i for i, t in enumerate(dtext) if "HOTEL PRINCESA DOS VALES" in t)
    end = next(i for i, t in enumerate(dtext) if "HOTEL ARAÇÁ LTDA" in t)
    hotel_text = " ".join(dtext[start : end + 1])
    hotel_text = hotel_text[: hotel_text.find("HOTEL ARAÇÁ LTDA")]
    for x in lista:
        check(
            f"Lista de Rafael item {x['item']} = escolhidos no PDF",
            x["valor"],
            esc.get(x["item"]),
            "lista x PDF Hotel p. 22–23",
        )
        shown = (
            f"{x['valor']:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        )
        check(
            f"Diretrizes p. {start + 1}–{end + 1}: valor do item {x['item']} impresso no bloco do Hotel",
            int(
                bool(
                    re.search(
                        r"(?<![\d.])" + re.escape(shown) + r"(?![\d])", hotel_text
                    )
                )
            ),
            1,
            "PDF Diretrizes",
        )

    comp = []
    for b in i25_05:
        k = b["item"]
        r = exe.get(k)
        conc = [
            n for n in rnotes if n["tipo"] == "concordância expressa" and n["item"] == k
        ]
        extra = [
            n["tipo"]
            for n in rnotes
            if n["item"] == k and n["tipo"] != "concordância expressa"
        ]
        cited = [
            f"p. {n['pagina']}"
            for n in rnotes
            if n["tipo"] == "concordância expressa"
            and k in re.findall(r"\d{2}", n["itens_citados"])
        ]
        if r is not None and r != b["valor"]:
            extra.append(
                "valor executado sem memória de cálculo"
                + (
                    "; ver faixas de mercado citadas"
                    if any(f["item"] == k for f in rr)
                    else ""
                )
                + (
                    "; componentes impressos = os do grupo"
                    if k == "01" and exe01_is_group
                    else ""
                )
            )
        if cited and not conc:
            extra.append(f"citado na concordância de outros itens ({', '.join(cited)})")
        comp.append(
            {
                "item": k,
                "descricao": b["descricao"][:120],
                "grupo": b["valor"],
                "rafael": r,
                "diferenca": (r - b["valor"]) if r is not None else None,
                "na_lista_rafael": "sim" if k in esc else "não",
                "rafael_avaliacao_p7_15": ava.get(k),
                "concordancia": f"sim, p. {conc[0]['pagina']}" if conc else "não",
                "faixas": len([f for f in rr if f["item"] == k]),
                "notas": "; ".join(extra),
            }
        )
    check(
        "Rafael − grupo, total da empresa 05",
        sum(c["diferenca"] for c in comp),
        exe["TOTAL"] - total(t25_05),
        "comparação item a item",
    )

    # ---- 2018 benchmark bases (M1): two bases, conciliated from the sources
    hotel18 = total(t18_13)
    base_liq = liquido + hotel18
    jp18 = next(x for x in fl18 if x[0].startswith("Jane"))
    inv_sold = (
        sold18 * jp18[2]
        + sum(x["valor"] for x in rigo[1:])
        + sum(x["valor"] for x in pell)
    )
    angeli = next(x for x in i18_17 if x["comprador"] == "ANGELI")
    inv_sold += angeli["valor"]
    trator = next(x for x in i18_17 if "VALTRA" in x["descricao"])
    lotes_sei = [x for x in i18_17 if x["item"] in ("03", "04", "05")]
    gain = recebido - inv_sold
    kept = (
        hotel18
        + fazenda[0]["valor"]
        - sold18 * jp18[2]
        + sum(x["valor"] for x in fazenda if not x["comprador"] and x["seq"] != 1)
        + next(x for x in i18_17 if x["matriculas"] == "7.391")["valor"]
        + next(x for x in i18_17 if "JTB 8585" in x["descricao"])["valor"]
    )
    wbg = openpyxl.load_workbook(SRC["planilha_2_0"], data_only=True)
    planilha_hotel = Decimal(str(wbg["GERAL"]["C28"].value))
    planilha_agro = Decimal(str(wbg["GERAL"]["C11"].value))
    planilha_base = planilha_hotel + planilha_agro
    check(
        "Base 2018 'itens que permaneceram' (Hotel 2018 + fazenda sem a parte vendida + 4 ha + benfeitorias + lote 52 + Mercedinha) = planilha GERAL C28 (seção Hotel)",
        kept,
        planilha_hotel,
        "docs 13 e 17 x planilha 2_0",
    )
    check(
        "Planilha GERAL C11 (seção Agropecuária) = trator Valtra (único item da 02 que a planilha mantém)",
        planilha_agro,
        trator["valor"],
        "doc 17 x planilha 2_0",
    )
    gado = brl(
        need(
            r"R\$ (" + MONEY + r") – 41 cabeças de gado", v18[1]["texto"], "gado"
        ).group(1)
    )
    bases = [
        {
            "linha": "A. Líquido do doc 17 + total do doc 13 (base 'documento')",
            "valor": base_liq,
        },
        {
            "linha": "B. Base da planilha 2_0 para as duas empresas (GERAL C28, seção Hotel + C11, seção Agropecuária)",
            "valor": planilha_base,
        },
        {
            "linha": "   B1. GERAL C28 (seção Hotel, = aba 5.Hotel C2)",
            "valor": planilha_hotel,
        },
        {
            "linha": "   B2. GERAL C11 (seção Agropecuária: trator Valtra)",
            "valor": planilha_agro,
        },
        {"linha": "B − A", "valor": planilha_base - base_liq},
        {
            "linha": "Decomposição que reproduz B − A (não prova que a planilha a usou):",
            "valor": None,
        },
        {"linha": "Preço recebido nas vendas de 2018 (doc 17)", "valor": recebido},
        {
            "linha": f"Valor de inventário 2018 das áreas vendidas (Rigo {br(sold18)} ha × R$ {br(jp18[2])}/ha, premissa da planilha C944, "
            f"+ parcela mat. 11.090; Pellegrini 4 parcelas; Angeli)",
            "valor": inv_sold,
        },
        {"linha": "Ganho na venda (preço − inventário)", "valor": gain},
        {
            "linha": "Lotes 12, 89 e 33 (2018 itens 03–05; em 2025 na SEI; fora da base B)",
            "valor": sum(x["valor"] for x in lotes_sei),
        },
        {
            "linha": "Ganho − lotes 03–05 (deve ser igual a B − A)",
            "valor": gain - sum(x["valor"] for x in lotes_sei),
        },
        {
            "linha": f"Nota: o preço Pellegrini inclui R$ {br(gado)} de gado; o preço Rigo inclui parcelas marcadas 'FALTA RECEBER' "
            "(ver Créditos de soja). A escolha da base é do Luís.",
            "valor": None,
        },
    ]

    # ---- Google sheets (M4)
    ranges = {
        "2.São Crispin": None,
        "5.Hotel Princesa": None,
        "Extras": None,
        # 229-230, 234-235, 241: SEI rural lots that were Agropecuária assets in 2018
        "GERAL": [
            (11, 15),
            (28, 69),
            (229, 230),
            (234, 235),
            (241, 241),
            (882, 882),
            (941, 950),
            (986, 987),
        ],
        "Resumo": [(16, 20), (73, 73), (78, 102), (190, 191)],
        "11.Empreendimentos": [(90, 90), (731, 731)],
        "Acertos e Acordos": [(36, 43)],
    }
    obs_cols = {
        "GERAL": {"F", "H"},
        "11.Empreendimentos": {"H", "I"},
        "Acertos e Acordos": {"I"},
    }
    gs = []
    for name, rngs in ranges.items():
        ws = wbg[name]
        hdr = {c.column_letter: c.value for c in ws[1]}
        rows = (
            range(2, ws.max_row + 1)
            if rngs is None
            else [r for lo, hi in rngs for r in range(lo, hi + 1)]
        )
        for r in rows:
            for c in ws[r]:
                if c.value is None and c.comment is None:
                    continue
                comment = c.comment.text if c.comment else ""
                if "Comment:" in comment:
                    comment = comment.split("Comment:", 1)[1].strip()
                is_obs = (
                    hdr.get(c.column_letter) == "OBSERVAÇÕES"
                    or bool(comment)
                    or (
                        c.column_letter in obs_cols.get(name, set())
                        and isinstance(c.value, str)
                    )
                    or (name == "Resumo" and c.coordinate in RESUMO_OBS)
                )
                gs.append(
                    {
                        "aba": name,
                        "celula": c.coordinate,
                        "coluna": hdr.get(c.column_letter),
                        "valor": c.value,
                        "observacao": "sim" if is_obs else "",
                        "comentario": comment,
                    }
                )
    wbf = openpyxl.load_workbook(SRC["planilha_2_0"], data_only=False)
    f73 = str(wbf["Resumo"]["C73"].value)
    ref = re.search(r"GERAL!\$?([A-Z]+)\$?(\d+)", f73)
    if ref and "Hotel" in str(wbf["Resumo"]["B73"].value):
        target = wbg["GERAL"][f"B{ref.group(2)}"].value
        aviso(
            "Planilha 2_0",
            f"Resumo B73 'Hotel Princesa dos Vales' aponta para GERAL linha {ref.group(2)} ({str(target)[:80]}).",
            "planilha 2_0, Resumo C73 = " + f73,
        )
    wbo = openpyxl.load_workbook(SRC["observacoes"], data_only=True)
    covered = {(g["aba"], g["celula"]) for g in gs if g["observacao"]}
    obs_prev = sum(
        1
        for row in wbo["Observacoes"].iter_rows(min_row=2, values_only=True)
        if str(row[0]).startswith("2_0") and (row[1], row[2]) in covered
    )
    n_obs = len(covered)
    aviso(
        "Planilha 2_0",
        f"{n_obs} observações nas faixas destas empresas; {obs_prev} delas estão em "
        "observacoes-planilhas-google.xlsx (sessão SEI). As demais só nesta extração (aba 'Planilhas Google').",
        "planilha 2_0 x observacoes-planilhas-google.xlsx",
    )

    # ---- SEI 2018 'other companies' (M7) and 2018 -> 2025 by matrícula
    wb = openpyxl.load_workbook(SRC["inv2018_18x"], data_only=True)
    outras = []
    for row in wb["Bens em outras empresas"].iter_rows(min_row=4, values_only=True):
        if row[2] in ("AGROPECUÁRIA", "HOTEL"):
            outras.append(
                {
                    "item_no_doc": row[0],
                    "descricao": row[1],
                    "empresa": row[2],
                    "valor": Decimal(str(row[3])),
                    "obs": row[4],
                }
            )
    for x in i18_13 + i18_17:
        label = "HOTEL" if x["empresa"].startswith("05") else "AGROPECUÁRIA"
        hit = [
            o
            for o in outras
            if o["empresa"] == label
            and o["item_no_doc"] == x["item"]
            and o["valor"] == x["valor"]
        ]
        x["tambem_sei_2018"] = "sim" if hit else ""
    check(
        "2018: linhas dos docs 13/17 também listadas pela SEI = linhas da aba da SEI",
        sum(1 for x in i18_13 + i18_17 if x["tambem_sei_2018"]),
        len(outras),
        "doc 13/17 x xlsx SEI 2018",
        nota=(
            f"Listados pela SEI em 2018: R$ {br(sum(o['valor'] for o in outras))}. Entram na base A (docs 13 + 17) "
            "inteiros, mas o Angeli foi vendido em 2018 (o preço de venda o anula); na base B só o lote 52 e a "
            f"Sede Campestre (R$ {br(sum(x['valor'] for x in i18_13 + i18_17 if x['tambem_sei_2018'] and (x['empresa'].startswith('05') or x['matriculas'] == '7.391')))})."
        ),
    )

    docs2025 = {
        "atual": {"inv2025_05": SRC["inv2025_05"], **OTHER_2025},
        "antigo": {"antigo_05": SRC["antigo_05"], **OLD_2025},
    }
    where = []
    rows2025 = []
    for version, docs in docs2025.items():
        for key, path in docs.items():
            for t_idx, rows in enumerate(docx_tables(path), 1):
                for r_idx, cells in enumerate(rows, 1):
                    text = " | ".join(cells)
                    rows2025.append(
                        (version, key, f"T{t_idx}R{r_idx}", text, set(matriculas(text)))
                    )
    # a matrícula shared by a sold parcel and a kept one (11.090) says nothing
    # about the sold parcel: flag it on the row itself
    kept_mats = {
        m
        for x in i18_13 + i18_17
        if "VENDIDA" not in x["situacao"] or "PARCIAL" in x["situacao"]
        for m in x["matriculas"].split(", ")
        if m
    }
    for x in i18_13 + i18_17:
        mats = {m for m in x["matriculas"].split(", ") if m}
        kw = [
            k
            for k in ("VALTRA", "JTB 8585", "Benfeitorias existentes")
            if k in x["descricao"]
        ]
        for version in ("atual", "antigo"):
            hits = [
                (k, r, t)
                for v, k, r, t, ms in rows2025
                if v == version and ((mats and mats & ms) or any(w in t for w in kw))
            ]
            areas25 = [area_text(t) for _, _, t in hits]
            values25 = [
                money_values(t)[-1] if money_values(t) else None for _, _, t in hits
            ]
            note = ""
            if x["situacao"] == "VENDIDA" and mats & kept_mats:
                note = "matrícula compartilhada com parcela não vendida; a parcela vendida não está em 2025"
            elif (
                len(hits) == 1
                and x["area"]
                and areas25[0]
                and area_ha(x["area"])
                and area_ha(areas25[0])
                and area_ha(x["area"]) != area_ha(areas25[0])
            ):
                gap = area_ha(x["area"]) - area_ha(areas25[0])
                surname = x["adquirido_de"].split()[-1] if x["adquirido_de"] else ""
                cand = [
                    (k, r, t)
                    for v, k, r, t, _ in rows2025
                    if v == version
                    and surname
                    and surname.upper() in t.upper()
                    and area_ha(area_text(t)) == gap
                    and (k, r, t) not in hits
                ]
                note = f"área {x['area']} em 2018 × {areas25[0]} em 2025 (diferença {br(gap, 1)} ha)"
                if cand:
                    note += "; mesma área e vendedor em " + "; ".join(
                        f"{k[-2:]} {r}: {t[:180]}…" for k, r, t in cand
                    )
                aviso(
                    "Área 2018 × 2025",
                    f"Seq. {x['seq']} ({x['matriculas']}, {version}): {note}.",
                    "docs 17 2018 e inventários 2025",
                )
            where.append(
                {
                    "versao_2025": version,
                    "empresa_2018": x["empresa"],
                    "seq": x["seq"],
                    "item_2018": x["item"],
                    "matriculas": x["matriculas"],
                    "area": x["area"],
                    "valor_2018": x["valor"],
                    "situacao_2018": x["situacao"],
                    "tambem_sei_2018": x["tambem_sei_2018"],
                    "onde_2025": "; ".join(
                        f"{k.split('_')[-1]} {r}" for k, r, _ in hits
                    )
                    or "não encontrado",
                    "area_2025": "; ".join(a or "—" for a in areas25),
                    "valor_2025": "; ".join(
                        br(v) if v is not None else "—" for v in values25
                    ),
                    "nota": note,
                    "texto_2025": " || ".join(t[:160] for _, _, t in hits),
                }
            )

    # ---- mentions (B2)
    mentions = []
    for key in ("inv2025_02", "inv2025_05", "inv2025_areas", "antigo_05", "inv2025_00"):
        for p in docx_paragraphs(SRC[key]):
            if re.search(r"ROV[ÍI]LIO|Rov[íi]lio|outubro/2025", p, re.IGNORECASE):
                mentions.append({"fonte": key, "texto": p[:400]})
    check(
        "Menção a Rovílio ou a outubro/2025 nos docs 02, 05 e docx de áreas",
        len([m for m in mentions if m["fonte"] != "inv2025_00"]),
        0,
        "docs 02, 05, antigo 05, docx áreas",
        nota="O doc 00 menciona (aba 'Menções').",
    )

    write_xlsx(
        fontes,
        i25_05 + i25a_05 + i25_02,
        diff_rows,
        i18_13 + i18_17,
        totals18,
        v18,
        va,
        sales_diff,
        soy,
        areas_rows,
        rv,
        rr,
        rnotes + [dict(n, item="Agro " + n["item"]) for n in rnotes2],
        comp,
        lista,
        gs,
        outras,
        where,
        resumo,
        bases,
        mentions,
    )

    print(f"fontes: {len(fontes)}")
    print(
        f"2025: 05 atual {len(i25_05)} itens; 05 antigo {len(i25a_05)}; 02 atual {len(i25_02)}"
    )
    print(
        f"2018: 05 {len(i18_13)} itens; 02 {len(i18_17)} linhas; vendas {len(v18)}; soja {len(soy)} linhas"
    )
    print(f"Rafael: {len(rv)} valores, {len(rr)} faixas, {len(rnotes)} notas")
    print(
        f"Planilhas Google: {len(gs)} células, {n_obs} observações; já no arquivo da SEI: {obs_prev}"
    )
    print(f"avisos: {len(avisos)}")
    for a in avisos:
        print("  AVISO", a["assunto"], "|", a["texto"])
    print("--- conferências")
    for c in checks:
        print(
            f"{c['ok']:8} {c['conferencia']}: extraído {c['extraido']} | declarado {c['declarado']}"
            + (f"  [{c['nota']}]" if c["ok"] != "OK" and c["nota"] else "")
        )
    print(f"total {len(checks)}, divergem {sum(c['ok'] != 'OK' for c in checks)}")


def write_xlsx(
    fontes,
    i25,
    diff_rows,
    i18,
    t18,
    v18,
    va,
    sales_diff,
    soy,
    areas_rows,
    rv,
    rr,
    rnotes,
    comp,
    lista,
    gs,
    outras,
    where,
    resumo,
    bases,
    mentions,
):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wide = {
        "descricao",
        "texto",
        "caminho",
        "contexto",
        "texto_valor",
        "nota",
        "onde_2025",
        "texto_2025",
        "linha",
        "notas",
        "doc17",
        "docx2025",
    }

    def sheet(name, rows, cols, note=None):
        ws = wb.create_sheet(name)
        r0 = 1
        if note:
            ws.cell(1, 1, note).font = Font(italic=True)
            r0 = 3
        for j, (key, title) in enumerate(cols, 1):
            c = ws.cell(r0, j, title)
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDEBF7")
        for i, row in enumerate(rows, r0 + 1):
            for j, (key, _) in enumerate(cols, 1):
                v = row.get(key)
                if isinstance(v, Decimal):
                    ws.cell(i, j, float(v)).number_format = "#,##0.00"
                else:
                    ws.cell(i, j, v)
        for j, (key, title) in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(j)].width = (
                60 if key in wide else max(12, len(title) + 2)
            )
        ws.freeze_panes = ws.cell(r0 + 1, 1)

    sheet(
        "Conferências",
        checks,
        [
            ("conferencia", "Conferência"),
            ("extraido", "Extraído"),
            ("declarado", "Declarado / esperado"),
            ("ok", "Resultado"),
            ("fonte", "Fonte"),
            ("nota", "Nota"),
        ],
        "Cada total ou área extraído por script comparado com o que o documento declara. DIVERGE = achado a tratar, não erro de leitura.",
    )
    sheet(
        "Avisos",
        avisos,
        [("assunto", "Assunto"), ("texto", "Texto"), ("fonte", "Fonte")],
        "Inconsistências de redação ou de referência encontradas nas fontes.",
    )
    sheet(
        "Fontes",
        fontes,
        [
            ("chave", "Chave"),
            ("papel", "Papel"),
            ("md5", "MD5"),
            ("caminho", "Caminho"),
        ],
    )
    item_cols = [
        ("empresa", "Empresa"),
        ("versao", "Versão"),
        ("seq", "Seq."),
        ("item", "Item"),
        ("matriculas", "Matrículas"),
        ("area", "Área total"),
        ("area_avaliada", "Área avaliada"),
        ("adquirido_de", "Adquirido de"),
        ("valor", "Valor (R$)"),
        ("situacao", "Situação"),
        ("comprador", "Comprador"),
        ("ref", "Tabela/linha"),
        ("texto_valor", "Texto da célula de valor"),
        ("descricao", "Descrição"),
    ]
    sheet("Itens 2025", i25, item_cols)
    sheet(
        "2025 antigo x atual",
        diff_rows,
        [
            ("item", "Item"),
            ("antigo", "Antigo (R$)"),
            ("atual", "Atual (R$)"),
            ("diferenca", "Diferença (R$)"),
            ("texto_mudou", "Texto mudou?"),
        ],
    )
    sheet(
        "Itens 2018",
        i18,
        item_cols[:12]
        + [
            ("bloco_fazenda", "Bloco da fazenda?"),
            ("tambem_sei_2018", "Também listado pela SEI em 2018?"),
        ]
        + item_cols[12:],
        "Doc 17 (Agropecuária) e doc 13 (Hotel) de 2018. O doc 17 repete o número de item 01 nas parcelas da fazenda; 'Seq.' é a ordem no documento.",
    )
    sheet(
        "Totais 2018",
        t18,
        [
            ("empresa", "Empresa"),
            ("rotulo", "Rótulo"),
            ("valor", "Valor (R$)"),
            ("ref", "Tabela/linha"),
        ],
    )
    sheet(
        "Vendas 2018",
        [dict(v, fonte="doc 17 2018") for v in v18]
        + [dict(v, fonte="docx áreas vendidas 2025") for v in va],
        [
            ("fonte", "Fonte"),
            ("comprador", "Comprador"),
            ("area_ha", "Área (ha)"),
            ("valor", "Valor (R$)"),
            ("falta_receber", "Parcelas 'FALTA RECEBER'"),
            ("ref", "Tabela/linha"),
            ("texto", "Texto"),
        ],
    )
    sheet(
        "Vendas docx x doc 17",
        sales_diff,
        [
            ("comprador", "Venda"),
            ("operacao", "Diferença"),
            ("doc17", "Texto no doc 17 (2018)"),
            ("docx2025", "Texto no docx (2025)"),
        ],
        "Diferenças de texto, palavra a palavra, entre as vendas do doc 17 de 2018 e as do docx 'áreas vendidas' de 2025.",
    )
    sheet(
        "Créditos de soja",
        soy,
        [
            ("origem", "Origem"),
            ("tipo", "Tipo"),
            ("ano", "Ano"),
            ("sacas", "Sacas"),
            ("empresa", "Empresa anotada"),
            ("texto", "Texto"),
        ],
        "Parágrafos OBSERVAÇÕES do docx 'áreas vendidas' (2025). Sem valor em R$ no documento.",
    )
    sheet(
        "Áreas da fazenda",
        areas_rows,
        [
            ("bloco", "Bloco"),
            ("chave", "Matrícula / área"),
            ("ha", "ha"),
            ("valor", "Valor (R$)"),
            ("situacao", "Situação 2018"),
            ("comprador", "Comprador"),
            ("ref", "Tabela/linha"),
        ],
    )
    sheet(
        "Base 2018",
        bases,
        [("linha", "Linha"), ("valor", "Valor (R$)")],
        "Duas bases possíveis para o benchmark 2018 das empresas 02 + 05, conciliadas pelas fontes. Só informação: a escolha é do Luís.",
    )
    sheet(
        "Rafael 05 valores",
        rv,
        [
            ("secao", "Seção do PDF"),
            ("item", "Item"),
            ("valor", "Valor (R$)"),
            ("pagina", "Página"),
            ("observacao", "Observação"),
            ("contexto", "Contexto"),
        ],
    )
    sheet(
        "Rafael 05 faixas",
        rr,
        [
            ("secao", "Seção do PDF"),
            ("item", "Item"),
            ("pagina", "Página"),
            ("texto", "Texto"),
        ],
        "Faixas de mercado e referências citadas por Rafael (argumento C1).",
    )
    sheet(
        "Rafael notas",
        rnotes,
        [
            ("item", "Item"),
            ("pagina", "Página"),
            ("tipo", "Tipo"),
            ("itens_citados", "Itens citados"),
            ("texto", "Texto"),
        ],
    )
    sheet(
        "Grupo x Rafael 05",
        comp,
        [
            ("item", "Item"),
            ("grupo", "Grupo 2025 (R$)"),
            ("rafael", "Rafael, balanço executado (R$)"),
            ("diferenca", "Rafael − grupo (R$)"),
            ("na_lista_rafael", "Na lista de Rafael?"),
            ("rafael_avaliacao_p7_15", "Valor na seção 'avaliação executada' p. 7–15"),
            ("concordancia", "Concordância expressa?"),
            ("faixas", "Faixas de mercado"),
            ("notas", "Notas"),
            ("descricao", "Descrição"),
        ],
    )
    sheet(
        "Lista de Rafael",
        lista,
        [
            ("n", "Nº"),
            ("empresa", "Empresa"),
            ("item", "Item"),
            ("categoria", "Categoria"),
            ("matriculas", "Matrículas"),
            ("area", "Área"),
            ("valor", "Valor (R$)"),
            ("pagina_diretrizes", "Página Diretrizes"),
            ("obs", "Observação"),
            ("descricao", "Descrição"),
        ],
    )
    sheet(
        "Planilhas Google",
        gs,
        [
            ("aba", "Aba"),
            ("celula", "Célula"),
            ("coluna", "Coluna"),
            ("observacao", "É observação?"),
            ("valor", "Valor / texto"),
            ("comentario", "Comentário"),
        ],
    )
    sheet(
        "SEI 2018 outras empresas",
        outras,
        [
            ("item_no_doc", "Item no doc"),
            ("empresa", "Relacionado em"),
            ("valor", "Valor (R$)"),
            ("obs", "Observação"),
            ("descricao", "Descrição"),
        ],
        "Inventário 2018 da SEI, aba 'Bens em outras empresas': linhas da Agropecuária e do Hotel.",
    )
    sheet(
        "2018 → 2025 por matrícula",
        where,
        [
            ("versao_2025", "Versão 2025"),
            ("empresa_2018", "Empresa 2018"),
            ("seq", "Seq."),
            ("item_2018", "Item 2018"),
            ("matriculas", "Matrículas"),
            ("area", "Área 2018"),
            ("valor_2018", "Valor 2018 (R$)"),
            ("situacao_2018", "Situação 2018"),
            ("tambem_sei_2018", "Também SEI 2018?"),
            ("onde_2025", "Onde aparece em 2025 (doc, tabela/linha)"),
            ("area_2025", "Área 2025"),
            ("valor_2025", "Valor 2025 (R$)"),
            ("nota", "Nota"),
            ("texto_2025", "Texto em 2025"),
        ],
        "Ocorrências da mesma matrícula (ou do mesmo veículo/benfeitoria) nos inventários 2025, versão atual e antiga. "
        "Só localização: o destino (P4) é decidido depois.",
    )
    sheet(
        "Resumos de capital",
        resumo,
        [
            ("fonte", "Fonte"),
            ("tipo", "Tipo"),
            ("texto", "Linha"),
            ("valor", "Valor (R$)"),
            ("ref", "Tabela/linha"),
        ],
    )
    sheet(
        "Menções",
        mentions,
        [("fonte", "Fonte"), ("texto", "Parágrafo")],
        "Parágrafos que citam Rovílio ou outubro/2025.",
    )
    wb.save(OUT_XLSX)


if __name__ == "__main__":
    main()
