"""Packet 1 (block rules) for companies 03, 04, 07, 09 and 10: which master rows each rule
covers, their totals and a sample, written to pacotes/pacote-01-tabelas.md. The narrative
packet (pacote-01.md) quotes these tables; numbers come only from here."""

from __future__ import annotations

import hashlib
from pathlib import Path

import openpyxl

BASE = Path(
    r"C:\entrega-cisao-siviero\02_trabalho\empresas\03-04-07-09-10-radio-ceval-baixadas"
)
MASTER = BASE / "planilha-mestre-03-04-07-09-10.xlsx"
OUT = BASE / "pacotes" / "pacote-01-tabelas.md"
SAMPLE = ["07-2025-02", "07-2025-03", "07-2025-04", "07-2025-05", "07-2025-07"]


def brl(v) -> str:
    if v in (None, ""):
        return "—"
    s = f"{v:,.2f}"
    return "R$ " + s.replace(",", "X").replace(".", ",").replace("X", ".")


def main():
    wb = openpyxl.load_workbook(MASTER, data_only=True)
    ws = wb["Mestre"]
    it = ws.iter_rows(values_only=True)
    head = next(it)
    rows = [dict(zip(head, r)) for r in it if r[0]]
    summ = {
        r[0]: dict(zip([c.value for c in wb["Resumo por empresa"][1]], r))
        for r in wb["Resumo por empresa"].iter_rows(min_row=2, values_only=True)
    }

    # R1 = same value on both sides, except when Rafael's balance value lies outside his own
    # market range: that "agreement" is only formal, and treating it is the user's decision.
    agree = [r for r in rows if r["Situação"] == "Concorda"]
    q_luis = [r for r in agree if "fora da faixa" in (r["Critério de Rafael"] or "")]
    r1 = [r for r in agree if r not in q_luis]
    n_notes = wb["Créditos e dívidas (notas)"].max_row - 1
    r3 = [r for r in rows if r["Situação"] == "Bem de outra empresa listado em 2018"]
    exc = [r for r in rows if r["Situação"] == "Discorda"]
    dest = [
        r
        for r in rows
        if r["Situação"]
        in (
            "Só em 2018 (ausente em 2025)",
            "Vendido antes de 2018 (crédito a receber em 2018)",
        )
    ]
    baixadas = [summ[k] for k in summ if k.startswith(("03", "09", "10"))]
    covered = {r["ID"] for r in r1 + q_luis + r3 + exc + dest}
    dest_goods = [r for r in dest if r["Situação"].startswith("Só em 2018")]
    dest_sold = [r for r in dest if r["Situação"].startswith("Vendido antes")]
    missing = [r["ID"] for r in rows if r["ID"] not in covered]

    out = [
        f"<!-- gerado por _scripts/pacote01.py a partir de planilha-mestre (md5 {hashlib.md5(MASTER.read_bytes()).hexdigest()}) -->",
        "",
    ]
    out += [
        "## R1 — bens de 2025 com o mesmo valor nas duas avaliações",
        "",
        "| ID | Empresa | Bem | Grupo 2025 | Rafael | Critério de Rafael | Lista C0? | Parte 6 |",
        "|---|---|---|---|---|---|---|---|",
    ]
    for r in r1:
        out.append(
            f"| {r['ID']} | {r['Empresa (2025)']} | {r['Descrição'][:70]} | {brl(r['Valor grupo 2025'])} | {brl(r['Valor Rafael'])} | "
            f"{r['Critério de Rafael'] or ''} | {r['Lista de Rafael (C0)?']} | {r['Parte 6'] or ''} |"
        )
    out += [
        "",
        (
            f"Total R1: {len(r1)} bens, {brl(sum(r['Valor grupo 2025'] for r in r1))} "
            f"(Rádio {brl(sum(r['Valor grupo 2025'] for r in r1 if r['ID'].startswith('07')))}; "
            f"Ceval {brl(sum(r['Valor grupo 2025'] for r in r1 if r['ID'].startswith('04')))})."
        ),
        "",
    ]
    out += [
        "## R2 — empresas baixadas sem bens em 2025 nas duas avaliações",
        "",
        "| Empresa | Grupo 2025 | Doc 00 | Rafael |",
        "|---|---|---|---|",
    ]
    for s in baixadas:
        out.append(
            f"| {s['Empresa']} | {brl(s['Soma grupo 2025'])} | {brl(s['Doc 00 (2025)'])} | {s['Confere Rafael?']} |"
        )
    out += [
        "",
        "## R3 — bem de outra empresa listado em 2018",
        "",
        "| ID | Bem | Valor 2018 | Em 2025 | Destino |",
        "|---|---|---|---|---|",
    ]
    for r in r3:
        out.append(
            f"| {r['ID']} | {r['Descrição'][:70]} | {brl(r['Valor 2018'])} | {r['Ref. 2025']} | {r['Destino 2018→2025'][:160]} |"
        )
    out += ["", "## Fora das regras", "", "Exceções (pacote 2):", ""]
    for r in exc:
        out.append(
            f"- {r['ID']} {r['Empresa (2025)']}: {r['Descrição'][:60]} — grupo {brl(r['Valor grupo 2025'])}, Rafael {brl(r['Valor Rafael'])}"
        )
    out += [
        "",
        "Concordância só formal (valor de Rafael fora da própria faixa) — pergunta ao Luís (pacote 2):",
        "",
    ]
    for r in q_luis:
        out.append(
            f"- {r['ID']} {r['Empresa (2025)']}: {r['Descrição'][:60]} — grupo = Rafael {brl(r['Valor grupo 2025'])}; "
            f"faixa de Rafael {r['Faixa de Rafael (C1)']}; outras referências {r['Outras referências de Rafael']}"
        )
    out += [
        "",
        f"Destino 2018 → 2025 (etapa 6): {len(dest_goods)} linhas de bens de 2018 "
        f"({brl(sum(r['Valor 2018'] for r in dest_goods))}): "
        + ", ".join(r["ID"] for r in dest_goods)
        + f"; e {len(dest_sold)} ônibus vendidos a prazo antes de 2018 ({brl(sum(r['Valor 2018'] for r in dest_sold))}, "
        "fora do total de 2018): " + ", ".join(r["ID"] for r in dest_sold) + ".",
        "",
        (
            f"Créditos, dívidas e intragrupo (aba 'Créditos e dívidas' da planilha-mestre): {n_notes} notas, "
            "nenhuma coberta pelas regras; vão para a etapa 6 e para a Parte 6."
        ),
        "",
    ]
    out += [
        f"Linhas da planilha sem enquadramento: {missing or 'nenhuma'} (de {len(rows)}).",
        "",
    ]
    out += [
        "## Amostra para conferir",
        "",
        "| ID | Bem | Grupo 2025 (fonte) | Rafael (fonte) | Faixa de Rafael |",
        "|---|---|---|---|---|",
    ]
    for sid in SAMPLE:
        r = next(x for x in rows if x["ID"] == sid)
        out.append(
            f"| {sid} | {r['Descrição'][:60]} | {brl(r['Valor grupo 2025'])} ({r['Ref. 2025']}) | {brl(r['Valor Rafael'])} ({r['Ref. Rafael']}) | "
            f"{r['Faixa de Rafael (C1)'] or r['Critério de Rafael'] or ''} |"
        )
    OUT.write_text("\n".join(out) + "\n", encoding="utf-8")
    print("\n".join(out))


if __name__ == "__main__":
    main()
