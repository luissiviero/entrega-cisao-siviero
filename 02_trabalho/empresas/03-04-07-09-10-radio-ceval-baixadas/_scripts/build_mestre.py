"""Master spreadsheet for companies 03, 04, 07, 09 and 10 (Siviero split study).

Input: the Phase 0 extraction (_extracao/fase0-extracao.json, built from the sources by
fase0_extracao.py), the approved decisions (decisoes.json) and the Google sheets (only to
read the full description of rows that carry Luís's notes). Output:
planilha-mestre-03-04-07-09-10.xlsx. Never edited by hand: rerun after every decision.

Pairs 2018 x 2025 are made only by registration keys (matrícula, placa, chassi, nota
fiscal), never by text similarity (method lesson 3). Pairs without a key come from
decisoes.json ("casamentos") and stay marked as proposed until approved.
Labels are in Portuguese (deliverable language); code is in English.
"""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(
    r"C:\entrega-cisao-siviero\02_trabalho\empresas\03-04-07-09-10-radio-ceval-baixadas"
)
EXTRACT = BASE / "_extracao" / "fase0-extracao.json"
DECISIONS = BASE / "decisoes.json"
OUT = BASE / "planilha-mestre-03-04-07-09-10.xlsx"
GSHEETS = {
    "2_0": Path(
        r"C:\entrega-cisao-siviero\01_fontes\planilhas-google\2_0_Comparativo Patrimonial Siviero 2018 x 2025.xlsx"
    ),
    "emp. e Rafael": Path(
        r"C:\entrega-cisao-siviero\01_fontes\planilhas-google\Comparativo Patrimonial Siviero empreendimentos e Rafael 2018 x 2025.xlsx"
    ),
}

NAMES = {
    "03": "03 Associação Cultural",
    "04": "04 Ceval",
    "07": "07 Rádio",
    "09": "09 Turismo",
    "10": "10 Siviero Transportes",
}
# File names of the 2018 inventories on the Drive (numbering differs from 2025).
FILE18 = {
    "03": "arq. 14.1 (Associação 2018)",
    "04": "arq. 11.1 (Ceval 2018)",
    "07": "arq. 9.1 (Rádio 2018)",
    "09": "arq. 8.1 (Turismo 2018)",
    "10": "arq. 7.1 (Siviero Transportes 2018)",
}
# Rafael's C0 list (Valter selection): only the Toyota among these companies (Phase 0 check).
C0_KEYS = {"placa IHW0285"}

# Facts read in the sources (Phase 0 and reviewer of the master), applied by registration key.
OTHER_OWNER = {
    "mat. 3.751": {
        "situacao": "Bem de outra empresa listado em 2018",
        "ref18": "também doc 18 (SEI 2018), tab. 31, item 06: 'Relacionado no Inventário da empresa: RÁDIO'",
        "destino": "Declarado como da SEI: doc 18 tab. 31 (bens da SEI relacionados em outras empresas); doc 25 par. 204 "
        "('imóvel que pertence à empresa Siviero Empreendimentos'); inventário 2025 da Rádio par. 15; "
        "arq. 1 p. 12 (cessão verbal da área). Em 2018 só listado no inventário da Rádio; em 2025 no doc 11, "
        "tab. 7, item 02.",
    }
}
# Destinations backed by sources, by item id (2018 ids).
DEST_FACTS = {
    "10-2018-05": "Vendido a prazo em dez/2015 por R$ 585.000 (doc 25 par. 249). Em 2025: processo da venda, acordo de "
    "R$ 115.000 (R$ 20.000 de despesas + parcelas de R$ 5.000 desde nov/2024) — ver notas 10-2025-N22; "
    "crédito em aberto em 31.12.2024 numa empresa baixada. Recebido até a baixa e saldo: pendente (P4).",
    "10-2018-06": "Vendido a prazo em dez/2016. Doc 25 par. 229–230 (2018): a receber 24 parcelas de R$ 10.000 "
    "('venda do ônibus para uma empresa de Santa Catarina') e R$ 100.000 'do Rodrigo'. Recebimento até a "
    "baixa: pendente (P4).",
    "10-2018-07": "Vendido a prazo em dez/2015. Doc 25 par. 229–230 (2018): recebíveis da Siviero Transportes. "
    "Recebimento até a baixa: pendente (P4).",
}
DEST_PREFIX = {
    "10-2018-01": "Financiado em 2018 (doc 25 par. 227: 18 parcelas de R$ 8.000 à GUIBOR 'do ônibus atual'). ",
}
# Part 6 candidates of this session (parte6-candidatos.md), by key and by id.
PART6_KEY = {
    "mat. 3.751": "S4-1",
    "NF 4946": "S4-2",
    "placa IKI4325": "S4-3",
    "placa IHW0285": "S4-8",
}
PART6_ID = {
    "07-2025-04": ["S4-1"],
    "07-2025-01": ["S4-6", "S4-9"],
    "07-2025-05": ["S4-9"],
    "10-2018-05": ["S4-7"],
    "10-2018-06": ["S4-7"],
    "10-2018-07": ["S4-7"],
}
NOTE_PART6 = [
    (r"Celito", "S4-4"),
    (
        r"Siviero Empreendimentos Imobiliários Ltda\., ainda|2\.003\.000|NÃO paga",
        "S4-5",
    ),
    (r"COLINA DO RETIRO", "S4-6"),
    (r"DESLOCAMENTO DA ANTENA|redução de 8", "S4-1"),
    (r"IXO 6000", "S4-7"),
    (r"redução de 8", "S4-10"),
]
REG_DATA_BASE_ID = {
    "07-2025-03": "S4-DB3: licença Anatel vigente em 31.12.2024 (validade 08/04/2027); renovação de 2026 é posterior",
    "10-2018-05": "S4-DB2: acordo do IXO 6000 (parcelas desde nov/2024; natureza, titular e saldo a confirmar — pergunta 10)",
}
REG_DATA_BASE = {
    "placa IHW0285": "Verificar titular em 31.12.2024 (ST baixada em 08/04/2024; sem documento da transferência)",
    "mat. 11.382": "Titularidade em 31.12.2024 a conferir no arq. 13 (matrícula)",
    "mat. 11.324": "Titularidade em 31.12.2024 a conferir no arq. 13 (matrícula)",
}
# Base of the exiting partner's quota (P1); Rádio decided by Luís in packet 2, item 7.
QUOTA_BASE = {"07": (0.165, "16,50 % (sem o 1 % de Telmo de Toni; pacote 2, item 7)")}
OBS_DONE = {"07-2018-06": "Respondida (R3 do pacote 1; S4-1)"}
# Open points found in packet 3 outside the inventory texts (approved 28/10/2026).
NOTES_EXTRA = [
    {"ID": "04-2024-ATIV", "Empresa": "04 Ceval", "Versão": "contas pagas 2023–2024", "Parágrafo": "Consolidado",
     "Valores": "R$ 446.244,50 (bases das notas)", "Natureza": "atividade não registrada no inventário",
     "Texto": "Notas da Ceval (CERAN NF 20 R$ 126.590 e NF 22 R$ 286.940; Luiz Argenta R$ 27.200; Renan Baronio NF 21; "
              "nota 23 R$ 2.362,50), folha, alvará e contabilidade, com tributos e despesas pagos pela SEI.",
     "Situação em 31.12.2024 / destino": "Pacote 3 item 4: caixa, recebíveis, passivos e saldo com a SEI não registrados "
                                          "— perguntas 8 e 17.", "Parte 6": "S4-12"},
    {"ID": "07-2023-CAIXA", "Empresa": "07 Rádio", "Versão": "contas pagas jul/2023", "Parágrafo": "p. 5 (imagem)",
     "Valores": "R$ 6.500 (parcela 28/28)", "Natureza": "intragrupo",
     "Texto": "Seção da Abastecedora, atribuído à SIVIERO TRANSPORTES LTDA: 'P. 28/28 - EMPRESTIMO CAIXA FEDERAL - RADIO C…' "
              "e tarifa 'PAGTO EMPRESTIMO RADIO-SIVIER…' (texto cortado na fonte).",
     "Situação em 31.12.2024 / destino": "Pacote 3 item 3: empréstimo bancário encerrado em jul/2023; devedor e conta entre "
                                          "Abastecedora, ST e Rádio a confirmar — pergunta 15. A S3 lê como quitado e sem efeito "
                                          "(alinhar na união, P13).", "Parte 6": "S4-11"},
]
RANGE = re.compile(r"R\$\s*([\d.]+(?:,\d{2})?)\s*(?:e|a)\s*R\$\s*([\d.]+(?:,\d{2})?)")


def md5(p: Path) -> str:
    return hashlib.md5(p.read_bytes()).hexdigest()


def brl(v: float | None, dec: int = 2) -> str:
    if v is None:
        return ""
    s = f"{v:,.{dec}f}"
    return "R$ " + s.replace(",", "X").replace(".", ",").replace("X", ".")


def keyset(chaves: str) -> set[str]:
    return {k.strip() for k in chaves.split(";") if k.strip()}


def item_no(s) -> str:
    d = re.sub(r"\D", "", str(s or ""))
    return d.zfill(2) if d else ""


def block_of(desc: str) -> str:
    u = desc.upper()
    if re.search(r"SALA COMERCIAL|BOX DE ESTACIONAMENTO|TERRENO|LOTE URBANO", u):
        return "Imóveis"
    if re.search(r"ESCAVADEIRA|RETROESCAVADEIRA", u):
        return "Máquinas"
    if re.search(r"VEÍCULO|CAMINHÃO|ÔNIBUS|TOYOTA", u):
        return "Veículos"
    if re.search(r"ANTENA|EQUIPAMENTOS|MÓVEIS|CANAL", u):
        return "Equipamentos, móveis e concessão"
    return "Outros"


def short(desc: str, n: int = 110) -> str:
    d = re.sub(r"\s+", " ", desc).strip()
    return d if len(d) <= n else d[: n - 1] + "…"


def main():
    data = json.loads(EXTRACT.read_text(encoding="utf-8"))
    dec = json.loads(DECISIONS.read_text(encoding="utf-8"))
    inv = data["inventories"]

    # --- items by version, with ids --------------------------------------------------
    items = {"2018": {}, "2025": {}}
    for code, d in inv.items():
        for ver in ("2018", "2025"):
            for it in d[ver]["items"]:
                if it["vazio"] or (
                    it["valor"] in (None, 0) and it["descricao"].strip(".- ") == ""
                ):
                    continue
                iid = f"{code}-{ver}-{item_no(it['item'])}"
                items[ver][iid] = {
                    **it,
                    "id": iid,
                    "empresa": code,
                    "keys": keyset(it["chaves"]),
                }
    rafael, faixas = {}, {}
    for code, d in inv.items():
        for it in d["raf"]["items"]:
            if it["secao"] == "balanco_valter":
                rafael[f"{code}-{it['item']:02d}"] = it
            elif it["secao"] == "avaliacao_valter":
                faixas[f"{code}-{it['item']:02d}"] = it

    # SEI 2025 rows that share a key with an asset of these companies (cross references).
    sei25 = []
    for x in data["cross_refs"]:
        if x["fonte"].startswith("Doc 11") and x["local"].startswith("tabela"):
            cells = [c.strip() for c in x["texto"].split("|")]
            m = re.search(r"Matr[íi]cula\s*n?º?\s*([\d.]+\d)", x["texto"])
            sei25.append(
                {
                    "ref": f"doc 11 (SEI 2025), {x['local']}, item {cells[0]}",
                    "valor": float(x["valores"].split(";")[-1].replace(",", ""))
                    if x["valores"]
                    else None,
                    "keys": {f"mat. {m.group(1)}"} if m else set(),
                }
            )

    # --- pairing 2018 x 2025 by key ----------------------------------------------------
    pair18 = {}  # id2018 -> (target, how); target = 2025 id, list of 2025 ids, or SEI dict
    used25 = set()
    for i18, a in items["2018"].items():
        for i25, b in items["2025"].items():
            common = a["keys"] & b["keys"]
            if common and i25 not in used25:
                pair18[i18] = (i25, "por chave: " + ", ".join(sorted(common)))
                used25.add(i25)
                break
        if i18 not in pair18:
            for s in sei25:
                common = a["keys"] & s["keys"]
                if common:
                    pair18[i18] = (
                        s,
                        "por chave (outra empresa): " + ", ".join(sorted(common)),
                    )
                    break
    split_of = {}  # 2025 id -> (2018 id, how) for the 2nd+ targets of a 1 -> n proposed pair
    for i18, c in dec.get("casamentos", {}).items():
        targets = c["com"] if isinstance(c["com"], list) else [c["com"]]
        if (
            i18 not in items["2018"]
            or i18 in pair18
            or not all(t in items["2025"] for t in targets)
        ):
            continue
        label = (
            "aprovado" if c["situacao"] == "aprovado" else "PROPOSTO, aguarda aprovação"
        )
        if len(targets) > 1:
            tot = sum(items["2025"][t]["valor"] or 0 for t in targets)
            how = (
                f"sem chave — {label}: {i18} → {' + '.join(targets)} (desmembramento; comparar "
                f"{brl(items['2018'][i18]['valor'])} × {brl(tot)}). {c['motivo']}"
            )
        else:
            how = f"sem chave — {label}: {c['motivo']}"
        main_t = targets[-1]
        pair18[i18] = (main_t, how)
        used25.update(targets)
        for t in targets[:-1]:
            split_of[t] = (i18, how)

    # --- Luís's notes from the Google sheets, attached by registration key -------------
    obs_by_key = {}
    wbs = {k: openpyxl.load_workbook(p, data_only=True) for k, p in GSHEETS.items()}
    for h in data["gsheet_notes"]:
        if h["tipo"] != "texto":
            continue
        ws = wbs[h["planilha"]][h["aba"]]
        full = str(ws.cell(int(re.sub(r"\D", "", h["celula"])), 2).value or "")
        m = re.search(r"Matr[íi]cula\s*n?º?\s*([\d.]+\d)", full)
        if m:
            obs_by_key.setdefault(f"mat. {m.group(1)}", []).append(
                f"{h['planilha']}, {h['aba']} {h['celula']}: {h['texto']}"
            )

    decisions = dec.get("decisoes", {})
    destinos = dec.get("destinos", {})

    def decision_cols(rid):
        d = decisions.get(rid)
        if not d:
            return "Pendente.", None, ""
        return (
            d["decisao"],
            d.get("valor_adotado"),
            f"{d.get('motivo', '')} (aprovado em {d.get('data', '?')})",
        )

    # --- rows ----------------------------------------------------------------------------
    rows = []

    def add_row(rid, it25, it18, sei, how, split18=None):
        base = it25 or it18
        desc = base["descricao"]
        split_state = "aprovado" if "— aprovado" in how else "proposto"
        keys = sorted((it25 or {}).get("keys", set()) | (it18 or {}).get("keys", set()))
        k25 = f"{it25['empresa']}-{item_no(it25['item'])}" if it25 else None
        raf = rafael.get(k25) if it25 else None
        fx = faixas.get(k25) if it25 else None
        g25 = it25["valor"] if it25 else None
        rv = raf["valor"] if raf else None
        diff = (rv - g25) if (rv is not None and g25 is not None) else None
        pct = (diff / g25) if (diff is not None and g25) else None
        sold18 = bool(it18 and it18["situacao_no_doc"])
        other = next((OTHER_OWNER[k] for k in keys if k in OTHER_OWNER), None)
        if it25 and raf:
            sit = "Concorda" if abs(diff) < 0.005 else "Discorda"
        elif it25:
            sit = "Só grupo"
        elif other:
            sit = other["situacao"]
        elif sold18:
            sit = "Vendido antes de 2018 (crédito a receber em 2018)"
        else:
            sit = "Só em 2018 (ausente em 2025)"
        # Rafael's criterion: say when his balance value falls outside his own range.
        crit = fx["criterio"] if fx else ""
        if crit.startswith("concordância expressa"):
            # Approved wording (packet 1): Rafael's agreement is conditioned (PDF Rádio p. 4–5).
            crit = "concordância condicionada à conferência de titularidade, existência e regularidade (PDF p. 4–5)"
        if (
            fx
            and fx.get("faixa_min") is not None
            and rv is not None
            and not (fx["faixa_min"] <= rv <= fx["faixa_max"])
        ):
            crit = (
                "manteve o valor do grupo, "
                if g25 is not None and abs(rv - g25) < 0.005
                else ""
            ) + "valor do balanço fora da faixa principal (ver outras referências)"
        fx_txt = ""
        if fx and fx.get("faixa_min") is not None:
            fx_txt = f"{brl(fx['faixa_min'], 0)} a {brl(fx['faixa_max'], 0)} (p. {fx['pagina_faixa']})"
        other_refs = ""
        if fx and fx.get("outras_referencias"):
            vals = []
            for ln in fx["outras_referencias"].split("\n"):
                m = re.match(r"p\. (\d+): R\$ ([\d.,]+)", ln)
                if m:
                    vals.append(f"R$ {m.group(2)} (p. {m.group(1)})")
            other_refs = "; ".join(dict.fromkeys(vals))
        # destination 2018 -> 2025
        dest = ""
        if it18:
            if rid in destinos:
                dest = destinos[rid]
            elif rid in DEST_FACTS:
                dest = DEST_FACTS[rid]
            elif other:
                dest = other["destino"]
            elif it25 and "PROPOSTO" in how:
                dest = "Permanece, se aprovado o par proposto sem chave (ver coluna 'Par 2018 × 2025')"
            elif it25 and it25["empresa"] != it18["empresa"]:
                dest = (
                    f"Mesma placa em {NAMES[it18['empresa']]} 2018 e {NAMES[it25['empresa']]} 2025 (identidade "
                    "confirmada); forma, data e contrapartida da transferência sem documento — pendente (P4)"
                )
            elif it25:
                dest = f"Permanece ({it25['id']})"
            else:
                dest = (
                    DEST_PREFIX.get(rid, "") + "Sem destino documentado — pendente (P4)"
                )
        elif split18:
            dest = f"Desmembrado do item {split18} de 2018 (par {split_state})"
        elif it25:
            dest = "Entrada depois de 2018 (sem par em 2018)"
        decisao, adot, motivo = decision_cols(rid)
        obs = "; ".join(o for k in keys for o in obs_by_key.get(k, []))
        p6 = sorted(
            {PART6_KEY[k] for k in keys if k in PART6_KEY}
            | set(PART6_ID.get(rid, []))
            | set(PART6_ID.get(it18["id"], []) if it18 else [])
        )
        alerts = [
            a for a in ((it18 or {}).get("alerta"), (it25 or {}).get("alerta")) if a
        ]
        alert_txt = ""
        if alerts:
            src = it18 if (it18 or {}).get("alerta") else it25
            cited = "; ".join(brl(float(v.replace(",", ""))) for v in src["valor_descricao"].split(";") if v.strip())
            alert_txt = f"{alerts[0]}: texto cita {cited}, coluna {brl(src['valor'])}"
        if it18:
            ref18 = f"{FILE18[it18['empresa']]}, tab. {it18['tabela']}, item {item_no(it18['item'])}"
            if other:
                ref18 += "; " + other["ref18"]
            exists18 = (
                "Não como bem: já vendido (crédito a receber)" if sold18 else "Sim"
            )
        elif split18:
            ref18 = f"{FILE18[items['2018'][split18]['empresa']]}, item {split18[-2:]} (dentro do item, par {split_state})"
            exists18 = f"Sim, dentro de outro item (desmembramento, par {split_state})"
        else:
            ref18, exists18 = "", "Não"
        rows.append(
            {
                "ID": rid,
                "Empresa (2025)": NAMES[it25["empresa"]]
                if it25
                else ("11 SEI" if sei else "—"),
                "Empresa (2018)": NAMES[it18["empresa"]]
                if it18
                else (NAMES[split18[:2]] if split18 else "—"),
                "Bloco": block_of(desc),
                "Descrição": short(desc),
                "Chave": "; ".join(keys) or "(sem chave de registro)",
                "Existe em 2018?": exists18,
                "Ref. 2018": ref18,
                "Existe em 2025 (grupo)?": "Sim"
                if it25
                else ("Sim, em outra empresa" if sei else "Não"),
                "Ref. 2025": f"Inventário 2025 doc {it25['empresa']}, item {item_no(it25['item'])}"
                if it25
                else (sei["ref"] if sei else ""),
                "Existe no Estudo 2 (Rafael)?": "Sim" if raf else "Não",
                "Ref. Rafael": f"PDF {it25['empresa']}, balanço executado por Valter, item {raf['item']}, p. {raf['pagina_valor']}"
                if raf
                else "",
                "Valor 2018": it18["valor"] if it18 else None,
                "Valor grupo 2025": g25,
                "Valor Rafael": rv,
                "Diferença R$": diff,
                "Diferença %": pct,
                "Critério de Rafael": crit,
                "Faixa de Rafael (C1)": fx_txt,
                "Outras referências de Rafael": other_refs,
                "Lista de Rafael (C0)?": "Sim" if set(keys) & C0_KEYS else "Não",
                "Situação": sit,
                "Par 2018 × 2025": how,
                "Alerta da fonte": alert_txt,
                "Decisão (P5)": decisao,
                "Valor adotado": adot,
                "Motivo da decisão": motivo,
                "Destino 2018→2025": dest,
                "Registro de data-base": REG_DATA_BASE_ID.get(rid)
                or next(
                    (REG_DATA_BASE[k] for k in keys if k in REG_DATA_BASE),
                    "Nenhum ajuste registrado",
                ),
                "Observação do Luís": obs,
                "Situação da observação": (OBS_DONE.get(rid) or "Pendente") if obs else "",
                "Parte 6": ", ".join(p6),
            }
        )

    done25 = set()
    for i18, it18 in items["2018"].items():
        p = pair18.get(i18)
        if p and isinstance(p[0], str):
            done25.add(p[0])
            add_row(p[0], items["2025"][p[0]], it18, None, p[1])
        elif p:
            add_row(i18, None, it18, p[0], p[1])
        else:
            add_row(i18, None, it18, None, "sem par")
    for i25, it25 in items["2025"].items():
        if i25 in done25:
            continue
        if i25 in split_of:
            add_row(i25, it25, None, None, split_of[i25][1], split18=split_of[i25][0])
        else:
            add_row(i25, it25, None, None, "sem par em 2018")
    # 2018 totals declared without any item (Associação): a row of its own, destination pending.
    for code in inv:
        d18 = inv[code]["2018"]
        if d18["declared"] and not any(i["valor"] for i in d18["items"]):
            rid = f"{code}-2018-TOTAL"
            rows.append(
                {k: "" for k in rows[0]}
                | {
                    "ID": rid,
                    "Empresa (2025)": "—",
                    "Empresa (2018)": NAMES[code],
                    "Bloco": "Outros",
                    "Descrição": "Total declarado em 2018 sem nenhum item preenchido (5 linhas pontilhadas)",
                    "Chave": "(sem chave de registro)",
                    "Existe em 2018?": "Sim (só o total)",
                    "Ref. 2018": f"{FILE18[code]}, linha de total; doc 27",
                    "Existe em 2025 (grupo)?": "Não",
                    "Existe no Estudo 2 (Rafael)?": "Não",
                    "Valor 2018": d18["declared"],
                    "Valor grupo 2025": None,
                    "Valor Rafael": None,
                    "Diferença R$": None,
                    "Diferença %": None,
                    "Valor adotado": None,
                    "Lista de Rafael (C0)?": "Não",
                    "Situação": "Só em 2018 (ausente em 2025)",
                    "Par 2018 × 2025": "sem par",
                    "Decisão (P5)": decision_cols(rid)[0],
                    "Motivo da decisão": decision_cols(rid)[2],
                    "Destino 2018→2025": destinos.get(
                        rid,
                        "Sem destino documentado — pendente (P4). Comparar com o capital "
                        "social da Associação em 2025 (R$ 10.000).",
                    ),
                    "Registro de data-base": "Nenhum ajuste registrado",
                }
            )
    order = {"07": 0, "04": 1, "03": 2, "09": 3, "10": 4}
    rows.sort(
        key=lambda r: (
            order.get(r["ID"][:2], 9),
            r["Existe em 2025 (grupo)?"] != "Sim",
            r["ID"],
        )
    )

    # --- notes block: credits, debts, lawsuits outside the asset table -----------------
    notes = []
    for n in data["note_values"]:
        if n["versao"] == "2025 antigo":
            continue
        nid = f"{n['empresa']}-{n['versao'][:4]}-N{n['paragrafo']:02d}"
        vals = [
            brl(float(v.replace(",", "")))
            for v in n["valores"].split("|")[0].split(";")
            if v.strip()
        ]
        pct = n["valores"].split("|")[1].strip() if "|" in n["valores"] else ""
        p6 = ", ".join(
            dict.fromkeys(c for pat, c in NOTE_PART6 if re.search(pat, n["texto"]))
        )
        notes.append(
            {
                "ID": nid,
                "Empresa": NAMES[n["empresa"]],
                "Versão": n["versao"],
                "Parágrafo": n["paragrafo"],
                "Valores": "; ".join(vals + ([pct] if pct else [])),
                "Natureza": n["natureza"],
                "Texto": n["texto"],
                "Situação em 31.12.2024 / destino": destinos.get(nid, "Pendente"),
                "Parte 6": p6,
                "Registro de data-base": "S4-DB1"
                if "venda de quotas" in n["natureza"]
                else ("S4-DB2" if "IXO 6000" in n["texto"] else ""),
            }
        )
    for i18, it in items["2018"].items():
        if re.search(r"financiado", it["descricao"], re.IGNORECASE):
            notes.append(
                {
                    "ID": f"{i18}-FIN",
                    "Empresa": NAMES[it["empresa"]],
                    "Versão": "2018",
                    "Parágrafo": f"item {i18[-2:]}",
                    "Valores": "",
                    "Natureza": "dívida (bem financiado)",
                    "Texto": it["descricao"]
                    + " — doc 25 par. 227: 18 parcelas de R$ 8.000 à GUIBOR 'do ônibus atual'.",
                    "Situação em 31.12.2024 / destino": destinos.get(
                        f"{i18}-FIN", "Pendente"
                    ),
                    "Parte 6": "",
                    "Registro de data-base": "",
                }
            )
    notes += [{**n, "Registro de data-base": ""} for n in NOTES_EXTRA]
    notes.sort(key=lambda n: n["ID"])

    # --- per-company summary with checks ----------------------------------------------
    summary = []
    for code in ("07", "04", "03", "09", "10"):
        name = NAMES[code]
        r25 = [r for r in rows if r["Empresa (2025)"] == name]
        r18 = [
            r
            for r in rows
            if r["Empresa (2018)"] == name and r["Valor 2018"] not in (None, "")
        ]
        g = sum(r["Valor grupo 2025"] or 0 for r in r25)
        rv = sum(r["Valor Rafael"] or 0 for r in r25)
        decl25 = inv[code]["2025"]["declared"]
        raf_bal = next(
            (
                t["valor"]
                for t in inv[code]["raf"]["totals"]
                if t["secao"] == "balanco_valter"
            ),
            None,
        )
        v18 = sum(r["Valor 2018"] for r in r18)
        sold = sum(
            r["Valor 2018"] for r in r18 if r["Situação"].startswith("Vendido antes")
        )
        other = sum(
            r["Valor 2018"]
            for r in r18
            if r["Situação"] == "Bem de outra empresa listado em 2018"
        )
        d18 = inv[code]["2018"]["declared"]
        doc27, doc00 = data["doc27"].get(code), data["doc00"].get(code)
        cnt = {}
        # A row belongs to its 2025 company; rows gone by 2025 to their 2018 company (no double count).
        ours = set(NAMES.values())
        for r in r25 + [
            x
            for x in rows
            if x["Empresa (2018)"] == name and x["Empresa (2025)"] not in ours
        ]:
            cnt[r["Situação"]] = cnt.get(r["Situação"], 0) + 1
        summary.append(
            {
                "Empresa": name,
                "Linhas (bens 2018 ou 2025)": sum(cnt.values()),
                "Por situação": "; ".join(f"{k}: {v}" for k, v in sorted(cnt.items())),
                "Soma grupo 2025": g,
                "Total declarado 2025": decl25,
                "Doc 00 (2025)": doc00,
                "Confere 2025 (inventário e doc 00)?": "sim"
                if decl25 is not None
                and abs(g - decl25) < 0.005
                and doc00 is not None
                and abs(g - doc00) < 0.005
                else "NÃO",
                "Soma Rafael": rv if raf_bal is not None else None,
                "Total Rafael (balanço de Valter)": raf_bal,
                "Confere Rafael?": ("sim" if abs(rv - raf_bal) < 0.005 else "NÃO")
                if raf_bal is not None
                else "sem balanço próprio de Rafael (só repete o grupo)",
                "Soma 2018 (bruta)": v18,
                "(−) vendidos a prazo antes de 2018": sold,
                "Soma 2018 líquida": v18 - sold,
                "Total declarado 2018": d18,
                "Doc 27 (2018)": doc27,
                "Confere 2018 (inventário e doc 27)?": "sim"
                if doc27 is not None
                and abs(v18 - sold - doc27) < 0.005
                and (d18 is None or abs(v18 - sold - d18) < 0.005)
                else "NÃO",
                "2018 sem bens de outra empresa": v18 - sold - other,
                "Itens com decisão": f"{sum(1 for r in r25 if r['Decisão (P5)'] != 'Pendente.')} de {len(r25)}",
                "Soma adotada": sum(r["Valor adotado"] or 0 for r in r25),
                "Base da quota": QUOTA_BASE.get(code, (1 / 6, "1/6 (P1)"))[1],
                "Quota (soma adotada × base)": round(sum(r["Valor adotado"] or 0 for r in r25) * QUOTA_BASE.get(code, (1 / 6, ""))[0], 2),
            }
        )

    # --- workbook ---------------------------------------------------------------------
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    hdr = PatternFill("solid", fgColor="DDEBF7")
    fills = {
        "Concorda": "E2EFDA",
        "Discorda": "FCE4D6",
        "Só grupo": "FFF2CC",
        "Só em 2018 (ausente em 2025)": "EDEDED",
        "Bem de outra empresa listado em 2018": "D9E1F2",
        "Vendido antes de 2018 (crédito a receber em 2018)": "EDEDED",
    }

    def sheet(title, dicts, widths):
        ws = wb.create_sheet(title)
        cols = list(dicts[0].keys())
        ws.append(cols)
        for c in ws[1]:
            c.font = bold
            c.fill = hdr
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for d in dicts:
            ws.append([d[c] if d[c] != "" else None for c in cols])
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    head = ws.cell(1, c.column).value
                    c.number_format = (
                        "0.0%"
                        if head == "Diferença %"
                        else (
                            "0"
                            if head in ("Parágrafo", "Linhas (bens 2018 ou 2025)")
                            else "#,##0.00"
                        )
                    )
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    wb.remove(wb.active)
    leg = [
        {
            "Campo": "Fonte dos números",
            "Conteúdo": f"_extracao/fase0-extracao.json (md5 {md5(EXTRACT)}), gerado por _scripts/fase0_extracao.py a partir das fontes; decisoes.json (md5 {md5(DECISIONS)}). Planilha gerada por _scripts/build_mestre.py; não editar à mão.",
        },
        {
            "Campo": "Pares 2018 × 2025",
            "Conteúdo": "Só por matrícula, placa, chassi ou nota fiscal (lição 3). Itens sem chave: par em decisoes.json, marcado 'PROPOSTO' até aprovação. O item 02 de 2018 da Rádio (equipamentos, móveis e canal) foi desmembrado nos itens 02 e 03 de 2025 e o item 03 (antena) corresponde ao 04 — pares aprovados em 28/10/2026.",
        },
        {
            "Campo": "Valor Rafael",
            "Conteúdo": "Balanço executado por Valter (seção final de cada PDF). A 'avaliação executada por Valter' repete o valor do grupo e traz a faixa de mercado (coluna C1) e outras referências (FIPE, anúncios). 03, 09 e 10 não têm balanço próprio de Rafael.",
        },
        {
            "Campo": "Situação",
            "Conteúdo": "Concorda / Discorda (grupo × Rafael); Só grupo; Só em 2018 (ausente em 2025); Vendido antes de 2018 (crédito a receber em 2018); Bem de outra empresa listado em 2018 (terreno 3.751, declarado como da SEI).",
        },
        {
            "Campo": "C0",
            "Conteúdo": "Bem da lista de Rafael (só a Toyota, IHW 0285) recebe o menor valor; os demais, o maior (P10). Decisões só depois de aprovadas: coluna Decisão = 'Pendente.' até lá.",
        },
        {
            "Campo": "Registro de data-base",
            "Conteúdo": "'Nenhum ajuste registrado' = nenhum fato fora de 31.12.2024 mudou o item. Texto = ponto a verificar; o registro formal fica em registro-data-base.md desta pasta.",
        },
        {
            "Campo": "Aba Créditos e dívidas",
            "Conteúdo": "Valores, processos, ônus e relações intragrupo que estão no texto dos inventários, fora da tabela de bens. 'Parágrafo' = nº de ordem entre os parágrafos não vazios do .docx (fora das tabelas).",
        },
        {
            "Campo": "Empresas baixadas",
            "Conteúdo": "03, 09 e 10 não têm bens em 2025 (linhas '-.-' R$ 0,00 não entram). Associação 2018: total de R$ 10.000 sem itens (linha 03-2018-TOTAL). Turismo 2018: sem itens e sem total; doc 27 dá R$ 0,00.",
        },
    ]
    sheet("Leia-me", leg, [22, 140])
    ws = sheet(
        "Mestre",
        rows,
        [
            12,
            16,
            16,
            16,
            44,
            24,
            14,
            30,
            10,
            30,
            10,
            30,
            14,
            14,
            14,
            14,
            9,
            24,
            24,
            30,
            9,
            18,
            40,
            30,
            12,
            14,
            30,
            50,
            26,
            40,
            12,
            12,
        ],
    )
    sit_col = list(rows[0].keys()).index("Situação")
    for row in ws.iter_rows(min_row=2):
        f = fills.get(row[sit_col].value)
        if f:
            row[sit_col].fill = PatternFill("solid", fgColor=f)
    sheet("Créditos e dívidas (notas)", notes, [16, 22, 10, 10, 30, 24, 100, 30, 12])
    sheet(
        "Resumo por empresa",
        summary,
        [22, 8, 50, 15, 15, 15, 12, 15, 15, 14, 15, 15, 15, 15, 15, 12, 15, 10],
    )
    wb.save(OUT)

    print("rows:", len(rows), "| notes:", len(notes))
    for s in summary:
        print(
            s["Empresa"],
            "|",
            s["Por situação"],
            "| 2025",
            s["Confere 2025 (inventário e doc 00)?"],
            "| Rafael",
            s["Confere Rafael?"],
            "| 2018",
            s["Confere 2018 (inventário e doc 27)?"],
            s["Soma 2018 (bruta)"],
            s["(−) vendidos a prazo antes de 2018"],
            s["Doc 27 (2018)"],
            "| sem outra empresa",
            s["2018 sem bens de outra empresa"],
        )
    for r in rows:
        print(
            r["ID"],
            r["Situação"],
            "|",
            r["Existe em 2018?"],
            "|",
            r["Valor 2018"],
            r["Valor grupo 2025"],
            r["Valor Rafael"],
            "|",
            r["Critério de Rafael"],
            "|",
            r["Parte 6"],
            "|",
            r["Alerta da fonte"][:50],
            "|",
            r["Destino 2018→2025"][:60],
        )
    print("wrote", OUT, md5(OUT))


if __name__ == "__main__":
    main()
