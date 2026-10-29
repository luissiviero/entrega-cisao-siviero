"""Search every source for the 2018 Hotel Araçá assets with no 2025 match.

Scans .docx (paragraphs and tables), .xlsx (all cells), .pdf (text layer),
.md and the reviewed OCR corpus of the replies to Dr. Pablo, in the Drive
folder Cisao_Siviero and in 01_fontes. Looks for the registrations, the
matrículas, the property codes and a few names tied to those assets, and
writes every hit with file, location and context to
_extracoes/busca-destinos-06.xlsx. It only finds text: each hit used in a
decision is read again at the source.

Usage: python busca_destinos_06.py
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl
import pymupdf

sys.path.insert(0, str(Path(__file__).resolve().parent))
import fase0_extracao_06 as f0

OUT = f0.BASE / "02_trabalho/empresas/06-hotel-araca/_extracoes/busca-destinos-06.xlsx"
ROOTS = [
    Path("G:/Meu Drive/02_Patrimonio_e_Sociedades/Cisao_Siviero"),
    f0.BASE / "01_fontes",
]
OCR = f0.BASE / "02_trabalho/ocr-respostas-docs/ocr.jsonl"

TERMS = {
    "insc 360.0019.004 (2018 item 20, Válter)": r"360[.\- ]?0019[.\- ]?004",
    "insc 360.0057.002 (2018 item 22, Raul)": r"360\.?0057\.?002",
    "insc 360.0057.001 (2018 item 23, Élcio)": r"360\.?0057\.?001",
    "insc 360.0074.030 (2018 item 24, Adélia)": r"360\.?0074\.?030",
    "insc 360.0074.031 (2018 item 25, Pedro)": r"360\.?0074\.?031",
    "mat 77.833 (item 24)": r"(?<![\d.])77\.?833(?![\d])",
    "mat 74.857 (item 25)": r"(?<![\d.])74\.?857(?![\d])",
    "cód 038999 / 039772 / 039773 / 039863 / 039864": r"0?3(?:8999|9772|9773|9863|9864)(?!\d)",
    "Busatto": r"(?i)busatto",
    "casa de veraneio": r"(?i)veraneio",
    "Rua da Divisa": r"(?i)rua da divisa",
    "Fiat Toro": r"(?i)\btoro\b",
    "Capão da Canoa": r"(?i)cap[ãa]o da canoa",
    "Hotel Araçá": r"(?i)hotel\s+ara[çc][áa]",
    "esquina": r"(?i)esquina",
    "casa ao lado / casa da esquina": r"(?i)casa (?:ao lado|da esquina)",
    "consignado": r"(?i)consignad",
    "Coletto": r"(?i)coletto",
    "dólar / cofre": r"(?i)d[óo]lar|cofre",
    "lote/quadra das inscrições (04/q19, 01-02/q57, 30-31/q74)": r"(?i)lote\s*(?:n[ºo°.]*\s*)?0?(?:4|1|2|30|31)\s*[-–,]?\s*(?:da\s+)?quadra\s*(?:n[ºo°.]*\s*)?0*(?:19|57|74)",
}


def md5_10(p: Path) -> str:
    return hashlib.md5(p.read_bytes()).hexdigest()[:10]


def units(path: Path):
    """Yield (location, text) units of one file."""
    suf = path.suffix.lower()
    if suf == ".docx":
        for n, (kind, b) in enumerate(f0.doc_blocks(path), 1):
            if kind == "p":
                yield f"bloco {n}", b
            else:
                for i, r in enumerate(b.rows, 1):
                    yield (
                        f"tabela (bloco {n}) linha {i}",
                        " || ".join(dict.fromkeys(c.text.strip() for c in r.cells)),
                    )
    elif suf == ".xlsx":
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        for ws in wb.worksheets:
            for i, row in enumerate(ws.iter_rows(values_only=True), 1):
                t = " | ".join(str(x) for x in row if x is not None)
                if t:
                    yield f"{ws.title}!{i}", t
    elif suf == ".pdf":
        for i, pg in enumerate(pymupdf.open(str(path)), 1):
            yield f"p. {i}", pg.get_text()
    elif suf == ".md":
        for i, line in enumerate(
            path.read_text(encoding="utf-8", errors="replace").splitlines(), 1
        ):
            yield f"linha {i}", line


def main():
    files = sorted(
        {
            p.resolve()
            for root in ROOTS
            for p in root.rglob("*")
            if p.suffix.lower() in (".docx", ".xlsx", ".pdf", ".md")
            and not p.name.startswith("~$")
        }
    )
    hits, errors, seen_md5 = [], [], {}
    for p in files:
        try:
            h = md5_10(p)
        except OSError as e:
            errors.append([str(p), f"leitura: {e}"])
            continue
        dup = seen_md5.get(h)
        seen_md5.setdefault(h, str(p))
        if dup:
            continue  # same content already scanned under another path
        try:
            for loc, text in units(p):
                flat = re.sub(r"\s+", " ", text)
                for label, pat in TERMS.items():
                    for m in re.finditer(pat, flat):
                        hits.append(
                            [
                                label,
                                str(p),
                                h,
                                loc,
                                flat[max(0, m.start() - 200) : m.end() + 200],
                            ]
                        )
        except Exception as e:  # noqa: BLE001 - record and continue
            errors.append([str(p), f"{type(e).__name__}: {e}"])
    for line in OCR.read_text(encoding="utf-8").splitlines():
        d = json.loads(line)
        t = re.sub(r"\s+", " ", d.get("texto") or d.get("text") or "")
        for label, pat in TERMS.items():
            for m in re.finditer(pat, t):
                hits.append(
                    [
                        label,
                        f"OCR: {d.get('arquivo') or d.get('fonte')}",
                        "ocr",
                        f"p. {d.get('pagina') or d.get('page')}",
                        t[max(0, m.start() - 200) : m.end() + 200],
                    ]
                )
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Achados"
    ws.append(["termo", "arquivo", "md5", "local", "contexto"])
    for r in hits:
        ws.append(r)
    we = wb.create_sheet("Erros")
    we.append(["arquivo", "erro"])
    for r in errors:
        we.append(r)
    wc = wb.create_sheet("Contagem")
    wc.append(["termo", "achados", "arquivos"])
    for label in TERMS:
        hs = [h for h in hits if h[0] == label]
        wc.append([label, len(hs), len({h[1] for h in hs})])
    wb.save(OUT)
    print(
        f"files {len(files)} unique {len(seen_md5)} errors {len(errors)} hits {len(hits)}"
    )
    for label in TERMS:
        hs = [h for h in hits if h[0] == label]
        print(f"  {label}: {len(hs)} in {len({h[1] for h in hs})} files")
    for e in errors:
        print("  ERR", e)


if __name__ == "__main__":
    main()
