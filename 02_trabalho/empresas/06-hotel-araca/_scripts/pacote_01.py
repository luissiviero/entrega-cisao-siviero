"""Packet 1 for company 06: block rules for the items where both sides agree.

Reads planilha-mestre-06.xlsx (read only: run build_mestre_06.py first when
the sources or decisions change), applies the rule criteria below from the
structured columns and writes pacotes/pacote-01-regras.md with counts, sums,
the item list of each rule, the caveats and a sample of rows. The md5 of the
master sheet used is written into the packet.

R1  same value; paired with 2018; no partner name in 2018; Rafael's express
    concurrence (column "Concordância de Rafael" = expressa).
R2  same value with an origin or ownership caveat: no 2018 pair, partner
    name in 2018, or concurrence by value only.
R3  rows with no value in the inventory nor in Rafael's total.
Items that disagree are left for the exception packets.

Usage: python pacote_01.py
"""

from __future__ import annotations

import hashlib
import re
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parents[1]
MESTRE = HERE / "planilha-mestre-06.xlsx"
OUT = HERE / "pacotes" / "pacote-01-regras.md"
SAMPLE_ITEMS = ["02", "05", "13", "18", "17", "22"]
TOL = 0.005


def br(v: float) -> str:
    return f"{v:,.2f}".replace(",", "#").replace(".", ",").replace("#", ".")


def md(s: str) -> str:
    """Keep table cells intact: the key uses '|' as separator."""
    return s.replace(" | ", "; ")


def item_no(ref):
    m = re.match(r"item (\d{2})", ref or "")
    return m.group(1) if m else None


def main():
    md5 = hashlib.md5(MESTRE.read_bytes()).hexdigest()[:10]
    wb = openpyxl.load_workbook(MESTRE, data_only=True)
    ws = wb["Mestre"]
    hdr = [c.value for c in ws[1]]
    rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    for r in rows:
        r["item"] = item_no(r["Ref. 2025"])
        r["notas"] = r["Notas de identificação"] or ""
    totais = {r[0]: r[1] for r in wb["Totais"].iter_rows(min_row=2, values_only=True)}
    gap = next(
        v for k, v in totais.items() if k.startswith("Diferença de soma de Rafael")
    )

    total_grupo = sum(
        r["Valor grupo 2025"] for r in rows if r["Valor grupo 2025"] is not None
    )
    conc = [r for r in rows if r["Situação"] == "Concorda"]
    disc = [r for r in rows if r["Situação"] == "Discorda"]
    r1 = [
        r
        for r in conc
        if r["Existe em 2018?"] == "Sim"
        and not r["Nome de sócio em 2018"]
        and (r["Concordância de Rafael"] or "").startswith("expressa")
    ]
    r2 = [r for r in conc if r not in r1]
    r3 = [r for r in rows if r["Situação"] in ("Só 2018", "Só Rafael", "Não abordado")]
    assert len(r1) + len(r2) == len(conc) and len(conc) + len(disc) + len(r3) == len(
        rows
    )
    for r in conc:
        assert abs(r["Valor grupo 2025"] - r["Valor Rafael"]) < TOL
    s1 = sum(r["Valor grupo 2025"] for r in r1)
    s2 = sum(r["Valor grupo 2025"] for r in r2)
    sd = sum(r["Valor grupo 2025"] for r in disc)
    assert abs(s1 + s2 + sd - total_grupo) < TOL
    gap_items = [r["item"] for r in conc if abs(r["Valor grupo 2025"] - gap) < TOL]

    def items(lst):
        return ", ".join(r["item"] for r in lst) or "nenhum"

    def where(lst, key):
        return [r for r in lst if key in r["notas"]]

    below = where(conc, "abaixo do piso da faixa por lote")
    below_stay = [r for r in below if r["Lista de Rafael (C0)"] != "Sim"]
    below_list = [r for r in below if r["Lista de Rafael (C0)"] == "Sim"]
    inside_m2 = [r for r in below if "dentro da faixa de Rafael por m²" in r["notas"]]
    floor = 380000.0
    gap_stay = sum(floor - r["Valor grupo 2025"] for r in below_stay)
    r2_17 = [r for r in r2 if r["Nome de sócio em 2018"]]
    r2_new = [r for r in r2 if r["Existe em 2018?"] == "Não"]
    lst_items = [r for r in conc if r["Lista de Rafael (C0)"] == "Sim"]
    casas = next(r for r in r3 if "Rua da Divisa" in r["Descrição"])

    lines = []
    w = lines.append
    w("# Hotel Araçá (empresa 06) — Pacote 1: regras em bloco")
    w("")
    w(
        f"Data: 28/10/2026. Gerado por `_scripts/pacote_01.py` a partir de `planilha-mestre-06.xlsx` (md5 `{md5}`). Revisado por revisor independente (`_diario/revisao-pacote-01.md`) antes desta versão. Nenhuma decisão está aplicada até a sua aprovação."
    )
    w("")
    w(
        f"Total do grupo 2025: R$ {br(total_grupo)} em {len(conc) + len(disc)} itens. Este pacote cobre os {len(conc)} itens em que grupo e Rafael dão o mesmo valor (R$ {br(s1 + s2)}). Os {len(disc)} itens que divergem ({items(disc)}; R$ {br(sd)} pelo grupo) vão para o pacote de exceções."
    )
    w("")
    w(
        "Nas regras abaixo aprova-se o **valor**, não a titularidade. A própria concordância de Rafael é condicionada à conferência de titularidade e registro (PDF p. 12–13 e 24)."
    )
    w("")
    w("## Regras propostas")
    w("")
    w("| Regra | Abrangência | Itens 2025 | Valor (grupo = Rafael) | Proposta |")
    w("|---|---|---|---|---|")
    w(
        f"| R1 | Mesmo valor nos dois lados. Bem identificado em 2018 pela inscrição imobiliária (e matrícula e código, quando o 2018 os traz). Coberto pela concordância expressa de Rafael aos itens 02 a 19 (PDF p. 12–13 e 16). | {len(r1)}: {items(r1)} | R$ {br(s1)} | Adotar o valor do grupo, igual ao de Rafael. |"
    )
    w(
        f'| R2 | Mesmo valor, com ressalva de origem ou titularidade. Item {items(r2_17)}: em 2018 era o lote "{r2_17[0]["Nome de sócio em 2018"] if r2_17 else ""}", sem valor. Itens {items(r2_new)}: não existem no inventário 2018, e Rafael concorda só por valor (p. 15 e 20); a concordância da p. 16 cita "itens 02 à 19". | {len(r2)}: {items(r2)} | R$ {br(s2)} | Adotar o valor do grupo, igual ao de Rafael. Titularidade (G11) e data e forma de aquisição (G10) seguem como perguntas. Se o bem não for da empresa em 31/12/2024, sai na Parte 4. |'
    )
    w(
        f"| R3 | Linhas sem valor no inventário nem no total de Rafael: {sum(1 for r in r3 if r['Situação'] == 'Só 2018')} lotes de 2018 sem par em 2025, {sum(1 for r in r3 if r['Situação'] == 'Só Rafael')} bens citados só por Rafael, {sum(1 for r in r3 if r['Situação'] == 'Não abordado')} não abordados. | — | — | Nenhum valor adotado neste pacote. Os lotes de 2018 vão para a Parte 4 (destino); os não abordados, para a Parte 5. As residências citadas por Rafael vão para o pacote de exceções (ver abaixo). Perguntas G1, G7, G8 e G9. |"
    )
    w("")
    w(
        f"Cobertura das regras de valor (R1 + R2): R$ {br(s1 + s2)} de R$ {br(total_grupo)}."
    )
    w("")
    w("## Ressalvas (não mudam valor)")
    w("")
    w(
        f"- **Lista de Rafael.** Os itens {items(lst_items)}, que viriam para nós, têm o mesmo valor nos dois lados. O C0 não tem o que escolher: não há valor menor."
    )
    w(
        f"- **Abaixo do piso de Rafael, bens que ficam na empresa** (itens {items(below_stay)}, R$ 360.000,00 cada). Rafael dá R$ 380.000 a 550.000 por lote de 360 m² nas quadras 37 e 38 (p. 10–11) e mesmo assim adota R$ 360.000. O C0 pediria o maior, mas sem reavaliação própria (P5) não há valor maior disponível. Por metro quadrado, os itens {items(inside_m2)} (307,80 m²) ficam a R$ 1.169,59/m², dentro da faixa dele de R$ 1.050 a 1.500/m² (p. 11). Os itens 04 a 10 dependem ainda da quadra real (G5). Se todos os {len(below_stay)} fossem ao piso de R$ 380.000, a empresa subiria R$ {br(gap_stay)}, ou R$ {br(gap_stay / 6)} em 1/6. Fica registrado como argumento, sem valor."
    )
    w(
        f"- **Abaixo do piso de Rafael, bens da lista** (itens {items(below_list)}, R$ 350.000,00 cada, R$ 972,22/m²). Estão abaixo do piso de Rafael por lote e por m². Como vêm para nós, isso favorece o C0: é argumento técnico (C1) para manter o valor igual dos dois lados."
    )
    w(
        f"- **Quadra no texto × inscrição.** Itens {items(where(conc, 'quadra 0071'))}: o texto diz quadra 37, a inscrição é da quadra 0071 (G5)."
    )
    w(
        f"- **Matrícula repetida.** Itens {items(where(conc, '73.482'))}: matrícula 73.482 nos dois lotes (G4)."
    )
    w(
        f"- **Comarca de Osório.** Item {items(where(conc, 'Osório'))}: matrícula 124.906 registrada em Osório."
    )
    w(
        f"- **Lotes no quarteirão do hotel.** Itens {items(where(conc, 'quarteirão do hotel'))} são os lotes 16, 15 e 17 da quadra 37, ao lado dos lotes do item 01. Se parte do hotel (piscinas, estacionamento) estiver sobre eles, o valor de lote avulso pode se sobrepor ao item 01. Junto com os lotes 13 e 14 (G1), formam o vão do quarteirão a unificar."
    )
    w(
        f"- **Diferença de soma de Rafael.** Os R$ {br(gap)} a mais no total dele não estão identificados. O valor é igual ao de cada um dos itens {', '.join(gap_items)}. A soma dos itens de Rafael confere com os valores destas regras. O tratamento fica para o pacote de exceções."
    )
    w(
        '- **Cabeçalho do PDF.** A p. 16 de Rafael abre com "GRUPO SIVIERO — SIVIERO EMPREENDIMENTOS IMOBILIÁRIOS LTDA." (erro de cópia). O conteúdo do balanço executado (p. 16–20) é o do Hotel Araçá, com as mesmas matrículas e inscrições do inventário 06.'
    )
    w("")
    w("## Amostra para conferir (6 linhas)")
    w("")
    w("| Item 2025 | Bem | Chave | 2018 | Grupo 2025 | Rafael | Regra | Notas |")
    w("|---|---|---|---|---|---|---|---|")
    by_item = {r["item"]: r for r in conc}
    for it in SAMPLE_ITEMS:
        r = by_item[it]
        rule = "R1" if r in r1 else "R2"
        if r["Valor 2018 (referência, P3)"]:
            v18 = f"item {item_no(r['Ref. 2018'])}, R$ {br(r['Valor 2018 (referência, P3)'])}"
        elif r["Existe em 2018?"] == "Não":
            v18 = "sem par"
        else:
            v18 = f"item {item_no(r['Ref. 2018'])}, sem valor ({r['Nome de sócio em 2018']})"
        w(
            f"| {it} | {r['Descrição']} | {md(r['Chave'])} | {v18} | R$ {br(r['Valor grupo 2025'])} | R$ {br(r['Valor Rafael'])} ({r['Ref. Rafael'].split(', ')[-1]}) | {rule} | {r['notas'].replace('; ', ' / ')} |"
        )
    w("")
    w(
        "Fontes da amostra: inventário 2025 (`01_fontes/inventario-2025/06. …docx`), inventário 2018 (Drive, doc 16), PDF de Rafael (`01_fontes/estudo-rafael/HOTEL ARAÇÁ LTDA - BELISE LTDA.pdf`)."
    )
    w("")
    w("## R3 em detalhe")
    w("")
    for r in r3:
        extra = ""
        if r is casas:
            extra = f' Rafael estima {r["Faixa de mercado de Rafael"]}, não soma ao total dele e escreve que eles "devem entrar no escopo negocial ou ter a sua destinação auferida e comprovada documentalmente" (p. 25–26). Podem ser as casas dos itens 22 e 23 de 2018 (Raul e Élcio). Vão ao pacote de exceções como "só Rafael", sem valor adotado até se saber se existem e de quem são.'
        pista = "; ".join(
            x
            for x in r["notas"].split("; ")
            if x.startswith(("lote vizinho", "matrícula ", "doc 6", "doc 25"))
        )
        w(
            f"- **{r['Situação']}:** {r['Descrição']}."
            + (f" Pistas: {pista}." if pista else "")
            + extra
        )
    w("")
    w("## Lista completa por regra")
    w("")
    for name, lst in (("R1", r1), ("R2", r2)):
        w(
            f"**{name}** ({len(lst)} itens, R$ {br(sum(r['Valor grupo 2025'] for r in lst))}):"
        )
        w("")
        for r in lst:
            w(
                f"- {r['item']}: {r['Descrição']} — R$ {br(r['Valor grupo 2025'])} — {md(r['Chave'])}"
                + (" — lista de Rafael" if r["Lista de Rafael (C0)"] == "Sim" else "")
            )
        w("")
    w("Perguntas citadas (G1 a G11): `perguntas-ao-grupo.md`, ainda não enviadas.")
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print(
        f"mestre md5 {md5}; R1 {len(r1)} = {s1:,.2f} [{items(r1)}]; R2 {len(r2)} = {s2:,.2f} [{items(r2)}]; R3 {len(r3)}; disc {len(disc)} = {sd:,.2f}"
    )
    print(
        f"below stay {items(below_stay)} gap {gap_stay:,.2f}; below list {items(below_list)}; inside m2 {items(inside_m2)}"
    )
    print("saved", OUT)


if __name__ == "__main__":
    main()
