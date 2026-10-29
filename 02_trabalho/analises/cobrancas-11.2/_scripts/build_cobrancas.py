"""Every charge in section 11.2 of doc 11 (SEI liabilities): was it already paid, how, and what to ask the group.

Charge values come from the GERAL tab (column E = group 2025). Payment evidence comes from file 17 (lots handed
over), the paid-bills ledger Jul/2023-Oct/2024, the group's answers to Dr. Pablo (file 1), the Rovílio quota contract
(file 10), the Susin terms (file 16), the Buriti memorandum, ACERTOS findings (S1-S13) and the Part 4 cross-check.
Ledger figures are recomputed here. Also lists the facts that favour the group (omitted liabilities), for honesty.

Writes cobrancas-11.2-quitacao.xlsx. Run: python -I _scripts/build_cobrancas.py
"""

from __future__ import annotations

import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")
HERE = Path(__file__).resolve().parent.parent
ROOT = HERE.parents[2]
GS = (
    ROOT
    / "01_fontes/planilhas-google/Comparativo Patrimonial Siviero empreendimentos e Rafael 2018 x 2025.xlsx"
)
LEDGER = ROOT / "01_fontes/outros/Contas_Pagas_2023-07_a_2024-10.xlsx"
OUT = HERE / "cobrancas-11.2-quitacao.xlsx"


def norm(s) -> str:
    s = (
        unicodedata.normalize("NFKD", str(s or ""))
        .encode("ascii", "ignore")
        .decode()
        .lower()
    )
    return re.sub(r"\s+", " ", s).strip()


def brl(v: float) -> str:
    return f"R$ {v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


# ---------------------------------------------------------------- ledger evidence (SEI only)
def ledger_evidence() -> dict:
    sheet = openpyxl.load_workbook(LEDGER, read_only=True, data_only=True)[
        "Consolidado"
    ]
    rows = list(sheet.iter_rows(values_only=True))[1:]
    sei = [r for r in rows if r[2] and "EMPREENDIMENTOS" in str(r[2])]

    def pick(pat, supplier_only=False):
        return [
            r
            for r in sei
            if re.search(pat, norm(r[3] if supplier_only else f"{r[3]} {r[8]}"))
        ]

    ev = {}
    m = pick(r"mercali")
    ev["mercali"] = (len(m), sum(r[7] for r in m))
    for who in ("ines", "samir"):
        p = pick(rf"{who} trevisan piccoli")
        last = max(r[0] for r in p)
        ev[f"{who}_principal_last"] = sum(r[6] for r in p if r[0] == last)
        ev[f"{who}_interest"] = sum(r[7] for r in p)
    esc = pick(r"pesa/cat|escavadeira hidraulica cat")
    ev["esc_cash"] = sum(
        r[7]
        for r in esc
        if "entrada" in norm(r[8]) or re.search(r"parc\. 0[12]/02", norm(r[8]))
    )
    ev["esc_parcelas"] = len([r for r in esc if "/48" in str(r[8])])
    bett = pick(r"paviment.*falavign")
    ev["bett_falavigna"] = (len(bett), sum(r[7] for r in bett))
    ev["fioreze"] = sum(r[7] for r in pick(r"lori fioreze", supplier_only=True))
    ev["zanrosso"] = sum(r[7] for r in pick(r"evandro mario zanrosso"))
    ev["boff"] = sum(r[7] for r in pick(r"dorval francisco boff"))
    ev["boff_2024_new"] = len(
        [r for r in pick(r"dorval francisco boff") if "/20" in str(r[8])]
    )
    ev["areas_rurais"] = sum(r[7] for r in pick(r"stuani|area nestor"))
    ev["months"] = len({r[0] for r in rows})
    ev["trator"] = len(pick(r"d6n"))
    return ev


# ---------------------------------------------------------------- charge values (GERAL, column E)
def charge_values() -> dict:
    ws = openpyxl.load_workbook(GS, data_only=True)["GERAL"]
    vals = {None: (None, None, 0.0)}
    for r in range(664, 744):
        a, b, e = ws[f"A{r}"].value, ws[f"B{r}"].value, ws[f"E{r}"].value
        vals[r] = (a, b, -float(e) if isinstance(e, (int, float)) else None)
    return vals


SIT_FILL = {
    "Duplicidade já abatida": "C6EFCE",
    "Quitada": "C6EFCE",
    "Paga em parte": "FFF2CC",
    "Em aberto, documentada": "DDEBF7",
    "Em aberto, sem comprovação": "F8CBAD",
    "Contingência (nada pago)": "EDEDED",
    "Posterior à data-base": "E4DFEC",
    "Sem valor cobrado": "FFFFFF",
}
NOT_DUP = "Não é duplicidade; multa contestável"


def build_rows(ev: dict, gv: dict) -> list[dict]:
    """One dict per charge. 'row' is the GERAL row the value is read from; 'abat' is what Luís already took out."""
    R = []

    def add(
        row,
        secao,
        item,
        assunto,
        natureza,
        situacao,
        como,
        fonte,
        dup,
        perg,
        decisao="",
        abat=0.0,
    ):
        R.append(
            {
                "row": row,
                "secao": secao,
                "item": item,
                "assunto": assunto,
                "valor": gv[row][2] or 0.0,
                "natureza": natureza,
                "situacao": situacao,
                "como": como,
                "fonte": fonte,
                "dup": dup,
                "perg": perg,
                "decisao": decisao,
                "abat": abat,
            }
        )

    multa = "Multa contratual estimada (20% do custo da obra), sem distrato"
    dec_multa = "Luís: multa só com distrato assinado; 11.2.1 = R$ 0 na proposta"

    def m(row):  # the whole multa is out of the proposal
        return gv[row][2]

    # 11.2.1 ------------------------------------------------------------------
    add(
        667,
        "11.2.1",
        "1",
        "Caminhos do Sol (Boff, Buriti)",
        "Sem multa (R$ 0)",
        "Sem valor cobrado",
        (
            f"Não há cobrança. A SEI pagou a Dorval Boff {brl(ev['boff'])} de jul/2023 a out/2024 ('contrato antecipação "
            f"valores'), com um contrato NOVO de 20 parcelas a partir de jul/2024 ({ev['boff_2024_new']} lançamentos). "
            "O grupo respondeu que 'nada mais foi investido desde 2022' (arq. 1, item 19). Crédito com os Boff: "
            "R$ 3.315.675,23 nominais (ACERTOS S7), fora do ativo."
        ),
        "Contas pagas; arq. 1, p. 30; ACERTOS S7",
        "—",
        "P05",
    )
    add(
        668,
        "11.2.1",
        "2",
        "Vale das Colinas (Piccoli, Gasparin, Zanrosso)",
        multa,
        "Contingência (nada pago)",
        (
            "Multa não paga. Já pago aos donos da área: Arlindo Piccoli R$ 300.000 (2018); dívida de Arlindo de R$ 270.000 "
            "paga com o lote San Marino 03/1342 (06/05/2024); advogada de Arlindo (nov/2023) e empréstimo a ele (out/2024); "
            "Zanrosso R$ 62.500 (2019), lote Alto das Graças 13/6372 (19/03/2024) e um contrato de R$ 150.000 de 2024 "
            f"('compra 01 lote conf. acordo'), com {brl(ev['zanrosso'])} pagos até out/2024. Depois da data-base: lote "
            "Falavigna 008/876 + R$ 165.000 a Fabiano Vanin (09/04/2026), 'para não incorrer em processo' contra a SEI e "
            "Arlindo Piccoli."
        ),
        "Arq. 17; contas pagas; ACERTOS S3",
        "Possível: pagamentos e indenizações já feitos não abatem a multa cheia",
        "P05",
        dec_multa,
        m(668),
    )
    add(
        669,
        "11.2.1",
        "3",
        "Vista da Colina (Zanandréa, Fioreze)",
        multa,
        "Contingência (nada pago)",
        (
            "Multa não paga. Já indenizado pelo atraso: 4 lotes a Fochesatto (22/02/2022, 'pelos lotes adquiridos da "
            "família Lori Fioreze pelo atraso da execução') e lote 05/6372 + R$ 139.500 a Angest (2019, 'pela não "
            "execução'). Fioreze R$ 148.000 + R$ 410.000 (ACERTOS S3). Contrato de R$ 170.000 com Lori Fioreze para COMPRA "
            f"de 1 ha dentro do loteamento ({brl(ev['fioreze'])} pagos até out/2024): é compra de área, não indenização. "
            "4 lotes do Termo B de Susin vêm de compras da família Susin junto a Fioreze, e 1 a Bovo/Zanandréa."
        ),
        "Arq. 17; arq. 16; contas pagas; ACERTOS S3",
        "Possível: atraso já indenizado com lotes e dinheiro, e a multa por não executar ainda cobrada",
        "P05",
        dec_multa,
        m(669),
    )
    add(
        670,
        "11.2.1",
        "4",
        "Pedrotti + Andreis + Lorandi",
        multa,
        "Contingência (nada pago)",
        (
            "Multa não paga. O 7º lote do Termo B de Susin vem de Pedrotti; crédito de R$ 270.000 a confirmar "
            "(ACERTOS S3). Em 2024 a SEI pagou a defesa de um 'processo Pedrotti venda lote'."
        ),
        "Arq. 16; contas pagas; ACERTOS S3",
        "Possível",
        "P05",
        dec_multa,
        m(670),
    )
    add(
        671,
        "11.2.1",
        "5",
        "Cecconello (parceira HECTA, 50%)",
        multa,
        "Contingência (nada pago)",
        (
            "Multa não paga. O próprio doc 11 diz que a parceira do loteamento é a HECTA (50%), não a SEI. A SEI só tem "
            "2,0 ha, comprados por R$ 600.000 em 2014."
        ),
        "Doc 11; diff antigo × atual; ACERTOS S4",
        "Não é duplicidade; multa de contrato da HECTA cobrada da SEI",
        "P06",
        dec_multa,
        m(671),
    )
    add(
        672,
        "11.2.1",
        "6",
        "Bonetto",
        multa,
        "Contingência (nada pago)",
        "Nada pago. Só a retificação de matrícula da área (jul/2023, R$ 68,50).",
        "Arq. 17; contas pagas",
        NOT_DUP,
        "P04",
        dec_multa,
        m(672),
    )
    add(
        673,
        "11.2.1",
        "7",
        "Pasqual",
        multa,
        "Contingência (nada pago)",
        (
            "Multa não paga. Em 19/11/2024, Darcy Pasqual recebeu o lote Altos da Julieta 03/1730 'com permuta de um lote "
            "no Loteamento Industrial Pasqual'. Em 2024 a SEI pagou topografia do 'Lot. Pascoal' (R$ 7.100) e "
            "reconhecimento de assinaturas de Darcy Pasqual. O grupo diz que não há 'nenhuma regularização' em andamento."
        ),
        "Arq. 17; contas pagas; arq. 1, p. 27",
        NOT_DUP + " (há atividade em 2024)",
        "P04",
        dec_multa,
        m(673),
    )
    add(
        674,
        "11.2.1",
        "8",
        "Industrial Dr. Letti",
        multa,
        "Contingência (nada pago)",
        (
            "Multa não paga. O Luís lista 'Dr. Letti' entre os loteamentos FINALIZADOS (regra de custos), e a SEI pagou em "
            "2023 o frete de 2 reservatórios para o 'Lot. Dr. Letti'. O doc 11 cobra o 'Industrial Dr. Letti' como "
            "contratado e sem projeto. Pode ser outro loteamento."
        ),
        "Contas pagas; regra de custos do Luís",
        NOT_DUP + " (pode ser o loteamento finalizado)",
        "P04",
        dec_multa,
        m(674),
    )
    add(
        675,
        "11.2.1",
        "9",
        "Sol Nascente III",
        multa,
        "Contingência (nada pago)",
        (
            "Multa não paga. Em 25/04/2024 o lote Sol Nascente II 11/C foi dado à Sustenseg 'pela prestação de serviços "
            "de topografia, licenciamento, aprovação e registro do Loteamento Sol Nascente III'. O doc 11 diz que não há "
            "projeto, e o grupo diz que não há regularização em andamento."
        ),
        "Arq. 17; arq. 1, p. 27",
        NOT_DUP + " (loteamento licenciado e pago com lote em 2024)",
        "P04",
        dec_multa,
        m(675),
    )

    # 11.2.2 ------------------------------------------------------------------
    lots = {  # GERAL row -> (label, kind d=left SEI stock in 2025 / c=third-party lot, facts, question)
        680: ("Dalzóchio 09/0910", "c", "Sem registro no arq. 17.", "P02"),
        681: ("Dalzóchio 03/1292", "c", "Sem registro no arq. 17.", "P02"),
        682: (
            "Dalzóchio 04/1292 (a planilha repete 09/0910)",
            "c",
            "Sem registro no arq. 17.",
            "P02",
        ),
        686: ("Jardim Itália 07/1741", "c", "Sem registro no arq. 17.", "P01"),
        687: ("Jardim Itália 09/1741", "c", "Sem registro no arq. 17.", "P01"),
        688: (
            "Jardim Itália 01/1747",
            "d",
            "Era da SEI em 31/12/2024; vendido a Erick de Almeida em 10/03/2025.",
            "",
        ),
        689: (
            "Jardim Itália 02/1747",
            "c",
            "Dado à Ambitec em 14/07/2021 em acordo de distrato (ACERTOS S5).",
            "P01",
        ),
        690: ("Jardim Itália 09/1747", "c", "Sem registro no arq. 17.", "P01"),
        694: (
            "Perini 11/1469",
            "c",
            "Dado aos Farina em 14/08/2019, no acordo judicial do Vicentino (ACERTOS S5).",
            "P01",
        ),
        695: ("Perini 12/1469", "c", "Sem registro no arq. 17.", "P03"),
        696: (
            "Perini 18/1469",
            "c",
            "Vendido à Efetiva em 10/10/2024 ('diferença de área constatada').",
            "P03",
        ),
        697: (
            "Perini 19/1469",
            "c",
            "Vendido à Efetiva em 10/10/2024 ('diferença de área constatada').",
            "P03",
        ),
        701: ("Primavera 05/1199", "c", "Sem registro no arq. 17.", "P01"),
        702: ("Primavera 06/1199", "c", "Sem registro no arq. 17.", "P01"),
        703: (
            "Primavera 10/1199",
            "d",
            "Era da SEI em 31/12/2024; dado à Safra em fev/2025 (arq. 1, p. 26).",
            "",
        ),
        704: (
            "Primavera 11/1199",
            "d",
            "Era da SEI em 31/12/2024; dado à Safra em 27/03/2025.",
            "",
        ),
        708: (
            "Vicentino 03/1301",
            "d",
            "Era da SEI em 31/12/2024; vendido em 19/09/2025.",
            "",
        ),
        709: ("Vicentino 04/1301", "c", "Sem registro no arq. 17.", "P01"),
        710: (
            "Vicentino 04/1302",
            "c",
            "Permutado com Alcides Basso Imóveis em 26/02/2014 (3 ônibus + R$ 552.000).",
            "P01",
        ),
        711: ("Vicentino 09/1302", "c", "Sem registro no arq. 17.", "P01"),
        712: ("Vicentino 10/1302", "c", "Sem registro no arq. 17.", "P01"),
        713: ("Vicentino 11/1302", "c", "Sem registro no arq. 17.", "P01"),
        714: ("Vicentino 01/1303", "c", "Sem registro no arq. 17.", "P01"),
        715: ("Vicentino 12/1303", "c", "Sem registro no arq. 17.", "P01"),
        716: ("Vicentino 14/1303", "c", "Sem registro no arq. 17.", "P01"),
        717: ("Vicentino 16/1303", "c", "Sem registro no arq. 17.", "P01"),
    }
    nature = "Indenização possível a terceiro (APP, Marco Regulatório de Farroupilha)"
    for row, (lab, kind, how, q) in lots.items():
        if kind == "d":
            add(
                row,
                "11.2.2",
                "lote",
                lab,
                nature,
                "Duplicidade já abatida",
                how
                + " O lote saiu do estoque e o mesmo valor é descontado como indenização a terceiro.",
                "Arq. 17; ACERTOS S1; revisão D4",
                "Confirmada",
                "",
                "Luís: abatido (R$ 1.887.000 nos 4 lotes)",
                gv[row][2],
            )
            continue
        if row in (689, 694):
            dup = "Possível: o acordo de 2019/2021 pode ter dado quitação"
        elif row in (680, 681, 682, 695, 696, 697):
            dup = "Possível: mesma área do processo do 11.2.4-2/3"
        else:
            dup = "A verificar (dono e data da compra)"
        add(
            row,
            "11.2.2",
            "lote",
            lab,
            nature,
            "Contingência (nada pago)",
            how + " Nenhuma reclamação de dono nas fontes.",
            "Arq. 17; arq. 1, p. 28",
            dup,
            q,
            "Comentário do Luís: não se desconta indenização só possível (pendente, pacote ambiental)",
        )

    # 11.2.3 ------------------------------------------------------------------
    add(
        720,
        "11.2.3",
        "1",
        "Colina do Retiro II: obras finais",
        "Estimativa de obra a concluir",
        "Paga em parte",
        (
            "O lote Carnevalli 07/A foi dado à CVC em 01/08/2024 'como forma de pagamento de pavimentação no Loteamento "
            "Colina do Retiro II'. Contas pagas: taxa CORSAN de aprovação do projeto de água (nov/2023, R$ 6.229,95) e "
            "projeto de arborização (mar/2024, R$ 1.300), que são componentes do orçamento; contrato CVC de R$ 54.000 em 4 "
            "cheques (abr-jul/2024), 'PAVIMENTACAO LO...', com o nome do loteamento truncado. Lotes antigos (2006 e 2009) "
            "foram dados a Frainer 'em troca de pavimentação'. O doc 11 não diz o que os R$ 3,8 mi já descontam (eram "
            "R$ 3,5 mi na versão antiga)."
        ),
        "Arq. 17; contas pagas; diff antigo × atual",
        "Confirmada só para o lote CVC; o resto depende da composição dos R$ 3,8 mi",
        "P07",
        "Luís: abatido o lote CVC (R$ 205.667)",
        205666.67,
    )

    # 11.2.4 ------------------------------------------------------------------
    add(
        722,
        "11.2.4",
        "1",
        "Rovílio Siviero: 3 lotes do Morada dos Ipês II 'à pagar'",
        "Preço de quotas (contrato de 28/08/2025) + comissão",
        "Em aberto, documentada",
        (
            "Contrato de quotas (arq. 10, cl. 3ª): já pagos o 13/02 (24/03/2023), o 10/02 (fev/2025), o 07/04 (comissão de "
            "Rovílio pela intermediação do Morada I e II; no arq. 17, 'Structura... COMISSÃO - Rovílio', 30/08/2021), os "
            "lotes da N. S. da Saúde (12/237 vendido em 01/11/2012; 15/252 trocado por um lote do Dalla Costa em "
            "25/10/2013), R$ 110.000 (apartamento de Rovílio usado no muro do condomínio, mai/2022) e R$ 25.500 de "
            "impostos. Falta entregar 3 lotes do Morada II: 2 de preço de quota (R$ 1.100.000) e 1 de comissão de "
            "intermediação (R$ 550.000)."
        ),
        "Arq. 10, p. 4-6; arq. 17",
        "Nenhuma (o 07/04 é a comissão, dentro do preço)",
        "P10",
        "Luís: Morada II prova valor de loteamento não finalizado (tarefa pendente)",
    )
    add(
        723,
        "11.2.4",
        "2",
        "Infração ambiental do Perini",
        "Contingência judicial (execução do MP de dez/2024, valor da causa R$ 4.663.325,88)",
        "Contingência (nada pago)",
        (
            "Nada pago. Embargos à execução: ilegitimidade (loteamento entregue ao Município em 08/05/2017), prescrição e "
            "não descumprimento do TAC. Perícia deferida, sem perito nomeado. O grupo cobra R$ 5,0 mi; o valor da causa é "
            "R$ 4.663.325,88."
        ),
        "Arq. 1, p. 23-24",
        "Possível: mesma área dos 4 lotes do Perini no 11.2.2",
        "P03",
    )
    add(
        724,
        "11.2.4",
        "3",
        "Fonte de água do Dalzóchio (MP, desde 2006)",
        "Contingência judicial",
        "Contingência (nada pago)",
        "Nada pago. Sem acordo. O valor subiu de R$ 500 mil (versão antiga) para R$ 2,2 mi.",
        "Doc 11; diff antigo × atual",
        "Possível: mesma área dos 3 lotes do Dalzóchio no 11.2.2",
        "P02",
    )
    add(
        725,
        "11.2.4",
        "4",
        "APP do Jardim Itália (canalização)",
        "Sem valor",
        "Sem valor cobrado",
        "O MP arquivou o inquérito civil da canalização em jul/2026 (não trata da APP dos lotes do 11.2.2).",
        "Arq. 1, p. 24",
        "—",
        "",
    )
    add(
        726,
        "11.2.4",
        "5",
        "ETE do San Marino (terreno R$ 230 mil + obra R$ 300 mil)",
        "Obrigação com o Município, sem projeto",
        "Em aberto, sem comprovação",
        "Nada pago. O doc 11 não diz qual lote seria doado; o Luís já anotou que não há lote de R$ 230 mil na relação.",
        "Doc 11; comentário do Luís",
        "Possível: se o lote doado já saiu do estoque, o terreno é descontado duas vezes",
        "P11",
    )
    add(
        727,
        "11.2.4",
        "6",
        "Lago Negro: PRAD, custas e honorários",
        "Estimativa",
        "Em aberto, sem comprovação",
        (
            "As indenizações foram pagas: lote 5/5 a Adriana Wasen (acordo de ago/2018, arq. 17 em 24/04/2019), lote 10/9 "
            "ao promotor Flach (abr/2019), lote 14/5 a Candiago (28/08/2025) e R$ 580.000 a Lucas Zucoloto (set/2025). Em "
            "jun/2026 o processo foi encerrado com gravame perpétuo de preservação nos lotes 15 a 21 da quadra 11B."
        ),
        "Arq. 1, p. 24-25; arq. 17",
        "Possível: o acordo de jun/2026 pode ter dispensado o PRAD",
        "P14",
    )
    add(
        728,
        "11.2.4",
        "7",
        "Família Acordi: atualização da multa + TAC/PRAD",
        "Estimativa",
        "Paga em parte",
        (
            "O próprio doc 11 diz que a multa principal (~R$ 80 mil) e a multa do Município (~R$ 70 mil) já foram pagas. O "
            "ACERTOS dá o TAC do MP como pago (R$ 88.457,52) e soma R$ 130.000 de acerto com o Município (S13)."
        ),
        "Doc 11; ACERTOS S13",
        "Possível: o ACERTOS dá como pago o que o doc 11 ainda cobra",
        "P13",
    )
    for row, txt in (
        (729, "Impostos na transmissão de imóveis dos sócios (5,93% a 7%)"),
        (731, "Impostos sobre imóveis vendidos e não transmitidos"),
        (732, "7% de impostos sobre vendas"),
        (733, "6% de comissão sobre vendas"),
    ):
        add(
            row,
            "11.2.4",
            str(gv[row][0]).replace(".0", ""),
            txt,
            "Percentual, sem valor",
            "Sem valor cobrado",
            "Não entra no total.",
            "Doc 11",
            "—",
            "P18",
        )
    add(
        730,
        "11.2.4",
        "9",
        "Regularização de áreas rurais compradas por contrato",
        "Estimativa",
        "Em aberto, sem comprovação",
        (
            "Não diz quais áreas. Nas contas pagas de jul/2023 a out/2024, só parecem ser de áreas o levantamento GEO "
            "M.61.874 (Stuani-Mast) e a topografia da 'área Nestor J. Colo...' (Topo Ambiental, 4 parcelas): "
            f"{brl(ev['areas_rurais'])}."
        ),
        "Doc 11; contas pagas",
        "A verificar",
        "P15",
    )
    add(
        734,
        "11.2.4",
        "13",
        "Celestino Rossi, arquiteto (3 lotes)",
        "Honorários; documento não apresentado",
        "Em aberto, sem comprovação",
        (
            "Nenhum lote entregue (arq. 17) nem pagamento (contas pagas). O doc 11 fala em 'acordos firmados'; o grupo diz "
            "que o compromisso é verbal (arq. 1, p. 27)."
        ),
        "Arq. 1; arq. 17; contas pagas",
        "A verificar: sem documento, não se sabe o que já foi pago",
        "P12",
    )
    add(
        735,
        "11.2.4",
        "14",
        "Solar Serra: placas solares do Hotel Araçá (1 lote)",
        "Dívida do Hotel Araçá",
        "Quitada",
        (
            "O lote Aditivo Medianeira 02/07 foi entregue à Solar Serra em 24/04/2024 'como forma de pagamento placas "
            "solares Hotel Araçá'. O lote continua no estoque de 2025 (R$ 165.000) e a dívida no passivo (R$ 150.000). "
            "Mário Henrique Bordignon (Solar Serra) assina o doc 11 com a família: parte relacionada."
        ),
        "Arq. 17; doc 11; revisão D5",
        "Não é duplicidade: estoque e passivo inflados. Corrigir os dois (−R$ 15 mil) e lançar crédito contra o Hotel Araçá",
        "",
        "Pendente: registrar crédito da SEI contra o Hotel Araçá",
    )
    add(
        736,
        "11.2.4",
        "15",
        "Alexandre Susin: comissão pelos contratos Buriti (+ 6 lotes do Termo B)",
        "Comissão sem documento escrito + lotes de acordo assinado",
        "Em aberto, sem comprovação",
        (
            "Termo A (13/07/2020): 17 lotes + R$ 200 mil quitam as intermediações até aquela data, inclusive com as famílias "
            "donas das áreas depois cedidas à Buriti (cl. 1.1); a quitação depende de cumprir o acordo (cl. 6) e faltam 2 "
            "lotes-equivalentes. Termo B (30/04/2020): 7 lotes, só o 12/1729 entregue; os '06 lotes' do doc 11 são os que "
            "faltam, dívida documentada (~R$ 600 mil no ACERTOS). Memorando Buriti: 14/04/2022, sem assinatura, DEPOIS do "
            "Termo A; a comissão 'pelos contratos com a Buriti' não está no Termo A, e a cl. 8 exige documento escrito. "
            "Nenhum pagamento a Susin nas contas pagas."
        ),
        "Arq. 16; memorando Buriti; arq. 17; ACERTOS S2/S9; revisão D1",
        "Não é cobrança dupla pela data: a comissão é cobrança sem documento; os 6 lotes são dívida real",
        "P09",
        "Luís: abatido como duplicidade (R$ 1.200.000); a base muda, ver P09",
        1200000.0,
    )
    add(
        737,
        "11.2.4",
        "19",
        "Reserva Falavigna: pavimentação, parque e galeria",
        "Estimativa de obra a concluir",
        "Paga em parte",
        (
            "Lotes já dados por obras do Falavigna: Alto das Graças 15/6371 à Bett (23/11/2022, pavimentação, contrato "
            "de R$ 465.000); 18/6373 à AM Cimento (08/11/2022, drenagem pluvial); 3 lotes à Florense (20/03/2025, aduelas "
            "da galeria) e 2 à Bett (10/07/2025, pavimentação), 'na condição de 50%' em espécie e 50% em permuta; 028/322 à "
            f"Florense (02/06/2026, tubos da galeria). Contas pagas: {brl(ev['bett_falavigna'][1])} à Bett em "
            f"{ev['bett_falavigna'][0]} lançamentos (jul/2023-mar/2024, parcelas até P.10/10). O valor subiu de R$ 950 mil "
            "(versão antiga) para R$ 2,0 mi."
        ),
        "Arq. 17; contas pagas; diff antigo × atual",
        "Comprovada só até o valor dos 5 lotes de 2025; o resto depende da composição dos R$ 2,0 mi",
        "P08",
        "Luís: abatido (R$ 2.000.000)",
        2000000.0,
    )

    # 11.2.5 ------------------------------------------------------------------
    n_m, s_m = ev["mercali"]
    loan = "Em aberto, documentada"
    add(
        739,
        "11.2.5",
        "16",
        "Luis Carlos Mercali: empréstimo (jun/2020)",
        "Empréstimo de pessoa física, 1% a.m.",
        loan,
        (
            f"Principal em aberto, juros em dia: {n_m} pagamentos em {ev['months']} meses do arquivo (falta set/2023), "
            f"{brl(s_m)}, sobre R$ 200.000."
        ),
        "Contas pagas",
        "Baixo",
        "P16",
    )
    add(
        740,
        "11.2.5",
        "17",
        "Inês Trevisan Piccoli: empréstimo",
        "Empréstimo de pessoa física, 1% a.m.",
        loan,
        (
            f"Principal em aberto, juros em dia sobre {brl(ev['ines_principal_last'])} em out/2024 (R$ 250.000 + "
            "R$ 100.000; a parte de R$ 100.000 aparece só a partir de ago/2023, e o doc 11 diz 'desde dez/2021'). Juros "
            f"pagos: {brl(ev['ines_interest'])}."
        ),
        "Contas pagas",
        "A verificar: ligação com o acerto de Arlindo Piccoli (Vale das Colinas)",
        "P16",
    )
    add(
        741,
        "11.2.5",
        "18",
        "Samir Trevisan Piccoli: empréstimo",
        "Empréstimo de pessoa física, 1% a.m.",
        loan,
        (
            f"Principal em aberto, juros em dia sobre {brl(ev['samir_principal_last'])} em out/2024 (5 partes; "
            f"R$ 150.000 novos em fev/2024). Juros pagos: {brl(ev['samir_interest'])}."
        ),
        "Contas pagas",
        "A verificar: ligação com o acerto de Arlindo Piccoli (Vale das Colinas)",
        "P16",
    )
    add(
        742,
        "11.2.5",
        "20",
        "Escavadeira CAT 320 (2024): saldo financiado",
        "Financiamento",
        loan,
        (
            f"Preço R$ 887.000: entrada e 2 parcelas à vista ({brl(ev['esc_cash'])}) + R$ 377.000 em 48 × R$ 10.534,91 "
            f"desde mai/2024 ({ev['esc_parcelas']} parcelas pagas até out/2024). Confere com as 33 restantes em jul/2025 "
            "do doc 11. Na data-base o saldo era maior (40 × R$ 10.534,91 = R$ 421.396,40); o doc 11 usa o de jul/2025."
        ),
        "Contas pagas",
        "Nenhuma",
        "",
    )
    add(
        743,
        "11.2.5",
        "21",
        "Trator de esteira D6N XL: saldo financiado",
        "Financiamento",
        "Posterior à data-base",
        (
            "O próprio doc 11 diz 'adquirido em setembro de 2025' (máquinas, item 14, R$ 450.000; 'ano 2014' lá e '2021' "
            f"nas observações). Nenhuma parcela nas contas pagas até out/2024 ({ev['trator']} lançamentos). Na data-base "
            "não existiam nem o trator nem a dívida."
        ),
        "Doc 11 (máquinas, item 14); contas pagas",
        "Não é duplicidade; sair da data-base: −R$ 320 mil de passivo e −R$ 450 mil de ativo (−R$ 130 mil líquidos)",
        "",
    )

    # announced but without value ---------------------------------------------
    add(
        None,
        "11.2.4",
        "22-24",
        "'Itens 1 ao 24' citados na tabela 50",
        "Sem valor",
        "Sem valor cobrado",
        "O doc 11 só tem os itens 1 a 21; a soma dos 21 dá os R$ 16.127.652,00 da tabela 50.",
        "Doc 11, T50.R1",
        "—",
        "P18",
    )
    add(
        None,
        "Final",
        "—",
        "Custos de regularização, legalização, cartório, registro e 'ônus de dissolução'",
        "Valores em branco",
        "Sem valor cobrado",
        "O parágrafo final diz que 'o valor real acima apontado não é líquido', com todos os valores pontilhados.",
        "Doc 11, observações finais",
        "—",
        "P18",
    )
    return R


QUESTIONS = [
    # id, items, question, already asked; the value at stake is computed from the rows that point to the question
    (
        "P01",
        "11.2.2 (16 lotes)",
        (
            "Para cada lote do 11.2.2 que continua cobrado: quem é o dono hoje, quando comprou e por quanto? Algum dono "
            "já pediu indenização (notificação, processo)? Em especial: algum lote do Vicentino é da família Farina, que "
            "já recebeu lotes do Dalzóchio e do Perini em 2019 num 'acordo judicial indenizatório do Industrial "
            "Vicentino', ou de Alcides Basso (permuta de 2014)? Os lotes Jardim Itália 02/1747 (Ambitec, 2021) e Perini "
            "11/1469 (Farina, 2019) foram dados em acordos com quitação: por que voltam como indenizáveis?"
        ),
        "Em parte: Dr. Pablo, dúvida 5; perguntas-sugeridas 4",
    ),
    (
        "P02",
        "11.2.4-3; 11.2.2 Dalzóchio",
        (
            "O processo da fonte de água do Dalzóchio e os 3 lotes do Dalzóchio no 11.2.2 são a mesma área? Se o processo "
            "terminar com indenização, ela já inclui os donos desses lotes? Por que o valor subiu de R$ 500 mil para "
            "R$ 2,2 mi entre as duas versões do doc 11?"
        ),
        "Em parte: Dr. Pablo, item 3 (cópia dos processos), sem resposta",
    ),
    (
        "P03",
        "11.2.4-2; 11.2.2 Perini",
        (
            "A execução do MP no Perini (valor da causa R$ 4,66 mi; o grupo cobra R$ 5,0 mi) e os 4 lotes do Perini no "
            "11.2.2 tratam da mesma área? Os lotes 18 e 19/1469 foram vendidos à Efetiva em out/2024, com o risco já "
            "conhecido: o comprador tem direito a indenização? Se a defesa (ilegitimidade, entrega ao Município em 2017) "
            "prevalecer, a cobrança cai inteira?"
        ),
        "Em parte: Dr. Pablo, dúvida 16",
    ),
    (
        "P04",
        "11.2.1 Dr. Letti, Sol Nascente III, Pasqual, Bonetto",
        (
            "O grupo diz que não há 'nenhuma regularização' em andamento nesses loteamentos (arq. 1, p. 27). Mas: o "
            "'Industrial Dr. Letti' é o mesmo 'Dr. Letti' que recebeu reservatórios em 2023 e que tratamos como "
            "finalizado? No Sol Nascente III, a Sustenseg recebeu um lote em abr/2024 pelo licenciamento, aprovação e "
            "registro: ele está aprovado? No Pasqual, houve topografia e laudo em 2024, e Darcy Pasqual recebeu um lote em "
            "nov/2024 em permuta por um lote do Industrial Pasqual: o loteamento está sendo executado? Se estão andando, "
            "por que há multa por não execução?"
        ),
        "Em parte: Dr. Pablo, dúvida 13 (contratos e motivo do distrato)",
    ),
    (
        "P05",
        "11.2.1 Vale das Colinas, Vista da Colina, Pedrotti; Caminhos do Sol",
        (
            "Os pagamentos já feitos aos donos e aos compradores dessas áreas abatem a multa? Vale das Colinas: R$ 300 mil "
            "a Arlindo Piccoli (2018), a dívida de R$ 270 mil paga com lote (2024), e a Zanrosso R$ 62,5 mil, 1 lote e um "
            "contrato de R$ 150 mil (2024). Vista da Colina: R$ 558 mil a Fioreze, 4 lotes a Fochesatto (2022) e lote + "
            "R$ 139,5 mil a Angest (2019) 'pelo atraso'. Pedrotti: R$ 270 mil. Caminhos do Sol: o grupo disse que 'nada "
            "mais foi investido desde 2022', mas a SEI pagou R$ 70 mil a Dorval Boff em 2023-24, com um contrato novo de "
            "20 parcelas em jul/2024. Qual é esse contrato e qual é o saldo do crédito com os Boff (R$ 3,3 mi nominais)?"
        ),
        "Sim: perguntas-sugeridas 5; Dr. Pablo, dúvida 11 (item 19, Boff)",
    ),
    (
        "P06",
        "11.2.1 Cecconello",
        "A parceira do Cecconello é a HECTA (50%). De quem é a HECTA? Por que a multa de um contrato da HECTA é descontada na SEI?",
        "Sim: perguntas-sugeridas 7",
    ),
    (
        "P07",
        "11.2.3 Colina do Retiro II",
        (
            "Qual é o orçamento dos R$ 3,8 mi, item por item? Quanto da pavimentação já foi paga (lote à CVC em ago/2024, "
            "lotes antigos a Frainer) e isso já está descontado? O contrato CVC de R$ 54 mil (4 cheques, abr-jul/2024) é "
            "deste loteamento? A elevatória, a rede de água e a arborização já foram pagas em parte?"
        ),
        "Sim: Dr. Pablo, dúvida 6 (sem resposta)",
    ),
    (
        "P08",
        "11.2.4-19 Falavigna",
        (
            "Qual é o orçamento dos R$ 2,0 mi? Os lotes já dados por obras do Falavigna (Bett e AM Cimento em 2022; "
            "Florense e Bett em 2025; Florense em jun/2026, tubos da galeria) e as parcelas do contrato da Bett estão dentro "
            "ou fora desse valor? Por que subiu de R$ 950 mil para R$ 2,0 mi?"
        ),
        "Sim: Dr. Pablo, dúvida 6 (sem resposta)",
    ),
    (
        "P09",
        "11.2.4-15 Susin",
        (
            "Qual é o documento escrito, exigido pela cl. 8 do Termo A, que dá a Susin comissão pelos contratos Buriti "
            "(memorando de 14/04/2022)? A comissão é sobre os contratos com as famílias donas das áreas, já quitados pela "
            "cl. 1.1 do Termo A, ou sobre o negócio novo com a Buriti? Os R$ 1,2 mi incluem os 6 lotes do Termo B ou são "
            "só a comissão? Os 2 lotes que faltam do Termo A foram entregues? Se as áreas da Buriti não serão executadas "
            "(por isso o 11.2.1 cobra multa), por que pagar comissão pelo fechamento?"
        ),
        "Sim: perguntas-sugeridas 1 e 2",
    ),
    (
        "P10",
        "11.2.4-1 Rovílio",
        (
            "Quais são os 3 lotes do Morada dos Ipês II prometidos a Rovílio, e qual é o valor desse condomínio, que não "
            "está no inventário? O lote de comissão (R$ 550 mil) paga um serviço de 2020-21: por que não estava no passivo "
            "antes do contrato de 2025?"
        ),
        "Sim: Dr. Pablo, item 2 e dúvida 17 (respondido com o arq. 10)",
    ),
    (
        "P11",
        "11.2.4-5 ETE San Marino",
        "Qual lote será doado para a ETE? Ele está no estoque do doc 11? Existe projeto, orçamento ou termo assinado com o Município?",
        "Em parte: Dr. Pablo, dúvida 16",
    ),
    (
        "P12",
        "11.2.4-13 Celestino Rossi",
        (
            "O doc 11 fala em 'acordos firmados' e a resposta diz 'compromisso verbal'. Existe algum documento? Quais projetos, "
            "quais loteamentos e em que anos? Algum lote ou valor já foi pago a ele?"
        ),
        "Sim: Dr. Pablo, itens 7 e 9 ('compromisso verbal')",
    ),
    (
        "P13",
        "11.2.4-7 Acordi",
        "Os R$ 88.457,52 do TAC e os R$ 130.000 de acerto com o Município (ACERTOS) já cobrem o TAC/PRAD da Obs. 7?",
        "Sim: perguntas-sugeridas 9; Dr. Pablo, dúvida 16",
    ),
    (
        "P14",
        "11.2.4-6 Lago Negro",
        (
            "Depois do acordo de jun/2026 (gravame perpétuo e processo encerrado), ainda há PRAD a executar? As custas e "
            "os honorários já foram pagos?"
        ),
        "Sim: Dr. Pablo, itens 3 e 18 (resposta 3.3, sem falar do PRAD depois do acordo)",
    ),
    (
        "P15",
        "11.2.4-9 áreas rurais",
        "Quais são as áreas? O levantamento GEO M.61.874 e a topografia da área Nestor J. Colo (2024) fazem parte delas?",
        "Em parte: Dr. Pablo, item 5; comentário do Luís na planilha",
    ),
    (
        "P16",
        "11.2.5-16/17/18 Mercali e Piccoli",
        (
            "Contratos e saldo em 31/12/2024 de cada empréstimo, com a data de cada parte (R$ 100 mil de Inês desde "
            "ago/2023; R$ 150 mil de Samir desde fev/2024). Algum foi quitado em 2025? Inês e Samir Trevisan Piccoli têm "
            "relação com o acerto de Arlindo Piccoli no Vale das Colinas?"
        ),
        "Sim: Dr. Pablo, item 7 (sem resposta)",
    ),
    (
        "P17",
        "Fora do 11.2: empréstimo Buriti",
        (
            "A Buriti adiantou à SEI o empréstimo de R$ 1.000.000 previsto na cl. 12 do memorando? Se sim, qual era o "
            "saldo em 31/12/2024? (Se existir, é passivo que falta no 11.2.5: favorece o grupo.)"
        ),
        "Novo",
    ),
    (
        "P18",
        "11.2.4 itens sem valor; parágrafo final",
        (
            "O grupo pretende cobrar os itens sem valor (7% de impostos, 6% de comissão, impostos de transmissão, "
            "regularização, cartório, 'ônus de dissolução')? A tabela 50 cita os 'itens 1 ao 24': o que seriam os itens "
            "22 a 24?"
        ),
        "Em parte: Dr. Pablo, item 5",
    ),
]

FAVOR_GRUPO = [
    # fact, effect for the exiting partners, source
    (
        "Susin: 6 lotes do Termo B não entregues e 2 do Termo A faltando",
        "Dívida documentada (~R$ 600 mil no ACERTOS para os 6). Se os R$ 1,2 mi forem só a comissão, é passivo omitido.",
        "Arq. 16; arq. 17; ACERTOS S2/S9",
    ),
    (
        "Rovílio: 1 lote do Morada II (R$ 550 mil) é comissão por intermediação de 2020-21",
        "Serviço anterior à data-base; não é preço de quota entre sócios.",
        "Arq. 10, cl. 3ª, item 8",
    ),
    (
        "Lago Negro: R$ 580.000 pagos a Lucas Zucoloto (set/2025), por processo de 2016",
        "Passivo que já existia na data-base e não está no 11.2, que cobra só R$ 80 mil.",
        "Arq. 1, p. 25",
    ),
    (
        "Falavigna 016/876: dação à G3 em 10/09/2025, por aditivo de 2010",
        "Lote já devido na data-base; não está no 11.2.",
        "Arq. 17; revisão do cruzamento",
    ),
    (
        "Vale das Colinas: lote 008/876 + R$ 165.000 a Fabiano Vanin (09/04/2026)",
        "A indenização pela não execução já está acontecendo; enfraquece o argumento de que a multa é só possível.",
        "Arq. 17",
    ),
    (
        "Memorando Buriti, cl. 6.2.1 e 6.2.4",
        "A SEI é a única responsável pelas obras e multas dos 18 contratos com os proprietários.",
        "Memorando Buriti",
    ),
    (
        "Memorando Buriti, cl. 12: empréstimo de R$ 1.000.000 da Buriti à SEI",
        "Se foi feito, é passivo que falta no 11.2.5 (pergunta P17).",
        "Memorando Buriti",
    ),
    (
        "Lote 'que era do Mercali', hipotecado à Receita Estadual por dívida de ICMS",
        "A dívida de ICMS não está no 11.2; de quem é e qual o saldo?",
        "Doc 11, T6.R5",
    ),
    (
        "Escavadeira: saldo na data-base R$ 421.396,40, e não R$ 347.652",
        "Passivo de R$ 73.744 maior na data-base.",
        "Contas pagas",
    ),
    (
        "Trator D6N comprado em set/2025",
        "Sair da data-base tira R$ 450 mil de ativo e R$ 320 mil de passivo: −R$ 130 mil líquidos.",
        "Doc 11, máquinas item 14",
    ),
    (
        "Solar Serra: lote 02/07 no estoque (R$ 165 mil) e dívida no passivo (R$ 150 mil)",
        "Corrigir os dois lados reduz R$ 15 mil (e cria o crédito contra o Hotel Araçá).",
        "Arq. 17; doc 11",
    ),
    (
        "Falavigna: galeria ainda paga com lote em jun/2026",
        "A obra não estava concluída; enfraquece o 'Falavigna finalizado' na regra de custos.",
        "Arq. 17",
    ),
]


def write(rows: list[dict]):
    wb = openpyxl.Workbook()
    hdr_font, hdr_fill = (
        Font(bold=True, color="FFFFFF"),
        PatternFill("solid", fgColor="1F4E78"),
    )
    wrap = Alignment(wrap_text=True, vertical="top")

    def header(ws, cols):
        ws.append([c for c, _ in cols])
        for i, (_, w) in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for c in ws[1]:
            c.font, c.fill, c.alignment = hdr_font, hdr_fill, wrap

    # Resumo
    ws = wb.active
    ws.title = "Resumo"
    tot = sum(r["valor"] for r in rows)
    by, cnt, ab = defaultdict(float), defaultdict(int), defaultdict(float)
    for r in rows:
        by[r["situacao"]] += r["valor"]
        cnt[r["situacao"]] += 1
        ab[r["situacao"]] += r["abat"]
    ws.append(["Situação das cobranças do 11.2 do doc 11 (SEI)"])
    ws.append(
        [
            "Situação",
            "Cobranças",
            "Valor cobrado pelo grupo (R$)",
            "Já retirado pelo Luís (R$)",
            "Fica na proposta (R$)",
        ]
    )
    for s in SIT_FILL:
        ws.append([s, cnt[s], by[s], ab[s], by[s] - ab[s]])
    tab = sum(ab.values())
    ws.append(["Total", sum(cnt.values()), tot, tab, tot - tab])
    ws.append([])
    ws.append(
        [
            (
                "Leitura: 'Contingência' = nada foi pago e não há reclamação nem decisão; 'Em aberto, sem comprovação' = o grupo "
                "diz que deve, mas não há documento nas fontes; 'Em aberto, documentada' = há contrato ou juros/parcelas pagos; "
                "'Paga em parte' = há pagamento já feito e não se sabe se o saldo o desconta. 'Já retirado pelo Luís' = 11.2.1 "
                "a R$ 0 e os abatimentos por duplicidade. Valores do grupo (coluna E da aba GERAL), sem correção."
            )
        ]
    )
    ws["A1"].font = Font(bold=True, size=12)
    for c in ws[2]:
        c.font, c.fill, c.alignment = hdr_font, hdr_fill, wrap
    for i, s in enumerate(SIT_FILL, start=3):
        ws[f"A{i}"].fill = PatternFill("solid", fgColor=SIT_FILL[s])
    for row in ws.iter_rows(min_row=3, max_row=3 + len(SIT_FILL)):
        for c in row[2:5]:
            c.number_format = "#,##0.00"
    ws[f"A{3 + len(SIT_FILL)}"].font = Font(bold=True)
    for col, w in zip("ABCDE", (32, 11, 22, 22, 22), strict=True):
        ws.column_dimensions[col].width = w

    # Cobranças
    ws = wb.create_sheet("Cobranças")
    header(
        ws,
        [
            ("Seção", 8),
            ("Item", 6),
            ("Cobrança", 38),
            ("Valor cobrado (R$)", 15),
            ("Natureza", 26),
            ("Situação", 20),
            ("Já foi pago? Como (fatos das fontes)", 85),
            ("Fonte", 24),
            ("Risco de duplicidade", 34),
            ("Pergunta", 9),
            ("Decisão já tomada", 30),
            ("Já retirado pelo Luís (R$)", 14),
            ("Linha GERAL", 8),
        ],
    )
    for r in rows:
        ws.append(
            [
                r["secao"],
                r["item"],
                r["assunto"],
                r["valor"],
                r["natureza"],
                r["situacao"],
                r["como"],
                r["fonte"],
                r["dup"],
                r["perg"],
                r["decisao"],
                r["abat"],
                r["row"],
            ]
        )
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap
        row[3].number_format = row[11].number_format = "#,##0.00"
        row[5].fill = PatternFill("solid", fgColor=SIT_FILL[row[5].value])
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    # Perguntas
    stake = defaultdict(float)
    nlinks = defaultdict(int)
    for r in rows:
        if r["perg"]:
            stake[r["perg"]] += r["valor"]
            nlinks[r["perg"]] += 1
    ws = wb.create_sheet("Perguntas ao grupo")
    header(
        ws,
        [
            ("Nº", 6),
            ("Valor cobrado ligado (R$)", 16),
            ("Cobranças ligadas", 10),
            ("Cobrança", 30),
            ("Pergunta", 100),
            ("Já perguntado?", 32),
        ],
    )
    for qid, items, text, asked in QUESTIONS:
        ws.append([qid, stake[qid], nlinks[qid], items, text, asked])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap
        row[1].number_format = "#,##0.00"
    ws.freeze_panes = "A2"

    # A favor do grupo
    ws = wb.create_sheet("A favor do grupo")
    header(ws, [("Fato", 55), ("Efeito para quem sai", 75), ("Fonte", 30)])
    for line in FAVOR_GRUPO:
        ws.append(list(line))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap

    # Fontes
    ws = wb.create_sheet("Fontes")
    for line in [
        (
            "Doc 11 (inventário SEI 31/12/2024, versão atual)",
            "01_fontes/inventario-2025/11. INVENTÁRIO ... .docx; texto em 02_trabalho/empresas/01-08-.../_extracao/2025-11-sei.txt",
        ),
        ("Aba GERAL (valores do grupo, coluna E)", str(GS.relative_to(ROOT))),
        (
            "Arquivo 17 (390 registros, quase todos 2018-out/2025; alguns de 1994-2014 e de 2026)",
            "01_fontes/respostas-docs-2026-08-13/17 - ... .xlsx",
        ),
        (
            "Arquivo 10 (contrato de quotas de Rovílio, 28/08/2025)",
            "02_trabalho/ocr-respostas-docs/10 - ... .md",
        ),
        (
            "Arquivo 16 (Termos de acordo Susin/AGI)",
            "02_trabalho/ocr-respostas-docs/16 - ... .md",
        ),
        (
            "Arquivo 1 (respostas ao Dr. Pablo)",
            "01_fontes/respostas-docs-2026-08-13/1 - ... .pdf, p. 17-30",
        ),
        (
            "Contas pagas jul/2023-out/2024 (sem set/2023), só SEI",
            str(LEDGER.relative_to(ROOT)),
        ),
        (
            "Memorando Buriti (14/04/2022, sem assinatura)",
            "01_fontes/outros/Contrato Siviero X Buriti ... .pdf e Emailing MEMORANDO ... .pdf",
        ),
        (
            "ACERTOS E ACORDOS (achados S1-S13)",
            "02_trabalho/analises/acertos-e-areas-vendidas/achados-para-sei.md",
        ),
        (
            "Cruzamento do 11.2 com o arq. 17 e revisão",
            "02_trabalho/analises/parte4-fluxo-caixa/ (cruzamento + _diario/revisao-cruzamento-compromissos.md)",
        ),
        (
            "Diff doc 11 antigo × atual",
            "02_trabalho/diff-inventario-2025-antigo-x-atual.md",
        ),
        (
            "Revisão independente desta planilha",
            "02_trabalho/analises/cobrancas-11.2/_diario/revisao-cobrancas.md",
        ),
    ]:
        ws.append(list(line))
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 110
    wb.save(OUT)
    return tot, by, cnt, ab, stake


def main():
    ev = ledger_evidence()
    gv = charge_values()
    rows = build_rows(ev, gv)
    tot, by, cnt, ab, stake = write(rows)
    check = gv[664][2]
    print(
        f"cobranças: {len(rows)}  total: {tot:,.2f}  (GERAL 11.2: {check:,.2f})  confere: {abs(tot - check) < 0.01}"
    )
    for s in SIT_FILL:
        print(f"  {s:30s} {cnt[s]:3d}  {by[s]:>16,.2f}  retirado {ab[s]:>14,.2f}")
    print(f"  retirado pelo Luís: {sum(ab.values()):,.2f}")
    print("valor ligado por pergunta:", {k: round(v) for k, v in sorted(stake.items())})
    print(
        "evidências:",
        {k: (round(v, 2) if isinstance(v, float) else v) for k, v in ev.items()},
    )


if __name__ == "__main__":
    main()
