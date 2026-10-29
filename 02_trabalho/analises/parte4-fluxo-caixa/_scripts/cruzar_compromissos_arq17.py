"""Part 4 / Part 6 (form 1): does a liability charged in doc 11 (items 11.2.2-11.2.5) also appear in file 17 as
lots already handed over? If so the group may be taking the lot out of the inventory AND charging the debt.

Reads file 17 (xlsx transcription of the PDF, every record) and the GERAL tab of the user's Google sheet export.
Writes cruzamento-compromissos-arq17.xlsx and prints a summary. Run: python -I _scripts/cruzar_compromissos_arq17.py
"""

from __future__ import annotations

import datetime as dt
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.stdout.reconfigure(encoding="utf-8")
HERE = Path(__file__).resolve().parent.parent
ROOT = HERE.parents[2]
F17 = next((ROOT / "01_fontes/respostas-docs-2026-08-13").glob("17*.xlsx"))
GS = (
    ROOT
    / "01_fontes/planilhas-google/Comparativo Patrimonial Siviero empreendimentos e Rafael 2018 x 2025.xlsx"
)
DATA_BASE = dt.date(2024, 12, 31)


def norm(s) -> str:
    s = (
        unicodedata.normalize("NFKD", str(s or ""))
        .encode("ascii", "ignore")
        .decode()
        .lower()
    )
    return re.sub(r"\s+", " ", s).strip()


# doc 11 item -> (GERAL row, pattern on the file-17 record: subdivision + buyer + remarks)
ITEMS = {
    "11.2.3 Colina do Retiro II — obras finais": (
        720,
        (
            r"colina do retiro ii.*(paviment|obra|servi|tubo|material)|"
            r"(paviment|obra).*colina do retiro ii"
        ),
    ),
    "11.2.4-1 Rovílio — quotas": (722, r"rovilio|quotas de capital"),
    "11.2.4-2 Perini — infração ambiental": (
        723,
        r"perini.*(infra|ambiental|multa|app)|(infra|multa).*perini",
    ),
    "11.2.4-3 Dalzóchio — fonte de água": (
        724,
        r"dalzochio.*(fonte|ministerio|processo|acordo|indeniza)|fonte de agua",
    ),
    "11.2.4-5 San Marino — ETE": (726, r"\bete\b|esgoto"),
    "11.2.4-6 Lago Negro — APP / Candiago": (727, r"candiago|prad"),
    "11.2.4-7 Acordi": (728, r"acordi"),
    "11.2.4-13 Celestino Rossi": (734, r"celestino|rossi"),
    "11.2.4-14 Solar Serra": (735, r"solar serra|bordignon|placas solares"),
    "11.2.4-15 Alexandre Susin": (736, r"susin|razzera|buriti|agi empreend"),
    "11.2.4-19 Reserva Falavigna — obras": (
        737,
        (
            r"falavigna.*(aduela|galeria|paviment|obra|tubo|material|parque)|"
            r"(aduela|galeria).*falavigna"
        ),
    ),
    "11.2.5-16 Mercali": (739, r"mercali"),
    "11.2.5-17/18 Piccoli": (740, r"piccoli|picoli"),
    "11.2.5-20 Caterpillar (escavadeira)": (742, r"escavadeira|caterpil"),
    "11.2.5-21 Paraná Equipamentos (trator)": (743, r"parana equip|trator"),
}


def lot_key(lot, quadra):
    def n(x):
        m = re.match(r"\d+", str(x or "").strip())
        return int(m.group()) if m else str(x).strip().lower()

    return n(lot), n(quadra)


def main():
    f17 = openpyxl.load_workbook(F17, data_only=True)
    recs = []
    for r in list(f17["Lotes vendidos (PDF)"].iter_rows(values_only=True))[1:]:
        if not r[0]:
            continue
        recs.append(
            {
                "lot": r[0],
                "data": r[1],
                "lote": r[2],
                "quadra": r[3],
                "comprador": r[4],
                "obs": r[5],
                "pag": r[6],
                "txt": norm(f"{r[0]} {r[4]} {r[5]}"),
            }
        )
    geral = openpyxl.load_workbook(GS, data_only=True)["GERAL"]

    out = []
    # 11.2.2: the 26 third-party lots, matched lot by lot
    for rr in range(678, 718):
        b, e = geral[f"B{rr}"].value, geral[f"E{rr}"].value
        m = re.search(r"lote:\s*(\w+)\s*\|\s*quadra:\s*(\w+)", norm(b)) if b else None
        if not (m and isinstance(e, (int, float))):
            continue
        key = lot_key(m.group(1), m.group(2))
        for rec in recs:
            if lot_key(rec["lote"], rec["quadra"]) == key:
                out.append(
                    (
                        "11.2.2 Lotes de terceiros indenizáveis",
                        rr,
                        float(e),
                        rec,
                        f"lote {m.group(1)} / q {m.group(2)}",
                    )
                )
    # 11.2.3-11.2.5: by counterparty and subject
    for item, (rr, pat) in ITEMS.items():
        for rec in recs:
            if re.search(pat, rec["txt"]):
                out.append((item, rr, float(geral[f"E{rr}"].value or 0), rec, ""))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cruzamento"
    hdr = [
        "Item do doc 11",
        "Linha GERAL",
        "Valor cobrado no item (R$)",
        "Chave",
        "Loteamento (arq. 17)",
        "Data",
        "Antes da data-base?",
        "Lote",
        "Quadra",
        "Comprador / destinatário",
        "Observação do arquivo 17",
        "Página",
    ]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(1, j, h)
        c.font, c.fill = (
            Font(bold=True, color="FFFFFF"),
            PatternFill("solid", fgColor="1F3864"),
        )
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, (item, rr, val, rec, key) in enumerate(out, 2):
        d = rec["data"] if isinstance(rec["data"], dt.datetime) else None
        row = [
            item,
            rr,
            val,
            key,
            rec["lot"],
            d.date() if d else rec["data"],
            ("sim" if d.date() <= DATA_BASE else "não") if d else "sem data",
            rec["lote"],
            rec["quadra"],
            rec["comprador"],
            rec["obs"],
            rec["pag"],
        ]
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v)
        ws.cell(i, 3).number_format = "#,##0;[Red]-#,##0"
    for j, w in enumerate([34, 10, 16, 16, 34, 11, 10, 6, 7, 34, 110, 7], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    path = HERE / "cruzamento-compromissos-arq17.xlsx"
    wb.save(path)
    for item, rr, val, rec, key in out:
        d = rec["data"].date() if isinstance(rec["data"], dt.datetime) else rec["data"]
        print(
            f"{item[:38]:38s} | {val:>12,.0f} | {key:16s} | {str(rec['lot'])[:28]:28s} | {d} | {rec['lote']}/{rec['quadra']} | "
            f"{str(rec['comprador'])[:30]:30s} | {str(rec['obs'])[:120]}"
        )
    assess(recs, geral, wb, path)
    print("registros:", len(recs), "cruzamentos:", len(out), "saved", path)


# Cases read one by one from the matches above. Each lists the file-17 records it rests on, as
# (pattern on the record text, lot, quadra); a lot "*" takes every record that matches the pattern.
CASES = [
    {
        "id": "D1",
        "item": "11.2.2 Lotes que eram da SEI e saíram em 2025",
        "row": 677,
        "tipo": "cobrança dupla confirmada",
        "lotes": [
            (r"jardim italia", "1", "1747"),
            (r"primavera", "11", "1199"),
            (r"vicentino", "3", "1301"),
        ],
        "valores": [200_000, 166_000, 1_355_000],
        "base": "valor do lote no 11.2.2 (GERAL)",
        "extra": [
            (
                "2025-02 · 10/1199 · Safra (dação; arquivo 1, p. 26, achado ACERTOS S1)",
                166_000,
            )
        ],
        "leitura": (
            "Quatro lotes eram da SEI em 31.12.2024 e saíram em 2025: Jardim Itália 01/1747, vendido em mar/2025; "
            "Primavera 10 e 11/1199, dados à Safra em 2025; e Vicentino 03/1301, vendido em set/2025. Eles saíram "
            "do estoque e o mesmo valor é descontado como 'lote de terceiro indenizável'."
        ),
        "proposta": "Retirar do 11.2.2 e devolver os 4 lotes ao patrimônio de 31.12.2024 (registro de data-base).",
    },
    {
        "id": "D2",
        "item": "11.2.4-14 Solar Serra (placas do Hotel Araçá)",
        "row": 735,
        "tipo": "dívida já paga com lote, e dívida de outra empresa",
        "lotes": [(r"solar serra", "2", "7")],
        "leitura": (
            "O doc 11 cobra R$ 150 mil (1 lote) à Solar Serra pelas placas solares do Hotel Araçá. O arquivo 17 diz "
            "que o lote 02/07 do Aditivo Medianeira já foi entregue em 24/04/2024 por esse motivo, mas o lote "
            "continua no estoque de 2025 (R$ 165 mil). Estoque e passivo se compensam. O ponto real é que a SEI "
            "pagou com lote dela uma dívida do Hotel Araçá."
        ),
        "proposta": "Retirar o lote do estoque e o passivo, e registrar um crédito da SEI contra o Hotel Araçá.",
    },
    {
        "id": "D3",
        "item": "11.2.4-19 Reserva Falavigna — obras",
        "row": 737,
        "tipo": "possível (depende da composição do saldo)",
        "lotes": [
            (r"falavigna", "20", "174"),
            (r"falavigna", "21", "174"),
            (r"falavigna", "23", "174"),
            (r"falavigna", "12", "175"),
            (r"falavigna", "13", "175"),
        ],
        "leitura": (
            "O doc 11 cobra R$ 2,0 mi para concluir pavimentação, parque e galeria. Em 2025, 5 lotes do próprio "
            "Falavigna foram entregues à Florense (aduelas e tubos) e à Bett (pavimentação), como 50% de cada "
            "fornecimento. Se os R$ 2,0 mi incluem o que esses lotes pagaram, a cobrança é dupla."
        ),
        "proposta": "Pedir a composição dos R$ 2,0 mi e o que já foi pago com lotes.",
    },
    {
        "id": "D4",
        "item": "11.2.3 Colina do Retiro II — obras finais",
        "row": 720,
        "tipo": "possível (depende da composição do saldo)",
        "lotes": [(r"carnevalli", "7", "a")],
        "leitura": (
            "O lote 07/A do Carnevalli foi entregue à CVC em ago/2024 como pagamento de pavimentação no Colina do "
            "Retiro II. O doc 11 cobra R$ 3,8 mi para concluir a pavimentação, sem dizer se já desconta esse "
            "pagamento."
        ),
        "proposta": "Pedir a composição dos R$ 3,8 mi e o que já foi pago com lotes.",
    },
    {
        "id": "D5",
        "item": "11.2.4-15 Alexandre Susin — comissão Buriti",
        "row": 736,
        "tipo": "a esclarecer (depende da data dos contratos Buriti)",
        "lotes": [(r"susin|agi emp", "*", "*")],
        "valores": [0] * 40,
        "base": "sem valor: os lotes pagaram outra dívida",
        "leitura": (
            "Em 2020, 15 lotes-equivalentes foram entregues a Susin e AGI pelo Termo A (13/07/2020), mais 1 lote em "
            "2019 por outra dívida. O Termo A quita só as intermediações feitas até aquela data, e a cl. 8 exige "
            "documento escrito para as novas. Os 6 lotes do Termo B seguem devidos e no estoque. A comissão de "
            "R$ 1,2 mi só é dupla se os contratos Buriti forem anteriores a 13/07/2020."
        ),
        "proposta": "Pedir a data dos contratos Buriti e o documento escrito da cl. 8. Sem ele, contestar os R$ 1,2 mi.",
    },
    {
        "id": "D6",
        "item": "11.2.4-1 Rovílio — quotas",
        "row": 722,
        "tipo": "inconsistência (sem cobrança dupla)",
        "lotes": [(r"morada", "7", "4"), (r"rovilio siviero", "*", "*")],
        "valores": [432_072, 587_510, 500_000],
        "base": "valor do lote no item 1 do 11.2.4 (doc 11)",
        "leitura": (
            "O doc 11 dá como 'PAGO' a Rovílio o lote 07/04 do Morada dos Ipês I (R$ 432.072), mas o arquivo 17 "
            "diz que esse lote foi para a Structura Corretora em 2021, como comissão: o mesmo lote em dois "
            "pagamentos. Os lotes 12/237 e 15/252 do N. S. da Saúde, também 'PAGO', não estão no arquivo 17 nem "
            "no estoque. A compra das quotas é de 28/08/2025, depois da data-base, e os R$ 1,65 mi 'à pagar' são 3 "
            "lotes do Morada dos Ipês II a R$ 550 mil cada, loteamento que o grupo avalia em R$ 0."
        ),
        "proposta": (
            "Pedir o esclarecimento do 07/04 e dos 2 lotes do N. S. da Saúde. Registro de data-base. Se o Morada II "
            "vale R$ 550 mil por lote, não vale R$ 0."
        ),
    },
]
DISCARDED = [
    (
        "11.2.2 lotes já de terceiros (Ambitec 02/1747, Farina 11/1469, Perini 18 e 19/1469 vendidos em 2024)",
        (
            "Revisão: o 11.2.2 é contingência sobre lotes já vendidos, não cobrança dupla. Continua valendo o comentário do "
            "Luís: não se desconta indenização apenas possível."
        ),
    ),
    (
        "11.2.4-3 Dalzóchio",
        "Os lotes encontrados são do acordo Farina/Vicentino, outro processo.",
    ),
    (
        "11.2.4-6 Lago Negro",
        "A permuta Candiago é a indenização citada no próprio item; o item cobra só PRAD e custas.",
    ),
    (
        "11.2.5 Piccoli",
        "O acordo de 2024 é com Arlindo Piccoli (Vale das Colinas), não com Inês ou Samir (empréstimos).",
    ),
    (
        "11.2.4-13 Celestino Rossi; Mercali; Caterpillar; Paraná",
        "Nenhum lote entregue encontrado no arquivo 17.",
    ),
    (
        "Palavras soltas",
        "'prad' em Antônio Prado; 'rossi' em Rossim; Rovílio Dal Más (outra pessoa).",
    ),
]
OTHER = [
    (
        "Lotes já entregues que seguem no estoque de 2025",
        (
            "Aditivo Medianeira 02/07 (Solar Serra), Alto das Graças 11/6371 "
            "(HECTA/Basso; também na lista de Rafael), Sol Nascente I 1/A e 2/A (Válter, arquivo 18). Aumentam o patrimônio da "
            "empresa; se algum estiver na lista de bens que virão para nós, precisa sair dela."
        ),
    ),
    (
        "Passivo omitido",
        "Dação G3 do lote 016/876 do Falavigna (set/2025): não aparece no passivo do doc 11.",
    ),
]


GENERIC = {
    "loteamento",
    "familia",
    "reserva",
    "condominio",
    "industrial",
    "farroupilha",
    "veranopolis",
    "caxias",
    "nova",
    "prata",
    "flores",
    "cunha",
    "passo",
    "fundo",
}


def sig(name):
    # singular form, so that "Altos das Graças" and "Alto das Graças I" share both words
    return {
        x.rstrip("s") for x in re.findall(r"[a-z]+", norm(name)) if len(x) > 3
    } - GENERIC


def lot_value_2025(lot_name, lote, quadra, wb_model, stock25):
    """Group's own value: the lot in the 2025 list if it is there, else the subdivision's 2025 average."""
    k = lot_key(lote, quadra)
    for name, q, lt, v in stock25:
        if (lt, q) == k and sig(name) & sig(lot_name):
            return v, "doc 11 (lote no estoque de 2025)"
    best, bw = None, 0
    for name, v25 in wb_model:
        w = len(sig(name) & sig(lot_name))
        if w > bw:
            best, bw = (name, v25), w
    return (best[1], f"média 2025 de {best[0]}") if best else (None, "")


def assess(recs, geral, wb, path):
    model = openpyxl.load_workbook(HERE / "fluxo-caixa-sei-v0.2.xlsx", data_only=True)[
        "1 Lotes por loteamento"
    ]
    wb_model = [
        (r[0], r[7])
        for r in model.iter_rows(min_row=2, values_only=True)
        if r[0] and isinstance(r[7], (int, float))
    ]
    s25 = openpyxl.load_workbook(
        ROOT / "02_trabalho/sei-2025-itens.xlsx", data_only=True
    )["Lotes"]
    stock25 = []
    for r in list(s25.iter_rows(values_only=True))[1:]:
        if r[0] and isinstance(r[5], (int, float)):
            lt, q = lot_key(r[3], r[2])  # columns: quadra = r[2], lote = r[3]
            stock25.append((norm(r[0]), q, lt, float(r[5])))
    ws = wb.create_sheet("Avaliação")
    hdr = [
        "Caso",
        "Item do doc 11",
        "Valor cobrado no item (R$)",
        "Tipo",
        "Lotes do arquivo 17 (data · lote/quadra · destinatário)",
        "Lotes",
        "Valor dos lotes pelo grupo (R$)",
        "Base do valor",
        "Leitura",
        "Proposta",
    ]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(1, j, h)
        c.font, c.fill = (
            Font(bold=True, color="FFFFFF"),
            PatternFill("solid", fgColor="1F3864"),
        )
        c.alignment = Alignment(wrap_text=True, vertical="top")
    r = 2
    for case in CASES:
        found = []
        for pat, lote, quadra in case["lotes"]:
            for rec in recs:
                if (
                    rec not in found
                    and re.search(pat, rec["txt"])
                    and (
                        lote == "*"
                        or lot_key(rec["lote"], rec["quadra"]) == lot_key(lote, quadra)
                    )
                ):
                    found.append(rec)
        if not found:
            raise SystemExit(f"{case['id']}: nenhum registro encontrado")
        vals, bases = [], set()
        for i, rec in enumerate(found):
            if "valores" in case:
                v, b = (
                    case["valores"][i] if i < len(case["valores"]) else 0,
                    case.get("base", ""),
                )
            else:
                v, b = lot_value_2025(
                    rec["lot"], rec["lote"], rec["quadra"], wb_model, stock25
                )
            vals.append(v or 0)
            bases.add(b)
        desc = "; ".join(
            f"{rec['data'].date() if isinstance(rec['data'], dt.datetime) else rec['data']} · "
            f"{rec['lote']}/{rec['quadra']} · {str(rec['comprador'])[:40]}"
            for rec in found
        )
        for text, v in case.get("extra", []):  # lots documented outside file 17
            desc += f"; {text}"
            vals.append(v)
        row = [
            case["id"],
            case["item"],
            float(geral[f"E{case['row']}"].value or 0),
            case["tipo"],
            desc,
            len(vals),
            sum(vals),
            "; ".join(sorted(bases)),
            case["leitura"],
            case["proposta"],
        ]
        for j, v in enumerate(row, 1):
            c = ws.cell(r, j, v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for j in (3, 7):
            ws.cell(r, j).number_format = "#,##0;[Red]-#,##0"
        case["n"], case["valor_lotes"], case["cobrado"] = len(vals), sum(vals), row[2]
        r += 1
    r += 1
    ws.cell(
        r, 1, "Descartados (palavra-chave casou, mas não é o mesmo pagamento)"
    ).font = Font(bold=True)
    for item, why in DISCARDED:
        r += 1
        ws.cell(r, 1, item)
        ws.cell(r, 2, why)
    r += 2
    ws.cell(r, 1, "Outros achados do cruzamento (não são cobrança dupla)").font = Font(
        bold=True
    )
    for item, why in OTHER:
        r += 1
        ws.cell(r, 1, item)
        ws.cell(r, 2, why)
    r += 2
    ws.cell(
        r,
        1,
        "Conferido por revisor independente: _diario/revisao-cruzamento-compromissos.md",
    ).font = Font(italic=True)
    for j, w in enumerate([6, 34, 16, 22, 70, 7, 16, 30, 80, 60], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    wb.save(path)
    for c in CASES:
        print(
            f"{c['id']} {c['item'][:40]:40s} cobrado {c['cobrado']:>12,.0f} | {c['n']:2d} lotes | valor lotes "
            f"{c['valor_lotes']:>12,.0f} | {c['tipo']}"
        )


if __name__ == "__main__":
    main()
