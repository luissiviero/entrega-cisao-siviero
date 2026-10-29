"""Map ACERTOS E ACORDOS counterparts to file-17 lots and cross-check each lot.

Prints a pipe-separated table; the final workbook is built by build_outputs.py.
"""
import re
import sys
import unicodedata

import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
ROOT = 'C:/entrega-cisao-siviero/'
F17 = ROOT + '01_fontes/respostas-docs-2026-08-13/17 - SIVIERO EMPREENDIMENTOS IMOBILIÁRIOS LTDA. - LOTES VENDIDOS DE JANEIRO DE 2018 A OUTUBRO DE 2025 - REV.13.08.2026.xlsx'
S25 = ROOT + '02_trabalho/sei-2025-itens.xlsx'
GERAL = ROOT + '01_fontes/outros/Lotes vendidos 2018-2025 x GERAL.xlsx'


def norm(s):
    s = unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode().upper()
    return re.sub(r'\s+', ' ', s).strip()


def num(s):
    """Lot/quadra key: strip leading zeros and spaces, keep letters.

    Doc 11 writes some quadras as 'H - 6371' (letter + municipal number); file 17 uses '6371'.
    """
    n = norm(s)
    m = re.fullmatch(r'[A-Z]\s*-\s*(\d+)', n)
    if m:
        n = m.group(1)
    s = n.replace(' ', '').replace('O', '0') if re.fullmatch(r'[O0-9 ]+', n) else n.replace(' ', '')
    return s.lstrip('0') or '0'


# loteamento name -> canonical key shared by file 17, stock 2025 and indemnifications
CANON = [
    ('ALTOS DA JULIETA', 'ALTOS DA JULIETA'), ('ALTOS DAS GRACAS', 'ALTO DAS GRACAS I'),
    ('ALTO DAS GRACAS', 'ALTO DAS GRACAS I'), ('MONTE VENETO', 'MONTE VENETO'),
    ('PERINI', 'PERINI'), ('VIERO', 'VIERO'), ('SAN MARINO', 'SAN MARINO'),
    ('DALZOCHIO', 'DALZOCHIO'), ('JARDIM ITALIA', 'JARDIM ITALIA'),
    ('INDUSTRIAL VICENTINO', 'INDUSTRIAL VICENTINO'), ('PEDRA LISA', 'PEDRA LISA'),
    ('PRIMAVERA', 'PRIMAVERA'), ('DELTA', 'DELTA-NILUMAR'), ('FALAVIGNA', 'RESERVA FALAVIGNA'),
    ('SALVADOR', 'SALVADOR FEMACA II'), ('FEMACA', 'FEMACA'), ('CARNEVALLI', 'CARNEVALLI'),
]


def canon(name):
    n = norm(name)
    for k, v in CANON:
        if k in n:
            if v == 'MONTE VENETO' and 'CARLOS BARBOSA' in n:
                return 'MONTE VENETO (CARLOS BARBOSA)'
            return v
    return n


# counterpart patterns per ACERTOS item (matched against buyer + observation of file 17)
ITEMS = {
    '01': [r'ANGEST'],
    '03/04': [r'SUSIN', r'\bAGI\b'],
    '06': [r'AMBITEC'],
    '08': [r'FARINA'],
    '09': [r'BARUFFI'],
    '11': [r'ALCEONES'],
    '12': [r'FANTINEL'],
    '13': [r'FOCHESAT'],
    '14': [r'VIA NOVA'],
    '17': [r'VERONA'],
    '21': [r'ANSOLIN'],
    '25': [r'COLUNATA', r'ANTONIAZZI'],
    '26': [r'RAZ+ERA'],
    '27': [r'COLOMBO'],
    '32': [r'ALAN MARTINS'],
}


def load_f17():
    ws = openpyxl.load_workbook(F17, data_only=True).worksheets[0]
    out = []
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        lot, d, l, q, b, o, p = r
        out.append(dict(row=i, lot=lot, key=(canon(lot), num(q), num(l)), date=d, buyer=b, obs=o, page=p,
                        text=norm(f'{b} {o}')))
    return out


def load_stock():
    wb = openpyxl.load_workbook(S25, read_only=True, data_only=True)
    stock = {}
    for r in wb['Lotes'].iter_rows(min_row=2, values_only=True):
        stock[(canon(f'{r[0]} {r[1]}'), num(r[2]), num(r[3]))] = (r[0], r[5], r[7], r[10])
    indem = {}
    for r in wb['Compromissos'].iter_rows(min_row=2, values_only=True):
        if str(r[0]).startswith('Indeniza'):
            lot = canon(r[3])
            for q, l, v in re.findall(r'Q\s*([0-9O]+)\s*L\s*([0-9O]+)\s*\(R\$\s*([\d\.]+),', r[4]):
                indem[(lot, num(q), num(l))] = (r[0], float(v.replace('.', '')))
    return stock, indem


# file 18 lots (Válter), transcribed from the docx table rows R003-R032; quadra/lote as in the document
F18 = [
    ('LAGO NEGRO', '8', '2', 'R003 Oneide Wiltgen -> Land Rover (URBAN), R$ 273.000 (2 lotes)'),
    ('LAGO NEGRO', '8', '3', 'R003 Oneide Wiltgen -> Land Rover (URBAN), R$ 273.000 (2 lotes)'),
    ('ALTO DAS GRACAS I', '6371', '6', 'R004 consignado - 2,5 ha Monte Bérico, parte de Alexandre Susin, em troca de 3 lotes'),
    ('ALTO DAS GRACAS I', '6371', '12', 'R004 consignado - 2,5 ha Monte Bérico, parte de Alexandre Susin, em troca de 3 lotes'),
    ('ALTO DAS GRACAS I', '6371', '5', 'R004 consignado - 2,5 ha Monte Bérico, parte de Alexandre Susin, em troca de 3 lotes'),
    ('ALTOS DA JULIETA', '1728', '3', 'R005 Pago Césio Verona - serviços prestados para a empresa Rochedo'),
    ('ALTO DAS GRACAS I', '6365', '1', 'R009 25.08.2014 R$ 170.000'),
    ('ALTO DAS GRACAS I', '6365', '2', 'R009 25.08.2014 R$ 170.000'),
    ('ALTO DAS GRACAS I', '6372', '16', 'R014 consignado'),
    ('ALTO DAS GRACAS I', '6373', '11', 'R015 consignado'),
    ('ALTO DAS GRACAS I', '6374', '8', 'R016 consignado'),
    ('ALTO DAS GRACAS I', '6365', '5', 'R027 10.06.2019 consignado'),
    ('ALTO DAS GRACAS I', '6365', '6', 'R027 10.06.2019 consignado'),
    ('ALTO DAS GRACAS I', '6368', '13', 'R027 10.06.2019 consignado'),
    ('ALTO DAS GRACAS I', '6369', '13', 'R027 10.06.2019 consignado'),
    ('ALTO DAS GRACAS I', '6365', '4', 'R028 consignado'),
    ('ALTO DAS GRACAS I', '6368', '10', 'R029 consignado'),
    ('ALTO DAS GRACAS I', '6369', '8', 'R032/P026 Loteamento Basso (HECTA) - 2 lotes R$ 260.000 (lote 08/Q6369 + lote 11/Q6371)'),
    ('ALTO DAS GRACAS I', '6371', '11', 'R032/P026 Loteamento Basso (HECTA) - 2 lotes R$ 260.000 (lote 08/Q6369 + lote 11/Q6371)'),
    ('SOL NASCENTE', 'A', '1', 'R007 3 salas e 3 box Caxias (HECTA, VMC, URBAN)'),
    ('SOL NASCENTE', 'A', '2', 'R007'), ('SOL NASCENTE', 'C', '5', 'R007'),
    ('SOL NASCENTE', 'C', '6', 'R007'), ('SOL NASCENTE', 'F', '8', 'R007'),
]
F18_IDX = {(a, num(b), num(c)): d for a, b, c, d in F18}


def main():
    f17 = load_f17()
    stock, indem = load_stock()
    seen = {}
    print('item|linha17|pag17|loteamento|quadra|lote|data|comprador|em estoque 2025?|indenizavel doc11?|arquivo 18?|obs17')
    for item, pats in ITEMS.items():
        for r in f17:
            if any(re.search(p, r['text']) for p in pats):
                k = r['key']
                seen.setdefault(k, []).append(item)
                st = stock.get(k)
                ind = indem.get(k)
                f18 = F18_IDX.get(k)
                d = r['date'].strftime('%Y-%m-%d') if hasattr(r['date'], 'strftime') else r['date']
                print(f"{item}|{r['row']}|{r['page']}|{k[0]}|{k[1]}|{k[2]}|{d}|{r['buyer']}|"
                      f"{st if st else ''}|{ind if ind else ''}|{f18 or ''}|{str(r['obs'])[:120]}")
    print('\n# lots matched to more than one item')
    for k, v in seen.items():
        if len(v) > 1:
            print(k, v)
    print('\n# file-17 lots that are in doc-11 indemnification lists')
    for r in f17:
        if r['key'] in indem:
            d = r['date'].strftime('%Y-%m-%d') if hasattr(r['date'], 'strftime') else r['date']
            print(r['row'], r['key'], d, r['buyer'], indem[r['key']], str(r['obs'])[:100])
    print('\n# file-17 lots that are also in 2025 stock')
    for r in f17:
        if r['key'] in stock:
            d = r['date'].strftime('%Y-%m-%d') if hasattr(r['date'], 'strftime') else r['date']
            print(r['row'], r['key'], d, r['buyer'], stock[r['key']], str(r['obs'])[:80])
    print('\n# file-18 lots that are in file 17 / stock / indem')
    by17 = {}
    for r in f17:
        by17.setdefault(r['key'], []).append(r)
    for k, d in F18_IDX.items():
        hits = by17.get(k, [])
        print(k, d[:60], '| f17:', [(h['row'], str(h['buyer'])[:30]) for h in hits], '| stock:', stock.get(k), '| indem:', indem.get(k))


if __name__ == '__main__':
    main()
