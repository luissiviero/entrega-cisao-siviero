"""Build extracoes.xlsx and cruzamentos.xlsx for the ACERTOS / Hotel docx analysis.

Transcribed values come from the page images of ACERTOS E ACORDOS.pdf (checked against the
gdoc and the 26/10 extraction) and from the Hotel docx text. Lot cross-checks are recomputed
from file 17, sei-2025-itens.xlsx and file 18 via match_lots.py.
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(__file__))
import match_lots as ml  # noqa: E402

OUT = 'C:/entrega-cisao-siviero/02_trabalho/analises/acertos-e-areas-vendidas/'
SRC_AC = 'ACERTOS E ACORDOS.pdf'
SRC_HT = 'HOTEL PRINCESA DOS VALES - ÁREA DE TERRAS VENDIDAS + LOTEAMENTOS CONTRATADOS PELA SEI - SEM OBRAS.docx'

# ---------------------------------------------------------------- ACERTOS p.1-3
# (nº, contraparte, motivo, lotes, loteamento/lote citado, espécie, valor de lote citado, texto exato, página)
ACORDOS = [
    ('1', 'Valdir Angest', 'indenização (distrato/lote não entregue)', 1, 'Altos das Graças 05/6372', None, None, 'VALDIR ANGEST – 01 lote (05/6372-Altos das Graças).', 1),
    ('2', 'Valdir Angest', 'indenização (distrato/lote não entregue)', 0, '', 270000.00, None, 'VALDIR ANGEST – R$ 270.000,00 – pago em dinheiro.', 1),
    ('3', 'Alexandre Susin', 'lotes comprados por Susin e família de proprietários (Bovo, C. Fioreze, L. Pedrotti); entrega atrasada (Termo B, 30/04/2020)', 7, '', None, None, 'ALEXANDRE SUSIN – 07 lotes em um acordo.', 1),
    ('4', 'Alexandre Susin', 'comissão/serviços de intermediação (Termo A, 13/07/2020)', 17, '', None, None, 'ALEXANDRE SUSIN – 17 lotes em um acordo.', 1),
    ('5', 'Alexandre Susin', 'comissão/serviços (acordo extrajudicial)', 0, '', 200000.00, None, 'ALEXANDRE SUSIN – R$ 200.000,00 – pago em dinheiro.', 1),
    ('6', 'Ambitec', 'serviço/distrato', 11, '', None, None, 'AMBITEC – 11 lotes em acordo.', 1),
    ('7', 'Ambitec', 'serviço/distrato', 0, '', 15000.00, None, 'AMBITEC – R$ 15.000,00 – pago em dinheiro.', 1),
    ('8', 'Família Farina', 'acordo judicial', 17, '', None, None, 'FAMÍLIA FARINA – 17 lotes – acordo judicial.', 1),
    ('9', 'Família Farina', 'honorários advocatícios', 1, 'Industrial Vicentino (trocado por 04 lotes em Farroupilha)', None, 830000.00, 'FAMÍLIA FARINA – 01 lote, de R$ 830.000,00 – honorários advocatícios (Foi feito um novo acordo, onde foi trocado este lote, do Loteamento Industrial Vicentino, por 04 lotes nos Loteamentos na cidade de Farroupilha.', 1),
    ('10', 'Família Farina', 'honorários advocatícios', 0, '', 220000.00, None, 'FAMÍLIA FARINA – R$ 220.000,00 – honorários advocatícios – pago em dinheiro.', 1),
    ('11', 'Alceones Salvador', 'serviços (acordo)', 4, '', None, None, 'ALCEONES SALVADOR – 04 lotes em acordo.', 1),
    ('12', 'Telmo Fantinel', 'serviços (acordo)', 1, 'Monte Vêneto', None, None, 'TELMO FANTINEL – 01 lote em acordo – Loteamento Monte Vêneto.', 1),
    ('13', 'Fochesato – Caxias do Sul', 'distrato/atraso (lotes negociados no Zanandréa)', 4, 'Zanandréa', None, None, 'FOCHESATO – Caxias do Sul – 04 lotes – em acordo, referente aos lotes negociados no Loteamento Zanandréa.', 1),
    ('14', 'Via Nova – Boschetti + Sandro + Rafael', 'serviço/distrato', 9, '', None, None, 'VIA NOVA – BOSCHETTI+SANDRO+RAFAEL – 09 lotes em acordo.', 1),
    ('15', 'Via Nova – Boschetti + Sandro + Rafael', 'serviço/distrato', 0, '', 90000.00, None, 'VIA NOVA – BOSCHETTI+SANDRO+RAFAEL – R$ 90.000,00 – pago em dinheiro.', 1),
    ('16', 'Flávia Mânica Siviero', 'acordo (em pagamento, 24/50 parcelas)', 0, '', 500000.00, None, 'FLÁVIA MÂNICA SIVIERO – R$ 500.000,00 – acordo que está sendo pago em dinheiro – 24/50 parcelas.', 1),
    ('17', 'Verona – Farroupilha', 'serviços (acordo)', 2, '01 San Marino + 01 Altos da Julieta', None, None, 'VERONA – Farroupilha – 02 lotes em acordo (01/San Marino + 01/Altos da Julieta).', 2),
    ('18', 'Celestino Rossi', 'honorários (a acertar)', 0, '', None, None, 'CELESTINO ROSSI – A ACERTAR OS HONORÁRIOS.', 2),
    ('19', 'Antônio Ansolin', 'aquisição (compra de 06 lotes, segundo o texto)', 0, '', 55000.00, None, 'ANTÔNIO ANSOLIN – R$ 55.000,00 – compra de 06 lotes.', 2),
    ('20', 'Antônio Ansolin', 'aquisição (compra de 7,0 ha)', 0, '', 156666.00, None, 'ANTÔNIO ANSOLIN – R$ 156.666,00 – em dinheiro referente a compra de 7,0 ha.', 2),
    ('21', 'Antônio Ansolin', 'aquisição (compra de 7,0 ha)', 1, 'Industrial Pedra Lisa', None, 375000.00, 'ANTÔNIO ANSOLIN – R$ 375.000,00 – 01 lote no Loteamento Industrial Pedra lisa, referente a compra de 7,0 ha.', 2),
    ('22', 'Lori Fioreze', 'aquisição (compra de 02 lotes)', 0, '', 120000.00, None, 'LORI FIOREZE – R$ 120.000,00 – compra de 02 lotes.', 2),
    ('23', 'Ezequiel Fioreze', 'aquisição (compra de 01 ha, em pagamento)', 0, '', 170000.00, None, 'EZEQUIEL FIOREZE – R$ 170.000,00 – compra de 01 ha, que está sendo pago.', 2),
    ('24', 'Cândido Fioreze', 'aquisição (compra de 02 lotes)', 0, '', 120000.00, None, 'CÂNDIDO FIOREZE – 120.000,00 – compra de 02 lotes.', 2),
    ('25', 'Asdrubal Antoniazzi – Colunata', 'serviços', 3, '01 Vale das Colinas + 02 Basso', None, None, 'ASDRUBAL ANTONIAZZI – COLUNATA – 01 lote no Vale das Colinas. / 02 lotes no Basso.', 2),
    ('26', 'Rodrigo Razera – Susin', 'comissão', 1, 'Viero – Farroupilha', None, None, 'RODRIGO RAZERA – SUSIN – 01 lote em Farroupilha – Viero.', 2),
    ('27', 'Inês Colombo', 'multa por rescisão de contrato (distrato)', 2, 'Altos da Julieta', None, None, 'INÊS COLOMBO – 02 lotes no Altos da Julieta – multa rescisão contrato.', 2),
    ('28', 'Edoardo Bovo – Farroupilha', 'acordo judicial', 0, '', 120000.00, None, 'EDOARDO BOVO – Farroupilha – R$ 120.000,00 – acordo judicial.', 2),
    ('29a', 'Família Acordi – Farroupilha', 'multa ambiental (TAC MP)', 0, '', 88457.52, None, 'FAMILIA ACORDI – Farroupilha: multa de +/- R$ 1.200.000,00, foi acordado da seguinte forma: TAC Ministério Público de Farroupilha – R$ 88.457,52', 2),
    ('29b', 'Família Acordi – Farroupilha', 'multa ambiental (Prefeitura)', 0, '', 61920.30, None, 'Multa na Prefeitura de Farroupilha – R$ 61.920,30', 2),
    ('29c', 'Família Acordi – Farroupilha', 'acerto com o Município', 0, '', 130000.00, None, 'Acerto com o Município – R$ 130.000,00', 2),
    ('30a', 'Família Basso', 'dívida na Prefeitura de Caxias do Sul (40 parcelas, 11 pagas)', 0, '', 156100.41, None, 'FAMÍLIA BASSO – Acerto da dívida na Prefeitura de Caxias do Sul: R$ 156.100,41, em 40 parcelas – Pago 11/40 parcelas de R$ 4.473,71.', 2),
    ('30b', 'Família Basso', 'dívida na Prefeitura de Caxias do Sul (40 parcelas, 11 pagas)', 0, '', 69060.32, None, 'R$ 69.060,32, em 40 parcelas – Pago 11/40 parcelas de R$ 1.971,21.', 2),
    ('31', 'Família Pedrotti – Casa Fochesatto – Nova Prata', 'não informado', 0, '', 270000.00, None, 'FAMÍLIA PEDROTTI / Casa Fochesatto – Nova Prata / R$ 270.000,00 (R$ 190.000,00 + R$ 80.000,00).', 2),
    ('32a', 'Alan – advogado das famílias do Monte Vêneto', 'honorários', 1, 'Monte Vêneto L.11/Q.09', None, None, 'ALAN – Advogado Famílias do Loteamento Monte Vêneto. 01 lote no Loteamento Monte Vêneto – L.11/Q.09.', 3),
    ('32b', 'Alan – advogado das famílias do Monte Vêneto', 'honorários', 0, '', 70000.00, None, 'R$ 70.000,00 – pago em dinheiro.', 3),
]
RESUMO = [('Lotes declarados', 82, 'p. 3'), ('80 lotes x R$ 100.000,00', 8000000.00, 'p. 3'),
          ('01 lote', 830000.00, 'p. 3'), ('01 lote', 375000.00, 'p. 3'),
          ('Em espécie', 2882204.55, 'p. 3'), ('TOTAL', 12087204.55, 'p. 3')]

# ---------------------------------------------------------------- ACERTOS p.4-6 (Boff)
BOFF = [
    ('B1', 'Ludovico Dante Boff (Família Boff)', 'Pagamento dos impostos das 03 áreas e adiantamentos (valor corrigido até março de 2010)', 983724.02, 'até mar/2010', '05 lotes de frente para a Rodovia (com B2; total firmado R$ 1.123.724,02)', 'mar/2010 até 2022', 4),
    ('B2', 'Ludovico Dante Boff (Família Boff)', 'Adiantamento para pagamento de despesas', 140000.00, 'não informada', '05 lotes de frente para a Rodovia (com B1)', 'mar/2010 até 2022', 4),
    ('B3', 'Ludovico Dante Boff (Família Boff)', 'Adiantamento de valores para pagamento de despesas', 100000.00, 'não informada', '01 lote na rua paralela à Rodovia, fundos com a Área Verde', 'mar/2010 até 2022', 4),
    ('B4', 'Prefeitura (IPTU das áreas da família Boff)', 'IPTU das 03 áreas escrituradas do loteamento, 2011 a 2022 (R$ 3,30/m²)', 1555898.94, '2011 a 2022', 'nenhuma informada', '2011 até 2022', 4),
    ('B5', 'Dorval Francisco Boff', 'Ação trabalhista em Gramado', 221052.27, 'fev/2013', '03 lotes na primeira quadra (Termo de Divisão)', 'fev/2013 até 2022', 5),
    ('B6', 'Dorval Francisco Boff', 'Compra de 02 lotes (contrato de 03/09/2012)', 165000.00, '03/09/2012', '02 lotes na divisa do colégio', 'set/2012 até 2022', 5),
    ('B7', 'Odir Marcos Boff', 'Compra de 01 lote (contrato de 23/02/2011)', 50000.00, '23/02/2011', '01 lote de 390,00 m² nos fundos da casa do Sr. Odir', 'fev/2011 até 2022', 6),
    ('B8', 'Odir Marcos Boff', 'Compra de 01 lote (contrato de 05/03/2012)', 100000.00, '05/03/2012', '01 lote de 360,00 m² na rua paralela à Rodovia, divisa com a Área Verde', 'mar/2012 até 2022', 6),
]
BOFF_TOTAL = 3315675.23

# ---------------------------------------------------------------- Hotel docx
VENDAS = [
    ('Família Rigo', 'Fazenda São Crispin', 190.00, 2018, 9448000.00, 49726.30, '118.100 sacas de soja a R$ 80,00 (cotação out/2018): R$ 3.000.000,00 à vista (37.500 sacas) + 20.150 sacas em maio de 2019, 2020, 2021 e 2022; 2021 e 2022 "FALTA RECEBER"', 'R005'),
    ('Família Pellegrini', 'Fazenda São Crispin', 85.32, 2018, 2100000.00, 23441.16, 'R$ 2.000.000,00 pela área + R$ 100.000,00 em 41 cabeças de gado; R$ 1.100.000,00 à vista + R$ 1.000.000,00 em jun/2019', 'R006'),
    ('Senhor Angeli', 'Parte do lote rural 24 e 24/E, Cotiporã (mat. 2.424 e 2.425; adquirida de Celito Pierozan)', 10.2, 2018, 510000.00, 50000.00, 'à vista', 'R007'),
]
VENDAS_TOTAL = 12058000.00
SOJA = [
    ('Família Rigo (venda da área)', 'maio/2021', 20150, 'Hotel Princesa dos Vales (vendedor da área, imóvel em nome do Hotel)', 'vendidas 2.606 sacas em 2024 para aquisição da escavadeira hidráulica', 'P010'),
    ('Família Rigo (venda da área)', 'maio/2021 – vendidas em 2024', -2606, 'Hotel Princesa dos Vales', 'para aquisição da escavadeira hidráulica', 'P010-P011'),
    ('Família Rigo (venda da área)', 'maio/2022', 20150, 'Hotel Princesa dos Vales', '', 'P012'),
    ('Família Lorenzet (arrendamento)', '2021', 2606, 'Hotel (texto: "HOTEL/OLEOPLAN")', '', 'P017'),
    ('Família Lorenzet (arrendamento)', '2022', 3206, 'Hotel (texto: "HOTEL/OLEOPLAN")', '', 'P018'),
    ('Família Lorenzet (arrendamento)', '2023', 3591, 'Hotel (texto: "HOTEL/OLEOPLAN")', '', 'P019'),
    ('Família Lorenzet (arrendamento)', '2024', 3682, 'Agropecuária Fazenda São Crispin (texto: "AGROPECUÁRIA/OLEPOPLAN")', '', 'P020'),
]
ARREND = [
    ('a) lavoura', 155.0, '18 sacas de soja por ha plantado', '2.790 sacas/ano'),
    ('b) batata', 30.0, '30 sacas de soja por ha plantado', '900 sacas/ano'),
    ('c) campo nativo (bovinos)', 157.5, '3.000 kg de boi por invernada de 100 ha, ao preço da época, pago em dezembro', '4.725 kg de boi/ano (cálculo: 157,5/100 × 3.000)'),
    ('d) matos, capões e pedras', 68.9, 'sem pagamento', '-'),
]
AREAS_SEM_OBRAS = [
    ('01', 'Loteamento Alto das Graças II', 'Caxias do Sul'), ('02', 'Loteamento Santa Corona III', 'Caxias do Sul'),
    ('03', 'Condomínio Morada dos Ipês II', 'Passo Fundo'), ('04', 'Loteamento da Congregação das Irmãos de São Carlos', 'Passo Fundo'),
    ('05', 'Loteamento Bella Visione', 'Nova Prata'), ('06', 'Loteamento Bonetto', 'Farroupilha'),
    ('07', 'Loteamento Pascoal', 'Farroupilha'), ('08', 'Loteamento Letti', 'Farroupilha'),
    ('09', 'Loteamento Colombo', 'Farroupilha'), ('10', 'Loteamento Champagnat', 'Veranópolis'),
    ('11', 'Loteamento Simioni', 'Veranópolis'), ('12', 'Loteamento Sol Nascente', 'Antônio Prado'),
]
# (doc 11 contratados, 2018 contratos/em execução, resposta do grupo item 18 do arquivo 1)
AREAS_X = {
    '01': ('não', '2018 contratados nº 09 (Alto das Graças II – Bellini)', '18.1: contrato firmado, sem urbanístico, LP em processo'),
    '02': ('não', '2018 em execução nº 02 (Santa Corona III – Paulety, 44 lotes Siviero)', '18.2: contrato firmado; ação indenizatória com acordo; em execução'),
    '03': ('não como contratado; doc 11 tem 3 lotes do Morada dos Ipês II a entregar a Rovílio (R$ 1.650.000)', '2018 contratados nº 19 (Zanella II) e em execução nº 05 (Zanella)', '18.3: percentual de 36 %; LI em vigor; empresa estima prejuízo'),
    '04': ('não', 'não', '18.4: só estudo de viabilidade'),
    '05': ('não', '2018 contratados nº 01 (Bella Visione – Truccolo)', '18.5: contrato firmado, 50 % de permuta, risco de rescisão'),
    '06': ('sim (Contratado 06, multa R$ 782.000)', '2018 contratados nº 20', '18.5'),
    '07': ('sim (Contratado 07 Pasqual, multa R$ 1.500.000)', '2018 contratados nº 21', '18.5'),
    '08': ('sim (Contratado 08 Industrial Dr. Letti, multa R$ 1.650.000)', '2018 contratados nº 16', 'não citado no 18'),
    '09': ('não', 'não (arquivo 17 e ACERTOS item 27: distrato com a família Colombo em 05/02/2018)', '18.5: "contratos firmados"'),
    '10': ('não', '2018 contratados nº 03 e 04 (Champagnat I e II – Maristas)', '18.5'),
    '11': ('não como contratado; doc 11 tem o Fracionamento Simioni executado, 4 lotes sem valor', '2018 executados nº 02 (Fracionamento Simioni, 4 lotes)', '18.6: "sem contrato firmado, apenas uma conversação"'),
    '12': ('sim (Contratado 09 Sol Nascente III, multa R$ 1.058.000)', '2018 contratados nº 22 (Sol Nascente III)', '18.5'),
}

# ---------------------------------------------------------------- helpers
HDR = Font(bold=True)
FILL = PatternFill('solid', fgColor='DDEBF7')
MONEY = '#,##0.00'


def sheet(wb, title, header, rows, money_cols=(), widths=None, note=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(italic=True)
        r0 = 3
    for j, h in enumerate(header, 1):
        c = ws.cell(r0, j, h)
        c.font, c.fill = HDR, FILL
        c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.alignment = Alignment(wrap_text=True, vertical='top')
            if j in money_cols and isinstance(v, (int, float)):
                c.number_format = MONEY
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    ws.auto_filter.ref = f'A{r0}:{get_column_letter(len(header))}{r0 + len(rows)}'
    for j, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    return ws


def build_extracoes():
    wb = Workbook()
    wb.remove(wb.active)
    rows = [(n, c, m, 'SEI (cabeçalho das p. 4-6; p. 1-3 sem empresa)', l or None, loc, e, v, 'não informada', t, p)
            for n, c, m, l, loc, e, v, t, p in ACORDOS]
    ws = sheet(wb, 'Acordos', ['nº', 'contraparte', 'motivo', 'empresa', 'lotes', 'loteamento / lote citado', 'valor em espécie (R$)',
                               'valor atribuído a lote no texto (R$)', 'data', 'texto exato', 'página'],
               rows, money_cols=(7, 8), widths=[6, 30, 30, 22, 7, 28, 16, 16, 12, 70, 7],
               note=f'Fonte: {SRC_AC}, p. 1-3, lidas como imagem (sem camada de texto). Conferido com o gdoc e com a extração de 26/10: sem divergência. Nenhum trecho ilegível.')
    n = ws.max_row + 2
    ws.cell(n, 1, 'Resumo dos acordos (p. 3)').font = HDR
    for i, (d, v, p) in enumerate(RESUMO, n + 1):
        ws.cell(i, 2, d)
        ws.cell(i, 3, v).number_format = MONEY if v > 1000 else '0'
        ws.cell(i, 4, p)

    sheet(wb, 'Contratos Boff', ['pagamento', 'beneficiário', 'descrição', 'valor (R$)', 'data', 'forma / contrapartida prevista', 'corrigir no período', 'página'],
          BOFF, money_cols=(4,), widths=[6, 30, 50, 16, 14, 45, 18, 7],
          note=f'Fonte: {SRC_AC}, p. 4-6 (cabeçalho: SIVIERO EMPREENDIMENTOS IMOBILIÁRIOS LTDA.). Forma: pagamentos e depósitos antecipados pela SEI, em espécie; valores nominais.')
    sheet(wb, 'Áreas vendidas', ['comprador', 'área', 'área (ha)', 'ano', 'preço (R$)', 'R$/ha', 'forma de pagamento', 'linha no docx'],
          VENDAS, money_cols=(5, 6), widths=[20, 40, 10, 8, 16, 12, 70, 10],
          note=f'Fonte: {SRC_HT}. Fazenda São Crispin: 686,75 ha − 190,00 ha (Rigo) − 85,32 ha (Pellegrini) = 411,43 ha (P001-P004). Total declarado recebido: R$ 12.058.000,00 (R008).')
    sheet(wb, 'Créditos soja', ['origem', 'safra', 'sacas', 'empresa credora', 'observação', 'linha no docx'],
          SOJA, widths=[28, 22, 10, 50, 50, 12],
          note='O docx deixa o preço em branco ("37.694 sacas de soja X R$ ......" e "13.085 sacas de soja X R$ ......"). Empresa credora: leitura do próprio texto; ver cruzamentos.')
    sheet(wb, 'Arrendamento', ['área', 'hectares', 'preço', 'produto por ano'], ARREND, widths=[28, 10, 60, 40],
          note='Arrendamento à Família Lorenzetti, 3 anos: 02/05/2023 a 30/04/2026 (última safra 2026), área total 411,4 ha (P022-P026). Mesmo texto nos inventários 02 e 05.')
    sheet(wb, 'Áreas sem obras', ['nº', 'nome', 'cidade', 'situação (texto do docx)'],
          [(a, b, c, 'contratada pela SEI para urbanização, sem obras / em estudo de viabilidade ambiental e/ou com obras em execução com saldo zero e possíveis prejuízos a apurar') for a, b, c in AREAS_SEM_OBRAS],
          widths=[5, 50, 16, 80], note='Fonte: segunda parte do docx do Hotel (cabeçalho SEI), P030-R048.')
    s_esp = sum(e or 0 for *_, e, v, t, p in [(a[0], a[5], a[6], a[7], a[8]) for a in ACORDOS])
    s_lotes = sum(a[3] for a in ACORDOS)
    s_boff = round(sum(b[3] for b in BOFF), 2)
    s_vendas = sum(v[4] for v in VENDAS)
    rig = 20150 + 20150 - 2606
    lor = 2606 + 3206 + 3591 + 3682
    conf = [
        ('Lotes dos acordos (itens 1-32, item 9 contado como 1)', 82, s_lotes, s_lotes - 82),
        ('Espécie dos acordos', 2882204.55, round(s_esp, 2), round(s_esp - 2882204.55, 2)),
        ('Total: 80 × 100.000 + 830.000 + 375.000 + espécie', 12087204.55, 8000000 + 830000 + 375000 + 2882204.55, 0.0),
        ('Contratos Boff', BOFF_TOTAL, s_boff, round(s_boff - BOFF_TOTAL, 2)),
        ('Boff: 983.724,02 + 140.000,00 = total firmado', 1123724.02, 983724.02 + 140000, 0.0),
        ('Áreas vendidas', VENDAS_TOTAL, s_vendas, s_vendas - VENDAS_TOTAL),
        ('Rigo: 37.500 + 4 × 20.150 sacas', 118100, 37500 + 4 * 20150, 0),
        ('Rigo: 118.100 sacas × R$ 80,00', 9448000.00, 118100 * 80.0, 0.0),
        ('Rigo: 190,00 ha × R$ 49.726,30', 9448000.00, round(190 * 49726.30, 2), round(190 * 49726.30 - 9448000, 2)),
        ('Pellegrini: 85,32 ha × R$ 23.441,16', 2000000.00, round(85.32 * 23441.16, 2), round(85.32 * 23441.16 - 2000000, 2)),
        ('Créditos Rigo (sacas)', 37694, rig, rig - 37694),
        ('Créditos Lorenzet (sacas)', 13085, lor, lor - 13085),
        ('São Crispin: 686,75 − 190,00 − 85,32 (ha)', 411.43, round(686.75 - 190 - 85.32, 2), round(686.75 - 190 - 85.32 - 411.43, 2)),
        ('Arrendamento: 155 + 30 + 157,5 + 68,9 (ha)', 411.4, round(155 + 30 + 157.5 + 68.9, 2), round(155 + 30 + 157.5 + 68.9 - 411.4, 2)),
        ('Arrendamento lavoura: 155 × 18 (sacas)', 2790, 155 * 18, 0),
        ('Arrendamento batata: 30 × 30 (sacas)', 900, 30 * 30, 0),
    ]
    sheet(wb, 'Conferência', ['total', 'declarado', 'somado/calculado', 'diferença'], conf, money_cols=(2, 3, 4), widths=[55, 18, 18, 14])
    wb.save(OUT + 'extracoes.xlsx')
    return conf


# ---------------------------------------------------------------- cross-checks
ITEM_OF_ROW = {}  # file-17 row -> ACERTOS item, resolving pattern overlaps


def lots_table():
    f17 = ml.load_f17()
    stock, indem = ml.load_stock()
    rows, count = [], {}
    for item, pats in ml.ITEMS.items():
        for r in f17:
            if not any(__import__('re').search(p, r['text']) for p in pats):
                continue
            # Baruffi rows mention "Família Farina" (item 08 pattern) but belong to item 09;
            # the Viero lot negotiated with Razzera (2019) is item 26, not the 2020 Susin terms.
            if item == '08' and 'BARUFFI' in r['text']:
                continue
            if item == '03/04' and r['row'] == 82:
                continue
            # 2016 Via Nova lot (swap for a Caminhos do Sol lot) is not part of the 2020 agreement.
            if item == '14' and r['row'] == 364:
                continue
            # Inês Ansolin Pessutto (4,9 ha, 2024) is a separate purchase, not ACERTOS item 21.
            if item == '21' and 'PESSUTTO' in r['text']:
                continue
            k = r['key']
            st = stock.get(k)
            own = st if st and not str(st[3]).startswith('tabela 47') else None
            ind = indem.get(k)
            f18 = ml.F18_IDX.get(k)
            d = r['date'].strftime('%d/%m/%Y') if hasattr(r['date'], 'strftime') else str(r['date'])
            after = 'sim' if isinstance(r['date'], dt.datetime) and r['date'] > dt.datetime(2024, 12, 31) else 'não'
            before18 = 'sim' if isinstance(r['date'], dt.datetime) and r['date'] <= dt.datetime(2018, 12, 31) else 'não'
            rows.append((item, r['lot'], k[1], k[2], d, r['buyer'], r['obs'], f'arq. 17, linha {r["row"]}, p. {r["page"]}',
                         'sim' if own else 'não', f'{ind[0]} – R$ {ind[1]:,.2f}' if ind else 'não', f18 or 'não', after, before18))
            count[item] = count.get(item, 0) + 1
    return rows, count


def build_cruzamentos(lot_rows, count):
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, 'Lotes dos acordos', ['item ACERTOS', 'loteamento (arq. 17)', 'quadra', 'lote', 'data', 'comprador', 'observação do arq. 17', 'fonte',
                                    'no estoque próprio 2025 (doc 11)?', 'em "lotes de terceiros indenizáveis" (doc 11, tabela 47)?', 'no arquivo 18 (Válter)?',
                                    'depois de 31/12/2024?', 'antes ou em 31/12/2018?'],
          lot_rows, widths=[8, 30, 8, 6, 11, 28, 60, 22, 12, 24, 30, 10, 10],
          note='Lotes do arquivo 17 ligados a cada item do ACERTOS pelo comprador e pela observação. Chave: loteamento + quadra + lote.')
    declared = {'01': 1, '03/04': 24, '06': 11, '08': 17, '09': '1 (trocado por 4)', '11': 4, '12': 1, '13': 4, '14': 9, '17': 2, '21': 1, '25': 3, '26': 1, '27': 2, '32': 1}
    cnt = [(i, declared[i], count.get(i, 0), {
        '03/04': 'Termo A (17 lotes): 14 lotes com data 13/07/2020, ou 15 equivalentes, porque os lotes 05 e 06/Q6364 "equivalem a 03 lotes"; faltam 2. Termo B (7 lotes): só o Q1729 L12 (30/04/2020); faltam 6, que são os "06 lotes" da Obs. 15 do doc 11. As datas são as dos termos, não as das escrituras.',
        '08': '16 lotes do arq. 17 com a observação "Acordo Processo Judicial Indenizatório do Loteamento Industrial Vicentino".',
        '09': '4 lotes a Baruffi Advogados em 16/05/2022 ("honorários – Família Farina"); o resumo do ACERTOS conta 1 lote de R$ 830.000.',
        '11': 'falta 1 lote no arq. 17', '14': 'arq. 17 tem 10 lotes em 07/07/2020 com o texto "num total de 09 lotes". O lote Q6373 L12 de 2016 (permuta por 1 lote, 2 % do Caminhos do Sol) ficou fora: não é deste acordo.',
        '25': 'arq. 17: Q6370 L06-07 pelo Loteamento Basso; Q6369 L08 e Falavigna Q322 L020 por serviços no Reserva Falavigna (não são deste item). Nenhum no Vale das Colinas. O arq. 18 dá Q6369 L08 e Q6371 L11 como os 2 lotes do Basso, e o Q6371 L11 está no estoque do doc 11.',
        '27': 'distrato de 05/02/2018, antes do inventário de 2018',
    }.get(i, '')) for i in declared]
    sheet(wb, 'Contagem de lotes', ['item', 'lotes no ACERTOS', 'lotes achados no arq. 17', 'observação'], cnt, widths=[8, 18, 14, 100])

    passivos = [
        ('Susin', 'ACERTOS itens 3, 4, 5 e 26 (24 lotes + R$ 200.000 + 1 lote)', 'Doc 11 Obs. 15: "restando ainda a entrega de 06 lotes ... e mais o acerto dos valores, que deverão ser pagos como forma de comissão pela intermediação dos contratos firmados com a empresa BURITI" – R$ 1.200.000,00',
         1200000.00, 'Termo B (30/04/2020, arq. 16 p. 5-9): 7 lotes por lotes que a família Susin comprou de Bovo, C. Fioreze e L. Pedrotti, com sub-rogação da SEI; só 1 entregue (Q1729 L12); os 6 pendentes são os da Obs. 15. Termo A (13/07/2020, p. 1-4): 17 lotes + R$ 200.000 por toda a intermediação até a data; cl. 1.1 cita Piccoli, Boff, Andreis, Pedrotti e Fiorese (famílias dos contratos Buriti); cl. 5.1 inclui a comissão "ajustada com a outra parte" desses negócios; cl. 6: quitação total "cumprido integralmente este acordo"; cl. 5: vencimento antecipado após 90 dias; última etapa no 48º mês (07/2024); cl. 8 exclui só intermediações novas, por escrito. Arq. 17: 14 lotes do Termo A (15 equivalentes). Minuta não assinada de 04/02/2020 em 01_fontes/outros. Versão antiga do doc 11: valores em branco.',
         'Os 6 lotes do Termo B são devidos e não entram na quitação. A comissão pelos contratos das famílias do Termo A está dentro do Termo A (cl. 1.1 e 5.1), cuja quitação depende de a SEI cumprir o acordo (faltam 2 lotes equivalentes). Sem documento escrito de intermediação nova (cl. 8), a parte de comissão não tem base.'),
        ('Farina', 'ACERTOS item 8 (17 lotes) – lote Perini Q1469 L11 (arq. 17, 14/08/2019)', 'Doc 11 tabela 47, Indenização 03 – Perini Q1469 L011 – R$ 160.000,00', 160000.00, 'O lote saiu do estoque como pagamento do acordo Farina e volta como dívida de indenização a "terceiros".', 'Pedir ao grupo quem é o "terceiro" indenizável e se o acordo Farina dava quitação.'),
        ('Ambitec', 'ACERTOS item 6 (11 lotes) – lote Jardim Itália Q1747 L02 (arq. 17, 14/07/2021)', 'Doc 11 tabela 47, Indenização 02 – Jardim Itália Q1747 L002 – R$ 200.000,00', 200000.00, 'Mesmo padrão: lote dado em acerto de distrato e depois descontado como indenizável.', 'Idem.'),
        ('Acordi', 'ACERTOS item 29: R$ 88.457,52 (TAC MP) + R$ 61.920,30 (Prefeitura) + R$ 130.000,00 (Município)', 'Doc 11 Obs. 7: multa de ~R$ 80.000 "já pagos" e ~R$ 70.000 ao Município "já quitada"; falta atualização + TAC/PRAD + custas: R$ 100.000,00', 100000.00, 'Doc 11 não cita os R$ 130.000 do "Acerto com o Município".', 'Perguntar se os R$ 130.000 já cobrem o TAC/PRAD; se sim, a Obs. 7 duplica.'),
        ('Celestino Rossi', 'ACERTOS item 18: "A ACERTAR OS HONORÁRIOS" (sem valor)', 'Doc 11 Obs. 13: 3 lotes urbanos = R$ 600.000,00', 600000.00, 'ACERTOS não soma valor; não há duplicidade. Arq. 1, p. 27: "compromisso verbal".', 'Sem duplicidade. O valor de R$ 600.000 depende de compromisso só verbal.'),
        ('Flávia Mânica Siviero', 'ACERTOS item 16: R$ 500.000 (24/50 parcelas pagas)', 'Doc 11: nenhum passivo com esse nome', 260000.00, 'Saldo de 26 parcelas (26 × R$ 10.000, cálculo) não está no doc 11. ACERTOS soma o valor cheio como pago.', 'Sem duplicidade; ACERTOS soma como pago um valor ainda não pago.'),
        ('Basso', 'ACERTOS item 30: R$ 225.160,73 em 40 parcelas, 11 pagas', 'Doc 11: nenhum passivo; arq. 18 lista custos do Loteamento Basso (contrato da HECTA) para "acerto futuro" com Válter', 225160.73, 'Ver aba "Válter e HECTA".', 'Ver achado da HECTA.'),
    ]
    sheet(wb, 'Acordos x passivos', ['contraparte', 'no ACERTOS', 'no doc 11', 'valor no doc 11 / em jogo (R$)', 'evidência', 'leitura'], passivos,
          money_cols=(4,), widths=[14, 45, 55, 16, 80, 50])

    third = [
        ('Industrial Vicentino', '1301', '003', 1355000.00, '19/09/2025', 'Gerson Fernando Fattori: "Vendido à vista, com parte permuta de outro imóvel, no valor de R$ 500.000,00"', 'arq. 17, linha 24, p. 3',
         'Em 31/12/2024 o lote era da SEI. O doc 11 não o lista no estoque próprio e o desconta como lote de terceiro indenizável. A permuta trouxe o lote Reserva Paradiso Q104 L623, que o doc 11 lista no ativo (R$ 500.000) e foi vendido em 07/01/2026 (arq. 17, linha 65).'),
        ('Jardim Itália', '1747', '001', 200000.00, '10/03/2025', 'Erick de Almeida: "Vendido à prazo, em 15x"', 'arq. 17, linha 162, p. 8', 'Em 31/12/2024 o lote era da SEI; vendido em 2025 e descontado como indenizável.'),
        ('Primavera', '1199', '010 e 011', 332000.00, 'fev/2025 (arq. 1) / 27/03/2025 (arq. 17, L11)', 'Safra – São Francisco: dação em pagamento no acordo de distrato', 'arq. 1, p. 26; arq. 17, linha 6, p. 1',
         'Em 31/12/2024 os dois lotes eram da SEI; dados em 2025 para pagar a Safra e descontados como indenizáveis (2 × R$ 166.000).'),
        ('Perini', '1469', '011', 160000.00, '14/08/2019', 'Ricardo e João Farina – acordo Industrial Vicentino', 'arq. 17, linha 18, p. 2', 'Lote do acordo Farina (ACERTOS item 8).'),
        ('Jardim Itália', '1747', '002', 200000.00, '14/07/2021', 'Ambitec – distrato', 'arq. 17, linha 163, p. 8', 'Lote do acordo Ambitec (ACERTOS item 6).'),
        ('Perini', '1469', '018 e 019', 320000.00, '10/10/2024', 'Efetiva de Imóveis: "Vendido à prazo ... Diferença de área constatada do lote"', 'arq. 17, linhas 20-21, p. 3', 'Vendidos antes da data-base; consistentes como lotes de terceiros, mas indenizáveis logo após a venda.'),
        ('Industrial Vicentino', '1302', '004', 1355000.00, '26/02/2014', 'Alcides Basso Imóveis', 'arq. 17, linha 26, p. 3', 'Vendido em 2014; consistente.'),
    ]
    sheet(wb, 'Indenizáveis x arq. 17', ['loteamento', 'quadra', 'lote', 'valor no doc 11, tabela 47 (R$)', 'data da saída', 'destino no arq. 17', 'fonte', 'leitura'], third,
          money_cols=(4,), widths=[18, 8, 10, 16, 18, 50, 22, 80],
          note='Lotes da tabela 47 do doc 11 ("lotes de terceiros passíveis de indenização por caracterização ambiental") que aparecem no arquivo 17.')

    boff = [
        ('ACERTOS p. 4-6', 'Pagamentos antecipados à família Boff, para compensar com 13 lotes (5 + 1 + 3 + 2 + 1 + 1)', BOFF_TOTAL, 'nominal; "corrigir até 2022"'),
        ('Inventário 2018 (investimentos pagos, itens 10-12)', 'Dorval 165.000 + 220.000; Ludovico 1.473.000', 1858000.00, 'tratados como investimento realizado'),
        ('Doc 11 (2025), Contratado 01', 'Caminhos do Sol: "R$ 00.000.000,00"; nenhum ativo e nenhuma multa', 0.0, 'o crédito Boff não aparece'),
        ('Rafael – Caminhos do Sol, Quadro_Resumo', '61 lotes × R$ 141.000 = R$ 8.601.000 (10 % da SEI)', 8601000.00, 'não inclui o crédito nem os 13 lotes'),
        ('Arq. 17, linha 364', 'lote Alto das Graças Q6373 L12 dado à Via Nova em 2016, "PERMUTADO com 01 lote do Loteamento Caminhos do Sol, no percentual de 2,00%"', None, 'mais um direito sobre os lotes do Caminhos do Sol'),
        ('Arq. 1, p. 30, resposta 19 do grupo', '"Os demais créditos futuros adquiridos de áreas e proprietários no município de Caxias do Sul foi e será acertado como moeda de postergação de contratos e compensação em negociação para não instauração de processos judiciais/indenizações."', None, 'o grupo trata os créditos como já usados ("moeda de postergação"), não como ativo; no exemplo Boff, prevê compensação a favor dos Boff'),
    ]
    sheet(wb, 'Boff x Caminhos do Sol', ['fonte', 'o que registra', 'valor (R$)', 'observação'], boff, money_cols=(3,), widths=[38, 90, 16, 40])

    credit = [
        ('Vale das Colinas (Piccoli + Gasparin + Zanrosso)', 'Multa R$ 3.560.000 + despesas R$ 651.035 = R$ 4.211.035', 'Inventário 2018: "Pago despesas do Senhor Arlindo Piccoli – dívidas bancárias e de terceiros em troca de lotes" R$ 300.000; arq. 17 linha 117: lote San Marino Q1342 L03 (2024, "Acordo – Arlindo Píccoli – Loteamento Vale das Colinas. Dívida de R$ 270.000,00"); linha 358: R$ 62.500 a Evandro Zanrosso "pela negociação de 01 lote no Loteamento Vale das Colinas" (2019); linha 357: lote Q6372 L13 a Zanrosso em dação (2024, sem valor)', 632500.00),
        ('Vista da Colina (Zanandréa + Fioreze)', 'Multa R$ 2.556.900 + despesas R$ 108.000 = R$ 2.664.900', 'Inventário 2018: "Compra de lotes da Família Fioreze" R$ 148.000; ACERTOS itens 22-24: Lori, Ezequiel e Cândido Fioreze R$ 410.000 ("compra"); Termo B de Susin (arq. 16): SEI sub-rogada nos lotes que C. Fioreze (4) e E. Bovo (contrato Zanandréa de 14/05/2008) venderam à família Susin, sem valor; ACERTOS itens 1-2 e 13 (Angest, Fochesatto): lotes dados pela não execução na área Fioreze/Zanandréa', 558000.00),
        ('Pedrotti + Andreis + Lorandi', 'Multa R$ 2.280.646 + despesas R$ 308.000 = R$ 2.588.646', 'ACERTOS item 31: Família Pedrotti R$ 270.000 (natureza não informada, a confirmar); Termo B de Susin: SEI sub-rogada no 7º lote comprado de Luiz Pedrotti, sem valor', 270000.00),
    ]
    sheet(wb, 'Créditos x multas', ['loteamento contratado', 'desconto no doc 11', 'créditos/pagamentos já feitos aos mesmos proprietários', 'créditos documentados (R$)'],
          credit, money_cols=(4,), widths=[34, 34, 110, 16],
          note='Critério: só valores em R$ escritos nas fontes. Arq. 1, p. 30: o grupo diz que esses créditos "foi e será acertado como moeda de postergação de contratos e compensação". O doc 11 não lista os créditos e desconta a multa cheia.')

    valter = [
        ('Loteamento Basso (contrato da HECTA)', 'arq. 18, R032: "Feito o DISTRATO, mas a empresa SIVIERO arcou com os seguintes custos": R$ 120.000 Colunata; R$ 260.000 em 2 lotes do Alto das Graças I (lote 08/Q6369 e lote 11/Q6371); R$ 5.000 Colunata; R$ 59.810 Garden; R$ 32.000 Garden-Aditrea; R$ 55.000 Garden (lote do Zanandréa); R$ 2.577,50 e R$ 7.260,19 taxas; R$ 2.500 Flávia Mânica; R$ 39.547,82 IPTU; "FALTA acertar 02 LOTES com a COLUNATA"', 583695.51,
         'ACERTOS item 25 (2 lotes "no Basso" a Colunata) e item 30 (dívida da família Basso na Prefeitura de Caxias, R$ 225.160,73)'),
        ('Lote Altos da Julieta Q1728 L03', 'arq. 18, R005: "Pago Césio Verona, de Farroupilha – serviços prestados para a empresa Rochedo"', None, 'ACERTOS item 17 (Verona) e arq. 17, linha 289: "acerto geral ... nos Loteamento da empresa"'),
        ('Lote Alto das Graças Q6371 L11', 'arq. 18 (Basso, custo 2): "R$ 260.000,00, representados por 02 lotes ... (Lote 08–Quadra 6369 / Lote 11–Quadra 6371)". O lote está no estoque próprio do doc 11 (tabela 38 linha 16, R$ 141.000) e na lista de Rafael (ativos-avaliados-valter.xlsx, Ativos, linha 101). Não está no arq. 17, que dá o Q6369 L08 à Colunata pelo Reserva Falavigna.', 141000.00, 'lote possivelmente já entregue ou comprometido, na lista dos bens que viriam para nós'),
        ('Lotes Alto das Graças Q6371 L05, L06, L12 (consignados)', 'arq. 18, R004: "Área de terras de 2,5 ha em Monte Bérico – Caxias do Sul, com Alexandre Susin. Adquirida a parte do Alexandre Susin em troca de 03 lotes." Nenhum dos 3 está no arq. 17, e o Monte Bérico não aparece em nenhum inventário de 2018 ou 2025.', None, 'lotes da SEI trocados por área que não aparece em inventário'),
        ('Salas 201-203 (HECTA, VMC, URBAN)', 'arq. 18, R007: 5 lotes do Sol Nascente pagos na aquisição de 3 salas comerciais e 3 box em Caxias do Sul: "201 - HECTA 202 - VMC 203 - URBAN"', None, 'indica a HECTA no grupo de empresas de Válter'),
        ('Loteamento Cecconello (HECTA 50 %)', 'Doc 11 Contratado 05: HECTA 50 % + proprietários 50 %; multa R$ 1.300.000 descontada na SEI; SEI tem 2 ha comprados por R$ 600.000 (área rural item 10)', 1300000.00, 'contrato da HECTA com multa na SEI'),
    ]
    sheet(wb, 'Válter e HECTA', ['assunto', 'evidência', 'valor (R$)', 'onde aparece como custo da SEI'], valter, money_cols=(3,), widths=[30, 110, 16, 60])

    acomp = [
        ('Ansolin (itens 19-21) – hipótese', 55000 + 156666 + 375000, 'Hipótese (o item 19 diz "compra de 06 lotes"): é a compra dos 7,0 ha que o doc 11 lista no ativo (área rural item 06: "área de 7,0 hectares ... pelo valor de R$ 586.666,00") e o arq. 17 (linha 153) dá o lote Pedra Lisa Q11 L01 "pela aquisição de 7,0 há ... R$ 586.666,00". 55.000 + 156.666 + 375.000 = 586.666.'),
        ('Fioreze (itens 22-24)', 410000.00, '"compra" de lotes/área dos proprietários do Vista da Colina; não está no ativo do doc 11 (ver aba Créditos x multas).'),
        ('Angest (itens 1-2)', 130500.00, 'arq. 17, linha 353: lote + "R$ 139.500,00, pagos ... em 18 parcelas = R$ 270.000,00". O ACERTOS conta o lote (R$ 100.000) e mais R$ 270.000 em espécie; a espécie paga foi R$ 139.500. Diferença: R$ 130.500.'),
        ('Flávia Mânica (item 16)', 260000.00, 'Somada inteira; 26 de 50 parcelas ainda a pagar (cálculo 26 × 10.000).'),
        ('Basso (item 30)', round(29 * 4473.71 + 29 * 1971.21, 2), 'Somada inteira; 29 de 40 parcelas a pagar = R$ 186.902,68. As 40 parcelas somam R$ 257.796,80, mais que os R$ 225.160,73 declarados; pela conta valor − 11 parcelas, R$ 154.266,61. Aproximação.'),
        ('Susin, 6 lotes do Termo B', 600000.00, 'Contados entre os 80 lotes × R$ 100.000 do resumo, mas não entregues (Obs. 15 do doc 11). Mais 2 lotes equivalentes do Termo A sem registro no arq. 17.'),
    ]
    sheet(wb, 'Composição do ACERTOS', ['item', 'valor (R$)', 'por que não é perda já realizada'], acomp, money_cols=(2,), widths=[24, 16, 120],
          note='Parcelas dos R$ 12.087.204,55 que não são acertos já pagos. O ACERTOS não é deduzido no total do doc 11. O lado consulente cita os 82 lotes do ACERTOS ao perguntar pela queda de 751 para 421 lotes (arq. 1, p. 18).')

    ar = [(n, nome, cid, *AREAS_X[n], 'conflito: distrato em 2018, listada como contratada' if n == '09' else 'só no docx do Hotel e em 2018' if AREAS_X[n][0] == 'não' and AREAS_X[n][1] != 'não' else ('só no docx do Hotel' if AREAS_X[n][0] == 'não' else 'conflito / ver texto' if not AREAS_X[n][0].startswith('sim') else 'no doc 11'))
          for n, nome, cid in AREAS_SEM_OBRAS]
    sheet(wb, 'Áreas sem obras', ['nº', 'área', 'cidade', 'doc 11 (9 contratados)', 'inventário 2018', 'resposta do grupo (arq. 1, item 18)', 'classificação'], ar,
          widths=[5, 40, 14, 45, 50, 50, 22],
          note='Doc 11 contratados que não estão no docx: Caminhos do Sol, Vale das Colinas, Vista da Colina, Pedrotti/Andreis/Lorandi, Cecconello.')

    sc = [
        ('Docx do Hotel (P001-P004)', '686,75 ha − 190,00 ha Rigo − 85,32 ha Pellegrini = 411,43 ha', 411.43),
        ('Doc 05 (inventário do Hotel), item 05', '491,0 ha (10 matrículas; a 00.760 aparece duas vezes, com 48 e 100 ha); vendidos 83,47 ha a Rigo da mat. 00.760; avaliados 407,53 ha × R$ 110.000 = R$ 44.828.300', 407.53),
        ('Doc 05, item 06', 'mat. 11.090, 4,0 ha "dentro de um todo maior", R$ 440.000', 4.0),
        ('Soma doc 05 itens 05 + 06', '407,53 + 4,0', 411.53),
        ('Diferença', '411,53 − 411,43', 0.10),
        ('Rigo: 190,00 (docx) − 83,47 (doc 05)', 'parte da venda Rigo fora das matrículas listadas em 2025', 106.53),
        ('686,75 − 491,0 − 4,0', 'área total do docx que não está nos itens 05 e 06', 191.75),
        ('106,53 + 85,32 (Pellegrini)', 'vendas fora das matrículas listadas', 191.85),
        ('Pellegrini pelas parcelas de 2018', 'sessão 02/05: as 4 parcelas do inventário de 2018 (doc 17 de 2018) somam 4,9 + 10,0 + 26,16 + 44,16 = 85,22 ha; 106,53 + 85,22 = 191,75 ha, igual a 686,75 − 491 − 4. A diferença de 0,10 ha vem dos 85,32 ha declarados.', 85.22),
        ('Doc 02 (Agropecuária)', 'bens: R$ 0,00; diz que as escrituras estão em nome do Hotel e de terceiros', None),
        ('Estudo de Rafael', 'repete o doc 05: 491 ha, 83,47 ha Rigo, 407,53 ha', None),
    ]
    sheet(wb, 'São Crispin', ['fonte', 'o que diz', 'ha'], sc, widths=[38, 110, 10])

    sj = [
        ('Rigo 2021 + 2022 (líquido)', 37694, 'Hotel Princesa dos Vales (provável)', 'título do docx; doc 02 diz que as escrituras estão em nome do Hotel; só 83,47 ha dos 190 ha saem de matrícula listada no doc 05. Natureza indefinida: a tabela diz "FALTA RECEBER", a observação diz "AINDA NÃO FORAM VENDIDOS". Coincidência: as 2.606 sacas vendidas = Lorenzet 2021', 37694 * 80.0),
        ('Lorenzet 2021-2023', 2606 + 3206 + 3591, 'Hotel Princesa dos Vales', 'texto "HOTEL/OLEOPLAN"; o arrendador nos contratos 8 e 9 é o Hotel (p. 1 de cada contrato, segundo o revisor)', (2606 + 3206 + 3591) * 80.0),
        ('Lorenzet 2024', 3682, 'conflito: o docx diz Agropecuária; o contrato 9 (2023-2026), p. 1, tem como ARRENDADORA o CNPJ 89.471.270/0001-94, que é o do Hotel (doc 05); leitura por OCR (Tesseract)', 'texto "AGROPECUÁRIA/OLEPOPLAN"', 3682 * 80.0),
        ('Total', 37694 + 13085, '', 'nenhum inventário de 2025 avalia esses créditos (busca por "soja"/"saca" nos docs 01 a 10: só o texto do arrendamento)', 50779 * 80.0),
        ('Vendidas em 2024 para a escavadeira', 2606, 'saíram do crédito do Hotel (Rigo 2021)', 'a única escavadeira de 2024 nos inventários é a da SEI (320, R$ 887.000; Obs. 20: saldo Caterpillar R$ 347.652,03); nenhum inventário registra crédito do Hotel contra a SEI', 2606 * 80.0),
    ]
    sheet(wb, 'Soja e escavadeira', ['crédito', 'sacas', 'empresa', 'base da atribuição', 'referência a R$ 80,00/saca (preço do contrato Rigo, out/2018) – NÃO é avaliação'], sj,
          money_cols=(5,), widths=[30, 10, 30, 80, 22])

    datas = [(r[0], f'{r[1]} Q{r[2]} L{r[3]}', r[4], r[11]) for r in lot_rows if r[11] == 'sim'] + [
        ('Indenizáveis', 'Industrial Vicentino Q1301 L003', '19/09/2025', 'sim'),
        ('Indenizáveis', 'Jardim Itália Q1747 L001', '10/03/2025', 'sim'),
        ('Indenizáveis', 'Primavera Q1199 L010 e L011', 'fev/2025', 'sim'),
        ('Reserva Paradiso', 'Q104 L623 (no ativo do doc 11)', 'entrada 19/09/2025; venda 07/01/2026', 'sim'),
        ('Fora do estoque do doc 11', 'Barbiero Q05 L03 e L08 (acordo judicial, arq. 17 linhas 53-54)', '23/05/2025', 'sim'),
        ('Fora do estoque do doc 11', 'Falavigna Q174 L020, L021, L023 (Indústria Florense, arq. 17 linhas 387-389)', '20/03/2025', 'sim'),
        ('Fora do estoque do doc 11', 'Falavigna Q175 L012, L013 (Bett, arq. 17 linhas 390-391)', '10/07/2025', 'sim'),
        ('Fora do estoque do doc 11', 'Falavigna Q876 L016 (dação G3, arq. 17 linha 383)', '10/09/2025', 'sim'),
        ('Doc 11 Obs. 15', 'Susin R$ 1.200.000 preenchido na versão de 26/10/2026', 'versão antiga com "......"', 'versão posterior'),
        ('Arrendamento Lorenzetti', 'safras 2025 e 2026', '02/05/2023 a 30/04/2026', 'parcialmente'),
        ('ACERTOS', 'documento sem data; Flávia 24/50 e Basso 11/40 parcelas', 'data do arquivo: nov/2025', 'a confirmar'),
    ]
    sheet(wb, 'Datas', ['assunto', 'bem', 'data', 'depois de 31/12/2024 (P2)?'], datas, widths=[18, 45, 30, 22],
          note='Nenhum lote dos acordos saiu depois de 31/12/2024. Os fatos posteriores estão na tabela 47 e no Reserva Paradiso.')
    wb.save(OUT + 'cruzamentos.xlsx')


if __name__ == '__main__':
    conf = build_extracoes()
    for c in conf:
        print('CONF', c)
    rows, count = lots_table()
    build_cruzamentos(rows, count)
    print('lots', len(rows), count)
    print('own-stock hits', [r[:4] for r in rows if r[8] == 'sim'])
    print('table-47 hits', [r[:4] for r in rows if r[9] != 'não'])
    print('file-18 hits', [r[:4] for r in rows if r[10] != 'não'])
