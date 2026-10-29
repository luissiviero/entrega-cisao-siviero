"""Build planilha-mestre-sei.xlsx from the three SEI snapshots.

Inputs (all under C:/entrega-cisao-siviero):
  02_trabalho/sei-2025-itens.xlsx           inventory 31.12.2024 (group), item + lot level
  02_trabalho/sei-2018-itens.xlsx           inventory dez/2018 (benchmark of existence)
  01_fontes/estudo-rafael/empreendimentos-xlsx/*.xlsx   Rafael's Estudo 2 (BELISE LTDA)
Output:
  02_trabalho/planilha-mestre-sei.xlsx
Run:  PYTHONIOENCODING=utf-8 python -I build_mestre_sei.py
"""

import json
import os
import re
import unicodedata
from collections import Counter, OrderedDict, defaultdict
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path("C:/entrega-cisao-siviero")
F2025 = BASE / "02_trabalho/sei-2025-itens.xlsx"
F2018 = BASE / "02_trabalho/sei-2018-itens.xlsx"
RAF_DIR = BASE / "01_fontes/estudo-rafael/empreendimentos-xlsx"
OUT = Path(os.environ.get("MESTRE_OUT", BASE / "02_trabalho/planilha-mestre-sei.xlsx"))

STOP = {
    "loteamento",
    "fracionamento",
    "condominio",
    "desmembramento",
    "rs",
    "localizado",
    "localizada",
    "na",
    "cidade",
    "de",
    "da",
    "do",
    "das",
    "dos",
    "em",
    "a",
    "o",
    "no",
    "com",
    "area",
    "total",
    "saldo",
    "lotes",
    "lote",
    "e",
    "familia",
    "familias",
    "fechado",
    "item",
    "reavaliacao",
    "valter",
    "azul",
    "pdf",
    "bem",
    "selecionado",
    "por",
    "p",
    "doc",
    "grafado",
    "letra",
    "industrial",
    "aditivo",
}
CITIES = [
    "antonio prado",
    "nova prata",
    "passo fundo",
    "veranopolis",
    "carlos barbosa",
    "farroupilha",
    "cotipora",
    "bento goncalves",
    "gramado",
    "caxias do sul",
    "guapore",
    "flores da cunha",
    "fagundes varela",
    "sao marcos",
    "vila flores",
    "sao jorge",
    "garibaldi",
    "serafina correa",
    "nova bassano",
    "protasio alves",
    "vacaria",
]
ALIAS = [  # (substring in normalised name, canonical tokens)
    ("salvador", ("salvador",)),
    ("femaca", ("salvador",)),
    ("falavig", ("falavigna",)),
    ("dalzochio", ("dalzochio",)),
    ("julieta", ("altos", "julieta")),
    ("jardim italia", ("jardim", "italia")),
    ("solar norte", ("solar", "norte")),
    ("pedra lisa", ("pedra", "lisa")),
    ("alto das gracas i", ("alto", "gracas", "i")),
    ("alto das gracas ii", ("alto", "gracas", "ii")),
    ("medianeira", ("medianeira",)),
    ("senhora da saude", ("saude",)),
    ("sol nascente i", ("sol", "nascente", "i")),
    ("sol nascente ii", ("sol", "nascente", "ii")),
    ("sol nascente iii", ("sol", "nascente", "iii")),
]


def norm(s):
    if s is None:
        return ""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"[^a-z0-9 ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def canon(name):
    """Canonical token tuple of a loteamento name (roman numerals kept)."""
    if not name:
        return ()
    raw = re.split(r"[\(\[]", str(name))[0]  # drop '(saldo de ...)' and '[azul ...]'
    raw = re.sub(r"\s[–—-]\s*item\s*\d+.*$", "", raw)
    n = norm(raw)
    for c in CITIES:
        n = n.replace(c, " ")
    n = re.sub(r"\s+", " ", n).strip()
    for sub, toks in ALIAS:
        # alias ending in a roman numeral must match at a word boundary on both
        # sides; the others are prefixes (Falavigna / Falavigha)
        tail = r"\b" if re.search(r" (i|ii|iii)$", sub) else ""
        if re.search(rf"\b{re.escape(sub)}{tail}", n):
            return toks
    toks = tuple(t for t in n.split() if t not in STOP and not t.isdigit())
    return toks


ROMANS = {"i", "ii", "iii", "iv"}


def variants(toks):
    """'monte veneto i ii' -> {('monte','veneto','i'), ('monte','veneto','ii')}"""
    base = tuple(t for t in toks if t not in ROMANS)
    nums = [t for t in toks if t in ROMANS]
    if len(nums) <= 1:
        return {toks}
    return {base + (n,) for n in nums}


def _match(sa, sb):
    short, long_ = (sa, sb) if len(sa) <= len(sb) else (sb, sa)
    if sa == sb:
        return True
    return set(short) <= set(long_) and (len(short) >= 2 or short[0] == long_[0])


def names_match(a, b):
    sa, sb = canon(a), canon(b)
    if not sa or not sb:
        return False
    return any(_match(x, y) for x in variants(sa) for y in variants(sb))


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("R$", "").replace(" ", "")
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})*(,\d+)?", s):
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def rows(path, sheet):
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)[sheet]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    out = []
    for r in it:
        if all(v is None for v in r):
            continue
        out.append(dict(zip(header, r)))
    return out


def col(d, prefix):
    for k in d:
        if k.startswith(prefix):
            return d[k]
    return None


# ---------------------------------------------------------------- 2025
i25 = rows(F2025, "Itens")
l25 = rows(F2025, "Lotes")
p25 = rows(F2025, "Passivos")

BLOCK25 = [
    ("veiculos", "VEIC"),
    ("maquinas", "MAQ"),
    ("lotes urbanos", "LURB"),
    ("areas urbanas", "AURB"),
    ("areas rurais", "ARUR"),
    ("lotes loteamentos", "LOT"),
    ("areas de terras contratadas", "CONTR"),
]


def bkey25(b):
    nb = norm(b)
    return next((v for k, v in BLOCK25 if k in nb), "?")


for r in i25:
    r["_bk"] = bkey25(r["bloco/seção"])
    r["_item"] = re.sub(r"[^0-9]", "", str(r["nº item"] or "")).zfill(2)
    r["_val"] = num(r["valor total (R$)"])
    r["_key"] = f"SEI-2025-{r['_bk']}-{r['_item']}"

# ---------------------------------------------------------------- 2018
i18 = rows(F2018, "Itens")


def bkey18(b):
    nb = norm(b)
    if "nao registrados" in nb and "execucao" not in nb:
        return "LOT-NR"
    if "em execucao" in nb:
        return "LOT-EX"
    if "contratados" in nb:
        return "CONTR"
    if "registrados" in nb:
        return "LOT"
    if "lotes urbanos" in nb:
        return "LURB"
    if "areas" in nb:
        return "ARUR"
    if "veicul" in nb:
        return "VEIC"
    if "maquin" in nb:
        return "MAQ"
    if "imoveis" in nb:
        return "IMOV"
    if "relacionados" in nb:
        return "OUTRAS"
    return "?"


for r in i18:
    r["_bk"] = bkey18(r["bloco/seção do documento"])
    r["_item"] = re.sub(r"[^0-9]", "", str(r["nº item (doc.)"] or "")).zfill(2)
    r["_val"] = num(col(r, "valor total"))
    r["_qty"] = num(col(r, "qtde lotes (saldo"))
    r["_key"] = (
        f"SEI-2018-OUTRAS-s{r['nº seq']}"  # item numbers repeat in this block
        if r["_bk"] == "OUTRAS"
        else f"SEI-2018-{r['_bk']}-{r['_item']}"
    )
    r["_used"] = False
    r["_name"] = r.get("loteamento") or r["descrição completa"]

# ---------------------------------------------------------------- Rafael
raf_items, raf_eval, raf_lots = [], [], []
RAF_BLOCK = [
    ("LOTES LOTEAMENTOS EXECUTADOS", "LOT"),
    ("LOTES URBANOS", "LURB"),
    ("AREAS RURAIS", "ARUR"),
    ("AREAS CONTRATADAS", "CONTR"),
    ("LOTEAMENTO BONETTO", "CONTR"),
    ("LOTEAMENTO COLINA DO RETIRO II", "LOT"),
    ("LOTES DE TERCEIROS", "PASS"),
    ("DEMIAS COMPROMISSOS", "PASS"),
]
for p in sorted(RAF_DIR.glob("*.xlsx")):
    if p.name.startswith("00 -"):
        continue
    doc = re.sub(r"[,\s-]*BELISE LTDA$", "", p.stem).strip()
    bk = next((v for k, v in RAF_BLOCK if norm(p.name).startswith(norm(k))), "?")
    sn = openpyxl.load_workbook(p, read_only=True).sheetnames
    for r in rows(p, "Balanco_Grupo"):
        item = (
            re.sub(r"[^0-9]", "", str(r.get("Item") or "")).zfill(2)
            if r.get("Item")
            else ""
        )
        desc = str(r.get("Descrição") or "")
        if desc.startswith("[Balanço patrimonial executado por Valter"):
            continue  # Rafael's own re-statement; handled through Avaliacao
        raf_items.append(
            {
                "doc": doc,
                "bk": bk,
                "item": item,
                "desc": desc,
                "area": num(r.get("Área (m²)")),
                "matr": r.get("Matrícula"),
                "val": num(r.get("Valor (R$)")),
                "sinal": r.get("Sinal"),
                "pag": r.get("Página"),
            }
        )
    if "Avaliacao" in sn:
        for r in rows(p, "Avaliacao"):
            it = str(r.get("Item") or "")
            item = re.sub(r"[^0-9]", "", it.split("(")[0]).zfill(2) if it else ""
            raf_eval.append(
                {
                    "doc": doc,
                    "bk": bk,
                    "item": item,
                    "desc": str(r.get("Descrição") or ""),
                    "vgrupo": num(r.get("Valor Grupo (R$)")),
                    "vmin": num(r.get("Faixa mín. (R$)")),
                    "vmax": num(r.get("Faixa máx. (R$)")),
                    "vaval": num(r.get("Valor avaliado (R$)")),
                    "obs": r.get("Observação"),
                    "pag": r.get("Página"),
                    "comp": "(comp" in it,
                }
            )
    if "Lotes" in sn:
        for r in rows(p, "Lotes"):
            raf_lots.append(
                {
                    "doc": doc,
                    "bk": bk,
                    "lot": r.get("Loteamento"),
                    "cidade": r.get("Cidade"),
                    "lote": r.get("Lote"),
                    "quadra": r.get("Quadra"),
                    "val": num(r.get("Valor (R$)")),
                    "pag": r.get("Página"),
                    "_used": False,
                }
            )

raf_by = defaultdict(list)
for r in raf_items:
    raf_by[(r["bk"], r["item"])].append(r)
rafeval_by = defaultdict(list)
for r in raf_eval:
    if not r["comp"]:
        rafeval_by[(r["bk"], r["item"])].append(r)


# ---------------------------------------------------------------- matching 2025 -> 2018
def sim(a, b):
    return SequenceMatcher(None, norm(a), norm(b)).ratio()


def match_desc18(desc25, cands, matr=None, thr=0.80):
    """Return (match, score). Below thr nothing is asserted; the best candidate is
    returned as (None, score, best) through the third element."""
    best, bs = None, 0
    for r in cands:
        if r["_used"]:
            continue
        if (
            matr
            and r.get("matrícula")
            and norm(matr)
            and norm(matr) == norm(r["matrícula"])
        ):
            return r, 1.0, r
        s = sim(desc25[:120], (r["descrição completa"] or "")[:120])
        if s > bs:
            best, bs = r, s
    return (best, bs, best) if bs >= thr else (None, bs, best)


by_bk18 = defaultdict(list)
for r in i18:
    by_bk18[r["_bk"]].append(r)

CASAMENTOS = json.loads(
    (BASE / "02_trabalho/decisoes-sei.json").read_text(encoding="utf-8")
).get("casamentos", {})
CASAMENTOS = {k: v for k, v in CASAMENTOS.items() if not k.startswith("_")}
I18_BY_KEY = {r["_key"]: r for r in i18}
for _k, _v in CASAMENTOS.items():
    if _v is not None:
        I18_BY_KEY[_v]["_used"] = (
            True  # reserve forced targets before automatic matching
        )
master = []
for r in i25:
    bk = r["_bk"]
    m = {
        "empresa": "SEI",
        "bloco": r["bloco/seção"],
        "item": r["_item"],
        "chave": r["_key"],
        "desc": r["descrição completa"],
        "lote": r.get("loteamento / denominação"),
        "cidade": r.get("cidade"),
        "matr": r.get("matrícula"),
        "area": r.get("área"),
        "un": r.get("unidade da área"),
        "qty25": r.get("quantidade de lotes"),
        "v25": r["_val"],
        "obs25": r.get("observação do documento / nota"),
        "fonte25": r.get("ordem no documento"),
        "_bk": bk,
    }
    cands = []
    if bk in ("LOT", "CONTR"):
        pool = (
            (by_bk18["LOT"] + by_bk18["LOT-NR"] + by_bk18["LOT-EX"])
            if bk == "LOT"
            else (by_bk18["CONTR"] + by_bk18["LOT-EX"])
        )
        name25 = m["lote"] or m["desc"]
        cands = [c for c in pool if names_match(name25, c["_name"])]
        if "monte veneto" in norm(name25):  # same name in two cities
            cands = [
                c
                for c in cands
                if not c.get("cidade") or norm(c["cidade"]) == norm(m["cidade"] or "")
            ]
    else:
        pool18 = {
            "VEIC": "VEIC",
            "MAQ": "MAQ",
            "LURB": "LURB",
            "AURB": "ARUR",
            "ARUR": "ARUR",
        }.get(bk)
        pool = (
            by_bk18.get(pool18, [])
            + (by_bk18["IMOV"] if bk == "AURB" else [])
            + (by_bk18["OUTRAS"] if bk in ("LURB", "AURB", "ARUR") else [])
        )
        c, s, best = match_desc18(m["desc"], pool, m["matr"])
        cands = [c] if c else []
        m["sim18"] = round(s, 2)
        if c and s < 0.9:
            m["revisar"] = "Revisar casamento 2018 (similaridade baixa)"
        elif not c and best is not None and s >= 0.6:
            m["revisar"] = (
                f"Sem casamento seguro; candidato 2018 mais parecido: "
                f"{best['bloco/seção do documento'][:25]} #{best['_item']} ({(best['descrição completa'] or '')[:70]}) sim {s:.2f}"
            )
    ov = CASAMENTOS.get(m["chave"], "auto")
    if ov != "auto":
        cands = [] if ov is None else [I18_BY_KEY[ov]]
        m["sim18"] = "par definido no pacote 4"
        m["revisar"] = None
    if cands:
        for c in cands:
            c["_used"] = True
        m["existe18"] = "Sim"
        m["item18"] = "; ".join(
            f"{c['bloco/seção do documento'][:28]} #{c['_item']}" for c in cands
        )
        q = [c["_qty"] for c in cands if c["_qty"] is not None]
        m["qty18"] = sum(q) if q else None
        v = [c["_val"] for c in cands if c["_val"] is not None]
        m["v18"] = sum(v) if v else None
        m["fonte18"] = "; ".join(
            str(c["ordem no documento (tabela-linha)"]) for c in cands
        )
        m["desc18"] = " | ".join((c["descrição completa"] or "")[:90] for c in cands)
        other = [c for c in cands if c["_bk"] == "OUTRAS"]
        if other:
            emp = other[0].get("relacionado no inventário de outra empresa")
            m["existe18"] = f"Sim, fora do total da SEI (inventário de {emp})"
            m["revisar"] = (
                f"Em 2018 este bem constava no inventário de {emp}, fora do total da SEI; "
                "em 2025 entra no total da SEI. Conferir dupla contagem na Parte 6."
            )
    else:
        m["existe18"] = "Não"
    # ---- Rafael
    ri = raf_by.get((bk, r["_item"]), [])
    re_ = rafeval_by.get((bk, r["_item"]), [])
    if ri or re_:
        m["existeR"] = "Sim"
        m["docR"] = "; ".join(sorted({x["doc"] for x in ri + re_}))
        pos = [
            x["val"] for x in ri if x["val"] is not None and (x["sinal"] or "+") == "+"
        ]
        m["vR_grupo"] = pos[0] if pos else None
        ev = [x["vaval"] for x in re_ if x["vaval"] is not None]
        m["vR_aval"] = ev[0] if ev else None
        m["obsR"] = " | ".join(str(x["obs"])[:200] for x in re_ if x["obs"])
        m["fonteR"] = "; ".join(f"{x['doc'][:30]} p.{x['pag']}" for x in (ri + re_)[:3])
        neg = [x for x in ri if (x["sinal"] or "") == "-"]
        if neg:
            m["passR"] = "; ".join(f"{x['desc'][:60]} = {x['val']}" for x in neg)
    else:
        m["existeR"] = "Não"
    v25, vg, va = m["v25"], m.get("vR_grupo"), m.get("vR_aval")
    if v25 is None and vg is None and va is None:
        sit = "Sem valor nas duas fontes"
    elif m["existeR"] == "Não":
        sit = "Não abordado por Rafael"
    elif v25 is None and vg is None and va is None:
        sit = "Sem valor nas duas fontes"
    elif va is not None:
        sit = "Concorda" if abs(va - (v25 or 0)) < 1 else "Discorda"
    elif vg is not None:
        sit = (
            "Transcrito por Rafael sem avaliação própria"
            if abs(vg - (v25 or 0)) < 1
            else "Transcrição de Rafael diverge do inventário"
        )
    else:
        sit = "Rafael só trata do passivo deste item"
    m["sit"] = sit
    m["dif"] = (va - v25) if (va is not None and v25 is not None) else None
    m["difp"] = (m["dif"] / v25) if (m["dif"] is not None and v25) else None
    if vg is not None and v25 is not None and abs(vg - v25) >= 1:
        m["transc"] = (
            f"Rafael transcreve R$ {vg:,.2f} como valor do grupo; inventário diz R$ {v25:,.2f}"
        )
    master.append(m)

# 2018 items that did not match any 2025 item
gone = []
for r in i18:
    if r["_used"]:
        continue
    if r["_bk"] == "OUTRAS":
        sit = "Listado em 2018 na SEI como bem de outra empresa; não consta no inventário 2025 da SEI"
    elif r["_bk"] == "CONTR":
        sit = "Contrato de 2018 não listado em 2025 (Parte 5)"
    elif (
        r["_bk"] in ("LOT", "LOT-NR", "LOT-EX")
        and (r["_qty"] or 0) == 0
        and (r["_val"] or 0) == 0
    ):
        sit = "Saldo zero já em 2018 (não é saída)"
    else:
        sit = "Existia em 2018 e não consta em 2025 (destino a apurar)"
    gone.append(
        {
            "empresa": "SEI",
            "bloco": f"[2018] {r['bloco/seção do documento']}",
            "item": r["_item"],
            "chave": r["_key"],
            "desc": r["descrição completa"],
            "lote": r.get("loteamento"),
            "cidade": r.get("cidade"),
            "matr": r.get("matrícula"),
            "qty18": r["_qty"],
            "v18": r["_val"],
            "existe18": "Sim",
            "existeR": "?",
            "sit": sit,
            "fonte18": r["ordem no documento (tabela-linha)"],
            "obs25": r.get("observação textual do documento"),
        }
    )

only_raf, seen = [], set()
for x in raf_items:
    k = (x["bk"], x["item"])
    if x["bk"] == "PASS" or not x["item"] or k in seen:
        continue
    if any(mm["_bk"] == x["bk"] and mm["_item"] == x["item"] for mm in i25):
        continue
    seen.add(k)
    only_raf.append(
        {
            "empresa": "SEI",
            "bloco": f"[Rafael] {x['doc']}",
            "item": x["item"],
            "chave": f"RAF-{x['bk']}-{x['item']}",
            "desc": x["desc"],
            "vR_grupo": x["val"],
            "existe18": "?",
            "existeR": "Sim",
            "sit": "Só no estudo de Rafael",
            "fonteR": f"{x['doc']} p.{x['pag']}",
        }
    )


# ---------------------------------------------------------------- lots
def lotkey(lot, quadra, lote):
    q = norm(re.split(r"[\[\(]", str(quadra or ""))[0]).lstrip("0") or "0"
    lt = norm(re.split(r"[\[\(]", str(lote or ""))[0])
    lt = re.sub(r"^o(?=\d)", "", lt).lstrip("0") or "0"
    return (canon(lot), q, lt)


raf_lot_by = defaultdict(list)
for x in raf_lots:
    raf_lot_by[lotkey(x["lot"], x["quadra"], x["lote"])].append(x)
lots_out = []
for r in l25:
    k = lotkey(r["loteamento"], r["quadra"], r["lote"])
    rl = raf_lot_by.get(k, [])
    for x in rl:
        x["_used"] = True
    v25 = num(r.get("valor (R$)"))
    vals = [x["val"] for x in rl if x["val"] is not None]
    reav = [
        x["val"] for x in rl if x["val"] is not None and "reavalia" in norm(x["lot"])
    ]
    vr = reav[0] if reav else (vals[0] if vals else None)
    if not rl:
        sit = "Só no inventário 2025"
    elif v25 is None and vr is None:
        sit = "Sem valor nas duas fontes"
    elif v25 is not None and vr is not None and abs(v25 - vr) < 1:
        sit = "Mesmo valor"
    elif v25 is None or vr is None:
        sit = "Valor em só uma das fontes"
    else:
        sit = "Valor diferente"
    extra = None
    if len(set(vals)) > 1:
        extra = "Rafael traz valores diferentes para o mesmo lote: " + "; ".join(
            f"{x['doc'][:25]}={x['val']}" for x in rl
        )
    lots_out.append(
        [
            r["loteamento"],
            r["cidade"],
            r["quadra"],
            r["lote"],
            r.get("matrícula"),
            v25,
            r.get("incluído no subtotal"),
            r.get("observação"),
            r.get("bloco"),
            r.get("item de origem"),
            r.get("ordem no documento"),
            "Sim" if rl else "Não",
            "; ".join(sorted({x["doc"][:40] for x in rl})) or None,
            vr,
            "; ".join(str(x["pag"]) for x in rl) or None,
            (vr - v25) if (vr is not None and v25 is not None) else None,
            sit,
            len(rl),
            extra,
        ]
    )
for x in raf_lots:
    if not x["_used"]:
        lots_out.append(
            [
                x["lot"],
                x["cidade"],
                x["quadra"],
                x["lote"],
                None,
                None,
                None,
                None,
                f"[Rafael] {x['bk']}",
                None,
                None,
                "Sim",
                x["doc"][:40],
                x["val"],
                x["pag"],
                None,
                "Só no estudo de Rafael",
                1,
                None,
            ]
        )

# ---------------------------------------------------------------- loteamentos 2018 x 2025 (counts)
cnt = OrderedDict()
for r, m in zip(i25, master):
    if r["_bk"] in ("LOT", "CONTR"):
        c = canon(m["lote"] or m["desc"])
        cnt[m["chave"]] = {
            "lote": m["lote"],
            "cidade": m["cidade"],
            "bloco": r["_bk"],
            "q25": num(m["qty25"]),
            "v25": m["v25"],
            "q18": m.get("qty18"),
            "v18": m.get("v18"),
            "item18": m.get("item18"),
            "qR": len(
                {
                    lotkey(x["lot"], x["quadra"], x["lote"])
                    for x in raf_lots
                    if x["bk"] in ("LOT", "CONTR")
                    and canon(x["lot"]) == c
                    and (
                        norm(x["cidade"] or "") == norm(m["cidade"] or "")
                        or "veneto" not in " ".join(c)
                    )
                }
            ),
            "l25": sum(
                1
                for x in l25
                if canon(x["loteamento"]) == c
                and norm(x["cidade"] or "") == norm(m["cidade"] or "")
                and str(x.get("bloco") or "").startswith("Lotes / Loteamentos")
            ),
            "obs": (
                "Quantidade 2025 = total do loteamento contratado, não a parte da SEI; não comparável com 2018"
                if r["_bk"] == "CONTR"
                else None
            ),
        }
for g in gone:
    if g["bloco"].startswith("[2018] Loteamentos"):
        cnt[g["chave"]] = {
            "lote": g["lote"],
            "cidade": g["cidade"],
            "bloco": g["bloco"][7:],
            "q25": None,
            "v25": None,
            "q18": g["qty18"],
            "v18": g["v18"],
            "item18": f"#{g['item']}",
            "qR": 0,
            "l25": 0,
            "obs": g["sit"],
        }

# ---------------------------------------------------------------- decisions
DEC = json.loads((BASE / "02_trabalho/decisoes-sei.json").read_text(encoding="utf-8"))
RULES = DEC["regras"]
R1_EXCL = {"Lago Negro", "Alto das Graças I"}
R3_EXCL = {"SEI-2025-AURB-04", "SEI-2025-VEIC-14", "SEI-2025-MAQ-14"}


def rule_for_item(m):
    k, bk, sit = m["chave"], m.get("_bk"), m["sit"]
    if k == "SEI-2025-VEIC-14":
        return "R2"  # Ford Ranger: same value in Rafael's PDF Diretrizes
    if sit == "Concorda" and bk == "LOT" and m.get("lote") not in R1_EXCL:
        return "R1"
    if sit == "Concorda" and bk in ("LURB", "ARUR"):
        return "R2"
    if sit == "Não abordado por Rafael" and k not in R3_EXCL:
        return "R3"
    if sit == "Sem valor nas duas fontes":
        return "R4"
    if sit == "Saldo zero já em 2018 (não é saída)":
        return "R5"
    return None


def rule_for_lot(row):
    sit, bloco, incl = row[16], str(row[8] or ""), row[6]
    if (
        sit == "Mesmo valor"
        and bloco.startswith("Lotes / Lot")
        and incl == "sim"
        and row[0] not in R1_EXCL
    ):
        return "R1"
    # APP lots were removed from R4 on 2026-10-28: Rafael values them at item level.
    return None


def apply(target, rid):
    if rid and rid in RULES:
        target["decisao"] = RULES[rid]["decisao"]
        target["motivo"] = (
            f"{RULES[rid]['motivo']} Aprovada em {RULES[rid]['aprovada_em']}."
        )
        if rid == "R5":
            target["destino"] = "Não se aplica (saldo zero em 2018)"


# Rafael's selection list (assets that will come to the exiting partners, P10/C0)
SEL_ITEMS, SEL_LOTS = set(), set()
SEL_VAL = {}  # chave -> value in the Diretrizes PDF
SEL_LOTS_TAB = set()  # lots only in the per-development 'Lotes escolhidos' tables
_sel = openpyxl.load_workbook(
    BASE / "01_fontes/outros/ativos-avaliados-valter.xlsx",
    read_only=True,
    data_only=True,
)
for r in _sel["Ativos"].iter_rows(values_only=True):
    if not (r[1] and "Siviero Empreendimentos" in str(r[1])):
        continue
    cat, item = str(r[2] or ""), str(r[3] or "").zfill(2)
    if cat == "Lote em loteamento":
        name = re.sub(r"^\d+\.\s*", "", str(r[5] or "")).strip()
        # spelling in the Diretrizes PDF: "Colina do Retiro" = item 07 (Colina do
        # Retiro I); "Altos da Graça I" = item 29 (Alto das Graças I)
        name = {
            "Colina do Retiro": "Colina do Retiro I",
            "Loteamento Altos da Graça I": "Alto das Graças I",
        }.get(name, name)
        SEL_LOTS.add(lotkey(name, r[7], r[6]))
    elif cat == "Veículo":
        SEL_ITEMS.add(f"SEI-2025-VEIC-{item}")
        SEL_VAL[f"SEI-2025-VEIC-{item}"] = num(r[10])
    elif cat == "Imóvel urbano":
        _k = {"02": "SEI-2025-ARUR-02", "04": "SEI-2025-AURB-04"}.get(
            item, f"SEI-2025-LURB-{item}"
        )
        SEL_ITEMS.add(_k)
        SEL_VAL[_k] = num(r[10])
for _p in RAF_DIR.glob("LOTES LOTEAMENTOS*.xlsx"):
    for _r in rows(_p, "Quadro_Resumo"):
        _q, _d = str(_r.get("Quadro") or ""), str(_r.get("Descrição") or "")
        _m = re.match(r"Lote (\S+), quadra (\S+)", _d)
        if _q.startswith("Lotes escolhidos") and _m:
            _lot = re.sub(r"^.*?[–-]\s*", "", _q, count=1)
            _key = lotkey(_lot, _m.group(2), _m.group(1))
            if _key not in SEL_LOTS:
                SEL_LOTS_TAB.add(_key)
# Lots taken off Rafael's list by an approved decision (e.g. S4 of the ACERTOS analysis)
SEL_SUSP = {}
for _k, _v in DEC.get("lista_rafael_suspensos", {}).items():
    if _k.startswith("_"):
        continue
    _lot, _q, _lt = _k.split("|")
    _key = lotkey(_lot, _q, _lt)
    assert _key in SEL_LOTS, _k
    SEL_LOTS.discard(_key)
    SEL_SUSP[_key] = _v
SEL_BY_LOT = Counter(k[0] for k in SEL_LOTS)
# Rafael values from the Diretrizes PDF for items his development workbooks do not value
for m in master:
    if (
        m["chave"] in SEL_VAL
        and m.get("vR_aval") is None
        and SEL_VAL[m["chave"]] is not None
    ):
        m["vR_aval"] = SEL_VAL[m["chave"]]
        m["existeR"] = "Sim"
        m["docR"] = "; ".join(x for x in (m.get("docR"), "PDF Diretrizes") if x)
        if m.get("v25") is not None:
            m["dif"] = m["vR_aval"] - m["v25"]
            m["difp"] = m["dif"] / m["v25"] if m["v25"] else None
            m["sit"] = "Concorda" if abs(m["dif"]) < 1 else "Discorda"
# Contracted loteamentos: Rafael's asset value sits in Quadro_Resumo ('Patrimônio total')
CONTR_DOC = [
    ("CAMINHOS DO SOL", "01"),
    ("VALE DAS COLINAS", "02"),
    ("VISTA DA COLINA", "03"),
    ("PEDROTTI", "04"),
    ("CECCONELLO", "05"),
    ("PASQUAL", "07"),
    ("DR LETTI", "08"),
]
for _p in RAF_DIR.glob("*.xlsx"):
    _n = norm(_p.name).upper()
    _it = next((i for k, i in CONTR_DOC if norm(k).upper() in _n), None)
    if not _it:
        continue
    _tot = [
        _r.get("Valor (R$)")
        for _r in rows(_p, "Quadro_Resumo")
        if str(_r.get("Descrição") or "").startswith("Patrimônio total")
    ]
    _items = (
        [_it, "09"] if _it == "08" else [_it]
    )  # Dr. Letti doc also covers Sol Nascente III
    for _i in _items:
        m = next((x for x in master if x["chave"] == f"SEI-2025-CONTR-{_i}"), None)
        if m is None or not _tot:
            continue
        m["vR_aval"] = num(_tot[0])
        m["existeR"] = "Sim"
        m["docR"] = _p.stem[:60]
        m["fonteR"] = f"{_p.stem[:40]} Quadro_Resumo, Patrimônio total"
        m["sit"] = (
            "Discorda" if (m["vR_aval"] or 0) > 0 else "Sem valor nas duas fontes"
        )
        m["obsR"] = (
            "Valor de referência de Rafael para a participação da SEI (não é laudo)."
        )


def valor_rule(rid, v25):
    return {"R1": v25, "R2": v25, "R3": v25, "R4": 0}.get(rid)


for m in master + gone + only_raf:
    rid = rule_for_item(m)
    apply(m, rid)
    m["vadot"] = valor_rule(rid, m.get("v25"))
    if m["chave"] in SEL_ITEMS:
        m["lista"] = "Sim"
    elif m.get("_bk") == "LOT" and SEL_BY_LOT.get(canon(m.get("lote") or "")):
        m["lista"] = f"Parcial: {SEL_BY_LOT[canon(m['lote'])]} lote(s)"
    else:
        m["lista"] = "Não" if m["chave"].startswith("SEI-2025") else None
    it = DEC["itens"].get(m["chave"])
    if it:
        m["decisao"], m["motivo"] = (
            it["decisao"],
            f"{it['motivo']} {'Aprovada em ' + it['aprovada_em'] + '.' if it.get('aprovada_em') else 'Pendente.'}",
        )
        if "valor_adotado" in it:
            m["vadot"] = it["valor_adotado"]
        for f in ("destino", "regdb"):
            if it.get(f):
                m[f] = it[f]
LOTDEC = {}
for k, v in DEC.get("lotes", {}).items():
    lot, q, lt = k.split("|")
    LOTDEC[lotkey(lot, q, lt)] = v
for row in lots_out:
    d = {}
    rid = rule_for_lot(row)
    apply(d, rid)
    vad = valor_rule(rid, row[5])
    k = lotkey(row[0], row[2], row[3])
    ld = LOTDEC.get(k)
    if ld:
        d["decisao"], d["motivo"] = (
            ld["decisao"],
            f"{ld['motivo']} {'Aprovada em ' + ld['aprovada_em'] + '.' if ld.get('aprovada_em') else 'Pendente.'}",
        )
        vad = ld.get("valor_adotado")
    elif not d and str(row[8] or "").startswith("Lotes / Lot"):
        # item-level decision that says its lots keep the group's value
        parent = next(
            (
                v
                for m in master
                if m.get("_bk") == "LOT"
                and canon(m.get("lote") or "") == k[0]
                and (v := DEC["itens"].get(m["chave"]))
                and v.get("lotes") == "valor do grupo"
            ),
            None,
        )
        if parent:
            d["decisao"] = parent["decisao"]
            d["motivo"] = (
                f"{parent['motivo']} {'Aprovada em ' + parent['aprovada_em'] + '.' if parent.get('aprovada_em') else 'Pendente.'}"
            )
            vad = row[5]
    row += [
        d.get("decisao"),
        d.get("motivo"),
        "Sim"
        if k in SEL_LOTS
        else (
            f"Suspenso da lista: {SEL_SUSP[k]}"
            if k in SEL_SUSP
            else (
                "Sim, só nas tabelas de lotes escolhidos (fora do PDF Diretrizes)"
                if k in SEL_LOTS_TAB
                else "Não"
            )
        ),
        vad,
    ]

# ---------------------------------------------------------------- write
wb = openpyxl.Workbook()
HDR = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F4E78")
WRAP = Alignment(wrap_text=True, vertical="top")
NUM = "#,##0.00"


def sheet(title, header, data, widths=None, numcols=()):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HDR, FILL, WRAP
    for row in data:
        ws.append(row)
    for j, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for j in numcols:
        for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
            row[0].number_format = NUM
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


wb.remove(wb.active)
MH = [
    "Empresa",
    "Bloco (inventário 2025)",
    "Nº item",
    "Chave",
    "Descrição (2025)",
    "Loteamento / denominação",
    "Cidade",
    "Matrícula",
    "Área",
    "Unid.",
    "Qtde lotes 2025",
    "Valor grupo 2025 (R$)",
    "Existe em 2018?",
    "Item 2018",
    "Descrição 2018",
    "Qtde lotes 2018",
    "Valor 2018 (R$)",
    "Similaridade descrição 2018",
    "Revisar",
    "Existe no Estudo 2 (Rafael)?",
    "Documento de Rafael",
    "Valor do grupo segundo Rafael (R$)",
    "Valor avaliado por Rafael (R$)",
    "Observação de Rafael",
    "Passivo associado segundo Rafael",
    "Diferença Rafael − grupo (R$)",
    "Diferença %",
    "Situação",
    "Na lista de Rafael (virá para nós)?",
    "Valor adotado (R$)",
    "Decisão (P5)",
    "Motivo da decisão",
    "Destino 2018→2025",
    "Registro data-base",
    "Observação do Luís",
    "Situação da observação",
    "Nota de transcrição",
    "Observação do inventário",
    "Fonte 2025",
    "Fonte 2018",
    "Fonte Rafael",
]
K = [
    "empresa",
    "bloco",
    "item",
    "chave",
    "desc",
    "lote",
    "cidade",
    "matr",
    "area",
    "un",
    "qty25",
    "v25",
    "existe18",
    "item18",
    "desc18",
    "qty18",
    "v18",
    "sim18",
    "revisar",
    "existeR",
    "docR",
    "vR_grupo",
    "vR_aval",
    "obsR",
    "passR",
    "dif",
    "difp",
    "sit",
    "lista",
    "vadot",
    "decisao",
    "motivo",
    "destino",
    "regdb",
    "obsL",
    "sitL",
    "transc",
    "obs25",
    "fonte25",
    "fonte18",
    "fonteR",
]
data = [[m.get(k) for k in K] for m in master + gone + only_raf]
ws = sheet(
    "Mestre",
    MH,
    data,
    widths=[
        7,
        26,
        6,
        20,
        48,
        24,
        14,
        10,
        10,
        6,
        8,
        16,
        8,
        24,
        40,
        8,
        16,
        8,
        14,
        8,
        30,
        16,
        16,
        40,
        30,
        16,
        8,
        30,
        22,
        16,
        14,
        24,
        20,
        10,
        30,
        12,
        30,
        30,
        16,
        16,
        30,
    ],
    numcols=(12, 17, 22, 23, 26, 30),
)
for row in ws.iter_rows(min_row=2, min_col=27, max_col=27):
    row[0].number_format = "0.0%"

LH = [
    "Loteamento (2025)",
    "Cidade",
    "Quadra",
    "Lote",
    "Matrícula",
    "Valor 2025 (R$)",
    "Incluído no subtotal 2025",
    "Observação 2025",
    "Bloco",
    "Item de origem 2025",
    "Fonte 2025",
    "Existe em Rafael?",
    "Documento de Rafael",
    "Valor Rafael (R$)",
    "Página Rafael",
    "Diferença (R$)",
    "Situação",
    "Nº de linhas de Rafael casadas",
    "Nota",
    "Decisão",
    "Motivo da decisão",
    "Na lista de Rafael (virá para nós)?",
    "Valor adotado (R$)",
]
sheet(
    "Lotes",
    LH,
    lots_out,
    widths=[
        30,
        14,
        8,
        8,
        10,
        14,
        12,
        24,
        28,
        18,
        16,
        8,
        32,
        14,
        8,
        12,
        26,
        8,
        40,
        26,
        50,
        12,
        16,
    ],
    numcols=(6, 14, 16, 23),
)

CH = [
    "Chave",
    "Loteamento",
    "Cidade",
    "Bloco",
    "Qtde lotes 2018 (inventário)",
    "Valor 2018 (R$)",
    "Item 2018",
    "Qtde lotes 2025 (item)",
    "Lotes listados 2025",
    "Valor 2025 (R$)",
    "Lotes listados por Rafael",
    "Variação 2018→2025 (lotes)",
    "Observação",
]
cdata = []
for k, c in cnt.items():
    var = (
        (c["q25"] - c["q18"])
        if (c["q25"] is not None and c["q18"] is not None and c["bloco"] != "CONTR")
        else None
    )
    cdata.append(
        [
            k,
            c["lote"],
            c["cidade"],
            c["bloco"],
            c["q18"],
            c["v18"],
            c["item18"],
            c["q25"],
            c["l25"],
            c["v25"],
            c["qR"],
            var,
            c.get("obs"),
        ]
    )
sheet(
    "Loteamentos 2018x2025",
    CH,
    cdata,
    widths=[22, 30, 14, 10, 10, 14, 24, 10, 10, 14, 10, 10, 40],
    numcols=(6, 10),
)

PH = [
    "Item (inventário 2025)",
    "Natureza",
    "Valor (R$)",
    "Tratamento no documento",
    "Fonte",
    "Documento de Rafael que trata do mesmo passivo",
    "Valor segundo Rafael (R$)",
    "Valor adotado (R$)",
    "Decisão",
]
PASSDEC = {k: v for k, v in DEC.get("passivos", {}).items() if not k.startswith("_")}
pdata = []
for r in p25:
    v = num(r.get("valor (R$)"))
    rr = [
        x
        for x in raf_items
        if (x["sinal"] or "") == "-"
        and x["val"] is not None
        and v is not None
        and abs(x["val"] - v) < 1
        and {w for w in norm(x["desc"]).split() if len(w) > 3}
        & {w for w in norm(r.get("item")).split() if len(w) > 3}
    ]
    pdata.append(
        [
            r.get("item"),
            r.get("natureza"),
            v,
            col(r, "tratamento"),
            r.get("fonte no documento"),
            "; ".join(sorted({x["doc"] for x in rr})) or None,
            rr[0]["val"] if rr else None,
            *next(
                (
                    [pd["valor_adotado"], f"{pd['decisao']}. {pd['motivo']} Aprovada em {pd['aprovada_em']}."]
                    for pk, pd in PASSDEC.items()
                    if str(r.get("item") or "").startswith(pk)
                ),
                [None, None],
            ),
        ]
    )
sheet("Passivos", PH, pdata, widths=[40, 36, 14, 40, 20, 36, 14, 14, 50], numcols=(3, 7, 8))

cs = Counter(m["sit"] for m in master + gone + only_raf)
rdata = [["Itens (aba Mestre)", None]] + [[s, n] for s, n in cs.most_common()]
cl = Counter(x[16] for x in lots_out)
rdata += [[None, None], ["Lotes (aba Lotes)", None]] + [
    [s, n] for s, n in cl.most_common()
]
tot25 = sum(m.get("v25") or 0 for m in master)
totad = sum(m.get("vadot") or 0 for m in master if m.get("vadot") is not None)
pend = [m["chave"] for m in master if m.get("vadot") is None]
rdata += [
    [None, None],
    ["Bens 2025 — valor do grupo (R$)", tot25],
    ["Bens 2025 — soma dos valores adotados (R$)", totad],
    ["Itens 2025 ainda sem valor adotado", len(pend)],
    *(
        [f"Passivo — {k}: valor adotado (R$)", v["valor_adotado"]]
        for k, v in DEC.get("passivos", {}).items()
        if not k.startswith("_")
    ),
    ["Chaves pendentes", "; ".join(pend)],
]
sheet("Resumo", ["Situação", "Quantidade"], rdata, widths=[50, 40])

LEG = [
    [
        "Mestre",
        "Uma linha por item do inventário 2025 da SEI (93), mais os itens do inventário 2018 que não casaram com nenhum item de 2025 e os itens que só existem no Estudo 2 de Rafael.",
    ],
    [
        "Valor grupo 2025",
        "Valor total do item no doc 11 (31.12.2024). Vazio = sem valor no documento (placeholder).",
    ],
    [
        "Existe em 2018? / Item 2018",
        "Casamento automático: loteamentos por nome (apelidos tratados); veículos, máquinas, lotes e áreas por matrícula ou semelhança da descrição (coluna Similaridade). A coluna Revisar marca os casos a conferir à mão.",
    ],
    [
        "Valor do grupo segundo Rafael",
        "Valor que Rafael transcreve como 'balanço apresentado pelo grupo'. Quando difere do inventário, a coluna Nota de transcrição aponta.",
    ],
    [
        "Valor avaliado por Rafael",
        "Valor que Rafael adota na aba Avaliação do seu documento. Vazio = Rafael não avaliou o item (só transcreveu ou só escolheu lotes).",
    ],
    [
        "Situação",
        "Concorda = valor avaliado por Rafael igual ao do grupo; Discorda = diferente; Transcrito sem avaliação = Rafael repete o valor do grupo sem avaliar; Não abordado por Rafael = item sem documento de Rafael (veículos, máquinas, áreas urbanas).",
    ],
    [
        "Decisão (P5) / Motivo",
        "A preencher: Adoto grupo / Adoto Rafael / Caso extremo, com motivo.",
    ],
    [
        "Destino 2018→2025",
        "A preencher na Parte 4: venda, permuta, dação, consignação, perda, sem destino.",
    ],
    [
        "Registro data-base",
        "Número da linha em registro-data-base.md quando o item sofreu ajuste de data.",
    ],
    ["Observação do Luís / Situação", "Trazidas da planilha Google 2_0 (Parte 7)."],
    [
        "Lotes",
        "Lote a lote: inventário 2025 (438 lotes dos loteamentos + 26 lotes de terceiros + 8 lotes do acerto Rovílio) cruzado com as listas lote a lote de Rafael por loteamento + quadra + lote. Rafael escreve o aviso de APP dentro da célula da quadra; isso foi limpo para o cruzamento.",
    ],
    [
        "Loteamentos 2018x2025",
        "Contagem de lotes por loteamento nas três fontes. Variação negativa = lotes que saíram e precisam de destino (Parte 4). 'Saldo zero já em 2018' não é saída.",
    ],
    [
        "Passivos",
        "Passivos e deduções do doc 11 e o documento de Rafael que trata do mesmo valor.",
    ],
    [
        "Premissas",
        "1/6 em todas as instâncias; data-base 31.12.2024; 2018 é benchmark de existência, não de valor; sem reavaliação própria.",
    ],
]
sheet("Legenda", ["Campo / aba", "Significado"], LEG, widths=[30, 140])
wb.save(OUT)
print("saved", OUT)
print(
    "master rows",
    len(master),
    "gone",
    len(gone),
    "only_raf",
    len(only_raf),
    "lots",
    len(lots_out),
)
for s, n in cs.most_common():
    print(f"  {n:4d}  {s}")
print("lots:", dict(cl))
print("unmatched 2018 by block:", Counter(g["bloco"] for g in gone))
