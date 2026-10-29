"""Build a searchable text corpus from every source used in the SEI study.

Each text unit records the file, a location (page, sheet/row, table/row or
paragraph) and the text. PDFs without a text layer are listed separately so
they can be read page by page as images (P9: no source is skipped).

Usage: python corpus_fontes.py OUT_DIR
Writes OUT_DIR/corpus.jsonl, OUT_DIR/sem_texto.json and OUT_DIR/fontes.json.
"""

from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

import docx
import fitz
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
DRIVE = Path("G:/Meu Drive/02_Patrimonio_e_Sociedades/Cisao_Siviero")

# Drive material not copied to 01_fontes (read in place).
DRIVE_EXTRA = [
    DRIVE / "Documentos",
    DRIVE / "Material Rafael" / "Estudo 1",
    DRIVE / "Material Siviero" / "2018",
    DRIVE / "Material Siviero" / "Relatórios Mensais",
    DRIVE / "Material Siviero" / "Material antigo" / "inventário-2015",
]
EXTS = {".docx", ".xlsx", ".pdf"}
# OCR of scanned pages (ocr_respostas_docs.py); used only where the PDF has no text.
OCR_FILE = ROOT / "02_trabalho" / "ocr-respostas-docs" / "ocr.jsonl"


def load_ocr() -> dict[tuple[str, str], dict]:
    if not OCR_FILE.exists():
        return {}
    units = {}
    for line in OCR_FILE.open(encoding="utf-8"):
        u = json.loads(line)
        units[(Path(u["arquivo"]).name, u["local"])] = u
    return units


OCR = load_ocr()


def md5_10(path: Path) -> str:
    return hashlib.md5(path.read_bytes()).hexdigest()[:10]


def iter_docx(path: Path):
    d = docx.Document(str(path))
    body = d.element.body
    p_i = t_i = 0
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p_i += 1
            text = "".join(
                n.text or "" for n in child.iter() if n.tag.endswith("}t")
            ).strip()
            if text:
                yield f"P{p_i}", text
        elif tag == "tbl":
            t_i += 1
            table = docx.table.Table(child, d)
            for r_i, row in enumerate(table.rows, 1):
                cells, seen = [], set()
                for c in row.cells:
                    if id(c._tc) in seen:
                        continue
                    seen.add(id(c._tc))
                    cells.append(c.text.strip())
                text = " | ".join(c for c in cells if c)
                if text:
                    yield f"T{t_i}.L{r_i}", text


def iter_xlsx(path: Path):
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    for ws in wb.worksheets:
        for r_i, row in enumerate(ws.iter_rows(values_only=True), 1):
            text = " | ".join(str(v).strip() for v in row if v not in (None, ""))
            if text:
                yield f"{ws.title}!L{r_i}", text


def iter_pdf(path: Path, empty: list):
    d = fitz.open(str(path))
    md5 = md5_10(path)
    blank = []
    for p_i, page in enumerate(d, 1):
        text = page.get_text().strip()
        if len(text) < 30:
            # Still listed as a page without text (P9); OCR text only helps to search.
            blank.append(p_i)
            ocr = OCR.get((path.name, f"p.{p_i}"))
            if ocr and ocr["md5"] == md5 and ocr["texto"]:
                yield f"p.{p_i}", ocr["texto"], {"metodo": f"OCR: {ocr['metodo']}"}
            continue
        yield f"p.{p_i}", " ".join(text.split())
    if blank:
        empty.append(
            {"arquivo": str(path), "paginas_sem_texto": blank, "total_paginas": len(d)}
        )


def sources() -> list[Path]:
    files = [p for p in (ROOT / "01_fontes").rglob("*") if p.suffix.lower() in EXTS]
    for folder in DRIVE_EXTRA:
        files += [p for p in folder.rglob("*") if p.suffix.lower() in EXTS]
    return sorted(set(files))


def main(out_dir: str) -> None:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    empty: list = []
    index = []
    with (out / "corpus.jsonl").open("w", encoding="utf-8") as fh:
        for path in sources():
            suffix = path.suffix.lower()
            try:
                if suffix == ".docx":
                    units = iter_docx(path)
                elif suffix == ".xlsx":
                    units = iter_xlsx(path)
                else:
                    units = iter_pdf(path, empty)
                n = 0
                for loc, text, *extra in units:
                    unit = {"arquivo": str(path), "local": loc, "texto": text}
                    unit.update(*extra)
                    fh.write(json.dumps(unit, ensure_ascii=False) + "\n")
                    n += 1
                index.append({"arquivo": str(path), "md5": md5_10(path), "unidades": n})
            except Exception as exc:  # noqa: BLE001 - report every unreadable source
                index.append({"arquivo": str(path), "erro": repr(exc)})
    (out / "sem_texto.json").write_text(
        json.dumps(empty, ensure_ascii=False, indent=1), encoding="utf-8"
    )
    (out / "fontes.json").write_text(
        json.dumps(index, ensure_ascii=False, indent=1), encoding="utf-8"
    )
    print(
        f"{len(index)} fontes; {sum(i.get('unidades', 0) for i in index)} unidades; "
        f"{len(empty)} PDFs com páginas sem texto; {sum('erro' in i for i in index)} erros"
    )


if __name__ == "__main__":
    main(sys.argv[1])
