"""Lots that left the SEI after 31.12.2024 (arquivo 17) and are absent from the 2025 inventory.

The 2025 inventory (doc 11) is dated 31.12.2024 but shows the stock of late 2025.
For every row of arquivo 17 dated after the base date, the script looks for the
lot in doc 11: own stock, third-party table 47 and the Rovílio lots (obs. 1).

Usage: python saidas_pos_database.py   (prints the list; writes saidas-pos-database.xlsx)
"""

from __future__ import annotations

import datetime as dt
import re
import unicodedata
from pathlib import Path

import openpyxl

BASE = Path("C:/entrega-cisao-siviero")
W = BASE / "02_trabalho"
ARQ17 = next((BASE / "01_fontes/respostas-docs-2026-08-13").glob("17 - *.xlsx"))
DB = dt.datetime(2024, 12, 31)  # noqa: DTZ001 - spreadsheet dates are naive


def norm(s) -> str:
    s = (
        unicodedata.normalize("NFKD", str(s or ""))
        .encode("ascii", "ignore")
        .decode()
        .lower()
    )
    return re.sub(r"[^a-z0-9 ]", " ", s)


def key_num(s) -> str:
    s = str(s or "").strip().upper().replace("O", "0")
    s = re.sub(r"^[A-Z]+\s*-\s*", "", s)  # "H - 6371" -> "6371"
    return s.lstrip("0") or "0"


def words(s) -> set[str]:
    stop = {"loteamento", "de", "do", "da", "dos", "das", "rs", "e"}
    return {w for w in norm(s).split() if (len(w) > 2 or w in {"i", "ii", "iii"}) and w not in stop}


ROMANS = {"i", "ii", "iii"}


def same_development(a: set[str], b: set[str]) -> bool:
    """Share a name word; if arquivo 17 gives a roman numeral, it must match exactly."""
    ra, rb = a & ROMANS, b & ROMANS
    return bool((a - ROMANS) & (b - ROMANS)) and (not ra or ra == rb)


def main() -> None:
    wb = openpyxl.load_workbook(
        W / "sei-2025-itens.xlsx", read_only=True, data_only=True
    )
    it = wb["Lotes"].iter_rows(values_only=True)
    head = next(it)
    inv = [dict(zip(head, r)) for r in it]
    ws = openpyxl.load_workbook(ARQ17, read_only=True, data_only=True)[
        "Lotes vendidos (PDF)"
    ]
    it = ws.iter_rows(values_only=True)
    h17 = next(it)
    rows17 = [dict(zip(h17, r)) for r in it]

    out = []
    for r in rows17:
        d = r["Data"]
        if not isinstance(d, dt.datetime) or d <= DB:
            continue
        lw = words(str(r["Loteamento (PDF)"]).split("–")[0])
        hits = [
            x
            for x in inv
            if key_num(x["quadra"]) == key_num(r["Quadra"])
            and key_num(x["lote"]) == key_num(r["Lote"])
            and same_development(lw, words(x["loteamento"]))
        ]
        where = (
            "; ".join(f"{x['bloco'][:30]} ({x['loteamento']})" for x in hits)
            or "ausente do doc 11"
        )
        out.append(
            (
                r["Loteamento (PDF)"],
                r["Quadra"],
                r["Lote"],
                d.date().isoformat(),
                r["Comprador"],
                r["Observações"],
                r["Página"],
                where,
            )
        )

    absent = [o for o in out if o[7] == "ausente do doc 11"]
    print(
        f"Linhas do arquivo 17 depois de 31.12.2024: {len(out)}; ausentes do doc 11: {len(absent)}"
    )
    for o in out:
        print(" | ".join(str(x)[:60] for x in o))
    wbo = openpyxl.Workbook()
    wso = wbo.active
    wso.title = "Saidas apos data-base"
    wso.append(
        [
            "Loteamento (arq. 17)",
            "Quadra",
            "Lote",
            "Data",
            "Comprador",
            "Observações",
            "Página arq. 17",
            "No doc 11",
        ]
    )
    for o in out:
        wso.append(list(o))
    wbo.save(W / "saidas-pos-database.xlsx")


if __name__ == "__main__":
    main()
