"""Numbers for the Part 3 closing sheet of the SEI, read from planilha-mestre-sei.xlsx.

Usage: python ficha_parte3.py   (prints markdown tables)
"""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

import openpyxl

W = Path("C:/entrega-cisao-siviero/02_trabalho")


def rows(sheet: str):
    ws = openpyxl.load_workbook(
        W / "planilha-mestre-sei.xlsx", read_only=True, data_only=True
    )[sheet]
    it = ws.iter_rows(values_only=True)
    head = [str(h) for h in next(it)]
    return [dict(zip(head, r)) for r in it if any(v is not None for v in r)]


def brl(v: float) -> str:
    s = f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {s}"


def main() -> None:
    mestre = [
        m for m in rows("Mestre") if str(m["Chave"] or "").startswith("SEI-2025-")
    ]
    by_block = defaultdict(lambda: [0, 0.0, 0.0, 0])
    for m in mestre:
        b = by_block[m["Bloco (inventário 2025)"]]
        b[0] += 1
        b[1] += m["Valor grupo 2025 (R$)"] or 0
        b[2] += m["Valor adotado (R$)"] or 0
        b[3] += m["Valor adotado (R$)"] is None
    print("| Bloco | Itens | Grupo 2025 | Adotado | Diferença | Sem valor adotado |")
    print("|---|---|---|---|---|---|")
    tg = ta = 0.0
    for blk, (n, g, a, p) in sorted(by_block.items(), key=lambda kv: -kv[1][1]):
        tg += g
        ta += a
        print(f"| {blk} | {n} | {brl(g)} | {brl(a)} | {brl(a - g)} | {p or ''} |")
    print(
        f"| **Total** | {len(mestre)} | **{brl(tg)}** | **{brl(ta)}** | **{brl(ta - tg)}** | |"
    )

    dec = defaultdict(lambda: [0, 0.0])
    for m in mestre:
        d = str(m["Decisão (P5)"] or "sem decisão")
        key = d.split(":")[0][:70]
        dec[key][0] += 1
        dec[key][1] += m["Valor adotado (R$)"] or 0
    print("\n| Decisão | Itens | Valor adotado |")
    print("|---|---|---|")
    for k, (n, v) in sorted(dec.items(), key=lambda kv: -kv[1][1]):
        print(f"| {k} | {n} | {brl(v)} |")

    print("\nPassivos com valor adotado:")
    for p in rows("Passivos"):
        if p.get("Valor adotado (R$)") is not None:
            print(
                f"- {p['Item (inventário 2025)'][:70]}: doc 11 {brl(p['Valor (R$)'])}; adotado {brl(p['Valor adotado (R$)'])}"
            )
    res = {r["Situação"]: r["Quantidade"] for r in rows("Resumo")}
    print(
        "\nResumo:",
        {
            k: v
            for k, v in res.items()
            if k
            and ("Bens" in k or "Passivo" in k or "pendentes" in k or "sem valor" in k)
        },
    )
    pend_lots = [
        lo
        for lo in rows("Lotes")
        if str(lo.get("Decisão") or "").startswith(
            ("Pendente", "Ativo da SEI pendente")
        )
    ]
    print(
        "\nLotes com decisão pendente:",
        len(pend_lots),
        sorted({lo["Loteamento (2025)"] for lo in pend_lots}),
    )

    # Provisional net equity: doc 11 deductions, with the approved changes
    # (third-party liability R6/S1; tractor debt removed with the tractor, data-base nº 6)
    ws = openpyxl.load_workbook(
        W / "sei-2025-itens.xlsx", read_only=True, data_only=True
    )
    tot = {str(r[0]): r[1] for r in ws["Totais"].iter_rows(values_only=True) if r[0]}
    comp = {
        str(r[0]): r[6] for r in ws["Compromissos"].iter_rows(values_only=True) if r[0]
    }

    def find(prefix):
        return next(v for k, v in tot.items() if k.startswith(prefix))

    contratados = find("TOTAL GERAL DOS BENS - LOTEAMENTOS CONTRATADOS")
    obras = comp["Obras 01"]
    demais = find("Demais compromissos a pagar (tabela 50)")
    liquido_doc = find("TOTAL LÍQUIDO PATRIMONIAL DA EMPRESA")
    terc_doc = 16240000.0
    terc = res[
        "Passivo — Lotes de terceiros passíveis de indenização: valor adotado (R$)"
    ]
    trator = comp["Obs. 21"]
    bens = res["Bens 2025 — soma dos valores adotados (R$)"]
    liq = bens - contratados - terc - obras - (demais - trator)
    print("\nPatrimônio líquido provisório:")
    print(
        f"- doc 11: {brl(liquido_doc)} (bens {brl(res['Bens 2025 — valor do grupo (R$)'])} - contratados {brl(contratados)} - terceiros {brl(terc_doc)} - obras {brl(obras)} - demais compromissos {brl(demais)})"
    )
    print(
        f"- adotado até a Parte 3: {brl(liq)} (bens {brl(bens)} - contratados {brl(contratados)} - terceiros {brl(terc)} - obras {brl(obras)} - demais compromissos sem o trator {brl(demais - trator)})"
    )
    print(
        f"- diferença: {brl(liq - liquido_doc)}; 1/6 = {brl((liq - liquido_doc) / 6)}; 1/6 do líquido adotado = {brl(liq / 6)}"
    )


if __name__ == "__main__":
    main()
