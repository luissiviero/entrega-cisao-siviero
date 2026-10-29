"""Environmental-restrictions survey for the SEI (Parte 3, passo 1).

Rebuilds the survey from the sources: every lot or item with an environmental
issue, what each source says about it, whether a process, decision or report
exists, and where the same problem is discounted more than once.

Numbers come from the extraction workbooks and the source spreadsheets.
Every quoted piece of evidence is checked against the text corpus built by
corpus_fontes.py; the script stops if a quote is not found in its source.

Usage: python levantamento_ambiental.py CORPUS_DIR
Writes 02_trabalho/levantamento-ambiental-sei.xlsx.
"""

from __future__ import annotations

import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

BASE = Path("C:/entrega-cisao-siviero")
W = BASE / "02_trabalho"
F = BASE / "01_fontes"
RAF = F / "estudo-rafael/empreendimentos-xlsx"
ARQ17 = (
    F / "respostas-docs-2026-08-13/17 - SIVIERO EMPREENDIMENTOS IMOBILIÁRIOS LTDA. - "
    "LOTES VENDIDOS DE JANEIRO DE 2018 A OUTUBRO DE 2025 - REV.13.08.2026.xlsx"
)
INV18 = next((F / "inventario-2018").glob("18. 1. SIVIERO*.xlsx"))
OUT = W / "levantamento-ambiental-sei.xlsx"
CONTAS = F / "outros/Contas_Pagas_2023-07_a_2024-10.xlsx"
DATA_BASE = "2024-12-31"

# Short source labels -> substring of the corpus file path.
SRC = {
    "doc11": "inventario-2025\\11. inv",
    "doc11-antigo": "inventario-2025-antigo\\antigo 11",
    "doc18": "inventario-2018\\18. 1. siviero empreendimentos imobiliários ltda. = i n v e n t á r i o = e m p r e s a   d a   s i v i e r o   -   dezembro de 2018.docx",
    "inv2015": "siviero empreendimentos imobiliários ltda. = i n v e n t á r i o = e m p r e s a   d a   s i v i e r o   -   dezembro de 2015.docx",
    "arq1": "1 - grupo siviero - encaminhamento",
    "arq17": "17 - siviero empreendimentos imobiliários ltda. - lotes vendidos de janeiro de 2018 a outubro de 2025 - rev.13.08.2026.xlsx",
    "acertos": "acertos e acordos (gdoc)",
    "raf-exec": "estudo-rafael\\lotes loteamentos executados e registrados, belise ltda.pdf",
    "raf-terc": "estudo-rafael\\empreendimentos-pdf\\lotes de terceiros indenizaveis",
    "raf-bonetto": "estudo-rafael\\empreendimentos-pdf\\loteamento bonetto",
    "raf-compr": "estudo-rafael\\empreendimentos-pdf\\demias compromissos",
    "contas": "01_fontes\\outros\\contas_pagas_2023-07_a_2024-10.xlsx",
    "obs-luis": "comparativo patrimonial siviero empreendimentos e rafael",
}


def load_corpus(d: Path) -> dict[str, list[tuple[str, str]]]:
    units: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for line in (d / "corpus.jsonl").open(encoding="utf-8"):
        u = json.loads(line)
        units[u["arquivo"].lower()].append((u["local"], u["texto"]))
    return units


def squash(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip().lower()


class Evidence:
    """Quote checked against the corpus; returns 'label, local: "quote"'."""

    def __init__(self, corpus):
        self.corpus = corpus
        self.checked = 0

    def __call__(self, src: str, quote: str, local: str | None = None) -> str:
        sub = SRC[src]
        q = squash(quote)
        for path, units in self.corpus.items():
            if sub not in path:
                continue
            for loc, text in units:
                if (local is None or loc == local) and q in squash(text):
                    self.checked += 1
                    return f'{src} [{loc}]: "{quote}"'
        raise SystemExit(f"Citação não encontrada em {src} ({local}): {quote}")


def num(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return float(str(v).replace(".", "").replace(",", "."))


def nlote(s) -> str:
    s = str(s).strip().upper().replace("O", "0")
    return str(int(s)) if s.isdigit() else s


def nquadra(s) -> str:
    s = str(s).upper().replace(" ", "")
    return str(int(s)) if s.isdigit() else s


def rows(path: Path, sheet: str):
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)[sheet]
    it = ws.iter_rows(values_only=True)
    head = [str(h) for h in next(it)]
    return [dict(zip(head, r)) for r in it if any(v is not None for v in r)]


# --------------------------------------------------------------------------
# 1. Lots with an environmental issue in the 2025 inventory (doc 11)
# --------------------------------------------------------------------------
LOTES25 = rows(W / "sei-2025-itens.xlsx", "Lotes")
TERC_BLOCO = "Lotes de terceiros passíveis de indenização"


def categoria(r) -> str | None:
    obs = str(r.get("observação") or "")
    if str(r["bloco"]).startswith(TERC_BLOCO):
        return "Lote de terceiro indenizável"
    if obs == "ÁREA DE APP":
        return "APP (grupo: sem valor)"
    if obs == "ÁREA DE POSSÍVEL APP":
        return "Possível APP"
    return None


LOTS = [dict(r, categoria=categoria(r)) for r in LOTES25 if categoria(r)]
assert len(LOTS) == 4 + 7 + 7 + 26, len(LOTS)

# arquivo 17: sales, exchanges and notes per development
ARQ17_KEY = {
    "Fracionamento Simioni": "FRACIONAMENTO SIMIONI",
    "Lago Negro": "LAGO NEGRO",
    "San Marino": "SAN MARINO",
    "Firminio Dalzóchio": "LOTEAMENTO DALZÓCHIO",
    "Primavera": "LOTEAMENTO PRIMAVERA –",
    "Industrial Vicentino": "INDUSTRIAL VICENTINO",
    "Perini": "LOTEAMENTO PERINI",
    "Jardim Itália I – II – III": "JARDIM ITÁLIA",
}
VEND = rows(ARQ17, "Lotes vendidos (PDF)")
TOT17 = rows(ARQ17, "Totais por loteamento")


def arq17_lot(lot_name: str, quadra, lote):
    key = ARQ17_KEY[lot_name]
    out = []
    for v in VEND:
        if (
            key in str(v["Loteamento (PDF)"])
            and nquadra(v["Quadra"]) == nquadra(quadra)
            and nlote(v["Lote"]) == nlote(lote)
        ):
            d = v["Data"]
            ds = d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else str(d)
            out.append((ds, v["Comprador"], v["Observações"], v["Página"]))
    # Exchanges are recorded on the lot received; the APP lot is only named in
    # the note ("Permuta, com o lote 20, da quadra 11B").
    for v in VEND:
        m = re.search(r"com o lote (\d+), da quadra (\w+ ?\w?)", str(v["Observações"]))
        if (
            m
            and key in str(v["Loteamento (PDF)"])
            and nquadra(m.group(2)) == nquadra(quadra)
            and nlote(m.group(1)) == nlote(lote)
        ):
            d = v["Data"]
            ds = d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else str(d)
            got = (
                f"recebe o lote {v['Lote']}/{v['Quadra']} em troca: {v['Observações']}"
            )
            out.append((ds, v["Comprador"], got, v["Página"]))
    for k, ev_ in EXTRA_EVENTS.items():
        if k == (lot_name, nquadra(quadra), nlote(lote)):
            out.append(ev_)
    return out


# Events not in the lot columns of arquivo 17 (date, party, fact, page)
EXTRA_EVENTS = {
    ("Lago Negro", "11B", "19"): (
        "2025-10-14",
        "Lucas Zucoloto Dalsotto",
        "indenizado em espécie, R$ 580.000,00 (arquivo 17, aba Totais, nota do Lago Negro; arquivo 1 p.25)",
        "8",
    ),
    ("Primavera", "1199", "10"): (
        "2025-02",
        "SAFRA – São Francisco Investimentos e Participações Ltda.",
        "dação em pagamento no acordo do processo 5004974-83.2023 (só no arquivo 1 p.26; o arquivo 17 registra apenas o lote 11)",
        "arq1 p.26",
    ),
}


def arq17_note(lot_name: str) -> str:
    key = ARQ17_KEY[lot_name]
    for t in TOT17:
        if key in str(t["Loteamento (PDF)"]):
            return f"{t['Totais do PDF']} {t['Notas'] or ''}".strip()
    return ""


# Rafael: per-lot values (executed developments) and third-party lots
RAF_EXEC = rows(
    RAF / "LOTES LOTEAMENTOS EXECUTADOS E REGISTRADOS, BELISE LTDA.xlsx", "Lotes"
)
RAF_TERC = rows(RAF / "LOTES DE TERCEIROS INDENIZAVEIS, BELISE LTDA.xlsx", "Lotes")


def raf_exec(lot_name, quadra, lote):
    """All Rafael rows for the lot (transcription and revaluation)."""
    out = []
    first = lot_name.split()[0]
    for r in RAF_EXEC:
        q = str(r["Quadra"]).split("[")[0].strip()
        if (
            first.lower() in str(r["Loteamento"]).lower()
            and nquadra(q) == nquadra(quadra)
            and nlote(r["Lote"]) == nlote(lote)
        ):
            m = re.search(r"R\$ ([\d.]+,\d\d) escrito ao lado", str(r["Quadra"]))
            val = num(r["Valor (R$)"]) or (num(m.group(1)) if m else None)
            out.append((r["Loteamento"], val, r["Página"]))
    return out


# Value Rafael implies for APP lots he lists without a per-lot value:
# Lago Negro 7 x 1.100.000 = 7.700.000 (p.81, 84); Simioni 520.000 / 4 (p.20-21).
RAF_IMPLICIT = {"Lago Negro": 1100000.0, "Fracionamento Simioni": 130000.0}


# Rafael's position per development, quoted from his PDF
RAF_POS = {
    "Fracionamento Simioni": (
        "raf-exec",
        "diferença nominal em seu favor no montante de R$ 520.000,00",
        "p.21",
    ),
    "Lago Negro": (
        "raf-exec",
        "diferença nominal em seu favor no montante de R$ 7.700.000,00",
        "p.84",
    ),
    "San Marino": (
        "raf-exec",
        "diferença nominal em seu favor no montante de R$ 7.700.000,00",
        "p.94",
    ),
    "Firminio Dalzóchio": (
        "raf-exec",
        "diferença nominal em seu favor no montante de R$ 131.000,00",
        "p.100",
    ),
    "Primavera": (
        "raf-exec",
        "diferença nominal em seu favor no montante de R$ 335.000,00",
        "p.106",
    ),
    "Industrial Vicentino": (
        "raf-exec",
        "diferença nominal em seu favor no montante de R$ 1.125.000,00",
        "p.116",
    ),
}

# --------------------------------------------------------------------------
# 2. Documents found per development (process, decision, report)
# --------------------------------------------------------------------------


def documents(ev):
    """Evidence per development: list of (tipo, texto)."""
    return {
        "Fracionamento Simioni": [
            (
                "2015",
                ev(
                    "inv2015",
                    "Loteamento Fracionamento Simioni, localizado na cidade de Veranópolis – RS, num saldo de 05 lotes, ao valor de R$ 100.000,00, cada lote.",
                ),
            ),
            (
                "2018",
                ev(
                    "doc18",
                    "num saldo de 04 lotes, ao valor de R$ 250.000,00, cada lote. Sobre os lotes há restrições ambientais – Fonte de Água e Mato",
                ),
            ),
            (
                "2025",
                ev(
                    "doc11",
                    "Sobre os lotes há restrições ambientais – Fonte de Água e Mato.",
                ),
            ),
            (
                "TAC (citado, sem cópia)",
                ev(
                    "arq17",
                    "Conforme TAC firmado com o Ministério Público é necessário avisar aos possíveis futuros proprietários dos lotes da existência de uma nascente de água",
                ),
            ),
            (
                "Resposta ao Dr. Pablo",
                "Nenhum processo citado para o Fracionamento Simioni. Inferência: o item 18.6 parece tratar de outro empreendimento (área contratada), não dos 4 lotes em estoque: "
                + ev(
                    "arq1",
                    "Loteamento Simioni: sem contrato firmado, apenas uma conversação da empresa com os proprietários.",
                ),
            ),
            (
                "Contas pagas dez/2023 (a verificar)",
                "ART de bióloga para projeto de construção em lotes; observação cortada no relatório: "
                + ev("contas", "ART - PROJETO DE CONSTRUÇÃO LOTES S")
                + ". Se for o Simioni, reforça que há construção possível com restrição. Conferir na imagem do relatório de dez./2023.",
            ),
            (
                "Inferência",
                "O TAC obriga a avisar os compradores; não há cópia. 'Não impede a venda' é leitura do texto do arquivo 17, a confirmar com o TAC.",
            ),
        ],
        "Lago Negro": [
            (
                "2015",
                ev(
                    "inv2015",
                    "Processo Judicial referente ao surgimento de nascentes em 02 (dois) lotes localizados no Loteamento Lago Negro",
                ),
            ),
            (
                "2018 processo 10",
                ev(
                    "doc18",
                    "07 lotes em área de APP, que deverão ser preservados (Prejuízo de aproximadamente R$ 1.750.000,00)",
                ),
            ),
            (
                "2018 processo 15",
                ev(
                    "doc18",
                    "Promovido acordo em 2018, com a troca de 01 (um) lote dentro do mesmo Loteamento",
                ),
            ),
            (
                "2025 obs. 6",
                ev(
                    "doc11",
                    "basicamente com decisão já proferido pelo Ministério Público, com acato da Juíza da decisão",
                ),
            ),
            (
                "Resposta ao Dr. Pablo 3.3",
                ev(
                    "arq1",
                    "Foi realizada pericia ambiental, que confirmou que existe uma APP na referida quadra, atingindo os referidos lotes.",
                ),
            ),
            (
                "Resposta ao Dr. Pablo 3.3",
                ev(
                    "arq1",
                    "indenizou o montante de R$ 580.000,00 porque ele não aceitou nenhum dos lotes livres dentro do Loteamento",
                ),
            ),
            (
                "Resposta ao Dr. Pablo 3.3 (após a data-base)",
                ev(
                    "arq1",
                    "seja averbado o gravame perpetuo de preservação, bem como o Município não poderá autorizar construção nas mesmas unidades",
                ),
            ),
            (
                "arquivo 17",
                ev(
                    "arq17",
                    "indenizou, em espécie, o Lote 19, da Quadra 11B, ao Senhor Lucas Zucoloto Dalsotto, pelo valor de R$ 580.000,00, no dia 14 de outubro de 2025",
                ),
            ),
            (
                "Contas pagas jul/2023",
                ev("contas", "VISTORIA PROCESSO NASCENTES 07 LOTES -LOT LAG"),
            ),
            (
                "Versão antiga do doc 11 (removido na atual)",
                ev(
                    "doc11-antigo",
                    "poderá ser compensado 03 (três) lotes em área de Uso Institucional do Município, dentro da mesma área do Loteamento Lago Negro",
                ),
            ),
            (
                "Resposta ao Dr. Pablo 3.3 (após a data-base)",
                ev(
                    "arq1",
                    "o Ministério Publico foi contra esse acordo alinhado com o Município",
                ),
            ),
            (
                "Contagem divergente",
                "Obs. 6: indenização de 02 lotes; arquivo 17: 'supriu 04 lotes' e 'indenizar 03 lotes'; arquivo 1: 4 casos (lotes 15, 18, 19, 20); Dr. Pablo: '-8 lotes APP'. "
                + ev(
                    "arq17",
                    "A empresa supriu 04 lotes de sua propriedade e obrigou-se a indenizar 03 lotes vendidos para terceiros pela Família Candiago",
                ),
            ),
        ],
        "_farroupilha_marco": [
            (
                "Resposta ao Dr. Pablo 12 e 14",
                ev(
                    "arq1",
                    "Em função da aprovação da Lei do Marco Regulatório das APP’s do Município de Farroupilha, o entendimento é que os lotes estão dentro de faixas de APP de nascentes, córregos e outro georreferenciados pelo município",
                ),
            ),
            (
                "Resposta ao Dr. Pablo 3",
                ev(
                    "arq1",
                    "habilitar camada “Layer” no Marco Regulatório das APP’s no Sistema GEOFarroupilha",
                ),
            ),
            (
                "Proposta do Dr. Pablo (item 12)",
                ev(
                    "arq1",
                    "Acreditamos que o melhor seria deixarmos o destino deles para depois da conclusão do processo de cisão parcial em uma conta garantidora (escrow account).",
                ),
            ),
            (
                "Fonte externa (não está nas fontes do grupo)",
                "Lei Municipal de Farroupilha nº 4.814/2023, 'Marco Regulatório para as APPs consolidadas junto aos corpos hídricos do perímetro urbano'; listada pelo MP-RS entre leis municipais que reduziram limites de APP (mprs.mp.br, municipiosquereduziramlimitesdeapps.pdf). A conferir no texto oficial.",
            ),
        ],
        "Firminio Dalzóchio": [
            (
                "2018 processo 08",
                ev(
                    "doc18",
                    "Ação Civil Pública – Ivanor Carlos Dalzóchio e Siviero Empreendimentos Imobiliários Ltda. – Fonte da Água.",
                ),
            ),
            (
                "2015 processo 08 (mesmo número, outro loteamento)",
                ev(
                    "inv2015",
                    "Processo Judicial do Ministério Público contra a empresa Siviero – Ação Civil Pública – Loteamento Perini.",
                ),
            ),
            (
                "2025 obs. 3",
                ev(
                    "doc11",
                    "que questionas uma FONTE DE ÁGUA na área de terras do LOTEAMENTO DALZÓCHIO, em discussão de forma judicial desde meados dos anos 2006",
                ),
            ),
            (
                "Versão antiga do doc 11",
                "Mesma observação com R$ 500.000,00 (diff-inventario-2025-antigo-x-atual.md, linha da obs. Dalzóchio).",
            ),
            (
                "Resposta ao Dr. Pablo",
                "Sem relato sobre o processo do Dalzóchio (a Dra. Sandra relata Perini, Jardim Itália, Lago Negro e SAFRA).",
            ),
        ],
        "Perini": [
            (
                "2025 obs. 2",
                ev(
                    "doc11",
                    "Existe em tramitação uma infração ambiental aplicada sobre a área de terras do LOTEAMENTO PERINI",
                ),
            ),
            (
                "Resposta ao Dr. Pablo 3.1",
                ev(
                    "arq1",
                    "Ação de Execução de Obrigação de Fazer sob nº 5009118-66.2024.8.21.0048",
                ),
            ),
            (
                "Resposta ao Dr. Pablo 3.1",
                ev("arq1", "Valor da causa em dezembro de 2024 R$ 4.663.325,88"),
            ),
            (
                "Resposta ao Dr. Pablo 3.1",
                ev(
                    "arq1",
                    "ilegitimidade passiva porque o loteamento Perini foi entregue para o Município de Farroupilha 08.05.2017",
                ),
            ),
            (
                "Resposta ao Dr. Pablo 3.1",
                ev(
                    "arq1",
                    "Foi deferido mas ainda não nenhum perito nomeado aceitou o encargo.",
                ),
            ),
            (
                "Resposta ao Dr. Pablo 3.1 (a verificar)",
                ev("arq1", "Foram indicados a penhora três imóveis: matricula 34892")
                + ". Não se verificou se as matrículas 34892, 24305 e 24306 são lotes do estoque (por exemplo Perini 002–004/1688).",
            ),
            (
                "Histórico",
                "TAC de 04.12.2012; 2015: ACP 'Loteamento Perini' nº 048/1.08.0001924-5; 2018: nenhum processo do Perini (o mesmo número aparece como Dalzóchio); doc 11 antigo: infração citada sem valor; doc 11 atual: R$ 5.000.000. "
                + ev(
                    "inv2015",
                    "Processo Judicial do Ministério Público contra a empresa Siviero – Ação Civil Pública – Loteamento Perini.",
                ),
            ),
        ],
        "Jardim Itália I – II – III": [
            (
                "2025 obs. 4",
                ev(
                    "doc11",
                    "está em trâmite um Processo Judicial impetrado pelo Ministério Público referente ao trecho da canalização – área de APP",
                ),
            ),
            (
                "Resposta ao Dr. Pablo 3.2 (após a data-base)",
                ev(
                    "arq1",
                    "O Ministério Público já determinou o arquivamento do inquérito civil em julho de 2026",
                ),
            ),
            (
                "Contradição",
                "O doc 11 fala em 'Processo Judicial'; o arquivo 1 trata de inquérito civil (IC 00771.001.451/2020).",
            ),
        ],
        "Primavera": [
            (
                "Resposta ao Dr. Pablo 3.4",
                ev(
                    "arq1",
                    "indenizando mediante dação em pagamento de dois lotes urbanizados de nº 10 e 11 da quadra 1199",
                ),
            ),
            (
                "Contas pagas ago/2023",
                ev("contas", "LAUDO AMBIENTAL PROCESSO JUDICIAL - SAFRA - F"),
            ),
        ],
        "Acordi": [
            (
                "2018 processo 16",
                ev(
                    "doc18",
                    "Ação Civil Pública – Danos ambientais na área de terras da Família Acordi, de Farroupilha.",
                ),
            ),
            (
                "2018 processo 17",
                ev(
                    "doc18",
                    "Processo Administrativo da Prefeitura Municipal de Farroupilha – SEMMA, contra a empresa Siviero – Danos ambientais na área de terras da Família Acordi",
                ),
            ),
            (
                "2025 obs. 7",
                ev(
                    "doc11",
                    "A multa que reportava um valor de aproximadamente R$ 1.000.000,00 foi reduzida para aproximadamente R$ 80.000,00 (já pagos)",
                ),
            ),
            (
                "ACERTOS E ACORDOS (gdoc), item 29",
                ev("acertos", "TAC Ministério Público de Farroupilha - R$ 88.457,52"),
            ),
            (
                "ACERTOS E ACORDOS (gdoc), item 29",
                ev(
                    "acertos",
                    "Multa na Prefeitura de Farroupilha - R$ 61.920,30 Acerto com o Município - R$ 130.000,00",
                ),
            ),
            (
                "Divergência",
                "Valores não fecham: ACERTOS 88.457,52 + 61.920,30 + 130.000,00 = 280.377,82; doc 11 ≈ 80.000 (MP) + ≈ 70.000 (Município); multa de origem ≈ 1.200.000 (ACERTOS), ≈ 1.000.000 (doc 11), 1.499.492,79 + 1.069.049,00 (2018). Pagamento dos R$ 130.000 não comprovado; a linha do gdoc vem depois dos marcadores 30.) e 31.), a atribuição ao item 29 é interpretação. O doc 11 diz que ainda falta firmar o TAC.",
            ),
        ],
        "Colina do Retiro II": [
            (
                "arquivo 17",
                ev(
                    "arq17",
                    "A redução ocorreu em virtude de que praticamente 01 quadra foi considerada de área de app, com a doação de lotes para a Prefeitura e Corsan.",
                ),
            ),
            (
                "2025, áreas rurais 11 e 12",
                ev(
                    "doc11",
                    "Parte da área foi destinada para compensação ambiental do Loteamento Colina do Retiro II.",
                ),
            ),
        ],
        "Bonetto": [
            (
                "2025 contratado 06",
                ev("doc11", "podendo inviabilizar a implantação do Loteamento."),
            ),
            (
                "Resposta ao Dr. Pablo 18.5",
                ev(
                    "arq1",
                    "com grande custo estimado de obra e sem nenhuma trâmite de regularização",
                ),
            ),
            (
                "Rafael",
                "Soma dois cenários que se excluem: ativo de 30,5 lotes e multa. "
                + ev(
                    "raf-bonetto",
                    "R$ 518.500,00 – R$ 130.333,33 = R$ 388.166,67",
                ),
            ),
            (
                "Versão antiga do doc 11",
                "Sem a frase de risco de APP e sem multa; ambas entraram na versão atual (diff, linhas 306-308).",
            ),
        ],
        "Pasqual/Letti": [
            (
                "Versão antiga do doc 11",
                "Sem a frase de risco de APP e sem multa; ambas entraram na versão atual (diff, linhas 306-308).",
            ),
            (
                "Contas pagas 2024",
                ev("contas", "SERVICOS TOPOGRAFIA LOT. PASCOAL - FARROUPIL")
                + "; "
                + ev("contas", "2,5H BOBCAT-SERVICOS TESTES SOLO LAUDO GEOLOG")
                + ". Contrasta com 'sem nenhuma trâmite de regularização' (arquivo 1, 18.5).",
            ),
        ],
        "Rural 09": [
            (
                "2025 áreas rurais 09",
                ev(
                    "doc11",
                    "Uma área de terras rurais de MATO NATIVO, denominada de LOTE Nº 33",
                ),
            ),
            (
                "Decisão aprovada",
                "Pacote 3: valor de Rafael, R$ 250.000 (C0). Restrição de mato nativo sem processo.",
            ),
        ],
        "Rafael obs": [
            (
                "Rafael, DEMAIS COMPROMISSOS p.3",
                ev("raf-compr", "R$ 16.127.652,00 ÷ 6 = R$ 2.687.942,00")
                + ". Aceita as obrigações à razão de 1/6, sujeitas a conferência documental. Nos lotes de terceiros faz o inverso (trata os R$ 16.240.000 como patrimônio e cobra crédito de 1/6); no mesmo PDF diverge entre R$ 2.706.666,67 e R$ 2.706.333,33.",
            ),
        ],
        "Cotiporã": [
            (
                "Contas pagas 2024",
                ev("contas", "P.01/36-MULTA AMB COLONIA SAO BR")
                + ". Não aparece no doc 11.",
            ),
        ],
        "Outras despesas ambientais": [
            (
                "Contas pagas 2023-2024",
                ev("contas", "1200 MUDAS ARAUCARIA PROC ADM LOT")
                + "; "
                + ev("contas", "RALATORIO PRAD REP FLORESTAL LOT ALTO GRACAS")
                + "; "
                + ev("contas", "CUSTAS RECURSO REF MULTA R$ 40.000,00-LOT COL")
                + ". Reposição florestal em processos administrativos e PRAD do Alto das Graças; multa de R$ 40.000 de natureza não identificada.",
            ),
        ],
    }


# --------------------------------------------------------------------------
# 3. Overlaps: the same environmental problem discounted in more than one place
# --------------------------------------------------------------------------


def overlaps(lots_by_key, comp):
    """Each overlap: (nº, forma, desconto 1, desconto 2, valor em risco, situação, fontes)."""

    def v(*keys):
        return sum(lots_by_key[k]["valor (R$)"] for k in keys)

    terc = "Lote de terceiro indenizável"
    dz_terc = [
        k
        for k in lots_by_key
        if k[0] == "Firminio Dalzóchio" and lots_by_key[k]["categoria"] == terc
    ]
    pe_terc = [k for k in lots_by_key if k[0] == "Perini"]
    own_at_db = [
        ("Jardim Itália I – II – III", "1747", "1"),
        ("Primavera", "1199", "10"),
        ("Primavera", "1199", "11"),
        ("Industrial Vicentino", "1301", "3"),
    ]
    poss = [k for k in lots_by_key if lots_by_key[k]["categoria"] == "Possível APP"]
    sm = [k for k in poss if k[0] == "San Marino"]
    o = []
    o.append(
        (
            1,
            "Lote próprio na data-base lançado como passivo de terceiro",
            "Jardim Itália 001/1747 (vendido em 10/03/2025), Primavera 010 e 011/1199 (dados à SAFRA em fev.–mar./2025) e Vicentino 003/1301 (vendido em 19/09/2025): em 31.12.2024 eram estoque da SEI",
            "Os mesmos 4 lotes no passivo de terceiros indenizáveis (doc 11, tabela 47)",
            v(*own_at_db),
            "O doc 11 retrata out./2025 (arquivo 1 p.28: termo final 31/10/2025), não a data-base. Na data-base os 4 lotes eram ativo (no máximo 'possível APP', como os 7 próprios) e não havia terceiro a indenizar. Em contrapartida, na data-base existia o passivo da ação SAFRA (multa de 20% + 5 lotes pedidos), que o doc 11 não traz porque o acordo de fev./2025 o quitou com os lotes 10 e 11. Registro de data-base: 4 lotes + passivo SAFRA.",
            "arq17 p.1, 3, 8; arq1 p.25-26 e p.28; doc11 tabela 47",
        )
    )
    o.append(
        (
            2,
            "Mesmo problema em três lugares",
            "Dalzóchio: lote próprio 002/1292 tirado do total (possível APP) e 3 lotes de terceiros indenizáveis, dois na mesma quadra 1292",
            "Obs. 3: processo do MP sobre a fonte de água do Loteamento Dalzóchio",
            v(("Firminio Dalzóchio", "1292", "2"), *dz_terc) + comp["Obs. 3"],
            "Candidato. Nenhum documento diz se a fonte de água e a APP do Marco Regulatório são a mesma área. Pedir a delimitação (mapa GEOFarroupilha ou perícia).",
            "doc11 tabelas 30 e 47, obs. 3; arq17 p.6",
        )
    )
    o.append(
        (
            3,
            "Mesmo problema em dois lugares",
            "Perini: 4 lotes de terceiros indenizáveis na quadra 1469",
            "Obs. 2: execução do TAC de 2012 (área sob termo de suspensão da PATRAM)",
            v(*pe_terc) + comp["Obs. 2"],
            "Candidato. A defesa diz que o diagnóstico ambiental localiza a área preservada; não se sabe se inclui a quadra 1469.",
            "doc11 tabela 47, obs. 2; arq1 p.23-24",
        )
    )
    o.append(
        (
            4,
            "Possível mesmo processo em duas observações",
            "Obs. 2 (Perini, R$ 5.000.000)",
            "Obs. 3 (Dalzóchio, R$ 2.200.000)",
            comp["Obs. 2"] + comp["Obs. 3"],
            "Candidato. O nº 048/1.08.0001924-5 é 'ACP – Loteamento Perini' em 2015 e 'ACP – Dalzóchio – Fonte da Água' em 2018; o Perini some em 2018 e volta em 2025 (TAC de 2012). Pedir certidão de objeto e pé do 048/1.08.0001924-5 e do 5009118-66.2024.",
            "inv2015 proc. 08; doc18 proc. 08; arq1 p.23; doc11 obs. 2 e 3",
        )
    )
    o.append(
        (
            5,
            "Lote vendido pela SEI e listado como inedificável",
            "Perini 018 e 19/1469: vendidos pela SEI em 10/10/2024 (à Efetiva de Imóveis, 'diferença de área constatada')",
            "Os mesmos lotes na lista de terceiros indenizáveis (R$ 160.000 cada)",
            v(("Perini", "1469", "18"), ("Perini", "1469", "19")),
            "Contradição a esclarecer com o grupo: se o lote era inedificável por APP, a venda de out./2024 é o fato gerador da indenização; se não era, a indenização não existe. (Sem usar a data da lei municipal, que vem de fonte externa.)",
            "arq17 p.3; doc11 tabela 47",
        )
    )
    o.append(
        (
            6,
            "Lotes usados para pagar dívidas e cobrados como indenização",
            "Perini 011/1469 (Farina, acordo do processo Vicentino, 2019) e Jardim Itália 002/1747 (Ambitec, pagamento de serviços, 2021)",
            "Os mesmos lotes na lista de terceiros indenizáveis",
            v(("Perini", "1469", "11"), ("Jardim Itália I – II – III", "1747", "2")),
            "Candidato. A dívida foi quitada com o lote (o estoque já caiu); a indenização ao novo dono seria um segundo custo pelo mesmo lote. Decidir na Parte 6.",
            "arq17 p.2 e p.8; doc11 tabela 47",
        )
    )
    o.append(
        (
            7,
            "Critério de APP diferente na mesma quadra",
            "Lotes próprios com valor cheio: Vicentino 007–010/1303 (R$ 1.355.000 cada); Jardim Itália 002–004/1741 e 015–016/1747; Dalzóchio quadra 1292",
            "Vizinhos na mesma quadra tratados como inedificáveis: terceiros Vicentino 001, 012, 014, 016/1303 (R$ 1.355.000 cada); Jardim Itália 007, 009/1741 e 001, 002, 009/1747; próprio Vicentino 013/1303 (possível APP)",
            None,
            "Divergência para o pacote 5: sem delimitação por lote, o mesmo critério teria de valer para os lotes próprios das mesmas quadras. Pedir o mapa GEOFarroupilha.",
            "doc11 tabelas 33, 34 e 47",
        )
    )
    o.append(
        (
            8,
            "Contingência + bem retirado (forma 8 do plano)",
            "Lago Negro: 7 lotes da quadra 11 B sem valor (R$ 0 aprovado); lotes de reposição entregues (05/05 em 2019, 10/09 em 2019, 14/05 em 28/08/2025); R$ 580.000 pagos em 14/10/2025",
            "Obs. 6: PRAD, custas e honorários, R$ 80.000",
            comp["Obs. 6"],
            "Sem duplicidade, mas com omissão: a obrigação com os lotes 18 (Candiago) e 19 (Zucoloto) já existia na data-base (ACP de 2016) e não tem provisão no doc 11, que a dá como cumprida (retrata out./2025): R$ 580.000 + o lote 14/05. O prejuízo de R$ 1.750.000 de 2018 não é descontado de novo. A compensação de 3 lotes com o Município (doc 11 antigo) caiu em ago./2025 por oposição do MP. Contagens divergentes: 2, 3, 4 ou 8 lotes conforme a fonte.",
            "doc18 proc. 10; doc11 obs. 6; doc11 antigo P129; arq17 p.8-9; arq1 p.24-25",
        )
    )
    o.append(
        (
            9,
            "Acordo pago e cobrado de novo",
            "Acordi: TAC R$ 88.457,52 + multa da Prefeitura R$ 61.920,30 + acerto com o Município R$ 130.000,00 (ACERTOS E ACORDOS; atribuição das duas últimas ao Acordi é interpretação)",
            "Obs. 7: R$ 100.000 de atualização, TAC/PRAD, custas e honorários",
            comp["Obs. 7"],
            "Provavelmente sem duplicidade: a obs. 7 diz que o principal foi pago e cobra só o saldo. Valores das fontes não fecham entre si (ver Itens). Conferir com a sessão do ACERTOS E ACORDOS.",
            "doc11 obs. 7; acertos item 29",
        )
    )
    o.append(
        (
            10,
            "Ativo e multa do mesmo empreendimento em cenários que se excluem",
            "Rafael, Bonetto: ativo de 30,5 lotes × R$ 102.000 = R$ 3.111.000",
            "Grupo e Rafael: multa de R$ 782.000 por inviabilidade (risco de APP)",
            comp["Contratado 06"],
            "Se o loteamento é inviável, há multa e não há lotes; se é viável, há lotes e não há multa. Rafael soma os dois (líquido R$ 388.167 para o 1/6). Decidir na Parte 5 (contratados). Verificar se Rafael faz o mesmo no Pasqual e no Dr. Letti.",
            "raf-bonetto p.3-6; doc11 contratado 06",
        )
    )
    o.append(
        (
            11,
            "Compensação ambiental de um loteamento em área rural",
            "Áreas rurais 11 e 12 (lote 47): parte destinada à compensação ambiental do Colina do Retiro II; grupo R$ 20.000 e R$ 30.000",
            "Valor adotado no pacote 3: R$ 380.000 cada (Rafael, C0)",
            None,
            "Não é desconto duplo, é o oposto: a área comprometida foi avaliada cheia. Rafael transcreve a compensação mas não a considera no valor (Áreas rurais p.21). Registrar como risco para quem ficar com a área.",
            "doc11 áreas rurais 11 e 12; Rafael áreas rurais p.20-24; decisoes-sei.json ARUR-11 e 12",
        )
    )
    o.append(
        (
            12,
            "Lote de possível APP: divergência de tratamento",
            "Rafael lista os 7 lotes logo abaixo dos lotes escolhidos, fora do total escolhido, e cobra o valor cheio à parte (100%, não 1/6)",
            "O grupo tira os 7 do total de bens, sem processo ou laudo por lote",
            v(*poss),
            f"Não é duplicidade, é divergência (pacote 5). Total pelos valores listados no doc 11: San Marino {v(*sm):,.0f} + Dalzóchio 131.000 + Primavera 335.000 + Vicentino 1.125.000. No San Marino, Rafael escreve 'R$ 7.700.000,00' (p.94), frase copiada do Lago Negro: erro aparente, os 3 lotes somam {v(*sm):,.0f}. Com o número escrito por Rafael, a soma seria 9.291.000.",
            "raf-exec p.93-94, 99-100, 106, 115-116; doc11",
        )
    )
    return o


# --------------------------------------------------------------------------
# 4. Build the workbook
# --------------------------------------------------------------------------
HDR = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F4E78")
WRAP = Alignment(wrap_text=True, vertical="top")


def sheet(wb, title, header, data, widths):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HDR, FILL, WRAP
    for r in data:
        ws.append(list(r))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
            if isinstance(c.value, float):
                c.number_format = "#,##0.00"
    ws.freeze_panes = "A2"
    return ws


def main(corpus_dir: str) -> None:
    ev = Evidence(load_corpus(Path(corpus_dir)))
    docs = documents(ev)
    comp = {
        c["nº"]: num(c["valor total (R$)"])
        for c in rows(W / "sei-2025-itens.xlsx", "Compromissos")
    }

    lots_by_key = {}
    casos = []
    for r in LOTS:
        name, q, lt = r["loteamento"], r["quadra"], r["lote"]
        key = (name, nquadra(q), nlote(lt))
        lots_by_key[key] = r
        sales = arq17_lot(name, q, lt)
        sale_txt = "; ".join(f"{d} {b}: {o} (p.{p})" for d, b, o, p in sales)
        after = any(d > DATA_BASE for d, *_ in sales)
        if r["categoria"] == "Lote de terceiro indenizável":
            raf = [
                (x["Loteamento"], num(x["Valor (R$)"]), x["Página"])
                for x in RAF_TERC
                if nquadra(x["Quadra"]) == key[1] and nlote(x["Lote"]) == key[2]
            ]
            raf_pos = "Transcreve o valor e trata os R$ 16.240.000 como patrimônio: crédito de 1/6 a seu favor (p.4-5). Pede relatório técnico de APP por lote (p.2-4)."
        else:
            raf = raf_exec(name, q, lt)
            s = RAF_POS.get(name)
            raf_pos = ev(*s[:2], local=s[2]) if s else ""
            if name == "San Marino":
                raf_pos += " — erro aparente de Rafael (frase copiada do Lago Negro): os 3 lotes de possível APP somam R$ 730.000. Rafael os lista abaixo dos escolhidos, fora do total escolhido, e cobra o valor cheio à parte."
            elif r["categoria"] == "Possível APP":
                raf_pos += " — lote listado abaixo dos escolhidos, fora do total escolhido; Rafael cobra o valor cheio à parte."
        imp = RAF_IMPLICIT.get(name)
        raf_val = "; ".join(
            f"{v:,.0f} (p.{p})"
            if v
            else (
                f"sem valor por lote (p.{p}); implícito {imp:,.0f}"
                if imp
                else f"sem valor (p.{p})"
            )
            for _, v, p in raf
        )
        doc_list = docs.get(name, [])
        if (
            name
            in (
                "San Marino",
                "Firminio Dalzóchio",
                "Primavera",
                "Industrial Vicentino",
                "Perini",
                "Jardim Itália I – II – III",
            )
            and r["cidade"] == "Farroupilha"
        ):
            doc_list = doc_list + docs["_farroupilha_marco"]
        casos.append(
            (
                r["categoria"],
                name,
                r["cidade"],
                q,
                lt,
                num(r["valor (R$)"]),
                r["incluído no subtotal"],
                r["ordem no documento"],
                raf_val,
                raf_pos,
                sale_txt,
                "sim" if after else "",
                arq17_note(name),
                "\n".join(f"[{t}] {x}" for t, x in doc_list),
            )
        )

    # Cotiporã environmental fine, paid in 36 instalments (contas pagas)
    cot = [
        (int(m.group(1)), num(r["Valor Dupl."]), num(r["Valor NF"]))
        for r in rows(CONTAS, "Consolidado")
        if (m := re.search(r"P\.(\d+)/36-MULTA AMB COLONIA", str(r["Observações"])))
    ]
    cot_last, cot_parc, cot_nf = max(cot)[0], cot[0][1], cot[0][2]
    cot_open = round((36 - cot_last - 2) * cot_parc, 2)  # Nov and Dec/2024 assumed paid
    cot_txt = (
        f"NF R$ {cot_nf:,.2f} em 36 parcelas de R$ {cot_parc:,.2f}; última vista P.{cot_last:02d} (out./2024). "
        f"Em aberto na data-base, se nov. e dez./2024 foram pagas: {36 - cot_last - 2} parcelas (estimativa)."
    )
    # Declared total of observations 1-24 (doc 11, tabela 50); obs. 1 counts only the part still to pay
    comp_total = next(
        num(t["valor declarado (R$)"])
        for t in rows(W / "sei-2025-itens.xlsx", "Totais")
        if str(t["descrição"]).startswith("Demais compromissos a pagar (tabela 50)")
    )

    # Item-level cases (no lot list)
    itens = [
        ("Contingência ambiental", "Perini", comp["Obs. 2"], "Obs. 2", "Perini"),
        (
            "Contingência ambiental",
            "Dalzóchio (fonte de água)",
            comp["Obs. 3"],
            "Obs. 3",
            "Firminio Dalzóchio",
        ),
        (
            "Contingência ambiental",
            "Jardim Itália (canalização, APP)",
            comp["Obs. 4"],
            "Obs. 4",
            "Jardim Itália I – II – III",
        ),
        (
            "Compromisso de saneamento",
            "San Marino (ETE + doação de lote)",
            comp["Obs. 5"],
            "Obs. 5",
            None,
        ),
        (
            "Contingência ambiental",
            "Lago Negro (PRAD, custas)",
            comp["Obs. 6"],
            "Obs. 6",
            "Lago Negro",
        ),
        (
            "Contingência ambiental",
            "Acordi (atualização, TAC/PRAD)",
            comp["Obs. 7"],
            "Obs. 7",
            "Acordi",
        ),
        (
            "Contratado com risco de APP",
            "Bonetto (multa)",
            comp["Contratado 06"],
            "Contratado 06",
            "Bonetto",
        ),
        (
            "Contratado com risco de APP",
            "Pasqual (multa)",
            comp["Contratado 07"],
            "Contratado 07",
            "Pasqual/Letti",
        ),
        (
            "Contratado com risco de APP",
            "Dr. Letti (multa)",
            comp["Contratado 08"],
            "Contratado 08",
            "Pasqual/Letti",
        ),
        (
            "Restrição de mato nativo",
            "Área rural 09 (lote 33, Cotiporã, 15,1 ha)",
            None,
            "áreas rurais 09",
            "Rural 09",
        ),
        (
            "Multa ambiental fora do doc 11",
            f"Prefeitura de Cotiporã, Colônia São Brás: {cot_txt}",
            cot_open,
            "não consta",
            "Cotiporã",
        ),
        (
            "Despesas ambientais pequenas",
            "Reposição florestal, PRAD Alto das Graças, multa de R$ 40.000 não identificada",
            None,
            "não consta",
            "Outras despesas ambientais",
        ),
        (
            "Posição de Rafael",
            "Observações 1 a 24 (inclui as ambientais 2 a 7)",
            comp_total,
            "tabela 50",
            "Rafael obs",
        ),
        (
            "Redução de lotes por APP",
            "Colina do Retiro II (105 → 99 lotes)",
            None,
            "arquivo 17",
            "Colina do Retiro II",
        ),
        (
            "Compensação ambiental",
            "Áreas rurais 11 e 12 (lote 47)",
            None,
            "áreas rurais",
            "Colina do Retiro II",
        ),
    ]
    item_rows = []
    for cat, nome, val, onde, dk in itens:
        d = "\n".join(f"[{t}] {x}" for t, x in docs.get(dk, [])) if dk else ""
        item_rows.append((cat, nome, val, onde, d))

    # Processes table: 2015, 2018 (with values from the 2018 spreadsheet), 2025
    # The Processos sheet has a title block; the header is the row starting "Item no doc."
    ws18 = openpyxl.load_workbook(INV18, read_only=True, data_only=True)["Processos"]
    raw18 = [r for r in ws18.iter_rows(values_only=True)]
    h_i = next(i for i, r in enumerate(raw18) if r[0] == "Item no doc.")
    head18 = [str(h) for h in raw18[h_i]]
    p18 = {
        str(r[0]): dict(zip(head18, r))
        for r in raw18[h_i + 1 :]
        if r[0] and re.fullmatch(r"\d\d", str(r[0]))
    }
    proc = [
        (
            "Lago Negro, APP lotes 15–21 q. 11B",
            "ACP (MP)",
            "101/1.16.0002642-2 (2018); 5000454-62.2016.8.21.0101 (arq1)",
            "2016",
            num(p18["10"]["Valor da causa (R$)"]),
            "Perícia confirmou APP; acordos homologados; jun./2026: gravame perpétuo e encerramento.",
            "doc18 proc. 10; arq1 p.24-25",
        ),
        (
            "Lago Negro, Adriana Wasen, lote 20",
            "Indenizatória",
            "101/1.15.0003484-9",
            "2015",
            num(p18["15"]["Valor da causa (R$)"]),
            "Acordo 2018/2019: permuta pelo lote 05/05.",
            "doc18 proc. 15; arq17 p.9",
        ),
        (
            "Dalzóchio, fonte de água",
            "ACP (MP)",
            "048/1.08.0001924-5",
            "2006/2008",
            num(p18["08"]["Valor da causa (R$)"]),
            "2025: R$ 2.200.000 (versão antiga: R$ 500.000). Sem relato atual. Em 2015 o mesmo número aparece como 'ACP – Loteamento Perini'.",
            "doc18 proc. 08; inv2015 proc. 08; doc11 obs. 3",
        ),
        (
            "Acordi, danos ambientais",
            "ACP (MP) / Inquérito civil",
            "IC 00.771.00015/2018",
            "2018",
            num(p18["16"]["Valor da causa (R$)"]),
            "Principal pago (TAC R$ 88.457,52); falta atualização e PRAD.",
            "doc18 proc. 16; acertos item 29; doc11 obs. 7",
        ),
        (
            "Acordi, auto de infração SEMMA",
            "Administrativo",
            "PA 1198/2018, AI 003/2018",
            "2018",
            num(p18["17"]["Valor da causa (R$)"]),
            "ACERTOS cita multa da Prefeitura R$ 61.920,30 e acerto com o Município R$ 130.000 (pagamento dos R$ 130.000 não comprovado; doc 11 fala em ≈ R$ 70.000 já quitados).",
            "doc18 proc. 17; acertos item 29",
        ),
        (
            "Perini, execução de TAC de 2012",
            "Execução de obrigação de fazer (MP)",
            "5009118-66.2024.8.21.0048; embargos 5001467.46.2025.8.21.0048",
            "2008/2012 (TAC); execução dez./2024",
            4663325.88,
            "Embargos com perícia deferida, sem perito. Penhora indicada sobre as matrículas 34892, 24305 e 24306.",
            "arq1 p.23-24; doc11 obs. 2",
        ),
        (
            "Jardim Itália, canalização",
            "Inquérito civil (MP) — o doc 11 diz 'Processo Judicial'",
            "IC 00771.001.451/2020",
            "2020",
            None,
            "Arquivado em jul./2026 (após a data-base).",
            "arq1 p.24; doc11 obs. 4",
        ),
        (
            "SAFRA, distrato de loteamento",
            "Rescisória + indenizatória",
            "5004974-83.2023.8.21.0048",
            "2023",
            None,
            "Acordo em fev./2025: dação dos lotes 10 e 11/1199 (Primavera).",
            "arq1 p.25-26; arq17 p.1",
        ),
        (
            "Simioni, nascente",
            "TAC (MP)",
            "sem número",
            "?",
            None,
            "Obriga a avisar compradores da nascente e das restrições (inferência, sem cópia do TAC: a venda seria possível com o aviso).",
            "arq17 p.2",
        ),
        (
            "Marco Regulatório das APPs de Farroupilha",
            "Lei municipal",
            "Lei 4.814/2023 (fonte externa)",
            "2023",
            None,
            "Base citada pelo grupo para os 7 lotes de possível APP e os 26 lotes de terceiros. Sem processo, laudo ou notificação por lote.",
            "arq1 p.26, 28",
        ),
    ]

    ov = overlaps(lots_by_key, comp)

    # Checks
    def soma(cat):
        return sum(num(r["valor (R$)"]) or 0 for r in LOTS if r["categoria"] == cat)

    terc = soma("Lote de terceiro indenizável")
    poss = soma("Possível APP")
    amb_obs = sum(comp[k] for k in ("Obs. 2", "Obs. 3", "Obs. 5", "Obs. 6", "Obs. 7"))
    raf_terc = sum(num(x["Valor (R$)"]) for x in RAF_TERC)
    checks = [
        (
            "26 lotes de terceiros (doc 11, tabela 48)",
            16240000.0,
            terc,
            terc - 16240000.0,
        ),
        ("26 lotes de terceiros (Rafael)", 16240000.0, raf_terc, raf_terc - 16240000.0),
        (
            "7 lotes de possível APP (valor listado, fora do total)",
            2321000.0,
            poss,
            poss - 2321000.0,
        ),
        (
            "'Multas ambientais' segundo o Dr. Pablo (obs. 2, 3, 5, 6, 7)",
            7910000.0,
            amb_obs,
            amb_obs - 7910000.0,
        ),
        (
            "Quantidade de lotes no levantamento (4+7+7+26)",
            44.0,
            float(len(LOTS)),
            len(LOTS) - 44.0,
        ),
        ("Citações conferidas no texto das fontes", None, float(ev.checked), None),
    ]
    sales_after = [(c[1], c[3], c[4], c[10]) for c in casos if c[11]]
    # Facts after the base date that are not lot sales (arquivo 1)
    sales_after += [
        (
            "Lago Negro",
            "11 B",
            "15–21",
            "ago./2025: MP recusa a compensação de 3 lotes com o Município de Gramado, prevista na versão antiga do doc 11 (arq1 p.25)",
        ),
        (
            "Lago Negro",
            "11 B",
            "15–21",
            "jun./2026: acordo homologado, gravame perpétuo de preservação nos lotes 15 a 21; processo encerrado após a averbação (arq1 p.25)",
        ),
        (
            "Jardim Itália I – II – III",
            "",
            "",
            "jul./2026: MP arquiva o IC 00771.001.451/2020 (arq1 p.24)",
        ),
        (
            "Primavera (SAFRA)",
            "1199",
            "10, 11",
            "fev./2025: acordo na ação 5004974-83.2023 quita com 2 lotes o passivo que existia na data-base (arq1 p.25-26)",
        ),
    ]

    fontes = [
        (
            "Doc 11 (inventário 2025) e versão antiga",
            "Lido",
            "Lotes, contingências, contratados, observações 1–24.",
        ),
        (
            "Doc 18 (inventário 2018), docx e xlsx",
            "Lido",
            "Lista de 19 processos (soma das causas 5.187.216,15 x total declarado 6.278.430,81); Simioni a R$ 1.000.000; P69 tira as infrações do Acordi do valor dos loteamentos, mas os processos 16 e 17 entram no total de ações.",
        ),
        (
            "Inventário 2015 da SEI",
            "Lido",
            "Processos 2015; Perini com o mesmo número do Dalzóchio de 2018.",
        ),
        (
            "Arquivo 1, resposta ao Dr. Pablo",
            "Lido inteiro",
            "Relato da Dra. Sandra: Perini, Jardim Itália, Lago Negro, SAFRA; Marco Regulatório; nada sobre Dalzóchio, Simioni ou os 26 lotes por lote.",
        ),
        (
            "Arquivo 17, lotes vendidos 2018–2025",
            "Lido por script",
            "Vendas e permutas dos lotes do levantamento; TAC do Simioni; notas do Lago Negro e do Colina do Retiro II.",
        ),
        (
            "Arquivo 16, termos Susin/AGI (escaneado, lido por OCR)",
            "OCR (ocr-respostas-docs/), busca por termos",
            "Nenhuma menção a APP, nascente ou aos loteamentos do levantamento; o mesmo vale para os arquivos 2 a 15.",
        ),
        (
            "ACERTOS E ACORDOS (gdoc)",
            "Busca por termos",
            "Item 29, Acordi. A análise completa é da sessão própria.",
        ),
        (
            "Estudo 2 de Rafael: executados, terceiros, Bonetto, Diretrizes, Introdução, compromissos",
            "Lido",
            "Pede relatório técnico de APP; cobra valor cheio; Simioni R$ 520.000.",
        ),
        (
            "Estudo 1 de Rafael (Drive)",
            "Busca por termos",
            "Só modelos de laudo; nada específico.",
        ),
        (
            "Documentos do Drive (Buriti, Susin, termo de acordo, retiradas, avaliação com partilha)",
            "Busca por termos",
            "Nada sobre APP.",
        ),
        (
            "Relatórios mensais (xlsx de contas pagas, jul/2023–out/2024)",
            "Lido por script",
            "Despesas ambientais pequenas: vistoria de nascentes do Lago Negro, laudo SAFRA, multa FEPAM do Pedra Lisa, mudas doadas à SEMMA, reposição florestal em processos administrativos (Boff, Palavro, Piccoli), PRAD do Alto das Graças, ART 'projeto de construção lotes S…' (Simioni?), custas de recurso de multa de R$ 40.000. Multa ambiental de Cotiporã em 36 parcelas, fora do doc 11 (aba Itens).",
        ),
        (
            "Planilhas Google (observações do Luís)",
            "Lido",
            "Nota: verificar se os 35 lotes de 2018 do Lago Negro incluem os 7 de APP.",
        ),
        (
            "Planilha de custos da Dra. Sandra (anexo citado no arquivo 1)",
            "Não encontrada",
            "Citada como anexo da resposta; não está nas fontes.",
        ),
        (
            "22. Sócios… reunião 2017/2019 (Drive 2018)",
            "Não abriu",
            "Arquivo do Drive não baixado localmente (erro ao abrir).",
        ),
        (
            "Considerações finais Luciano 2019 (Drive 2018, escaneado)",
            "Não lido nesta etapa",
            "5 páginas sem texto; pode tratar dos processos de 2018.",
        ),
    ]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet(
        wb,
        "Lotes",
        [
            "Categoria",
            "Loteamento",
            "Cidade",
            "Quadra",
            "Lote",
            "Valor no doc 11 (R$)",
            "Incluído no total do grupo",
            "Onde no doc 11",
            "Valor de Rafael (página)",
            "Posição de Rafael",
            "Arquivo 17: venda, permuta ou dação",
            "Após a data-base?",
            "Arquivo 17: nota do loteamento",
            "Documentos encontrados (fonte [local]: citação)",
        ],
        casos,
        [18, 20, 12, 7, 6, 14, 18, 14, 18, 40, 45, 8, 40, 90],
    )
    sheet(
        wb,
        "Itens",
        [
            "Categoria",
            "Caso",
            "Valor no doc 11 (R$)",
            "Onde no doc 11",
            "Documentos encontrados",
        ],
        item_rows,
        [24, 34, 16, 14, 110],
    )
    sheet(
        wb,
        "Processos",
        [
            "Caso",
            "Tipo",
            "Número",
            "Início",
            "Valor da causa (R$)",
            "Situação conhecida",
            "Fontes",
        ],
        proc,
        [34, 22, 34, 10, 16, 60, 30],
    )
    sheet(
        wb,
        "Sobreposicoes",
        [
            "Nº",
            "Forma",
            "Desconto 1",
            "Desconto 2",
            "Valor envolvido no doc 11 (R$)",
            "Situação",
            "Fontes",
        ],
        ov,
        [5, 26, 50, 40, 16, 60, 30],
    )
    sheet(
        wb,
        "Apos data-base",
        ["Loteamento", "Quadra", "Lote", "Arquivo 17"],
        sales_after,
        [24, 8, 6, 100],
    )
    sheet(
        wb, "Fontes verificadas", ["Fonte", "Como", "Resultado"], fontes, [50, 20, 90]
    )
    sheet(
        wb,
        "Conferencias",
        ["Conferência", "Declarado", "Recalculado", "Diferença"],
        checks,
        [60, 16, 16, 12],
    )
    wb.save(OUT)
    print(
        f"{OUT.name}: {len(casos)} lotes, {len(item_rows)} itens, {len(proc)} processos, {len(ov)} sobreposições, {len(sales_after)} lotes com fato após a data-base; {ev.checked} citações conferidas"
    )
    for c in checks:
        print("  ", c)


if __name__ == "__main__":
    main(sys.argv[1])
